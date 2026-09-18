import argparse
import ctypes
import difflib
import json
import os
import re
import shlex
import shutil
import subprocess
import threading
import time
import unicodedata
import webbrowser
from ctypes import wintypes
from datetime import datetime
from pathlib import Path
from textwrap import wrap

APP_ROOT = Path(__file__).resolve().parent

CREATE_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)


def habilitar_dpi_awareness():
    if os.name != "nt" or not hasattr(ctypes, "windll"):
        return
    try:
        # Per-monitor v2: evita mistura de coordenadas entre monitores com escala diferente.
        ctypes.windll.user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4))
        return
    except Exception:
        pass
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(2)
        return
    except Exception:
        pass
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass


habilitar_dpi_awareness()

import pyautogui
from openpyxl import load_workbook


def popen_silencioso(argumentos, **kwargs):
    if CREATE_NO_WINDOW:
        kwargs["creationflags"] = kwargs.get("creationflags", 0) | CREATE_NO_WINDOW
    return subprocess.Popen(argumentos, **kwargs)

try:
    from PIL import Image, ImageDraw, ImageEnhance, ImageGrab, ImageOps

    PIL_DISPONIVEL = True
except Exception:
    Image = None
    ImageDraw = None
    ImageEnhance = None
    ImageGrab = None
    ImageOps = None
    PIL_DISPONIVEL = False

try:
    import cv2  # noqa: F401
    import numpy as np

    OPENCV_DISPONIVEL = True
except Exception:
    np = None
    OPENCV_DISPONIVEL = False

try:
    import pytesseract

    OCR_DISPONIVEL = PIL_DISPONIVEL
except Exception:
    pytesseract = None
    OCR_DISPONIVEL = False

try:
    from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
    from playwright.sync_api import sync_playwright

    PLAYWRIGHT_DISPONIVEL = True
except Exception:
    sync_playwright = None

    class PlaywrightTimeoutError(Exception):
        pass

    PLAYWRIGHT_DISPONIVEL = False

try:
    import tkinter as tk

    TK_DISPONIVEL = True
except Exception:
    tk = None
    TK_DISPONIVEL = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    PDF_RELATORIO_DISPONIVEL = True
except Exception:
    A4 = None
    canvas = None
    PDF_RELATORIO_DISPONIVEL = False


def localizar_tesseract_executavel():
    candidatos = [
        APP_ROOT / "Tesseract-OCR" / "tesseract.exe",
        APP_ROOT / "tesseract" / "tesseract.exe",
        APP_ROOT / "ocr" / "tesseract.exe",
        APP_ROOT / "tesseract.exe",
        APP_ROOT / "tauri_gerador" / "src-tauri" / "installer_support" / "Tesseract-OCR" / "tesseract.exe",
        APP_ROOT.parent / "tauri_gerador" / "src-tauri" / "installer_support" / "Tesseract-OCR" / "tesseract.exe",
    ]
    for candidato in candidatos:
        if candidato.exists():
            return candidato

    encontrado = shutil.which("tesseract")
    return Path(encontrado) if encontrado else None


def caminho_windows_curto(caminho):
    texto = str(caminho)
    if os.name != "nt":
        return texto
    try:
        buffer = ctypes.create_unicode_buffer(32768)
        resultado = ctypes.windll.kernel32.GetShortPathNameW(str(caminho), buffer, len(buffer))
        if resultado:
            return buffer.value
    except Exception:
        pass
    return texto


def configurar_tesseract_local():
    if pytesseract is None:
        return None

    executavel = localizar_tesseract_executavel()
    if not executavel:
        return None

    try:
        pytesseract.pytesseract.tesseract_cmd = caminho_windows_curto(executavel)
        tessdata = executavel.parent / "tessdata"
        if tessdata.exists():
            os.environ.setdefault("TESSDATA_PREFIX", caminho_windows_curto(tessdata))
    except Exception:
        return None
    return executavel


TESSERACT_EXECUTAVEL = configurar_tesseract_local()


ARQUIVO_MAPEAMENTO = Path("mapeamento_teclado.json")
ARQUIVO_PLANILHA = Path("db.xlsx")
PASTA_PLANILHAS = Path("planilhas")
PASTA_IMAGENS = Path("imagens")
PASTA_DIAGNOSTICOS_RECONHECIMENTO = Path("diagnosticos_reconhecimento")
PASTA_DIAGNOSTICOS_CHROME = Path("diagnosticos_chrome")
PASTA_RELATORIOS_LOOP = Path("relatorios_loop")
ARQUIVO_CANCELAMENTO = Path(".parar_automacao.flag")
INTERVALO_TEXTO = 0.01
INTERVALO_ENTRE_ITENS_PLANILHA = 0.5
TEMPO_ESTABILIZACAO_APOS_CLIQUE = 0.24
TEMPO_RESPIRO_ANTES_DIGITACAO = 0.05
TIPOS_SEQUENCIA_PLANILHA = {"sequencia_planilha"}
TIPOS_ESCRITA_PLANILHA = {"escrever_planilha"}
TIPOS_COMPARACAO_PLANILHA = {"comparar_planilha_tela"}
TIPOS_MARCADORES_LOOP = {"inicio_loop", "fim_loop"}
TIPOS_MOUSE = {"mouse_click"}
TIPOS_IMAGEM = {"clicar_imagem"}
TIPOS_OCR = {"procurar_texto_ocr"}
TIPOS_VERIFICACAO = {"verificar_estado_botao"}
TIPOS_CHROME = {"chrome_elemento_html"}
TIPOS_UIA = {"uia_elemento"}
TIPOS_PREPARACAO_TELA = {"abrir_url", "abrir_programa"}
TIPOS_SUPORTADOS_EXECUTOR = {
    "texto_livre",
    "tecla",
    "atalho",
    *TIPOS_SEQUENCIA_PLANILHA,
    *TIPOS_ESCRITA_PLANILHA,
    *TIPOS_COMPARACAO_PLANILHA,
    *TIPOS_MARCADORES_LOOP,
    *TIPOS_MOUSE,
    *TIPOS_IMAGEM,
    *TIPOS_OCR,
    *TIPOS_VERIFICACAO,
    *TIPOS_CHROME,
    *TIPOS_UIA,
    *TIPOS_PREPARACAO_TELA,
}
MAPA_PYAUTOGUI = {
    "space": "space",
    "enter": "enter",
    "tab": "tab",
    "backspace": "backspace",
    "esc": "esc",
    "shift": "shift",
    "shift_r": "shift",
    "ctrl": "ctrl",
    "ctrl_r": "ctrl",
    "alt": "alt",
    "alt_r": "alt",
    "cmd": "win",
    "cmd_r": "win",
    "up": "up",
    "down": "down",
    "left": "left",
    "right": "right",
    "delete": "delete",
    "insert": "insert",
    "home": "home",
    "end": "end",
    "page_up": "pageup",
    "page_down": "pagedown",
    "f8": "f8",
    "f9": "f9",
}


def safe_print(*args, **kwargs):
    try:
        print(*args, **kwargs)
    except (OSError, ValueError, BrokenPipeError):
        pass


class ExecutorAutomacao:
    def __init__(
        self,
        arquivo_mapeamento=ARQUIVO_MAPEAMENTO,
        arquivo_planilha=ARQUIVO_PLANILHA,
        arquivo_cancelamento=None,
        janela_alvo=None,
        target_context=None,
        modo_teste_reconhecimento=False,
    ):
        self.arquivo_mapeamento = Path(arquivo_mapeamento)
        self.arquivo_planilha = Path(arquivo_planilha)
        self.arquivo_cancelamento = Path(arquivo_cancelamento) if arquivo_cancelamento else None
        self.janela_alvo = str(janela_alvo or "").strip()
        self._janela_alvo_logada = ""
        self._janela_alvo_ativada_logada = ""
        self._ocr_indisponivel_logado = False
        self._janela_pendente_logada = False
        self._needle_maior_logada = False
        self._janela_contexto = ""
        self._janela_fixada = None
        self._ultima_regiao_reconhecimento = None
        self._ultimo_template_debug = None
        self._ultimo_diagnostico_reconhecimento = None
        self._cache_template_imagem = {}
        self._historico_reconhecimento = {}
        self._ultimo_clique_visual_em = 0.0
        self._ultima_posicao_clique_visual = None
        self._ultima_origem_captura_logada = ""
        self._ultima_captura_visual_info = None
        self._ultima_busca_template_info = None
        self._ultima_busca_ocr_info = None
        self._ultimo_ocr_global_em = 0.0
        self._cache_ocr_curto = {}
        # A leitura composta usa os mesmos pixels para todos os alvos OCR. Manter
        # a leitura por poucos instantes evita executar o Tesseract uma vez por
        # alvo, sem reaproveitar uma tela que ja pode ter mudado.
        self._cache_ocr_screenshot = {}
        self._diagnosticos_ocr_salvos = set()
        self._janela_alvo_perdida_logada = False
        self._conflitos_janela_evento_logados = set()
        self.contexto_variaveis = {}
        self.contexto_execucao_alvo = self.normalizar_contexto_execucao(target_context)
        self.modo_teste_reconhecimento = bool(modo_teste_reconhecimento)
        self._playwright_runtime = None
        self._chrome_browser = None
        self._chrome_context = None
        self._chrome_page = None
        self._chrome_source = ""
        self._chrome_launched_by_executor = False
        self._relatorio_loop = None
        self.cancelar = threading.Event()
        if modo_teste_reconhecimento:
            self.payload = {"eventos": [], "sequencias": {}}
            self.eventos = []
            self.sequencias = {}
            self.sequencia_inicial = "principal"
            self.configuracoes_sequencia = {}
            self.planilha_cache = {"linhas": [], "cabecalho": [], "colunas": [], "config": {}}
        else:
            self.payload = self.carregar_payload_mapeamento()
            self.eventos = self.payload.get("eventos", [])
            self.sequencias = self.agrupar_eventos_por_sequencia()
            self.sequencia_inicial = self.resolver_sequencia_inicial(self.payload)
            self.configuracoes_sequencia = self.carregar_configuracoes_sequencia(self.payload)
            self.planilha_cache = self.carregar_planilha()

    def normalizar_nome_sequencia(self, nome):
        texto = str(nome or "").strip()
        return texto or "principal"

    def agrupar_eventos_por_sequencia(self):
        sequencias = {}
        for evento in self.eventos:
            nome = self.normalizar_nome_sequencia(evento.get("sequencia"))
            evento["sequencia"] = nome
            sequencias.setdefault(nome, []).append(evento)
        return sequencias

    def carregar_payload_mapeamento(self):
        if not self.arquivo_mapeamento.exists():
            raise FileNotFoundError(f"Arquivo nao encontrado: {self.arquivo_mapeamento.resolve()}")

        dados = json.loads(self.arquivo_mapeamento.read_text(encoding="utf-8"))
        eventos = dados.get("eventos")
        if not isinstance(eventos, list) or not eventos:
            raise ValueError("Nenhum evento valido foi encontrado no roteiro salvo.")
        tipos_nao_suportados = sorted(
            {
                str((evento or {}).get("tipo") or "").strip()
                for evento in eventos
                if str((evento or {}).get("tipo") or "").strip()
                and str((evento or {}).get("tipo") or "").strip() not in TIPOS_SUPORTADOS_EXECUTOR
            }
        )
        if tipos_nao_suportados:
            raise ValueError(
                "Esta automacao possui blocos que este Agent ainda nao suporta: "
                + ", ".join(tipos_nao_suportados)
                + ". Atualize o runtime do Agent antes de executar."
            )
        return dados

    def resolver_sequencia_inicial(self, dados):
        preferida = self.normalizar_nome_sequencia(dados.get("sequencia_inicial"))
        if preferida in self.sequencias:
            return preferida
        if "principal" in self.sequencias:
            return "principal"
        return next(iter(self.sequencias.keys()))

    def carregar_configuracoes_sequencia(self, dados):
        bruto = dados.get("configuracoes_sequencia")
        if not isinstance(bruto, dict):
            return {}

        configuracoes = {}
        for nome_sequencia, configuracao in bruto.items():
            sequencia = self.normalizar_nome_sequencia(nome_sequencia)
            config = configuracao if isinstance(configuracao, dict) else {}
            modo_execucao = str(
                config.get("modo_execucao") or config.get("execution_mode") or "normal"
            ).strip().lower()
            if modo_execucao not in {"normal", "loop_planilha"}:
                modo_execucao = "normal"
            sequencia_apos_loop = str(
                config.get("sequencia_apos_loop") or config.get("after_loop_sequence") or ""
            ).strip()
            sequencia_apos_loop = (
                self.normalizar_nome_sequencia(sequencia_apos_loop) if sequencia_apos_loop else ""
            )
            if sequencia_apos_loop == sequencia:
                sequencia_apos_loop = ""
            if modo_execucao != "loop_planilha":
                sequencia_apos_loop = ""
            if modo_execucao != "normal" or sequencia_apos_loop:
                configuracoes[sequencia] = {
                    "modo_execucao": modo_execucao,
                    "sequencia_apos_loop": sequencia_apos_loop,
                }
        return configuracoes

    def obter_configuracao_sequencia(self, sequencia):
        nome = self.normalizar_nome_sequencia(sequencia)
        configuracao = self.configuracoes_sequencia.get(nome, {})
        return {
            "modo_execucao": str(configuracao.get("modo_execucao") or "normal").strip().lower() or "normal",
            "sequencia_apos_loop": str(configuracao.get("sequencia_apos_loop") or "").strip(),
        }

    def obter_sequencias_loop_configuradas(self):
        return {
            self.normalizar_nome_sequencia(nome_sequencia)
            for nome_sequencia, configuracao in self.configuracoes_sequencia.items()
            if str(configuracao.get("modo_execucao") or "").strip().lower() == "loop_planilha"
        }

    def localizar_configuracao_loop_planilha(self):
        sequencias_loop = self.obter_sequencias_loop_configuradas()
        if not sequencias_loop:
            return None

        evento_planilha = next(
            (
                evento
                for evento in self.eventos
                if self.normalizar_nome_sequencia(evento.get("sequencia")) in sequencias_loop
                and evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA
            ),
            None,
        )
        if evento_planilha is None:
            paginas = ", ".join(sorted(sequencias_loop))
            raise ValueError(
                f"As paginas em loop ({paginas}) precisam ter ao menos um bloco de planilha no grupo."
            )
        return {
            "sequencias_loop": sequencias_loop,
            "evento_planilha": evento_planilha,
        }
        return None

    def localizar_faixa_loop_por_blocos(self):
        marcador_inicio = None
        marcador_fim = None
        encontrou_marcador = False

        for nome_sequencia, eventos in self.sequencias.items():
            for indice, evento in enumerate(eventos):
                tipo = str(evento.get("tipo") or "").strip().lower()
                if tipo not in TIPOS_MARCADORES_LOOP:
                    continue
                encontrou_marcador = True
                if tipo == "inicio_loop" and marcador_inicio is None:
                    marcador_inicio = {
                        "sequencia": nome_sequencia,
                        "indice": indice,
                        "evento": evento,
                    }
                    continue
                if (
                    tipo == "fim_loop"
                    and marcador_inicio is not None
                    and nome_sequencia == marcador_inicio["sequencia"]
                    and indice > marcador_inicio["indice"]
                ):
                    marcador_fim = {
                        "sequencia": nome_sequencia,
                        "indice": indice,
                        "evento": evento,
                    }
                    break
            if marcador_fim is not None:
                break

        if not encontrou_marcador:
            return None

        if marcador_inicio is None or marcador_fim is None:
            raise ValueError(
                "Use os blocos Inicio do loop e Fim do loop na mesma pagina para limitar o trecho em loop."
            )

        sequencia_loop = self.normalizar_nome_sequencia(marcador_inicio["sequencia"])
        eventos_intervalo = self.sequencias.get(sequencia_loop, [])[
            marcador_inicio["indice"] + 1:marcador_fim["indice"]
        ]
        if not eventos_intervalo:
            raise ValueError("Adicione ao menos um bloco entre Inicio do loop e Fim do loop.")

        evento_planilha = next(
            (evento for evento in eventos_intervalo if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA),
            None,
        )
        if evento_planilha is None:
            raise ValueError(
                "O trecho entre Inicio do loop e Fim do loop precisa ter ao menos um bloco de planilha."
            )

        return {
            "sequencia_loop": sequencia_loop,
            "inicio_indice": marcador_inicio["indice"],
            "fim_indice": marcador_fim["indice"],
            "inicio_ordem": int(marcador_inicio["evento"].get("ordem", 0) or 0),
            "fim_ordem": int(marcador_fim["evento"].get("ordem", 0) or 0),
            "evento_planilha": evento_planilha,
        }

    def normalizar_id_planilha(self, arquivo):
        texto = str(arquivo or ARQUIVO_PLANILHA.name).strip().replace("\\", "/")
        return texto or ARQUIVO_PLANILHA.name

    def resolver_planilha(self, arquivo):
        workbook_id = self.normalizar_id_planilha(arquivo)
        caminho = Path(workbook_id)
        if caminho.is_absolute():
            return caminho, workbook_id
        if caminho.exists():
            return caminho, workbook_id
        if "/" not in workbook_id and "\\" not in workbook_id and workbook_id != ARQUIVO_PLANILHA.name:
            alternativa = PASTA_PLANILHAS / workbook_id
            if alternativa.exists():
                return alternativa, alternativa.as_posix()
        return caminho, workbook_id

    def resolver_imagem(self, imagem):
        imagem_id = str(imagem or "").strip().replace("\\", "/")
        if not imagem_id:
            return Path(""), ""
        caminho = Path(imagem_id)
        if caminho.is_absolute():
            return caminho, imagem_id
        if caminho.exists():
            return caminho, imagem_id
        if "/" not in imagem_id and "\\" not in imagem_id:
            alternativa = PASTA_IMAGENS / imagem_id
            if alternativa.exists():
                return alternativa, alternativa.as_posix()
        return caminho, imagem_id

    def limpar_saida_status(self, valor):
        return str(valor).replace("|", "/").replace("\n", " ").replace("\r", " ").strip()

    def aplicar_text_overrides(self, raw_overrides):
        bruto = str(raw_overrides or "").strip()
        if not bruto:
            return 0
        try:
            dados = json.loads(bruto)
        except Exception as erro:
            raise ValueError(f"Overrides de texto invalidos: {erro}") from erro
        if not isinstance(dados, dict):
            raise ValueError("Overrides de texto precisam ser um objeto JSON, exemplo: {\"4\":\"01072026\"}.")

        overrides = {
            str(chave).strip(): "" if valor is None else str(valor)
            for chave, valor in dados.items()
            if str(chave).strip()
        }
        if not overrides:
            return 0

        aplicados = 0
        for evento in self.eventos:
            if str(evento.get("tipo") or "").strip() != "texto_livre":
                continue
            ordem = str(evento.get("ordem", "") or "").strip()
            if ordem not in overrides:
                continue
            evento["texto"] = overrides[ordem]
            evento["text"] = overrides[ordem]
            aplicados += 1
        if aplicados:
            safe_print(f"Textos editaveis aplicados: {aplicados} bloco(s).", flush=True)
        return aplicados

    def corrigir_texto_corrompido(self, texto):
        bruto = str(texto or "")
        if not bruto:
            return ""
        marcadores = ("Ã", "Â", "â")
        melhor = bruto
        # Um texto pode passar por mais de uma camada de serializacao antes de
        # chegar ao worker. Corrigimos ate tres camadas, sempre aceitando apenas
        # uma conversao que reduza os marcadores de texto corrompido.
        for _ in range(3):
            marcadores_atuais = sum(melhor.count(marcador) for marcador in marcadores)
            if not marcadores_atuais:
                break
            proximo = melhor
            for origem in ("latin-1", "cp1252"):
                try:
                    convertido = melhor.encode(origem, errors="strict").decode("utf-8", errors="strict")
                except Exception:
                    continue
                marcadores_convertidos = sum(convertido.count(marcador) for marcador in marcadores)
                if convertido and marcadores_convertidos < marcadores_atuais:
                    proximo = convertido
                    break
            if proximo == melhor:
                break
            melhor = proximo
        return melhor

    def normalizar_texto_busca(self, texto):
        texto_corrigido = self.corrigir_texto_corrompido(texto)
        base = unicodedata.normalize("NFKD", str(texto_corrigido or ""))
        sem_acento = "".join(caractere for caractere in base if not unicodedata.combining(caractere))
        somente_validos = "".join(
            caractere.lower() if caractere.isalnum() else " "
            for caractere in sem_acento
        )
        return " ".join(somente_validos.split())

    def normalizar_texto_ocr(self, texto):
        return self.normalizar_texto_busca(texto)

    def normalizar_contexto_execucao(self, target_context):
        bruto = target_context
        if isinstance(bruto, str):
            texto = bruto.strip()
            if not texto:
                bruto = {}
            else:
                try:
                    bruto = json.loads(texto)
                except Exception:
                    bruto = {}
        if not isinstance(bruto, dict):
            bruto = {}

        rect = bruto.get("target_rect") or {}
        if not isinstance(rect, dict):
            rect = {}

        hwnd = int(bruto.get("target_hwnd", 0) or 0)
        titulo = str(bruto.get("target_title") or "").strip()
        classe = str(bruto.get("target_class") or "").strip()
        rect_normalizado = {
            "left": int(rect.get("left", 0) or 0),
            "top": int(rect.get("top", 0) or 0),
            "width": max(0, int(rect.get("width", 0) or 0)),
            "height": max(0, int(rect.get("height", 0) or 0)),
        }
        contexto = {
            "target_enabled": bool(bruto.get("target_enabled") and hwnd > 0),
            "target_hwnd": hwnd,
            "target_title": titulo,
            "target_class": classe,
            "target_rect": dict(rect_normalizado),
            "target_rect_atual": dict(rect_normalizado),
        }
        if contexto["target_enabled"] and titulo:
            self.janela_alvo = titulo
            self._janela_contexto = titulo
        return contexto

    def alvo_execucao_ativo(self):
        return bool(
            (self.contexto_execucao_alvo or {}).get("target_enabled")
            and int((self.contexto_execucao_alvo or {}).get("target_hwnd", 0) or 0) > 0
        )

    def descricao_alvo_execucao(self):
        if not self.alvo_execucao_ativo():
            return ""
        contexto = self.contexto_execucao_alvo or {}
        titulo = str(contexto.get("target_title") or self.janela_alvo or "").strip() or "Janela alvo"
        hwnd = int(contexto.get("target_hwnd", 0) or 0)
        return f"{titulo} / hwnd {hwnd}"

    def registrar_origem_captura(self, origem, janela=None):
        if origem == "janela_alvo":
            descricao = self.descricao_alvo_execucao() or str((janela or {}).get("title") or self.janela_alvo or "").strip()
            chave = f"janela_alvo::{descricao}"
            if self._ultima_origem_captura_logada != chave:
                safe_print(f"Captura usando janela alvo: {descricao}", flush=True)
                self._ultima_origem_captura_logada = chave
            return

        if self._ultima_origem_captura_logada != "tela_inteira":
            safe_print("Captura usando tela inteira.", flush=True)
            self._ultima_origem_captura_logada = "tela_inteira"

    def obter_janela_alvo_execucao_atual(self, obrigatoria=False):
        if not self.alvo_execucao_ativo():
            return None

        hwnd = int((self.contexto_execucao_alvo or {}).get("target_hwnd", 0) or 0)
        janela = self.descrever_janela(hwnd)
        if not janela:
            if not self._janela_alvo_perdida_logada:
                safe_print("Janela alvo perdida. Execucao pausada.", flush=True)
                self._janela_alvo_perdida_logada = True
            if obrigatoria:
                raise RuntimeError("Janela alvo perdida. Execucao pausada.")
            return None

        self._janela_alvo_perdida_logada = False
        contexto = self.contexto_execucao_alvo or {}
        contexto["target_title"] = str(janela.get("title") or contexto.get("target_title") or "").strip()
        contexto["target_class"] = str(janela.get("class_name") or contexto.get("target_class") or "").strip()
        contexto["target_rect_atual"] = {
            "left": int(janela.get("left", 0) or 0),
            "top": int(janela.get("top", 0) or 0),
            "width": max(0, int(janela.get("width", 0) or 0)),
            "height": max(0, int(janela.get("height", 0) or 0)),
        }
        self._janela_fixada = janela
        self._janela_contexto = contexto["target_title"]
        self.janela_alvo = contexto["target_title"]
        return janela

    def obter_janela_companheira_execucao(self, janela_principal=None):
        """Devolve passivamente um dialogo em primeiro plano do mesmo processo.

        Durante a sequencia, alguns sistemas abrem um modal separado da janela
        inicialmente escolhida. O alvo principal continua sendo a unica janela
        autorizada para clique e digitacao; esta janela complementar serve
        somente para capturar pixels e localizar o OCR/imagem no dialogo visivel.
        """
        if not self.alvo_execucao_ativo():
            return None

        principal = janela_principal or self.obter_janela_alvo_execucao_atual(obrigatoria=False)
        if not isinstance(principal, dict):
            return None
        ativa = self.obter_janela_ativa()
        if not isinstance(ativa, dict) or self.janela_parece_interface_flowdetect(ativa):
            return None

        hwnd_principal = int(principal.get("hwnd", 0) or 0)
        hwnd_ativa = int(ativa.get("hwnd", 0) or 0)
        pid_principal = int(principal.get("pid", 0) or 0)
        pid_ativa = int(ativa.get("pid", 0) or 0)
        if (
            hwnd_ativa <= 0
            or hwnd_ativa == hwnd_principal
            or pid_principal <= 0
            or pid_ativa != pid_principal
            or int(ativa.get("width", 0) or 0) < 120
            or int(ativa.get("height", 0) or 0) < 80
        ):
            return None
        return ativa

    def resolver_janela_evento(self, titulo_evento=None, registrar_conflito=False, priorizar_evento=False):
        titulo_evento = str(titulo_evento or "").strip()
        if priorizar_evento and titulo_evento:
            return titulo_evento
        if self.alvo_execucao_ativo():
            contexto = self.contexto_execucao_alvo or {}
            titulo_execucao = str(contexto.get("target_title") or self.janela_alvo or "").strip()
            if registrar_conflito and titulo_evento:
                conflito = self.normalizar_texto_busca(titulo_evento) != self.normalizar_texto_busca(titulo_execucao)
                chave_conflito = (titulo_evento, titulo_execucao)
                if conflito and chave_conflito not in self._conflitos_janela_evento_logados:
                    safe_print(
                        f"Bloco tentou usar a janela {titulo_evento} fora da janela alvo {titulo_execucao}. Mantendo a janela alvo da execucao.",
                        flush=True,
                    )
                    self._conflitos_janela_evento_logados.add(chave_conflito)
            return titulo_execucao
        return titulo_evento or self._janela_contexto or self.janela_alvo or ""

    def ajustar_coordenadas_para_janela_alvo(self, x, y):
        destino_x = int(x or 0)
        destino_y = int(y or 0)
        if not self.alvo_execucao_ativo():
            return destino_x, destino_y

        janela_atual = self.obter_janela_alvo_execucao_atual(obrigatoria=True)
        rect_base = (self.contexto_execucao_alvo or {}).get("target_rect") or {}
        base_left = int(rect_base.get("left", janela_atual.get("left", 0)) or 0)
        base_top = int(rect_base.get("top", janela_atual.get("top", 0)) or 0)
        delta_x = int(janela_atual.get("left", 0) or 0) - base_left
        delta_y = int(janela_atual.get("top", 0) or 0) - base_top
        destino_x += delta_x
        destino_y += delta_y

        esquerda = int(janela_atual.get("left", 0) or 0)
        topo = int(janela_atual.get("top", 0) or 0)
        direita = esquerda + max(0, int(janela_atual.get("width", 0) or 0))
        base = topo + max(0, int(janela_atual.get("height", 0) or 0))
        if not (esquerda <= destino_x <= direita and topo <= destino_y <= base):
            safe_print("Bloco tentou executar fora da janela alvo. Execucao pausada.", flush=True)
            raise RuntimeError("Bloco tentou executar fora da janela alvo. Execucao pausada.")

        return destino_x, destino_y

    def descrever_janela(self, hwnd):
        if not hasattr(ctypes, "windll"):
            return None

        user32 = ctypes.windll.user32
        if not user32.IsWindowVisible(hwnd):
            return None

        tamanho = user32.GetWindowTextLengthW(hwnd)
        if tamanho <= 0:
            return None

        buffer = ctypes.create_unicode_buffer(tamanho + 1)
        user32.GetWindowTextW(hwnd, buffer, tamanho + 1)
        titulo = buffer.value.strip()
        if not titulo:
            return None

        rect = wintypes.RECT()
        if not user32.GetWindowRect(hwnd, ctypes.byref(rect)):
            return None

        largura = int(rect.right - rect.left)
        altura = int(rect.bottom - rect.top)
        if largura <= 0 or altura <= 0:
            return None

        pid = wintypes.DWORD(0)
        try:
            user32.GetWindowThreadProcessId(wintypes.HWND(hwnd), ctypes.byref(pid))
        except Exception:
            pid = wintypes.DWORD(0)

        classe_buffer = ctypes.create_unicode_buffer(256)
        try:
            user32.GetClassNameW(hwnd, classe_buffer, 255)
            classe_janela = str(classe_buffer.value or "").strip()
        except Exception:
            classe_janela = ""

        return {
            "hwnd": int(hwnd),
            "title": titulo,
            "class_name": classe_janela,
            "left": int(rect.left),
            "top": int(rect.top),
            "width": largura,
            "height": altura,
            "pid": int(pid.value or 0),
        }

    def listar_janelas_visiveis(self):
        if not hasattr(ctypes, "windll"):
            return []

        user32 = ctypes.windll.user32
        janelas = []

        def callback(hwnd, _lparam):
            dados = self.descrever_janela(hwnd)
            if dados:
                janelas.append(dados)
            return True

        enum_proc = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)(callback)
        user32.EnumWindows(enum_proc, 0)
        return janelas

    def obter_janela_ativa(self):
        if not hasattr(ctypes, "windll"):
            return None
        user32 = ctypes.windll.user32
        hwnd = int(user32.GetForegroundWindow() or 0)
        if not hwnd:
            return None
        return self.descrever_janela(hwnd)

    def obter_janela_fixada(self):
        dados = self._janela_fixada or {}
        hwnd = int(dados.get("hwnd", 0) or 0)
        if not hwnd:
            return None
        janela = self.descrever_janela(hwnd)
        if janela:
            self._janela_fixada = janela
        return janela

    def localizar_janela_alvo(self, titulo_janela=None):
        if self.alvo_execucao_ativo():
            return self.obter_janela_alvo_execucao_atual(obrigatoria=False)

        titulo_consulta = str(titulo_janela or self._janela_contexto or self.janela_alvo or "").strip()
        termo = self.normalizar_texto_busca(titulo_consulta) if titulo_consulta else ""

        janela_fixada = self.obter_janela_fixada()
        if janela_fixada:
            titulo_fixado_normalizado = self.normalizar_texto_busca(janela_fixada.get("title", ""))
            if not termo:
                self._janela_contexto = str(janela_fixada.get("title") or "").strip()
                return janela_fixada
            if (
                termo == titulo_fixado_normalizado
                or termo in titulo_fixado_normalizado
                or difflib.SequenceMatcher(None, termo, titulo_fixado_normalizado).ratio() >= 0.52
            ):
                # Um Enter pode abrir um modal com titulo diferente. Quando ele
                # pertence ao mesmo processo, lemos esse dialogo visivel sem
                # ativar a janela principal nem trocar de aplicacao.
                janela_ativa = self.obter_janela_ativa()
                if (
                    isinstance(janela_ativa, dict)
                    and not self.janela_parece_interface_flowdetect(janela_ativa)
                    and int(janela_ativa.get("hwnd", 0) or 0) != int(janela_fixada.get("hwnd", 0) or 0)
                    and int(janela_ativa.get("pid", 0) or 0) > 0
                    and int(janela_ativa.get("pid", 0) or 0) == int(janela_fixada.get("pid", 0) or -1)
                ):
                    janela_fixada = janela_ativa
                    self._janela_fixada = janela_ativa
                self._janela_contexto = str(janela_fixada.get("title") or "").strip()
                return janela_fixada

        if not termo:
            return None

        melhor = None
        melhor_pontuacao = 0.0
        for janela in self.listar_janelas_visiveis():
            titulo_normalizado = self.normalizar_texto_busca(janela["title"])
            if not titulo_normalizado:
                continue

            if termo == titulo_normalizado:
                pontuacao = 1.0
            elif termo in titulo_normalizado:
                pontuacao = 0.96
            else:
                pontuacao = difflib.SequenceMatcher(None, termo, titulo_normalizado).ratio()

            if pontuacao > melhor_pontuacao:
                melhor = janela
                melhor_pontuacao = pontuacao

        if melhor and melhor_pontuacao >= 0.52:
            janela_ativa = self.obter_janela_ativa()
            if (
                janela_ativa
                and int(janela_ativa.get("pid", 0) or 0) > 0
                and int(janela_ativa.get("pid", 0) or 0) == int(melhor.get("pid", 0) or -1)
                and int(janela_ativa.get("hwnd", 0) or 0) != int(melhor.get("hwnd", 0) or 0)
            ):
                melhor = janela_ativa
            chave_log = f"{titulo_consulta}|{melhor['title']}"
            if self._janela_alvo_logada != chave_log:
                safe_print(f"Janela alvo capturada: {melhor['title']}", flush=True)
                self._janela_alvo_logada = chave_log
            self._janela_contexto = str(melhor.get("title") or "").strip()
            self._janela_fixada = melhor
            self._janela_pendente_logada = False
            return melhor

        if not self._janela_pendente_logada:
            safe_print(f"Aguardando janela alvo: {titulo_consulta}", flush=True)
            self._janela_pendente_logada = True
        return None

    def obter_regiao_busca(self, titulo_janela=None, margem=0):
        janela = (
            self.obter_janela_alvo_execucao_atual(obrigatoria=True)
            if self.alvo_execucao_ativo()
            else self.localizar_janela_alvo(titulo_janela)
        )
        if not janela:
            return None
        margem = max(0, int(margem or 0))
        return (
            int(janela["left"]) - margem,
            int(janela["top"]) - margem,
            int(janela["width"]) + margem * 2,
            int(janela["height"]) + margem * 2,
        )

    def ativar_janela_alvo(self, titulo_janela=None):
        if not hasattr(ctypes, "windll"):
            return None

        janela = (
            self.obter_janela_alvo_execucao_atual(obrigatoria=True)
            if self.alvo_execucao_ativo()
            else self.localizar_janela_alvo(titulo_janela)
        )
        if not janela:
            return None

        hwnd = int(janela.get("hwnd") or 0)
        if not hwnd:
            return janela

        user32 = ctypes.windll.user32
        kernel32 = ctypes.windll.kernel32
        SW_RESTORE = 9
        janela_ativa = self.obter_janela_ativa()

        if (
            janela_ativa
            and int(janela_ativa.get("pid", 0) or 0) > 0
            and int(janela_ativa.get("pid", 0) or 0) == int(janela.get("pid", 0) or -1)
        ):
            return janela_ativa

        try:
            if user32.IsIconic(wintypes.HWND(hwnd)):
                user32.ShowWindow(wintypes.HWND(hwnd), SW_RESTORE)
            try:
                foreground_hwnd = int(user32.GetForegroundWindow() or 0)
                foreground_thread = int(user32.GetWindowThreadProcessId(wintypes.HWND(foreground_hwnd), None) or 0)
                target_thread = int(user32.GetWindowThreadProcessId(wintypes.HWND(hwnd), None) or 0)
                current_thread = int(kernel32.GetCurrentThreadId() or 0)
                if foreground_thread and current_thread and foreground_thread != current_thread:
                    user32.AttachThreadInput(current_thread, foreground_thread, True)
                if target_thread and foreground_thread and target_thread != foreground_thread:
                    user32.AttachThreadInput(target_thread, foreground_thread, True)
            except Exception:
                foreground_thread = 0
                target_thread = 0
                current_thread = 0
            user32.SetForegroundWindow(wintypes.HWND(hwnd))
            user32.SetActiveWindow(wintypes.HWND(hwnd))
            user32.SetFocus(wintypes.HWND(hwnd))
            user32.BringWindowToTop(wintypes.HWND(hwnd))
            try:
                if foreground_thread and current_thread and foreground_thread != current_thread:
                    user32.AttachThreadInput(current_thread, foreground_thread, False)
                if target_thread and foreground_thread and target_thread != foreground_thread:
                    user32.AttachThreadInput(target_thread, foreground_thread, False)
            except Exception:
                pass
        except Exception:
            return janela

        chave_log = f"{self.resolver_janela_evento(titulo_janela)}|{janela['title']}"
        if self._janela_alvo_ativada_logada != chave_log:
            safe_print(f"Janela alvo ativada: {janela['title']}", flush=True)
            self._janela_alvo_ativada_logada = chave_log
        self._janela_fixada = janela
        return janela

    def preparar_janela_para_acao(self, titulo_janela=None, priorizar_evento=False):
        titulo = str(
            self.resolver_janela_evento(
                titulo_janela,
                registrar_conflito=True,
                priorizar_evento=priorizar_evento,
            ) or ""
        ).strip()
        if not titulo:
            return None
        janela = self.ativar_janela_alvo(titulo)
        if janela and str(janela.get("title") or "").strip():
            self._janela_contexto = str(janela.get("title") or "").strip()
        return janela

    def maximizar_janela_por_hwnd(self, hwnd):
        if not hwnd or not hasattr(ctypes, "windll"):
            return False
        try:
            user32 = ctypes.windll.user32
            SW_MAXIMIZE = 3
            user32.ShowWindow(wintypes.HWND(int(hwnd)), SW_MAXIMIZE)
            user32.BringWindowToTop(wintypes.HWND(int(hwnd)))
            user32.SetForegroundWindow(wintypes.HWND(int(hwnd)))
            return True
        except Exception:
            return False

    def janela_parece_interface_flowdetect(self, janela):
        texto = self.normalizar_texto_busca(
            f"{janela.get('title', '')} {janela.get('class_name', '')}"
            if isinstance(janela, dict)
            else ""
        )
        if not texto:
            return False
        termos_bloqueados = [
            "flowdetect studio",
            "flowdetect agent",
            "visual studio code",
            "windows terminal",
            "powershell",
            "cmd.exe",
            "prompt de comando",
            "terminal",
        ]
        return any(termo in texto for termo in termos_bloqueados)

    def janela_parece_navegador(self, janela):
        texto = self.normalizar_texto_busca(
            f"{janela.get('title', '')} {janela.get('class_name', '')}"
            if isinstance(janela, dict)
            else ""
        )
        if not texto:
            return False
        termos_navegador = [
            "google chrome",
            "chrome_widgetwin",
            "microsoft edge",
            "mozilla firefox",
            "brave",
            "opera",
        ]
        return any(termo in texto for termo in termos_navegador)

    def escolher_janela_ocr_automatica(self, texto_busca=""):
        if self.janela_alvo or self.alvo_execucao_ativo() or self._janela_contexto:
            return None

        texto_alvo = self.normalizar_texto_busca(texto_busca)
        termos_alvo = [termo for termo in re.split(r"\s+", texto_alvo) if len(termo) >= 3]
        janelas = [
            janela
            for janela in self.listar_janelas_visiveis()
            if int(janela.get("width", 0) or 0) >= 280
            and int(janela.get("height", 0) or 0) >= 180
            and not self.janela_parece_interface_flowdetect(janela)
        ]

        melhor = None
        melhor_score = -1
        for janela in janelas:
            titulo = self.normalizar_texto_busca(janela.get("title", ""))
            classe = self.normalizar_texto_busca(janela.get("class_name", ""))
            score = 0
            if self.janela_parece_navegador(janela):
                score += 40
            if "gov" in titulo or "fgts" in titulo:
                score += 35
            for termo in termos_alvo:
                if termo in titulo:
                    score += 8
            if "chrome_widgetwin" in classe:
                score += 5
            if score > melhor_score:
                melhor = janela
                melhor_score = score

        if melhor and melhor_score >= 35:
            titulo = str(melhor.get("title") or "").strip()
            safe_print(f"OCR inteligente selecionou janela automaticamente: {titulo}", flush=True)
            self._janela_contexto = titulo
            self._janela_fixada = melhor
            try:
                self.preparar_janela_para_acao(titulo)
            except Exception:
                pass
            return melhor
        return None

    def adotar_janela_ativa_para_ocr(self):
        if self.janela_alvo or self.alvo_execucao_ativo() or self._janela_contexto:
            return None
        janela = self.obter_janela_ativa()
        if not janela:
            return None
        if self.janela_parece_interface_flowdetect(janela):
            return None
        if int(janela.get("width", 0) or 0) < 280 or int(janela.get("height", 0) or 0) < 180:
            return None
        titulo = str(janela.get("title") or "").strip()
        if not titulo:
            return None
        safe_print(f"OCR inteligente usando janela ativa: {titulo}", flush=True)
        self._janela_contexto = titulo
        self._janela_fixada = janela
        return janela

    def adotar_janela_visivel_passiva_para_ocr(self):
        """Escolhe passivamente a janela externa mais proxima do topo para OCR.

        E um ultimo contexto seguro quando a execucao foi iniciada sem alvo.
        Nunca ativa, maximiza ou altera a ordem das janelas: apenas impede que
        o Tesseract precise analisar toda a area de trabalho para encontrar um
        dialogo pequeno.
        """
        if self.alvo_execucao_ativo() or self.janela_alvo or self._janela_contexto:
            return None
        try:
            candidatas = [
                janela
                for janela in self.listar_janelas_visiveis()
                if not self.janela_parece_interface_flowdetect(janela)
                and int(janela.get("hwnd", 0) or 0) > 0
                and int(janela.get("width", 0) or 0) >= 260
                and int(janela.get("height", 0) or 0) >= 180
            ]
        except Exception:
            candidatas = []
        if not candidatas:
            return None
        # EnumWindows preserva a ordem Z: a primeira janela valida e a mais
        # proxima do usuario. Nao usamos tamanho para evitar capturar outro
        # programa grande que esteja atras do sistema automatizado.
        candidata = candidatas[0]
        if not self.memorizar_contexto_visual_passivo(candidata):
            return None
        safe_print(
            f"OCR passivo adotou a janela visivel: {candidata.get('title', 'janela atual')}",
            flush=True,
        )
        return candidata

    def memorizar_contexto_visual_passivo(self, janela=None):
        """Guarda a primeira janela visual valida sem ativar ou mover nada.

        Quando a execucao e iniciada sem uma janela selecionada, o primeiro
        bloco hibrido ja sabe qual sistema esta em primeiro plano. Persistir
        essa referencia permite que os proximos OCRs continuem no mesmo sistema
        depois de um Enter abrir um dialogo, em vez de voltar para a tela inteira.
        """
        if self.alvo_execucao_ativo() or self.janela_alvo or self._janela_contexto:
            return None
        candidata = janela if isinstance(janela, dict) else self.obter_janela_ativa()
        if (
            not isinstance(candidata, dict)
            or self.janela_parece_interface_flowdetect(candidata)
            or int(candidata.get("hwnd", 0) or 0) <= 0
            or int(candidata.get("width", 0) or 0) < 260
            or int(candidata.get("height", 0) or 0) < 180
        ):
            return None
        titulo = str(candidata.get("title") or "").strip()
        if not titulo:
            return None
        self._janela_contexto = titulo
        self._janela_fixada = dict(candidata)
        safe_print(f"Contexto visual passivo memorizado: {titulo}", flush=True)
        return candidata

    def rolar_para_varredura_visual(self, titulo_janela=None, direcao="baixo", pixels=520):
        if self.cancelamento_solicitado():
            return False
        try:
            self.preparar_janela_para_acao(titulo_janela)
        except Exception:
            pass
        direcao = str(direcao or "baixo").strip().lower()
        pixels = max(120, int(abs(float(pixels or 520))))
        passos = max(1, min(8, int(round(pixels / 120))))
        sentido = -1 if direcao in {"baixo", "down", "descer"} else 1
        try:
            pyautogui.scroll(sentido * passos)
            safe_print(
                f"Varredura visual: rolando para {'baixo' if sentido < 0 else 'cima'} e tentando reconhecer novamente.",
                flush=True,
            )
            time.sleep(0.32)
            return True
        except Exception as erro:
            safe_print(f"Varredura visual nao conseguiu rolar a tela: {erro}", flush=True)
            return False

    def bounds_reconhecimento(self, local):
        if isinstance(local, dict):
            esquerda = int(local.get("left", local.get("x", 0)))
            topo = int(local.get("top", local.get("y", 0)))
            largura = max(24, int(local.get("width", 48)))
            altura = max(24, int(local.get("height", 32)))
            return esquerda, topo, largura, altura

        esquerda = int(getattr(local, "left", 0))
        topo = int(getattr(local, "top", 0))
        largura = max(24, int(getattr(local, "width", 48)))
        altura = max(24, int(getattr(local, "height", 32)))
        return esquerda, topo, largura, altura

    def normalizar_modo_area_busca(self, modo):
        texto = str(modo or "").strip().lower()
        aliases = {
            "janela": "target_window",
            "janela_alvo": "target_window",
            "janela_alvo_inteira": "target_window",
            "window": "target_window",
            "target_window": "target_window",
            "area": "custom_area",
            "area_personalizada": "custom_area",
            "custom": "custom_area",
            "custom_area": "custom_area",
            "ultima": "last_found",
            "ultima_posicao": "last_found",
            "cache": "last_found",
            "last_found": "last_found",
        }
        return aliases.get(texto, "target_window")

    def formatar_regiao_log(self, regiao):
        if not regiao:
            return "tela_inteira"
        try:
            esquerda, topo, largura, altura = [int(valor) for valor in regiao]
        except Exception:
            return str(regiao)
        return f"x={esquerda} y={topo} w={largura} h={altura}"

    def limitar_regiao(self, regiao, limite=None):
        if not regiao:
            return None
        try:
            esquerda, topo, largura, altura = [int(valor) for valor in regiao]
        except Exception:
            return None
        largura = max(0, largura)
        altura = max(0, altura)
        if largura <= 0 or altura <= 0:
            return None
        if limite:
            try:
                limite_esquerda, limite_topo, limite_largura, limite_altura = [int(valor) for valor in limite]
            except Exception:
                limite_esquerda = limite_topo = 0
                limite_largura = limite_altura = 0
            limite_direita = limite_esquerda + max(0, limite_largura)
            limite_base = limite_topo + max(0, limite_altura)
            direita_original = esquerda + largura
            base_original = topo + altura
            esquerda = max(limite_esquerda, esquerda)
            topo = max(limite_topo, topo)
            direita = min(limite_direita, direita_original)
            base = min(limite_base, base_original)
            largura = max(0, direita - esquerda)
            altura = max(0, base - topo)
        if largura <= 0 or altura <= 0:
            return None
        return (int(esquerda), int(topo), int(largura), int(altura))

    def construir_regiao_personalizada(self, area_config=None, janela_info=None):
        area = area_config or {}
        if isinstance(area, (list, tuple)) and len(area) >= 4:
            area = {
                "x": area[0],
                "y": area[1],
                "width": area[2],
                "height": area[3],
            }
        if not isinstance(area, dict):
            return None
        largura = max(0, int(area.get("width", area.get("largura", area.get("w", 0))) or 0))
        altura = max(0, int(area.get("height", area.get("altura", area.get("h", 0))) or 0))
        if largura <= 0 or altura <= 0:
            return None
        deslocamento_x = int(area.get("x", area.get("left", area.get("esquerda", 0))) or 0)
        deslocamento_y = int(area.get("y", area.get("top", area.get("topo", 0))) or 0)
        if janela_info:
            esquerda = int(janela_info.get("left", 0) or 0) + deslocamento_x
            topo = int(janela_info.get("top", 0) or 0) + deslocamento_y
            limite = (
                int(janela_info.get("left", 0) or 0),
                int(janela_info.get("top", 0) or 0),
                int(janela_info.get("width", 0) or 0),
                int(janela_info.get("height", 0) or 0),
            )
            return self.limitar_regiao((esquerda, topo, largura, altura), limite=limite)
        return self.limitar_regiao((deslocamento_x, deslocamento_y, largura, altura))

    def chave_historico_reconhecimento(self, imagem_id, janela_alvo=None, bloco_id=None):
        imagem_norm = str(imagem_id or "").strip().replace("\\", "/").lower()
        janela_norm = self.normalizar_texto_busca(janela_alvo or self._janela_contexto or self.janela_alvo or "")
        bloco_norm = str(bloco_id or "").strip().lower() or "default"
        return f"{bloco_norm}::{imagem_norm}::{janela_norm}"

    def registrar_reconhecimento_bem_sucedido(self, imagem_id, local, janela_alvo=None, bloco_id=None, janela_info=None):
        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        chave = self.chave_historico_reconhecimento(imagem_id, janela_alvo=janela_alvo, bloco_id=bloco_id)
        dados = {
            "left": esquerda,
            "top": topo,
            "width": largura,
            "height": altura,
            "updated_at": time.time(),
        }
        janela_base = janela_info or (
            self.obter_janela_alvo_execucao_atual(obrigatoria=False)
            if self.alvo_execucao_ativo()
            else None
        )
        if janela_base:
            dados["window_left"] = int(janela_base.get("left", 0) or 0)
            dados["window_top"] = int(janela_base.get("top", 0) or 0)
            dados["window_width"] = int(janela_base.get("width", 0) or 0)
            dados["window_height"] = int(janela_base.get("height", 0) or 0)
            dados["left_rel"] = int(esquerda - int(janela_base.get("left", 0) or 0))
            dados["top_rel"] = int(topo - int(janela_base.get("top", 0) or 0))
        self._historico_reconhecimento[chave] = dados

    def obter_regiao_historica_reconhecimento(self, imagem_id, janela_alvo=None, bloco_id=None, janela_info=None):
        chave = self.chave_historico_reconhecimento(imagem_id, janela_alvo=janela_alvo, bloco_id=bloco_id)
        dados = self._historico_reconhecimento.get(chave)
        if not dados:
            return None
        janela_base = janela_info or (
            self.obter_janela_alvo_execucao_atual(obrigatoria=False)
            if self.alvo_execucao_ativo()
            else None
        )
        if janela_base and "left_rel" in dados and "top_rel" in dados:
            esquerda = int(janela_base.get("left", 0) or 0) + int(dados.get("left_rel", 0) or 0)
            topo = int(janela_base.get("top", 0) or 0) + int(dados.get("top_rel", 0) or 0)
        else:
            esquerda = int(dados.get("left", 0))
            topo = int(dados.get("top", 0))
        largura = max(24, int(dados.get("width", 0) or 24))
        altura = max(24, int(dados.get("height", 0) or 24))
        margem_x = max(40, int(largura * 1.8))
        margem_y = max(28, int(altura * 1.8))
        limite = None
        if janela_base:
            limite = (
                int(janela_base.get("left", 0) or 0),
                int(janela_base.get("top", 0) or 0),
                int(janela_base.get("width", 0) or 0),
                int(janela_base.get("height", 0) or 0),
            )
        return self.limitar_regiao((
            esquerda - margem_x,
            topo - margem_y,
            largura + margem_x * 2,
            altura + margem_y * 2,
        ), limite=limite)

    def marcar_area_reconhecida(
        self,
        local,
        duracao=1.1,
        bloqueante=False,
        tema="visual",
        etiqueta=None,
        margem=16,
    ):
        if not TK_DISPONIVEL:
            safe_print("Marcador visual indisponivel: Tk nao esta carregado.", flush=True)
            return False

        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        exibido = self.marcar_regiao_absoluta(
            esquerda,
            topo,
            largura,
            altura,
            duracao=duracao,
            bloqueante=bloqueante,
            tema=tema,
            etiqueta=etiqueta,
            margem=margem,
        )
        # A primeira criacao da janela transparente pode ser descartada pelo
        # Windows enquanto o Studio esta minimizado para o botao Procurar.
        # Repetimos uma unica vez para o usuario sempre ver a confirmacao.
        if exibido is False and self.feedback_visual_habilitado():
            safe_print("Marcador visual nao abriu na primeira tentativa; repetindo.", flush=True)
            time.sleep(0.06)
            exibido = self.marcar_regiao_absoluta(
                esquerda,
                topo,
                largura,
                altura,
                duracao=duracao,
                bloqueante=bloqueante,
                tema=tema,
                etiqueta=etiqueta,
                margem=margem,
            )
        return bool(exibido)

    def feedback_visual_habilitado(self):
        """Indica se os contornos de OCR/reconhecimento podem ser exibidos.

        O feedback faz parte da execucao assistida: o usuario precisa conseguir
        ver o que foi lido em azul e o que foi confirmado em verde. Os
        marcadores sao configurados como janelas passivas mais abaixo, portanto
        nao precisam mais ficar restritos ao botao de teste.
        """
        return bool(TK_DISPONIVEL and not bool(getattr(self, "ocultar_feedback_visual", False)))

    def duracao_confirmacao_visual(self, duracao):
        """Mantem a confirmacao do botao Procurar visivel tempo suficiente."""
        try:
            valor = max(0.12, float(duracao or 0.24))
        except (TypeError, ValueError):
            valor = 0.24
        # O Studio retorna logo apos o worker responder. Uma borda de 0,36 s
        # podia desaparecer durante a restauracao da janela e parecer que a
        # confirmacao nao tinha sido desenhada.
        if bool(getattr(self, "modo_teste_reconhecimento", False)):
            return max(0.72, valor)
        return valor

    def aplicar_estilo_marcador_passivo(self, janela):
        """Torna um marcador Tk visivel e sem ativacao no Windows.

        O contorno e apenas uma camada de indicacao. Ele nunca pode receber
        clique, assumir o foco nem alterar a ordem da janela que esta sendo
        automatizada. Os mesmos estilos sao usados pelos overlays persistentes
        do FlowDetect.
        """
        if os.name != "nt" or not hasattr(ctypes, "windll") or janela is None:
            return False
        try:
            janela.update_idletasks()
            hwnd = int(janela.winfo_id() or 0)
            if hwnd <= 0:
                return False

            user32 = ctypes.windll.user32
            GWL_EXSTYLE = -20
            WS_EX_TRANSPARENT = 0x00000020
            WS_EX_TOOLWINDOW = 0x00000080
            WS_EX_NOACTIVATE = 0x08000000
            HWND_TOPMOST = -1
            SWP_NOSIZE = 0x0001
            SWP_NOMOVE = 0x0002
            SWP_NOACTIVATE = 0x0010
            SWP_SHOWWINDOW = 0x0040
            SW_SHOWNOACTIVATE = 4

            # Definir os prototipos evita que ctypes reduza o HWND em processos
            # de 64 bits antes de aplicar o estilo sem ativacao.
            get_window_long = getattr(user32, "GetWindowLongPtrW", None) or user32.GetWindowLongW
            set_window_long = getattr(user32, "SetWindowLongPtrW", None) or user32.SetWindowLongW
            get_window_long.argtypes = [wintypes.HWND, ctypes.c_int]
            get_window_long.restype = ctypes.c_ssize_t
            set_window_long.argtypes = [wintypes.HWND, ctypes.c_int, ctypes.c_ssize_t]
            set_window_long.restype = ctypes.c_ssize_t
            user32.SetWindowPos.argtypes = [
                wintypes.HWND, wintypes.HWND, ctypes.c_int, ctypes.c_int,
                ctypes.c_int, ctypes.c_int, ctypes.c_uint,
            ]
            user32.SetWindowPos.restype = wintypes.BOOL
            user32.ShowWindow.argtypes = [wintypes.HWND, ctypes.c_int]
            user32.ShowWindow.restype = wintypes.BOOL

            hwnd_win = wintypes.HWND(hwnd)
            estilo_atual = int(get_window_long(hwnd_win, GWL_EXSTYLE) or 0)
            estilo_passivo = estilo_atual | WS_EX_TRANSPARENT | WS_EX_TOOLWINDOW | WS_EX_NOACTIVATE
            set_window_long(hwnd_win, GWL_EXSTYLE, estilo_passivo)
            user32.SetWindowPos(
                hwnd_win,
                wintypes.HWND(HWND_TOPMOST),
                0,
                0,
                0,
                0,
                SWP_NOMOVE | SWP_NOSIZE | SWP_NOACTIVATE | SWP_SHOWWINDOW,
            )
            user32.ShowWindow(hwnd_win, SW_SHOWNOACTIVATE)
            return True
        except Exception:
            return False

    def marcar_regiao_absoluta(
        self,
        esquerda,
        topo,
        largura,
        altura,
        duracao=1.1,
        bloqueante=False,
        tema="visual",
        etiqueta=None,
        margem=16,
    ):
        # Os marcadores usam uma janela temporaria, mas recebem estilo nativo
        # click-through/no-activate. Assim eles voltam a aparecer tambem na
        # execucao assistida sem cobrir a interacao nem trocar o foco do alvo.
        if not self.feedback_visual_habilitado():
            return False

        esquerda = int(esquerda)
        topo = int(topo)
        largura = max(1, int(largura))
        altura = max(1, int(altura))
        tema_config = {
            "visual": {
                "cor_borda": "#2ecc71",
                "cor_flash": "#9bffbf",
                "cor_texto": "#f5fff8",
                "cor_fundo_etiqueta": "#1c9f5b",
                "etiqueta": "Reconhecido",
            },
            "ocr": {
                "cor_borda": "#38BDF8",
                "cor_flash": "#93C5FD",
                "cor_texto": "#E6F7FF",
                "cor_fundo_etiqueta": "#1D4ED8",
                "etiqueta": "OCR confirmado",
            },
            "ocr_busca": {
                "cor_borda": "#38BDF8",
                "cor_flash": "#7DD3FC",
                "cor_texto": "#E6F7FF",
                "cor_fundo_etiqueta": "#075985",
                "etiqueta": "OCR candidato",
            },
            "ocr_busca_1": {
                "cor_borda": "#38BDF8",
                "cor_flash": "#7DD3FC",
                "cor_texto": "#E6F7FF",
                "cor_fundo_etiqueta": "#075985",
                "etiqueta": "OCR 1",
            },
            "ocr_busca_2": {
                "cor_borda": "#A78BFA",
                "cor_flash": "#C4B5FD",
                "cor_texto": "#F5F3FF",
                "cor_fundo_etiqueta": "#6D28D9",
                "etiqueta": "OCR 2",
            },
            "ocr_busca_3": {
                "cor_borda": "#86EFAC",
                "cor_flash": "#BBF7D0",
                "cor_texto": "#F0FDF4",
                "cor_fundo_etiqueta": "#15803D",
                "etiqueta": "OCR 3",
            },
            "ocr_busca_4": {
                "cor_borda": "#F0ABFC",
                "cor_flash": "#F5D0FE",
                "cor_texto": "#FDF4FF",
                "cor_fundo_etiqueta": "#A21CAF",
                "etiqueta": "OCR 4",
            },
            "ocr_busca_5": {
                "cor_borda": "#FDBA74",
                "cor_flash": "#FED7AA",
                "cor_texto": "#FFF7ED",
                "cor_fundo_etiqueta": "#C2410C",
                "etiqueta": "OCR 5",
            },
            "ocr_busca_6": {
                "cor_borda": "#67E8F9",
                "cor_flash": "#A5F3FC",
                "cor_texto": "#ECFEFF",
                "cor_fundo_etiqueta": "#0E7490",
                "etiqueta": "OCR 6",
            },
            "ocr_alvo_falha": {
                "cor_borda": "#FB7185",
                "cor_flash": "#FDA4AF",
                "cor_texto": "#FFF1F2",
                "cor_fundo_etiqueta": "#BE123C",
                "etiqueta": "Falha OCR",
            },
            "ocr_falha": {
                "cor_borda": "#38BDF8",
                "cor_flash": "#93C5FD",
                "cor_texto": "#E6F7FF",
                "cor_fundo_etiqueta": "#1D4ED8",
                "etiqueta": "OCR lido aqui",
            },
        }.get(str(tema or "visual"), {})
        area_esquerda = int(esquerda - margem)
        area_topo = int(topo - margem)
        area_largura = int(largura + (margem * 2))
        area_altura = int(altura + (margem * 2))

        def exibir():
            try:
                cor_borda = tema_config.get("cor_borda", "#2ecc71")
                cor_flash = tema_config.get("cor_flash", "#9bffbf")
                cor_texto = tema_config.get("cor_texto", "#f5fff8")
                cor_fundo_etiqueta = tema_config.get("cor_fundo_etiqueta", "#1c9f5b")
                texto_marcador = str(etiqueta or tema_config.get("etiqueta") or "Reconhecido")
                fundo_transparente = "#ff00ff"
                raio = 16
                espessura = 3
                largura_minima_etiqueta = max(82, min(150, 22 + len(texto_marcador) * 7))
                altura_etiqueta = 20
                padding_overlay = 8
                raiz = tk.Tk()
                raiz.withdraw()
                centro_original_x = int(esquerda + (largura / 2))
                centro_original_y = int(topo + (altura / 2))
                monitor_esquerda, monitor_topo, monitor_largura, monitor_altura = self.obter_metricas_monitor_ponto(
                    centro_original_x,
                    centro_original_y,
                )
                monitor_direita = monitor_esquerda + max(1, monitor_largura)
                monitor_base = monitor_topo + max(1, monitor_altura)

                def pontos_retangulo_arredondado(x1, y1, x2, y2, r):
                    r = max(4, min(r, int((x2 - x1) / 2), int((y2 - y1) / 2)))
                    return [
                        x1 + r, y1,
                        x1 + r, y1,
                        x2 - r, y1,
                        x2 - r, y1,
                        x2, y1,
                        x2, y1 + r,
                        x2, y1 + r,
                        x2, y2 - r,
                        x2, y2 - r,
                        x2, y2,
                        x2 - r, y2,
                        x2 - r, y2,
                        x1 + r, y2,
                        x1 + r, y2,
                        x1, y2,
                        x1, y2 - r,
                        x1, y2 - r,
                        x1, y1 + r,
                        x1, y1 + r,
                        x1, y1,
                    ]

                area_direita = area_esquerda + area_largura
                area_base = area_topo + area_altura
                area_visivel_esquerda = max(monitor_esquerda, area_esquerda)
                area_visivel_topo = max(monitor_topo, area_topo)
                area_visivel_direita = min(monitor_direita, area_direita)
                area_visivel_base = min(monitor_base, area_base)

                if area_visivel_direita <= area_visivel_esquerda:
                    area_visivel_esquerda = max(monitor_esquerda, min(centro_original_x - 18, monitor_direita - 36))
                    area_visivel_direita = min(monitor_direita, area_visivel_esquerda + 36)
                if area_visivel_base <= area_visivel_topo:
                    area_visivel_topo = max(monitor_topo, min(centro_original_y - 18, monitor_base - 36))
                    area_visivel_base = min(monitor_base, area_visivel_topo + 36)

                largura_visivel = max(24, int(area_visivel_direita - area_visivel_esquerda))
                altura_visivel = max(24, int(area_visivel_base - area_visivel_topo))
                largura_overlay_desejada = max(
                    largura_visivel + (padding_overlay * 2),
                    largura_minima_etiqueta + 20,
                )
                altura_overlay_desejada = max(
                    altura_visivel + (padding_overlay * 2),
                    altura_etiqueta + 22,
                )

                overlay_esquerda = int(
                    min(
                        max(monitor_esquerda, area_visivel_esquerda - padding_overlay),
                        max(monitor_esquerda, monitor_direita - largura_overlay_desejada),
                    )
                )
                overlay_topo = int(
                    min(
                        max(monitor_topo, area_visivel_topo - padding_overlay),
                        max(monitor_topo, monitor_base - altura_overlay_desejada),
                    )
                )
                overlay_direita = int(min(monitor_direita, overlay_esquerda + largura_overlay_desejada))
                overlay_base = int(min(monitor_base, overlay_topo + altura_overlay_desejada))
                overlay_largura = int(max(40, overlay_direita - overlay_esquerda))
                overlay_altura = int(max(40, overlay_base - overlay_topo))

                overlay = tk.Toplevel(raiz)
                # Mantemos a camada criada e transparente desde o inicio. Se o
                # ajuste nativo falhar em alguma versao do Windows/Tk, ela ainda
                # podera aparecer nos frames seguintes em vez de ficar presa em
                # withdraw e sumir por completo.
                overlay.overrideredirect(True)
                overlay.attributes("-alpha", 0.0)
                try:
                    overlay.wm_attributes("-transparentcolor", fundo_transparente)
                except Exception:
                    pass
                overlay.configure(bg=fundo_transparente)
                overlay.geometry(
                    f"{overlay_largura}x{overlay_altura}+{int(overlay_esquerda)}+{int(overlay_topo)}"
                )

                canvas_overlay = tk.Canvas(
                    overlay,
                    width=overlay_largura,
                    height=overlay_altura,
                    bg=fundo_transparente,
                    highlightthickness=0,
                    bd=0,
                )
                canvas_overlay.pack(fill="both", expand=True)

                # A borda fica no topo apenas para ser vista; os estilos
                # passivos garantem que nem o mouse nem o foco sejam retidos
                # pela janela temporaria.
                try:
                    overlay.deiconify()
                    overlay.lift()
                    overlay.attributes("-topmost", True)
                    overlay.update_idletasks()
                    overlay.update()
                except Exception:
                    pass
                self.aplicar_estilo_marcador_passivo(overlay)

                retangulo_x1 = max(6, int(area_visivel_esquerda - overlay_esquerda) + 4)
                retangulo_y1 = max(6, int(area_visivel_topo - overlay_topo) + 4)
                retangulo_x2 = min(overlay_largura - 6, int(area_visivel_direita - overlay_esquerda) - 4)
                retangulo_y2 = min(overlay_altura - 6, int(area_visivel_base - overlay_topo) - 4)
                if retangulo_x2 <= retangulo_x1:
                    retangulo_x2 = min(overlay_largura - 6, retangulo_x1 + max(18, largura_visivel - 8))
                if retangulo_y2 <= retangulo_y1:
                    retangulo_y2 = min(overlay_altura - 6, retangulo_y1 + max(18, altura_visivel - 8))

                if tema.startswith("ocr_busca"):
                    retangulo_x1 = int(area_visivel_esquerda - overlay_esquerda)
                    retangulo_y1 = int(area_visivel_topo - overlay_topo)
                    retangulo_x2 = int(area_visivel_direita - overlay_esquerda)
                    retangulo_y2 = int(area_visivel_base - overlay_topo)
                    raio = min(4, max(1, altura / 4))

                retangulo_principal = canvas_overlay.create_polygon(
                    pontos_retangulo_arredondado(retangulo_x1, retangulo_y1, retangulo_x2, retangulo_y2, raio),
                    smooth=True,
                    splinesteps=24,
                    outline=cor_borda,
                    fill="",
                    width=espessura,
                )

                largura_etiqueta = min(104, max(largura_minima_etiqueta, overlay_largura - 20))
                etiqueta_x1 = min(
                    max(10, retangulo_x1 + 10),
                    max(10, overlay_largura - largura_etiqueta - 10),
                )
                etiqueta_y1 = retangulo_y1 - altura_etiqueta - 8
                if etiqueta_y1 < 6:
                    etiqueta_y1 = min(
                        overlay_altura - altura_etiqueta - 6,
                        retangulo_y2 + 8,
                    )
                etiqueta_x2 = etiqueta_x1 + largura_etiqueta
                etiqueta_y2 = etiqueta_y1 + altura_etiqueta

                etiqueta_canvas = canvas_overlay.create_polygon(
                    pontos_retangulo_arredondado(etiqueta_x1, etiqueta_y1, etiqueta_x2, etiqueta_y2, 10),
                    smooth=True,
                    splinesteps=24,
                    outline="",
                    fill=cor_fundo_etiqueta,
                )
                texto_etiqueta = canvas_overlay.create_text(
                    int((etiqueta_x1 + etiqueta_x2) / 2),
                    int((etiqueta_y1 + etiqueta_y2) / 2),
                    text=texto_marcador,
                    fill=cor_texto,
                    font=("Segoe UI", 8, "bold"),
                )

                try:
                    overlay.update_idletasks()
                except Exception:
                    pass

                flash_fases = [
                    (0.12, cor_flash, 5),
                    (0.92, cor_flash, 5),
                    (0.35, cor_borda, 3),
                ]

                for alpha, cor_atual, largura_borda in flash_fases:
                    try:
                        overlay.attributes("-alpha", alpha)
                        canvas_overlay.itemconfigure(retangulo_principal, outline=cor_atual, width=largura_borda)
                        canvas_overlay.itemconfigure(etiqueta_canvas, fill=cor_atual)
                        canvas_overlay.itemconfigure(texto_etiqueta, fill=cor_texto)
                        overlay.update_idletasks()
                        overlay.update()
                    except Exception:
                        pass
                    time.sleep(0.04)

                fim = time.perf_counter() + max(0.12, float(duracao))
                while time.perf_counter() < fim:
                    try:
                        overlay.attributes("-alpha", 0.72)
                        overlay.update_idletasks()
                        overlay.update()
                    except Exception:
                        pass
                    time.sleep(0.02)

                try:
                    overlay.destroy()
                except Exception:
                    pass
                try:
                    raiz.destroy()
                except Exception:
                    pass
                # Em execucao assistida os marcadores sempre sao sincronizados
                # com a automacao. Esperar um unico frame apos destruí-los evita
                # que uma captura imediatamente seguinte leia a borda colorida.
                if not bool(getattr(self, "modo_teste_reconhecimento", False)):
                    time.sleep(0.03)
                safe_print(f"Marcador visual exibido: {texto_marcador}.", flush=True)
                return True
            except Exception as erro:
                safe_print(f"Falha ao exibir marcador visual: {erro}", flush=True)
                return False

        # Marcadores assincronos podem permanecer visiveis enquanto a proxima
        # captura OCR/template e feita. Fora do botao Procurar, executamos o
        # pequeno feedback de forma sincronizada: ele aparece, fecha e so entao
        # a automacao continua a capturar pixels.
        if bloqueante or not bool(getattr(self, "modo_teste_reconhecimento", False)):
            return exibir()

        threading.Thread(target=exibir, daemon=True).start()
        return True

    def marcar_regiao_ocr(self, regiao, confirmado=True, duracao=0.9):
        if not regiao:
            return
        try:
            esquerda, topo, largura, altura = [int(valor) for valor in regiao]
        except Exception:
            return
        self.marcar_regiao_absoluta(
            esquerda,
            topo,
            largura,
            altura,
            duracao=duracao,
            bloqueante=True,
            tema="ocr" if confirmado else "ocr_falha",
            etiqueta="OCR confirmado" if confirmado else "OCR lido aqui",
            margem=4,
        )

    def escalar_offset_referencia(self, local, offset_x=0, offset_y=0, dimensoes_referencia=None):
        if not dimensoes_referencia:
            return int(offset_x or 0), int(offset_y or 0)
        try:
            ref_largura, ref_altura = dimensoes_referencia
            _esquerda, _topo, largura, altura = self.bounds_reconhecimento(local)
            if int(ref_largura or 0) <= 0 or int(ref_altura or 0) <= 0:
                return int(offset_x or 0), int(offset_y or 0)
            escala_x = float(largura or ref_largura) / float(ref_largura)
            escala_y = float(altura or ref_altura) / float(ref_altura)
            return (
                int(round(float(offset_x or 0) * escala_x)),
                int(round(float(offset_y or 0) * escala_y)),
            )
        except Exception:
            return int(offset_x or 0), int(offset_y or 0)

    def centro_area_reconhecida(self, local, offset_x=0, offset_y=0):
        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        centro_x = esquerda + int(largura / 2)
        centro_y = topo + int(altura / 2)
        return int(centro_x + int(offset_x or 0)), int(centro_y + int(offset_y or 0))

    def construir_regiao_ocr_relativa(self, local, leitura_relativa=None, dimensoes_referencia=None):
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        escala_x = 1.0
        escala_y = 1.0
        if dimensoes_referencia:
            try:
                ref_largura, ref_altura = dimensoes_referencia
                if int(ref_largura or 0) > 0 and int(ref_altura or 0) > 0:
                    escala_x = float(largura or ref_largura) / float(ref_largura)
                    escala_y = float(altura or ref_altura) / float(ref_altura)
            except Exception:
                escala_x = 1.0
                escala_y = 1.0
        largura_ocr = max(16, int(leitura.get("largura", leitura.get("width", 140)) or 140))
        altura_ocr = max(12, int(leitura.get("altura", leitura.get("height", 40)) or 40))
        offset_x = int(leitura.get("offset_x", leitura.get("offsetX", 0)) or 0)
        offset_y = int(leitura.get("offset_y", leitura.get("offsetY", 0)) or 0)
        largura_ocr = max(8, int(round(largura_ocr * escala_x)))
        altura_ocr = max(8, int(round(altura_ocr * escala_y)))
        offset_x = int(round(offset_x * escala_x))
        offset_y = int(round(offset_y * escala_y))
        centro_x = esquerda + int(largura / 2) + offset_x
        centro_y = topo + int(altura / 2) + offset_y
        limite_reconhecido = (int(esquerda), int(topo), int(largura), int(altura))
        return self.limitar_regiao((
            int(centro_x - largura_ocr / 2),
            int(centro_y - altura_ocr / 2),
            largura_ocr,
            altura_ocr,
        ), limite=limite_reconhecido)

    def construir_regiao_ocr_referencia(self, leitura_relativa=None, dimensoes_referencia=None):
        if not dimensoes_referencia:
            return None
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        try:
            ref_largura, ref_altura = [int(valor) for valor in dimensoes_referencia]
        except Exception:
            return None
        if ref_largura <= 0 or ref_altura <= 0:
            return None
        largura_ocr = max(16, int(leitura.get("largura", leitura.get("width", 140)) or 140))
        altura_ocr = max(12, int(leitura.get("altura", leitura.get("height", 40)) or 40))
        offset_x = int(leitura.get("offset_x", leitura.get("offsetX", 0)) or 0)
        offset_y = int(leitura.get("offset_y", leitura.get("offsetY", 0)) or 0)
        centro_x = int(ref_largura / 2) + offset_x
        centro_y = int(ref_altura / 2) + offset_y
        return self.limitar_regiao(
            (
                int(centro_x - largura_ocr / 2),
                int(centro_y - altura_ocr / 2),
                largura_ocr,
                altura_ocr,
            ),
            limite=(0, 0, ref_largura, ref_altura),
        )

    def confirmar_ocr_visual_pela_referencia(self, captura_ocr, caminho_referencia=None, leitura_relativa=None, dimensoes_referencia=None):
        if not OPENCV_DISPONIVEL or np is None or Image is None or captura_ocr is None or not caminho_referencia:
            return False
        try:
            caminho = Path(caminho_referencia)
            if not caminho.exists():
                return False
            regiao_ref = self.construir_regiao_ocr_referencia(leitura_relativa, dimensoes_referencia=dimensoes_referencia)
            if not regiao_ref:
                return False
            with Image.open(caminho) as imagem_ref:
                x_ref, y_ref, w_ref, h_ref = [int(valor) for valor in regiao_ref]
                recorte_ref = imagem_ref.convert("RGB").crop((x_ref, y_ref, x_ref + w_ref, y_ref + h_ref))
            tela = cv2.cvtColor(np.array(captura_ocr.convert("RGB")), cv2.COLOR_RGB2GRAY)
            ref = cv2.cvtColor(np.array(recorte_ref), cv2.COLOR_RGB2GRAY)
            if tela.size <= 0 or ref.size <= 0:
                return False
            if tela.shape[:2] != ref.shape[:2]:
                tela = cv2.resize(tela, (ref.shape[1], ref.shape[0]), interpolation=cv2.INTER_AREA)
            diff = cv2.absdiff(tela, ref)
            similaridade_tons = 1.0 - (float(diff.mean()) / 255.0)
            correlacao = float(cv2.matchTemplate(tela, ref, cv2.TM_CCOEFF_NORMED)[0][0]) if min(ref.shape[:2]) >= 3 else 0.0
            ref_bin = cv2.threshold(ref, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
            tela_bin = cv2.threshold(tela, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
            bin_diff = cv2.absdiff(tela_bin, ref_bin)
            similaridade_binaria = 1.0 - (float(bin_diff.mean()) / 255.0)
            confirmado = (
                similaridade_tons >= 0.88
                and similaridade_binaria >= 0.88
                and correlacao >= 0.62
            )
            safe_print(
                "OCR visual pela referencia "
                f"{'confirmou' if confirmado else 'nao confirmou'} "
                f"(tons={similaridade_tons:.2f}, bin={similaridade_binaria:.2f}, corr={correlacao:.2f}).",
                flush=True,
            )
            return confirmado
        except Exception as erro:
            safe_print(f"OCR visual pela referencia falhou: {erro}", flush=True)
            return False

    def localizar_referencia_por_ancora_ocr(
        self,
        caminho_referencia,
        texto_ocr_esperado="",
        region=None,
        leitura_relativa=None,
        dimensoes_referencia=None,
    ):
        esperado = str(texto_ocr_esperado or "").strip()
        if not esperado or not OCR_DISPONIVEL or not dimensoes_referencia:
            return None
        regiao_ref_ocr = self.construir_regiao_ocr_referencia(
            leitura_relativa,
            dimensoes_referencia=dimensoes_referencia,
        )
        if not regiao_ref_ocr:
            return None
        try:
            ref_largura, ref_altura = [int(valor) for valor in dimensoes_referencia]
            ref_x, ref_y, ref_w, ref_h = [int(valor) for valor in regiao_ref_ocr]
        except Exception:
            return None
        if ref_largura <= 0 or ref_altura <= 0 or ref_w <= 0 or ref_h <= 0:
            return None

        regiao_busca = region
        if regiao_busca is None:
            try:
                tela_largura, tela_altura = pyautogui.size()
                regiao_busca = (0, 0, int(tela_largura), int(tela_altura))
            except Exception:
                return None

        local_ocr = self.procurar_texto_por_ocr(esperado, region=regiao_busca)
        if not local_ocr:
            return None

        try:
            largura_lida = max(1, int(local_ocr.get("width", 0) or 0))
            altura_lida = max(1, int(local_ocr.get("height", 0) or 0))
            escala_x = largura_lida / max(1, ref_w)
            escala_y = altura_lida / max(1, ref_h)
            if not 0.65 <= escala_x <= 1.55:
                escala_x = 1.0
            if not 0.65 <= escala_y <= 1.55:
                escala_y = 1.0
            centro_ref_x = ref_x + (ref_w / 2.0)
            centro_ref_y = ref_y + (ref_h / 2.0)
            centro_ocr_x = float(local_ocr.get("left", 0) or 0) + (largura_lida / 2.0)
            centro_ocr_y = float(local_ocr.get("top", 0) or 0) + (altura_lida / 2.0)
            largura_estimativa = max(8, int(round(ref_largura * escala_x)))
            altura_estimativa = max(8, int(round(ref_altura * escala_y)))
            esquerda = int(round(centro_ocr_x - (centro_ref_x * escala_x)))
            topo = int(round(centro_ocr_y - (centro_ref_y * escala_y)))
            limite = self.limitar_regiao((esquerda, topo, largura_estimativa, altura_estimativa))
            if limite:
                esquerda, topo, largura_estimativa, altura_estimativa = limite
            local = {
                "x": int(esquerda + largura_estimativa / 2),
                "y": int(topo + altura_estimativa / 2),
                "left": int(esquerda),
                "top": int(topo),
                "width": int(largura_estimativa),
                "height": int(altura_estimativa),
                "score": float(local_ocr.get("score", 0.0) or 0.0),
                "search_stage": "ocr_anchor",
                "ocr_anchor_confirmed": True,
                "ocr_text": str(local_ocr.get("texto", "") or esperado),
                "ocr_bounds": {
                    "left": int(local_ocr.get("left", 0) or 0),
                    "top": int(local_ocr.get("top", 0) or 0),
                    "width": largura_lida,
                    "height": altura_lida,
                },
            }
            safe_print(
                "Imagem localizada por ancora OCR "
                f"(texto='{esperado}', capturado='{local.get('ocr_text', '')}', "
                f"area={self.formatar_regiao_log(regiao_busca)}, score={local['score']:.2f}).",
                flush=True,
            )
            return local
        except Exception as erro:
            safe_print(f"Falha ao reconstruir referencia pela ancora OCR: {erro}", flush=True)
            return None

    def tema_alvo_ocr(self, indice):
        """Mantem a mesma sequencia de cores usada pelos cards do Studio."""
        try:
            indice_normalizado = max(0, int(indice or 0))
        except (TypeError, ValueError):
            indice_normalizado = 0
        return f"ocr_busca_{(indice_normalizado % 6) + 1}"

    def marcar_alvos_ocr_em_sequencia(self, local, leitura_relativa=None, dimensoes_referencia=None):
        config = leitura_relativa or {}
        if not config.get("_mostrar_marcadores_ocr", True):
            return
        layouts = config.get("ocr_target_layouts") or []
        bounds = local.get("ocr_related_bounds") or [local.get("ocr_bounds") or {}]
        try:
            indice_inicial = max(0, int(config.get("ocr_target_marker_index", 0) or 0))
        except (TypeError, ValueError):
            indice_inicial = 0
        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        ref_w, ref_h = dimensoes_referencia or (largura, altura)
        scale_x, scale_y = largura / max(1, ref_w), altura / max(1, ref_h)
        for index, detected in enumerate(bounds):
            if self.cancelamento_solicitado():
                return
            if layouts:
                layout = layouts[index] if index < len(layouts) else {}
                if index == 0:
                    width, height = config.get("largura", 140), config.get("altura", 40)
                    ox, oy = config.get("offset_x", 0), config.get("offset_y", 0)
                else:
                    width, height = layout.get("width") or 140, layout.get("height") or 40
                    ox = layout.get("offsetX")
                    oy = layout.get("offsetY")
                    ox = ox if ox is not None else config.get("offset_x", 0) + index * .032 * ref_w
                    oy = oy if oy is not None else config.get("offset_y", 0) + index * .072 * ref_h
                w, h = max(1, round(width * scale_x)), max(1, round(height * scale_y))
                x = round(esquerda + largura / 2 + ox * scale_x - w / 2)
                y = round(topo + altura / 2 + oy * scale_y - h / 2)
            else:
                x, y, w, h = self.bounds_reconhecimento(detected)
            # Mantem a confirmacao visivel, mas sem introduzir quase meio
            # segundo de espera para cada OCR de um bloco composto.
            indice_visual = indice_inicial + index
            self.marcar_regiao_absoluta(x, y, w, h, tema=self.tema_alvo_ocr(indice_visual),
                                       etiqueta=f"OCR {indice_visual + 1}", duracao=0.18,
                                       bloqueante=True, margem=0)

    def regioes_alvos_ocr_por_layout(self, local, leitura_relativa=None, dimensoes_referencia=None, quantidade=0):
        config = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = config.get("ocr_target_layouts") if isinstance(config.get("ocr_target_layouts"), list) else []
        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        ref_w, ref_h = dimensoes_referencia or (largura, altura)
        scale_x, scale_y = largura / max(1, ref_w), altura / max(1, ref_h)
        total = max(int(quantidade or 0), len(layouts), 1)
        regioes = []
        for index in range(total):
            layout = layouts[index] if index < len(layouts) and isinstance(layouts[index], dict) else {}
            if index == 0:
                width = config.get("largura", layout.get("width", 140))
                height = config.get("altura", layout.get("height", 40))
                ox = config.get("offset_x", layout.get("offsetX", 0))
                oy = config.get("offset_y", layout.get("offsetY", 0))
            else:
                width = layout.get("width") or 140
                height = layout.get("height") or 40
                ox = layout.get("offsetX")
                oy = layout.get("offsetY")
                ox = ox if ox is not None else config.get("offset_x", 0) + index * 0.032 * ref_w
                oy = oy if oy is not None else config.get("offset_y", 0) + index * 0.072 * ref_h
            w = max(1, round(float(width or 140) * scale_x))
            h = max(1, round(float(height or 40) * scale_y))
            x = round(esquerda + largura / 2 + float(ox or 0) * scale_x - w / 2)
            y = round(topo + altura / 2 + float(oy or 0) * scale_y - h / 2)
            regioes.append((int(x), int(y), int(w), int(h)))
        return regioes

    def regiao_referencia_alvo_ocr(self, index, leitura_relativa=None, dimensoes_referencia=None):
        config = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = config.get("ocr_target_layouts") if isinstance(config.get("ocr_target_layouts"), list) else []
        ref_w, ref_h = dimensoes_referencia or (0, 0)
        if ref_w <= 0 or ref_h <= 0:
            return None
        layout = layouts[index] if index < len(layouts) and isinstance(layouts[index], dict) else {}
        if index == 0:
            width = config.get("largura", layout.get("width", 140))
            height = config.get("altura", layout.get("height", 40))
            ox = config.get("offset_x", layout.get("offsetX", 0))
            oy = config.get("offset_y", layout.get("offsetY", 0))
        else:
            width = layout.get("width") or 140
            height = layout.get("height") or 40
            ox = layout.get("offsetX")
            oy = layout.get("offsetY")
            ox = ox if ox is not None else config.get("offset_x", 0) + index * 0.032 * ref_w
            oy = oy if oy is not None else config.get("offset_y", 0) + index * 0.072 * ref_h
        w = max(1, int(round(float(width or 140))))
        h = max(1, int(round(float(height or 40))))
        x = int(round(ref_w / 2 + float(ox or 0) - w / 2))
        y = int(round(ref_h / 2 + float(oy or 0) - h / 2))
        x = max(0, min(x, max(0, ref_w - 1)))
        y = max(0, min(y, max(0, ref_h - 1)))
        w = max(1, min(w, ref_w - x))
        h = max(1, min(h, ref_h - y))
        return x, y, w, h

    def comparar_recorte_alvo_ocr_visual(self, captura, caminho_referencia, index, leitura_relativa=None, dimensoes_referencia=None):
        if not OPENCV_DISPONIVEL or np is None or cv2 is None or captura is None or not caminho_referencia:
            return 0.0
        regiao_ref = self.regiao_referencia_alvo_ocr(index, leitura_relativa, dimensoes_referencia)
        if not regiao_ref:
            return 0.0
        try:
            with Image.open(caminho_referencia) as imagem_ref:
                rx, ry, rw, rh = regiao_ref
                ref_crop = imagem_ref.convert("RGB").crop((rx, ry, rx + rw, ry + rh))
            live_crop = captura.convert("RGB").resize(ref_crop.size)
            ref_gray = cv2.cvtColor(np.array(ref_crop), cv2.COLOR_RGB2GRAY)
            live_gray = cv2.cvtColor(np.array(live_crop), cv2.COLOR_RGB2GRAY)
            resultado = cv2.matchTemplate(live_gray, ref_gray, cv2.TM_CCOEFF_NORMED)
            return float(resultado[0][0]) if resultado is not None and resultado.size else 0.0
        except Exception:
            return 0.0

    def confirmar_alvos_ocr_por_layout_visual(self, local, texto_esperado="", leitura_relativa=None, dimensoes_referencia=None, caminho_referencia=None, marcar=True):
        spec = self.interpretar_texto_ocr_composto(texto_esperado)
        partes = self.partes_alvos_ocr_esperados(texto_esperado, leitura_relativa)
        quantidade = self.quantidade_alvos_ocr_esperados(texto_esperado, leitura_relativa)
        if len(partes) < quantidade:
            safe_print(
                "OCR multiplo bloqueado: ha mais caixas de leitura que textos alvo configurados.",
                flush=True,
            )
            return None
        regioes = self.regioes_alvos_ocr_por_layout(local, leitura_relativa, dimensoes_referencia, quantidade)
        if len(regioes) < quantidade:
            return None
        bounds = []
        textos_lidos = []
        try:
            indice_inicial = max(0, int((leitura_relativa or {}).get("ocr_target_marker_index", 0) or 0))
        except (TypeError, ValueError):
            indice_inicial = 0
        for index, esperado in enumerate(partes[:quantidade]):
            regiao = self.limitar_regiao(regioes[index])
            if not regiao:
                return None
            if marcar:
                self.marcar_regiao_absoluta(
                    int(regiao[0]), int(regiao[1]), int(regiao[2]), int(regiao[3]),
                    tema=self.tema_alvo_ocr(indice_inicial + index),
                    etiqueta=f"OCR {indice_inicial + index + 1}",
                    duracao=0.12,
                    bloqueante=True,
                    margem=0,
                )
            try:
                captura, _base_left, _base_top = self.capturar_tela_pil(region=regiao)
            except Exception:
                return None
            textos = self.extrair_textos_ocr_rapido(captura, texto_esperado=esperado)
            if not textos:
                textos = [
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 6"),
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 7"),
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 8"),
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 13"),
                ]
            capturado = " ".join(str(item or "").strip() for item in textos if str(item or "").strip()).strip()
            confirmado = self.texto_ocr_confirma_alvo(esperado, capturado, limite_ratio=0.86)
            score_visual = 0.0
            if not confirmado and caminho_referencia:
                score_visual = self.comparar_recorte_alvo_ocr_visual(
                    captura,
                    caminho_referencia,
                    index,
                    leitura_relativa=leitura_relativa,
                    dimensoes_referencia=dimensoes_referencia,
                )
                confirmado = score_visual >= 0.72
                if confirmado and not capturado:
                    capturado = esperado
            safe_print(
                f"OCR alvo {index + 1}/{quantidade}: esperado='{esperado}' | capturado='{capturado}' "
                f"| visual={score_visual:.2f} | {'OK' if confirmado else 'falhou'}",
                flush=True,
            )
            if not confirmado:
                return None
            bounds.append({
                "parte": esperado,
                "texto": capturado or esperado,
                "left": int(regiao[0]),
                "top": int(regiao[1]),
                "width": int(regiao[2]),
                "height": int(regiao[3]),
            })
            textos_lidos.append(capturado or esperado)
        confirmado = dict(local)
        confirmado.update({
            "search_stage": "ocr_visual_hybrid",
            "ocr_anchor_confirmed": True,
            "ocr_layout_targets_marked": bool(marcar),
            "ocr_composto": True,
            "ocr_relacional": bool((spec or {}).get("relacional")),
            "ocr_partes": partes[:quantidade],
            "ocr_textos": textos_lidos,
            "ocr_text": " | ".join(textos_lidos),
            "ocr_related_bounds": bounds,
            "ocr_bounds": bounds[0] if bounds else {},
        })
        return confirmado

    def quantidade_alvos_ocr_esperados(self, texto_esperado="", leitura_relativa=None):
        partes = self.partes_alvos_ocr_esperados(texto_esperado, leitura_relativa)
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = leitura.get("ocr_target_layouts") if isinstance(leitura.get("ocr_target_layouts"), list) else []
        return max(1, len([parte for parte in partes if str(parte or "").strip()]), len(layouts))

    def partes_alvos_ocr_esperados(self, texto_esperado="", leitura_relativa=None):
        spec = self.interpretar_texto_ocr_composto(texto_esperado)
        partes_texto = [
            str(parte or "").strip()
            for parte in ((spec or {}).get("partes", []) if spec else [texto_esperado])
            if str(parte or "").strip()
        ]
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = leitura.get("ocr_target_layouts") if isinstance(leitura.get("ocr_target_layouts"), list) else []
        partes_layout = []
        for layout in layouts:
            if not isinstance(layout, dict):
                partes_layout.append("")
                continue
            texto_layout = (
                layout.get("texto")
                or layout.get("text")
                or layout.get("target_text")
                or layout.get("ocrText")
                or layout.get("value")
                or ""
            )
            partes_layout.append(str(texto_layout or "").strip())

        total = max(len(partes_texto), len(partes_layout))
        partes = []
        for indice in range(total):
            texto = partes_texto[indice] if indice < len(partes_texto) else ""
            if not texto and indice < len(partes_layout):
                texto = partes_layout[indice]
            if texto:
                partes.append(texto)
        return partes

    def ocr_hibrido_tem_todos_alvos(self, local, texto_esperado="", leitura_relativa=None):
        if not isinstance(local, dict):
            return False
        quantidade = self.quantidade_alvos_ocr_esperados(texto_esperado, leitura_relativa)
        if quantidade <= 1:
            return bool(local.get("ocr_anchor_confirmed") or local.get("ocr_visual_anchor"))
        bounds = local.get("ocr_related_bounds")
        if not isinstance(bounds, list):
            return False
        validos = [
            bound for bound in bounds
            if isinstance(bound, dict)
            and int(bound.get("width", 0) or 0) > 0
            and int(bound.get("height", 0) or 0) > 0
        ]
        if len(validos) < quantidade:
            return False
        partes = self.partes_alvos_ocr_esperados(texto_esperado, leitura_relativa)
        if len(partes) < quantidade:
            return False
        for esperado, bound in zip(partes[:quantidade], validos[:quantidade]):
            if not self.texto_ocr_confirma_alvo(esperado, bound.get("texto") or bound.get("text") or "", limite_ratio=0.86):
                return False
        return True

    def estimar_referencia_por_alvos_ocr(self, bounds, leitura_relativa=None, dimensoes_referencia=None):
        """Reconstrói a caixa da referência usando os alvos OCR já confirmados.

        As posições gravadas no editor são relativas ao centro da imagem de
        referência. Comparando essas posições com os textos encontrados na tela,
        obtemos a escala atual sem depender da resolução, DPI ou monitor.
        """
        if not isinstance(bounds, list) or not bounds or not dimensoes_referencia:
            return None
        try:
            ref_largura, ref_altura = [int(valor) for valor in dimensoes_referencia]
        except Exception:
            return None
        if ref_largura <= 0 or ref_altura <= 0:
            return None

        layouts = self.montar_layouts_alvos_ocr(
            leitura_relativa,
            dimensoes_referencia=dimensoes_referencia,
            quantidade=len(bounds),
        )
        observados = []
        for indice, bound in enumerate(bounds):
            if indice >= len(layouts):
                break
            retangulo = self._retangulo_texto_ocr(bound)
            if not retangulo:
                continue
            left, top, width, height = retangulo
            observados.append({
                "centro_x": left + (width / 2.0),
                "centro_y": top + (height / 2.0),
                "layout": layouts[indice],
            })
        if not observados:
            return None

        def mediana(valores, padrao=1.0):
            valores_validos = sorted(float(valor) for valor in valores if valor is not None)
            if not valores_validos:
                return float(padrao)
            meio = len(valores_validos) // 2
            if len(valores_validos) % 2:
                return valores_validos[meio]
            return (valores_validos[meio - 1] + valores_validos[meio]) / 2.0

        escalas_x = []
        escalas_y = []
        for indice, origem in enumerate(observados):
            for destino in observados[indice + 1:]:
                esperado_x = float(destino["layout"].get("offset_x", 0.0) or 0.0) - float(origem["layout"].get("offset_x", 0.0) or 0.0)
                esperado_y = float(destino["layout"].get("offset_y", 0.0) or 0.0) - float(origem["layout"].get("offset_y", 0.0) or 0.0)
                atual_x = destino["centro_x"] - origem["centro_x"]
                atual_y = destino["centro_y"] - origem["centro_y"]
                if abs(esperado_x) >= 8:
                    escala_x = atual_x / esperado_x
                    if 0.45 <= escala_x <= 2.40:
                        escalas_x.append(escala_x)
                if abs(esperado_y) >= 8:
                    escala_y = atual_y / esperado_y
                    if 0.45 <= escala_y <= 2.40:
                        escalas_y.append(escala_y)

        escala_x = mediana(escalas_x, 1.0)
        escala_y = mediana(escalas_y, escala_x if escalas_x else 1.0)
        if not escalas_x and escalas_y:
            escala_x = escala_y
        if not escalas_y and escalas_x:
            escala_y = escala_x
        escala_x = max(0.50, min(2.00, escala_x))
        escala_y = max(0.50, min(2.00, escala_y))

        centros_referencia_x = [
            item["centro_x"] - float(item["layout"].get("offset_x", 0.0) or 0.0) * escala_x
            for item in observados
        ]
        centros_referencia_y = [
            item["centro_y"] - float(item["layout"].get("offset_y", 0.0) or 0.0) * escala_y
            for item in observados
        ]
        centro_x = mediana(centros_referencia_x, observados[0]["centro_x"])
        centro_y = mediana(centros_referencia_y, observados[0]["centro_y"])
        largura = max(8, int(round(ref_largura * escala_x)))
        altura = max(8, int(round(ref_altura * escala_y)))
        esquerda = int(round(centro_x - (largura / 2.0)))
        topo = int(round(centro_y - (altura / 2.0)))
        erros = [
            abs(item["centro_x"] - (centro_x + float(item["layout"].get("offset_x", 0.0) or 0.0) * escala_x))
            + abs(item["centro_y"] - (centro_y + float(item["layout"].get("offset_y", 0.0) or 0.0) * escala_y))
            for item in observados
        ]
        return {
            "left": esquerda,
            "top": topo,
            "width": largura,
            "height": altura,
            "x": int(round(centro_x)),
            "y": int(round(centro_y)),
            "ocr_layout_scale_x": float(escala_x),
            "ocr_layout_scale_y": float(escala_y),
            "ocr_layout_error": float(mediana(erros, 0.0)),
        }

    def texto_ocr_confirma_alvo(self, esperado, capturado, limite_ratio=0.90):
        texto_esperado = str(esperado or "").strip()
        texto_capturado = str(capturado or "").strip()
        if not texto_esperado:
            return True
        comparacao = self.comparar_texto_reconhecimento_preciso(texto_esperado, texto_capturado)
        if comparacao.get("reference") == comparacao.get("captured"):
            return True
        limite_ratio = float(limite_ratio or 0.90)
        if float(comparacao.get("ratio", 0.0) or 0.0) >= limite_ratio:
            return True

        referencia = str(comparacao.get("reference") or "")
        leitura = str(comparacao.get("captured") or "")
        # Fontes pequenas de programas desktop confundem uma letra isolada com
        # frequencia (por exemplo, "Codigo" vira "Cedigo"). Nos alvos OCR
        # compostos a posicao relativa e a imagem ainda precisam confirmar o
        # bloco, entao aceitamos apenas uma troca interna em um rotulo curto.
        # O caminho padrao continua estrito e codigos numericos nunca entram
        # nessa tolerancia.
        if (
            limite_ratio <= 0.87
            and 5 <= len(referencia) <= 10
            and len(referencia) == len(leitura)
            and not re.search(r"\d", referencia + leitura)
            and referencia[:1] == leitura[:1]
            and referencia[-1:] == leitura[-1:]
            and sum(origem != destino for origem, destino in zip(referencia, leitura)) == 1
        ):
            return True
        esperado_digitos = re.sub(r"\D+", "", texto_esperado)
        if esperado_digitos and esperado_digitos == texto_esperado.replace(" ", ""):
            grupos = re.findall(r"\d+", texto_capturado)
            if esperado_digitos in grupos:
                return True
            if comparacao.get("has_digits") and float(comparacao.get("digit_ratio", 0.0) or 0.0) >= 0.999:
                return True
        return False

    def modo_acao_alvos_ocr(self, leitura_relativa=None):
        """Usa cada alvo OCR como uma alternativa independente.

        Os cards salvos pela aba Acoes por OCR registram a chave ``action``
        em cada alvo. Neles, cada texto e confirmado com sua propria area e
        imagem: encontrar um alvo ja e suficiente para executar sua acao,
        inclusive a acao normal do bloco. Layouts legados, sem essa chave,
        continuam com a confirmacao conjunta para nao alterar automacoes ja
        existentes.
        """
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = leitura.get("ocr_target_layouts") if isinstance(leitura.get("ocr_target_layouts"), list) else []
        return "individual" if any(
            isinstance(layout, dict) and ("action" in layout or "acao" in layout)
            for layout in layouts
        ) else "todos"

    def regra_acao_individual_ocr(self, config=None):
        """Normaliza as escolhas feitas para um alvo OCR individual.

        A configuracao e deliberadamente guardada junto do proprio alvo: assim
        cada OCR pode decidir se segue para outro OCR, se pula para um bloco
        ou se interrompe a automacao quando nao for encontrado.
        """
        dados = config if isinstance(config, dict) else {}
        acao = str(dados.get("action", dados.get("acao", "continuar")) or "continuar").strip().lower()
        aliases = {
            "click": "clicar",
            "clique": "clicar",
            "bloco": "confirmar_ir_para_bloco",
            "navegar": "confirmar_ir_para_bloco",
            "ir_para_bloco": "confirmar_ir_para_bloco",
            "clicar_e_ir_para_bloco": "clicar_ir_para_bloco",
            "click_and_jump": "clicar_ir_para_bloco",
            "click_goto": "clicar_ir_para_bloco",
            "confirmar_e_ir_para_bloco": "confirmar_ir_para_bloco",
            "confirm_and_jump": "confirmar_ir_para_bloco",
            "confirm_goto": "confirmar_ir_para_bloco",
        }
        acao = aliases.get(acao, acao)
        if acao not in {"continuar", "clicar", "clicar_ir_para_bloco", "confirmar_ir_para_bloco"}:
            acao = "continuar"
        falha = str(
            dados.get("failureAction", dados.get("failure_action", dados.get("acao_falha", "continuar")))
            or "continuar"
        ).strip().lower()
        if falha in {"stop", "interromper"}:
            falha = "parar"
        if falha not in {"continuar", "parar"}:
            falha = "continuar"
        return {
            "acao": acao,
            "falha": falha,
            "bloco": str(dados.get("block", dados.get("bloco", "")) or "").strip(),
            "clicks": max(1, int(dados.get("clicks", dados.get("cliques", 1)) or 1)),
        }

    def reler_alvo_ocr_na_referencia(self, local, indice, esperado, leitura_relativa=None, dimensoes_referencia=None):
        """Confere apenas o texto que a leitura ampla perdeu, na referencia ja localizada."""
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = leitura.get("ocr_target_layouts") or []
        if (
            not isinstance(local, dict) or not local.get("visual_confirmed")
            or indice >= len(layouts) or not dimensoes_referencia
            or self.cancelamento_solicitado()
        ):
            return None
        regioes = self.regioes_alvos_ocr_por_layout(
            local, leitura, dimensoes_referencia, quantidade=len(layouts),
        )
        regiao = self.limitar_regiao(regioes[indice], limite=self.bounds_reconhecimento(local))
        if not regiao:
            return None
        safe_print(f"OCR {indice + 1}: conferindo somente seu recorte ({self.formatar_regiao_log(regiao)}).", flush=True)
        try:
            captura, _left, _top = self.capturar_tela_pil(region=regiao)
            # Uma linha ampliada dispensa a varredura da janela inteira e
            # conserva acentos/espacos, sem a whitelist da leitura generica.
            ampliada = self.ampliar_imagem_para_ocr(captura.convert("L"), fator=2)
            primeira_leitura = self.extrair_texto_ocr_imagem_pil(ampliada, config="--psm 7")
            textos = [primeira_leitura] if primeira_leitura else []
            if not self.texto_ocr_confirma_alvo(esperado, primeira_leitura, limite_ratio=0.86):
                textos.extend(self.extrair_textos_ocr_rapido(captura, texto_esperado=esperado))
        except Exception as erro:
            safe_print(f"OCR {indice + 1}: falha na leitura do recorte: {erro}", flush=True)
            return None
        if self.cancelamento_solicitado():
            return None
        # Cada variante e uma leitura independente. Concatenar as alternativas
        # duplicava palavras e fazia uma leitura correta reprovar na comparacao.
        capturado = next((str(texto).strip() for texto in textos
                          if self.texto_ocr_confirma_alvo(esperado, texto, limite_ratio=0.86)), "")
        if not capturado:
            amostra = " | ".join(str(texto).strip() for texto in textos[:3])
            safe_print(f"OCR {indice + 1}: recorte nao confirmou '{esperado}' | lido='{amostra}'.", flush=True)
            return None
        bounds = dict(zip(("left", "top", "width", "height"), regiao))
        bounds["texto"] = capturado
        confirmado = dict(local)
        confirmado.update({
            "ocr_anchor_confirmed": True,
            "ocr_text": capturado,
            "ocr_bounds": bounds,
            "ocr_related_bounds": [bounds],
            "ocr_partes": [esperado],
            "ocr_textos": [capturado],
            "ocr_composto": False,
            "ocr_relacional": False,
            "ocr_layout_targets_marked": True,
            "search_stage": "ocr_visual_hybrid",
        })
        if leitura.get("_mostrar_marcadores_ocr", True):
            self.marcar_regiao_absoluta(
                *regiao, tema=self.tema_alvo_ocr(indice), etiqueta=f"OCR {indice + 1}",
                duracao=0.18, bloqueante=True, margem=0,
            )
        safe_print(f"OCR {indice + 1} confirmado no recorte: '{capturado}'.", flush=True)
        return confirmado

    def localizar_referencia_hibrida_por_primeiro_ocr(
        self,
        caminho_referencia,
        texto_ocr_esperado="",
        region=None,
        dimensoes_referencia=None,
        confianca_ocr=0.72,
        confianca_visual=0.8,
        perfil_busca="balanced",
        leitura_relativa=None,
        _escala_ocr=None,
    ):
        """Confirma a referencia a partir do primeiro alvo OCR disponivel.

        Cada alvo usa sua propria geometria salva no print de referencia. Isso
        preserva a busca independente de monitor/DPI e permite que cada OCR
        tenha uma acao diferente no fluxo.
        """
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        partes = self.partes_alvos_ocr_esperados(texto_ocr_esperado, leitura)
        layouts = leitura.get("ocr_target_layouts") if isinstance(leitura.get("ocr_target_layouts"), list) else []
        geometrias = self.montar_layouts_alvos_ocr(
            leitura,
            dimensoes_referencia=dimensoes_referencia,
            quantidade=len(partes),
        )
        testar_todos = bool(leitura.get("_testar_todos_alvos_ocr"))
        resultados = []
        ultimo_confirmado = None
        for indice, esperado in enumerate(partes):
            if self.cancelamento_solicitado():
                return None
            if not str(esperado or "").strip():
                continue
            geometria = geometrias[indice] if indice < len(geometrias) else {}
            layout_origem = layouts[indice] if indice < len(layouts) and isinstance(layouts[indice], dict) else {}
            layout_unico = dict(layout_origem)
            layout_unico["texto"] = str(esperado).strip()
            leitura_alvo = {
                **leitura,
                "_ocr_individual_resolvido": True,
                "ocr_target_index": indice,
                # A leitura individual recebe apenas sua propria geometria,
                # mas o marcador ainda deve manter a cor e o numero originais
                # do OCR dentro do bloco.
                "ocr_target_marker_index": indice,
                "ocr_target_action": layout_origem,
                "offset_x": float(geometria.get("offset_x", 0) or 0),
                "offset_y": float(geometria.get("offset_y", 0) or 0),
                "largura": max(16, float(geometria.get("width", 140) or 140)),
                "altura": max(12, float(geometria.get("height", 40) or 40)),
                "ocr_target_layouts": [layout_unico],
            }
            safe_print(
                f"OCR alternativo {indice + 1}/{len(partes)}: procurando '{esperado}'.",
                flush=True,
            )
            confirmado = self.localizar_referencia_hibrida_ocr_visual(
                caminho_referencia,
                texto_ocr_esperado=str(esperado).strip(),
                region=region,
                dimensoes_referencia=dimensoes_referencia,
                confianca_ocr=confianca_ocr,
                confianca_visual=confianca_visual,
                perfil_busca=perfil_busca,
                leitura_relativa=leitura_alvo,
                ocr_confirmado=None,
                _escala_ocr=_escala_ocr,
            )
            if not confirmado and ultimo_confirmado:
                confirmado = self.reler_alvo_ocr_na_referencia(
                    ultimo_confirmado, indice, esperado, leitura, dimensoes_referencia,
                )
            status = (
                "confirmed" if confirmado and confirmado.get("visual_confirmed")
                else "visual_failed" if confirmado and confirmado.get("ocr_anchor_confirmed")
                else "not_found"
            )
            resultados.append({"index": indice, "text": str(esperado).strip(), "status": status})
            if status != "confirmed":
                motivo = "imagem nao confirmou" if status == "visual_failed" else "texto nao encontrado"
                safe_print(f"Falha OCR {indice + 1}: {motivo} ('{esperado}').", flush=True)
                # Se outro alvo ja localizou a referencia, podemos mostrar
                # onde este OCR falhou sem inventar uma posicao na tela.
                if not confirmado and ultimo_confirmado and leitura.get("_mostrar_marcadores_ocr", True):
                    regioes = self.regioes_alvos_ocr_por_layout(
                        ultimo_confirmado, leitura, dimensoes_referencia, quantidade=len(partes),
                    )
                    regiao_falha = self.limitar_regiao(regioes[indice])
                    if regiao_falha:
                        self.marcar_regiao_absoluta(
                            *regiao_falha, tema="ocr_alvo_falha",
                            etiqueta=f"Falha OCR {indice + 1}", duracao=0.28, bloqueante=True, margem=0,
                        )
            if confirmado and bool(confirmado.get("visual_confirmed")):
                ultimo_confirmado = confirmado
                confirmado["ocr_target_index"] = indice
                confirmado["ocr_target_action"] = layout_origem
                confirmado["ocr_target_individual"] = True
                if testar_todos:
                    continue
                regra = self.regra_acao_individual_ocr(layout_origem)
                # "Continuar" em um OCR intermediario significa avaliar o
                # proximo alvo; somente o ultimo libera a acao normal do bloco.
                if regra["acao"] == "continuar" and indice < len(partes) - 1:
                    safe_print(
                        f"OCR {indice + 1} confirmado. Regra: continuar para OCR {indice + 2}.",
                        flush=True,
                    )
                    continue
                safe_print(
                    f"OCR alternativo confirmado: OCR {indice + 1} ('{esperado}').",
                    flush=True,
                )
                return confirmado
            if testar_todos:
                continue
            regra = self.regra_acao_individual_ocr(layout_origem)
            if regra["falha"] == "parar":
                safe_print(
                    f"OCR {indice + 1} nao foi confirmado. Regra: parar automacao.",
                    flush=True,
                )
                return {
                    "ocr_target_individual": True,
                    "ocr_target_index": indice,
                    "ocr_target_action": layout_origem,
                    "ocr_target_flow": {"acao": "parar"},
                }
            if indice < len(partes) - 1:
                safe_print(
                    f"OCR {indice + 1} nao foi confirmado. Regra: procurar OCR {indice + 2}.",
                    flush=True,
                )
                continue
            # No ultimo OCR, "continuar" encerra este bloco sem erro e deixa
            # a sequencia seguir para o proximo comando da automacao.
            safe_print(
                f"OCR {indice + 1} nao foi confirmado. Regra: continuar automacao.",
                flush=True,
            )
            return {
                "ocr_target_individual": True,
                "ocr_target_index": indice,
                "ocr_target_action": layout_origem,
                "ocr_target_flow": {"acao": "continuar"},
            }
        if testar_todos:
            # Procurar e um diagnostico: percorre todos os alvos uma vez e
            # conserva inclusive as falhas. Nao executa cliques nem saltos.
            return {
                **(ultimo_confirmado or {}),
                "ocr_validation": {
                    "targets": resultados,
                    "visual_confirmed": ultimo_confirmado is not None,
                    "success": bool(resultados) and all(item["status"] == "confirmed" for item in resultados),
                },
            }
        return None

    def acao_individual_ocr_configurada(self, local, leitura_relativa=None):
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        if not bool(local.get("ocr_target_individual")):
            return None
        config = local.get("ocr_target_action") if isinstance(local.get("ocr_target_action"), dict) else {}
        if not config:
            config = leitura.get("ocr_target_action") if isinstance(leitura.get("ocr_target_action"), dict) else {}
        regra = self.regra_acao_individual_ocr(config)
        acao = regra["acao"]
        if acao not in {"clicar", "clicar_ir_para_bloco", "confirmar_ir_para_bloco"}:
            return None
        return {
            "acao": acao,
            "bloco": regra["bloco"],
            "clicks": regra["clicks"],
            "indice": int(local.get("ocr_target_index", 0) or 0),
        }

    def aplicar_acao_individual_ocr(self, local, leitura_relativa=None):
        acao = self.acao_individual_ocr_configurada(local, leitura_relativa)
        if not acao:
            return None
        if acao["acao"] == "confirmar_ir_para_bloco":
            if not acao["bloco"]:
                safe_print("OCR individual encontrou a regra de salto sem bloco de destino; mantendo o fluxo normal.", flush=True)
                return None
            local["ocr_target_flow"] = {"acao": "ir_para_bloco", "bloco": acao["bloco"]}
            safe_print(
                f"OCR {acao['indice'] + 1}: redirecionando para o bloco {acao['bloco']}.",
                flush=True,
            )
            return "ir_para_bloco"

        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        # A bolinha "Clique" mostrada em cada preview continua sendo o ponto
        # de clique do bloco. O OCR escolhido apenas decide qual regra sera
        # seguida; ele nao desloca o clique para o centro do texto encontrado.
        # Isso permite, por exemplo, ler o rotulo e clicar no botao ao lado.
        dimensoes = leitura.get("_dimensoes_referencia")
        offset_x, offset_y = self.escalar_offset_referencia(
            local,
            offset_x=leitura.get("_click_offset_x", 0),
            offset_y=leitura.get("_click_offset_y", 0),
            dimensoes_referencia=dimensoes,
        )
        destino_x, destino_y = self.centro_area_reconhecida(
            local,
            offset_x=offset_x,
            offset_y=offset_y,
        )
        self.marcar_regiao_absoluta(
            int(destino_x) - 9,
            int(destino_y) - 9,
            18,
            18,
            tema=self.tema_alvo_ocr(acao['indice']),
            etiqueta=f"Clique OCR {acao['indice'] + 1}",
            duracao=0.18,
            bloqueante=True,
            margem=0,
        )
        if self.cancelamento_solicitado():
            return "cancelado"
        time.sleep(0.04)
        pyautogui.click(x=destino_x, y=destino_y, clicks=acao["clicks"], interval=0.08)
        safe_print(
            f"OCR {acao['indice'] + 1}: clique individual executado em X/Y: {destino_x}/{destino_y} | clicks={acao['clicks']}",
            flush=True,
        )
        local["ocr_target_action_applied"] = "clicar"
        if acao["acao"] == "clicar_ir_para_bloco":
            if not acao["bloco"]:
                safe_print("OCR individual clicou, mas a regra nao possui bloco de destino; mantendo o fluxo normal.", flush=True)
                return "clicar"
            local["ocr_target_flow"] = {"acao": "ir_para_bloco", "bloco": acao["bloco"]}
            safe_print(
                f"OCR {acao['indice'] + 1}: clique concluido. Indo para bloco {acao['bloco']}.",
                flush=True,
            )
            return "ir_para_bloco"
        return "clicar"

    def localizar_referencia_hibrida_ocr_visual(
        self,
        caminho_referencia,
        texto_ocr_esperado="",
        region=None,
        dimensoes_referencia=None,
        confianca_ocr=0.72,
        confianca_visual=0.8,
        perfil_busca="balanced",
        leitura_relativa=None,
        ocr_confirmado=None,
        _escala_ocr=None,
    ):
        esperado = str(texto_ocr_esperado or "").strip()
        if not esperado or not OCR_DISPONIVEL or not dimensoes_referencia:
            return None
        try:
            ref_largura, ref_altura = [int(valor) for valor in dimensoes_referencia]
        except Exception:
            return None
        if ref_largura <= 0 or ref_altura <= 0:
            return None

        leitura_configurada = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        if (
            not leitura_configurada.get("_ocr_individual_resolvido")
            and self.modo_acao_alvos_ocr(leitura_configurada) == "individual"
        ):
            return self.localizar_referencia_hibrida_por_primeiro_ocr(
                caminho_referencia,
                texto_ocr_esperado=esperado,
                region=region,
                dimensoes_referencia=dimensoes_referencia,
                confianca_ocr=confianca_ocr,
                confianca_visual=confianca_visual,
                perfil_busca=perfil_busca,
                leitura_relativa=leitura_configurada,
                _escala_ocr=_escala_ocr,
            )

        spec_composto = self.interpretar_texto_ocr_composto(esperado)
        quantidade_esperada = self.quantidade_alvos_ocr_esperados(esperado, leitura_relativa)
        partes_obrigatorias = self.partes_alvos_ocr_esperados(esperado, leitura_relativa)
        partes_compostas = partes_obrigatorias or (spec_composto.get("partes", []) if spec_composto else [])
        if quantidade_esperada > 1 and len([parte for parte in partes_compostas if str(parte or "").strip()]) < quantidade_esperada:
            safe_print(
                "Hibrido OCR + imagem bloqueado: existem varios alvos OCR, mas o texto de busca chegou incompleto. "
                f"Recebido='{esperado}'.",
                flush=True,
            )
            return None
        reutilizar_ocr = (
            isinstance(ocr_confirmado, dict)
            and self.ocr_hibrido_tem_todos_alvos(ocr_confirmado, esperado, leitura_relativa)
        )
        if reutilizar_ocr:
            candidatos_ocr = [ocr_confirmado]
        else:
            spec_busca = spec_composto
            if quantidade_esperada > 1 and not spec_busca:
                spec_busca = {
                    "partes": partes_compostas,
                    "relacional": any(re.search(r"\d", str(parte or "")) for parte in partes_compostas[1:]),
                    "separador": "layout",
                }
            layout_alvos = self.montar_layouts_alvos_ocr(
                leitura_relativa,
                dimensoes_referencia=dimensoes_referencia,
                quantidade=quantidade_esperada,
            ) if spec_busca else []
            if spec_busca and quantidade_esperada >= 2:
                candidatos_ocr = self.localizar_texto_composto_por_ocr_screenshot(
                    esperado,
                    spec_busca,
                    confianca=max(0.55, min(float(confianca_ocr or 0.72), 0.78)),
                    janela_alvo=None,
                    region=region,
                    retornar_candidatos=True,
                    marcar_candidatos=False,
                    limite_candidatos=3 if bool(getattr(self, "modo_teste_reconhecimento", False)) else 6,
                    layout_alvos=layout_alvos,
                    _escala_ocr=_escala_ocr,
                    _tentativa_ampliada=bool(_escala_ocr is not None),
                )
            else:
                candidatos_ocr = self.localizar_texto_por_ocr_screenshot(
                    esperado,
                    confianca=max(0.55, min(float(confianca_ocr or 0.72), 0.78)),
                    janela_alvo=None,
                    region=region,
                    retornar_candidatos=True,
                    marcar_candidatos=False,
                    limite_candidatos=3 if bool(getattr(self, "modo_teste_reconhecimento", False)) else 6,
                    _escala_ocr=_escala_ocr,
                    _tentativa_ampliada=bool(_escala_ocr is not None),
                )
        if not candidatos_ocr:
            return None

        partes = partes_compostas or [esperado]
        for candidato in candidatos_ocr:
            if self.cancelamento_solicitado():
                return None
            # A similaridade da busca gera candidatos; so texto confirmado libera o bloco.
            bounds = candidato.get("ocr_related_bounds") or [candidato]
            if len(bounds) < quantidade_esperada:
                continue
            if not all(
                self.texto_ocr_confirma_alvo(
                    parte, bound.get("texto") or bound.get("text") or "", limite_ratio=0.86
                )
                for parte, bound in zip(partes, bounds)
            ):
                continue

            ancora = bounds[0]
            local = self.estimar_referencia_por_alvos_ocr(
                bounds,
                leitura_relativa=leitura_relativa,
                dimensoes_referencia=dimensoes_referencia,
            )
            if not local:
                ax, ay, aw, ah = self._retangulo_texto_ocr(ancora)
                layouts = self.montar_layouts_alvos_ocr(
                    leitura_relativa, dimensoes_referencia, quantidade=quantidade_esperada
                )
                esquerda = round(ax + aw / 2 - layouts[0]["offset_x"] - ref_largura / 2)
                topo = round(ay + ah / 2 - layouts[0]["offset_y"] - ref_altura / 2)
                local = {
                    "left": esquerda, "top": topo,
                    "width": ref_largura, "height": ref_altura,
                    "x": round(esquerda + ref_largura / 2),
                    "y": round(topo + ref_altura / 2),
                    "ocr_layout_scale_x": 1.0,
                    "ocr_layout_scale_y": 1.0,
                }
            local["score"] = float(candidato.get("score", 0.0) or 0.0)
            local.update({
                "search_stage": "ocr_visual_hybrid",
                "ocr_anchor_confirmed": True,
                "visual_confirmed": False,
                "ocr_text": str(candidato.get("text") or candidato.get("texto") or esperado),
                "ocr_bounds": dict(ancora),
                "ocr_layout_targets_marked": bool(candidato.get("ocr_layout_targets_marked")),
            })
            for chave in (
                "ocr_composto", "ocr_relacional", "ocr_partes", "ocr_textos",
                "ocr_related_bounds", "ocr_anchor_bounds", "ocr_value_bounds",
                "ocr_union_bounds", "ocr_relation_score",
            ):
                if chave in candidato:
                    local[chave] = candidato[chave]
            if not self.ocr_hibrido_tem_todos_alvos(local, esperado, leitura_relativa):
                continue
            if not local.get("ocr_layout_targets_marked"):
                self.marcar_alvos_ocr_em_sequencia(local, leitura_relativa, dimensoes_referencia)
                if self.cancelamento_solicitado():
                    return None
                local["ocr_layout_targets_marked"] = True
                safe_print(
                    f"Hibrido OCR + imagem: etapa 1/2 concluida com {quantidade_esperada} OCR(s) confirmados.",
                    flush=True,
                )

            esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
            # A geometria dos OCRs ja calcula onde a referencia deve estar. A
            # confirmacao visual comeca exatamente nesse retangulo para impedir
            # que detalhes genericos da janela (bordas, campos cinza etc.)
            # ganhem a busca ampliada antes da imagem que corresponde aos OCRs.
            regiao_visual_exata = self.limitar_regiao(
                (esquerda, topo, largura, altura),
                limite=region,
            )
            margem_x = max(46, round(largura * 0.42))
            margem_y = max(34, round(altura * 0.42))
            regiao_visual_ampliada = self.limitar_regiao(
                (esquerda - margem_x, topo - margem_y,
                 largura + margem_x * 2, altura + margem_y * 2),
                limite=region,
            )
            # A primeira passagem da busca balanceada e rapida. O reconhecimento
            # hibrido precisa respeitar esse perfil; antes ele voltava para
            # "balanced" aqui e fazia a varredura visual pesada mesmo quando a
            # etapa externa ja tinha escolhido "fast".
            perfil_visual = str(perfil_busca or "balanced").strip().lower()
            if perfil_visual not in {"fast", "balanced", "precise"}:
                perfil_visual = "balanced"
            escala_visual = (
                float(local.get("ocr_layout_scale_x", 1.0) or 1.0)
                + float(local.get("ocr_layout_scale_y", 1.0) or 1.0)
            ) / 2.0
            regioes_visuais = []
            for nome_regiao, regiao_candidata in (
                ("exata", regiao_visual_exata),
                ("ampliada", regiao_visual_ampliada),
            ):
                if not regiao_candidata:
                    continue
                regiao_candidata = tuple(regiao_candidata)
                if any(regiao_candidata == regiao_existente for _, regiao_existente in regioes_visuais):
                    continue
                regioes_visuais.append((nome_regiao, regiao_candidata))

            for indice_regiao, (nome_regiao, regiao_visual) in enumerate(regioes_visuais):
                if indice_regiao == 0 and nome_regiao == "exata":
                    safe_print(
                        "Hibrido OCR + imagem: etapa 2/2 conferindo a imagem na regiao exata estimada pelos OCRs "
                        f"(escala estimada={escala_visual:.2f}).",
                        flush=True,
                    )
                elif indice_regiao == 0:
                    safe_print(
                        "Hibrido OCR + imagem: etapa 2/2 procurando a imagem proxima aos OCRs "
                        f"(escala estimada={escala_visual:.2f}).",
                        flush=True,
                    )
                else:
                    safe_print(
                        "Hibrido OCR + imagem: sem confirmacao na regiao exata; ampliando a busca ao redor dos OCRs.",
                        flush=True,
                    )

                visual = self.procurar_imagem_por_template(
                    caminho_referencia,
                    confianca=max(0.55, min(float(confianca_visual or 0.8), 0.9)),
                    region=regiao_visual,
                    perfil_busca=perfil_visual,
                    escala_referencia=escala_visual,
                    # Os OCRs da tela acima ja confirmaram todos os alvos do bloco.
                    # Nao reler a imagem de referencia evita dezenas de chamadas ao
                    # Tesseract, inclusive quando o usuario escolhe o perfil preciso.
                    incluir_ocr_referencia=False,
                )
                if self.cancelamento_solicitado():
                    return None
                if not visual:
                    continue

                vx, vy, vw, vh = self.bounds_reconhecimento(visual)
                centro_visual_x = vx + (vw / 2.0)
                centro_visual_y = vy + (vh / 2.0)
                centro_ocr_x = esquerda + (largura / 2.0)
                centro_ocr_y = topo + (altura / 2.0)
                proximo_ao_ocr = (
                    abs(centro_visual_x - centro_ocr_x) <= max(34, largura * 0.62, vw * 0.62)
                    and abs(centro_visual_y - centro_ocr_y) <= max(28, altura * 0.62, vh * 0.62)
                )
                if proximo_ao_ocr:
                    confirmado = dict(visual)
                    confirmado.update(local)
                    confirmado.update({
                        "left": int(vx),
                        "top": int(vy),
                        "width": int(vw),
                        "height": int(vh),
                        "x": int(round(centro_visual_x)),
                        "y": int(round(centro_visual_y)),
                        "ocr_score": float(local.get("score", 0.0) or 0.0),
                        "score": float(visual.get("rank_score", visual.get("score", 0.0)) or 0.0),
                        "visual_confirmed": True,
                        "visual_scale": float(visual.get("scale", escala_visual) or escala_visual),
                    })
                    if self.modo_acao_alvos_ocr(leitura_configurada) == "individual":
                        confirmado["ocr_target_individual"] = True
                        confirmado["ocr_target_index"] = int(leitura_configurada.get("ocr_target_index", 0) or 0)
                        confirmado["ocr_target_action"] = (
                            leitura_configurada.get("ocr_target_action")
                            if isinstance(leitura_configurada.get("ocr_target_action"), dict)
                            else (leitura_configurada.get("ocr_target_layouts") or [{}])[0]
                        )
                    safe_print(
                        f"Hibrido OCR + imagem: etapa 2/2 confirmada pela imagem (escala={confirmado['visual_scale']:.2f}).",
                        flush=True,
                    )
                    return confirmado
                safe_print(
                    "Hibrido OCR + imagem: a imagem encontrada ficou fora da regiao confirmada pelos OCRs; descartada.",
                    flush=True,
                )
            safe_print(
                "Hibrido OCR + imagem: os OCRs foram confirmados, mas a imagem ainda nao confirmou a referencia.",
                flush=True,
            )
            return local
        # A leitura nativa e muito mais rapida, mas labels pequenos podem
        # gerar um candidato parecido sem atingir a confirmacao estrita. Nesse
        # caso fazemos uma unica releitura ampliada antes de repetir o ciclo;
        # nao deixamos o Tesseract rodar continuamente e travar o mouse.
        if _escala_ocr is None:
            safe_print(
                "OCR candidato parcial; conferindo uma vez a mesma janela ampliada.",
                flush=True,
            )
            return self.localizar_referencia_hibrida_ocr_visual(
                caminho_referencia,
                texto_ocr_esperado=texto_ocr_esperado,
                region=region,
                dimensoes_referencia=dimensoes_referencia,
                confianca_ocr=confianca_ocr,
                confianca_visual=confianca_visual,
                perfil_busca=perfil_busca,
                leitura_relativa=leitura_relativa,
                ocr_confirmado=None,
                _escala_ocr=1.35,
            )
        safe_print("OCR: nem todos os textos obrigatorios confirmaram; bloco aguardando.", flush=True)
        return None

    def localizar_referencia_por_recorte_ocr_visual(
        self,
        caminho_referencia,
        region=None,
        leitura_relativa=None,
        dimensoes_referencia=None,
        confianca=0.78,
    ):
        if not OPENCV_DISPONIVEL or np is None or Image is None or not caminho_referencia or not dimensoes_referencia:
            return None
        regiao_ref_ocr = self.construir_regiao_ocr_referencia(
            leitura_relativa,
            dimensoes_referencia=dimensoes_referencia,
        )
        if not regiao_ref_ocr:
            return None
        try:
            ref_largura, ref_altura = [int(valor) for valor in dimensoes_referencia]
            ref_x, ref_y, ref_w, ref_h = [int(valor) for valor in regiao_ref_ocr]
            if ref_w <= 4 or ref_h <= 4:
                return None
            with Image.open(caminho_referencia) as imagem_ref:
                recorte_ref = imagem_ref.convert("RGB").crop((ref_x, ref_y, ref_x + ref_w, ref_y + ref_h))
            captura, base_left, base_top = self.capturar_tela_pil(region=region)
            tela_gray = cv2.cvtColor(np.array(captura.convert("RGB")), cv2.COLOR_RGB2GRAY)
            ref_gray = cv2.cvtColor(np.array(recorte_ref), cv2.COLOR_RGB2GRAY)
            if tela_gray.shape[0] < ref_gray.shape[0] or tela_gray.shape[1] < ref_gray.shape[1]:
                return None
            resultado = cv2.matchTemplate(tela_gray, ref_gray, cv2.TM_CCOEFF_NORMED)
            _min_val, max_val, _min_loc, max_loc = cv2.minMaxLoc(resultado)
            score = float(max_val or 0.0)
            if score < float(confianca or 0.78):
                return None

            ocr_left = int(base_left + max_loc[0])
            ocr_top = int(base_top + max_loc[1])
            centro_ocr_x = ocr_left + (ref_w / 2.0)
            centro_ocr_y = ocr_top + (ref_h / 2.0)
            centro_ref_x = ref_x + (ref_w / 2.0)
            centro_ref_y = ref_y + (ref_h / 2.0)
            esquerda = int(round(centro_ocr_x - centro_ref_x))
            topo = int(round(centro_ocr_y - centro_ref_y))
            limite = self.limitar_regiao((esquerda, topo, ref_largura, ref_altura))
            if limite:
                esquerda, topo, ref_largura, ref_altura = limite
            safe_print(
                "Referencia localizada pelo recorte visual do Ler aqui "
                f"(score={score:.2f}, area={self.formatar_regiao_log(region)}).",
                flush=True,
            )
            return {
                "x": int(esquerda + ref_largura / 2),
                "y": int(topo + ref_altura / 2),
                "left": int(esquerda),
                "top": int(topo),
                "width": int(ref_largura),
                "height": int(ref_altura),
                "score": score,
                "search_stage": "ocr_visual_anchor",
                "ocr_visual_anchor": True,
                "ocr_bounds": {
                    "left": int(ocr_left),
                    "top": int(ocr_top),
                    "width": int(ref_w),
                    "height": int(ref_h),
                },
            }
        except Exception as erro:
            safe_print(f"Falha ao localizar recorte visual do Ler aqui: {erro}", flush=True)
            return None

    def confirmar_ocr_relativo(self, local, texto_esperado="", leitura_relativa=None, ordem_bloco=None, dimensoes_referencia=None, marcar_ocr=True, caminho_referencia=None):
        esperado = str(texto_esperado or "").strip()
        leitura = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        salvar_em = str(leitura.get("salvar_em", "") or "").strip()
        if not esperado and not salvar_em:
            return True
        if not OCR_DISPONIVEL:
            safe_print("OCR nao confirmou: recurso OCR indisponivel.", flush=True)
            return not esperado
        quantidade_alvos = self.quantidade_alvos_ocr_esperados(esperado, leitura_relativa)
        if esperado and quantidade_alvos > 1:
            confirmado_multi = self.confirmar_alvos_ocr_por_layout_visual(
                local,
                texto_esperado=esperado,
                leitura_relativa=leitura_relativa,
                dimensoes_referencia=dimensoes_referencia,
                caminho_referencia=None,
                marcar=marcar_ocr,
            )
            if not confirmado_multi:
                safe_print(
                    f"OCR multiplo nao confirmou todos os {quantidade_alvos} alvo(s); bloco bloqueado.",
                    flush=True,
                )
                return False
            capturado_multi = str(confirmado_multi.get("ocr_text") or esperado).strip()
            if capturado_multi:
                self.salvar_variavel_contexto("ultimo_ocr", capturado_multi)
                if ordem_bloco:
                    self.salvar_variavel_contexto(f"ocr_bloco_{ordem_bloco}", capturado_multi)
                    self.salvar_variavel_contexto(f"bloco_{ordem_bloco}_ocr", capturado_multi)
                if salvar_em:
                    self.salvar_variavel_contexto(salvar_em, capturado_multi)
            safe_print(
                f"OCR multiplo confirmou todos os {quantidade_alvos} alvo(s): {capturado_multi}",
                flush=True,
            )
            return True
        regiao = self.construir_regiao_ocr_relativa(local, leitura, dimensoes_referencia=dimensoes_referencia)
        safe_print(f"OCR lendo area: {self.formatar_regiao_log(regiao)}", flush=True)

        def ler_texto_regiao(regiao_ocr):
            try:
                captura, _base_left, _base_top = self.capturar_tela_pil(region=regiao_ocr)
            except Exception as erro:
                safe_print(f"OCR nao confirmou: falha ao capturar area de leitura ({erro}).", flush=True)
                return "", None
            textos = self.extrair_textos_ocr_rapido(captura, texto_esperado=esperado)
            if not textos:
                textos = [
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 6"),
                    self.extrair_texto_ocr_imagem_pil(captura, config="--psm 7"),
                ]
            return " ".join(texto for texto in textos if texto).strip(), captura

        def texto_confirma(capturado_atual):
            if not esperado:
                return True
            comparacao_atual = self.comparar_texto_reconhecimento_preciso(esperado, capturado_atual)
            return (
                comparacao_atual.get("reference") == comparacao_atual.get("captured")
                or (
                    comparacao_atual.get("has_digits")
                    and float(comparacao_atual.get("digit_ratio", 0.0) or 0.0) >= 0.999
                )
                or float(comparacao_atual.get("ratio", 0.0) or 0.0) >= 0.96
            )

        capturado, captura_ocr = ler_texto_regiao(regiao)
        confirmado = texto_confirma(capturado)

        if not confirmado:
            limite_reconhecido = self.bounds_reconhecimento(local)
            try:
                esquerda, topo, largura, altura = [int(valor) for valor in regiao]
                local_esquerda, local_topo, local_largura, local_altura = limite_reconhecido
                margem_x = max(14, int(largura * 0.28))
                margem_y = max(8, int(altura * 0.45))
                limite_ocr = self.limitar_regiao((
                    local_esquerda - 16,
                    local_topo - 10,
                    local_largura + 32,
                    local_altura + 20,
                ))
                regiao_ampliada = self.limitar_regiao(
                    (
                        esquerda - margem_x,
                        topo - margem_y,
                        largura + margem_x * 2,
                        altura + margem_y * 2,
                    ),
                    limite=limite_ocr or limite_reconhecido,
                )
            except Exception:
                regiao_ampliada = None
            if regiao_ampliada and regiao_ampliada != regiao:
                safe_print(f"OCR tentando area ampliada: {self.formatar_regiao_log(regiao_ampliada)}", flush=True)
                capturado_ampliado, captura_ampliada = ler_texto_regiao(regiao_ampliada)
                if texto_confirma(capturado_ampliado):
                    regiao = regiao_ampliada
                    capturado = capturado_ampliado
                    captura_ocr = captura_ampliada
                    confirmado = True
                elif not capturado and capturado_ampliado:
                    regiao = regiao_ampliada
                    capturado = capturado_ampliado
                    captura_ocr = captura_ampliada

        if not confirmado and caminho_referencia:
            confirmado_visual = self.confirmar_ocr_visual_pela_referencia(
                captura_ocr,
                caminho_referencia=caminho_referencia,
                leitura_relativa=leitura_relativa,
                dimensoes_referencia=dimensoes_referencia,
            )
            if confirmado_visual:
                capturado = capturado or esperado
                confirmado = True

        if capturado:
            self.salvar_variavel_contexto("ultimo_ocr", capturado)
            if ordem_bloco:
                self.salvar_variavel_contexto(f"ocr_bloco_{ordem_bloco}", capturado)
                self.salvar_variavel_contexto(f"bloco_{ordem_bloco}_ocr", capturado)
            if salvar_em:
                self.salvar_variavel_contexto(salvar_em, capturado)
        if not esperado:
            return True
        safe_print(
            f"OCR {'confirmou' if confirmado else 'nao confirmou'}: esperado='{esperado}' | capturado='{capturado}'",
            flush=True,
        )
        if marcar_ocr:
            self.marcar_regiao_ocr(
                regiao,
                confirmado=confirmado,
                duracao=0.9 if confirmado else 1.15,
            )
        return confirmado

    def confirmar_localizacao_imagem(self, caminho_imagem, local, confianca=0.8, perfil_busca="balanced"):
        perfil_busca = str(perfil_busca or "balanced").strip().lower()
        if perfil_busca not in {"fast", "balanced", "precise"}:
            perfil_busca = "balanced"
        modo_estrito = float(confianca or 0.8) >= 0.99

        if perfil_busca == "fast" and isinstance(local, dict):
            score_local = float(local.get("score", 0.0) or 0.0)
            if score_local >= max(0.9, float(confianca or 0.8) - 0.04):
                return local

        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        margem = 6 if modo_estrito else (12 if perfil_busca == "fast" else 24)
        regiao = (
            int(esquerda - margem),
            int(topo - margem),
            int(largura + margem * 2),
            int(altura + margem * 2),
        )
        confianca_confirmacao = max(0.88, float(confianca or 0.8))

        try:
            local_template = self.procurar_imagem_por_template(
                caminho_imagem,
                confianca=confianca_confirmacao,
                region=regiao,
                perfil_busca=perfil_busca,
            )
            if local_template and self.validar_correspondencia_visual(
                caminho_imagem,
                local_template,
                confianca_confirmacao,
                perfil_busca=perfil_busca,
            ):
                return local_template
        except Exception:
            pass

        if perfil_busca == "fast" or modo_estrito:
            return None

        try:
            if OPENCV_DISPONIVEL:
                local_confirmado = pyautogui.locateOnScreen(
                    str(caminho_imagem),
                    confidence=confianca_confirmacao,
                    grayscale=False,
                    region=regiao,
                )
            else:
                local_confirmado = pyautogui.locateOnScreen(
                    str(caminho_imagem),
                    grayscale=False,
                    region=regiao,
                )
            if local_confirmado and self.validar_correspondencia_visual(
                caminho_imagem,
                local_confirmado,
                confianca_confirmacao,
                perfil_busca=perfil_busca,
            ):
                return local_confirmado
        except Exception:
            return None
        return None

    def clicar_area_reconhecida(
        self,
        local,
        offset_x=0,
        offset_y=0,
        button="left",
        clicks=1,
        janela_alvo=None,
        preservar_janelas=True,
    ):
        destino_x, destino_y = self.centro_area_reconhecida(local, offset_x=offset_x, offset_y=offset_y)
        safe_print(f"Clicando em X/Y: {destino_x}/{destino_y}", flush=True)
        # O ponto veio da captura que ja esta visivel. Preservar as janelas
        # evita reativar um contexto antigo entre o reconhecimento e o clique.
        janela_preparada = None
        if not bool(preservar_janelas):
            janela_preparada = self.preparar_janela_para_acao(
                janela_alvo,
                priorizar_evento=bool(str(janela_alvo or "").strip()),
            )
        if janela_preparada:
            time.sleep(0.08)
        total_clicks = max(1, int(clicks or 1))
        if hasattr(ctypes, "windll"):
            user32 = ctypes.windll.user32
            eventos = {
                "left": (0x0002, 0x0004),
                "right": (0x0008, 0x0010),
                "middle": (0x0020, 0x0040),
            }
            down_flag, up_flag = eventos.get(str(button or "left").lower(), (0x0002, 0x0004))
            user32.SetCursorPos(int(destino_x), int(destino_y))
            time.sleep(0.06)
            for indice in range(total_clicks):
                user32.mouse_event(down_flag, 0, 0, 0, 0)
                time.sleep(0.05)
                user32.mouse_event(up_flag, 0, 0, 0, 0)
                if indice < total_clicks - 1:
                    time.sleep(0.09)
        else:
            pyautogui.moveTo(destino_x, destino_y, duration=0.12)
            time.sleep(0.06)
            for indice in range(total_clicks):
                pyautogui.mouseDown(x=destino_x, y=destino_y, button=button)
                time.sleep(0.04)
                pyautogui.mouseUp(x=destino_x, y=destino_y, button=button)
                if indice < total_clicks - 1:
                    time.sleep(0.09)
        self._ultimo_clique_visual_em = time.perf_counter()
        self._ultima_posicao_clique_visual = {
            "x": int(destino_x),
            "y": int(destino_y),
            "button": str(button or "left"),
            "clicks": int(total_clicks),
        }
        time.sleep(0.12)
        safe_print(f"Clique executado em X/Y: {destino_x}/{destino_y} | clicks={total_clicks}", flush=True)
        return {"x": destino_x, "y": destino_y}

    def liberar_teclas_modificadoras(self):
        for tecla in ("shift", "ctrl", "alt", "win"):
            try:
                pyautogui.keyUp(tecla)
            except Exception:
                continue

    def preparar_foco_para_digitacao(self):
        if self.cancelamento_solicitado():
            return False

        if self._ultimo_clique_visual_em:
            decorrido = time.perf_counter() - float(self._ultimo_clique_visual_em or 0.0)
            restante = max(0.0, TEMPO_ESTABILIZACAO_APOS_CLIQUE - decorrido)
            if restante > 0 and not self.esperar(restante):
                return False

            dados_clique = self._ultima_posicao_clique_visual or {}
            if dados_clique:
                safe_print(
                    "Aguardando foco estabilizar apos clique visual "
                    f"em X/Y {dados_clique.get('x', 0)}/{dados_clique.get('y', 0)}.",
                    flush=True,
                )

        self.liberar_teclas_modificadoras()
        if TEMPO_RESPIRO_ANTES_DIGITACAO > 0 and not self.esperar(TEMPO_RESPIRO_ANTES_DIGITACAO):
            return False
        return True

    def extrair_texto_referencia(self, caminho_imagem):
        if not OCR_DISPONIVEL:
            return ""

        try:
            with Image.open(caminho_imagem) as imagem:
                texto = pytesseract.image_to_string(imagem, config="--psm 6")
        except Exception:
            return ""

        return self.normalizar_texto_ocr(texto)

    def extrair_texto_ocr_imagem_pil(self, imagem, config="--psm 7"):
        if not OCR_DISPONIVEL or imagem is None:
            return ""
        try:
            texto = pytesseract.image_to_string(imagem, config=config)
        except Exception:
            return ""
        return self.normalizar_texto_ocr(texto)

    def ampliar_imagem_para_ocr(self, imagem, fator=3):
        if imagem is None or Image is None:
            return imagem
        try:
            largura = max(1, int(imagem.width * fator))
            altura = max(1, int(imagem.height * fator))
            resampling = getattr(Image, "Resampling", None)
            metodo = getattr(resampling, "LANCZOS", getattr(Image, "LANCZOS", 1))
            return imagem.resize((largura, altura), metodo)
        except Exception:
            return imagem

    def gerar_variantes_imagem_ocr(self, imagem, fator=2, incluir_cor=True):
        if imagem is None or Image is None:
            return []
        variantes = []
        vistos = set()

        def adicionar(nome, img):
            if img is None:
                return
            chave = (nome, getattr(img, "mode", ""), int(getattr(img, "width", 0) or 0), int(getattr(img, "height", 0) or 0))
            if chave in vistos:
                return
            vistos.add(chave)
            variantes.append((nome, img))

        try:
            base_rgb = imagem.convert("RGB")
            base_l = imagem.convert("L")
            ampliada_l = self.ampliar_imagem_para_ocr(base_l, fator=fator)
            adicionar("cinza", ampliada_l)
            if incluir_cor:
                adicionar("cor", self.ampliar_imagem_para_ocr(base_rgb, fator=fator))
            if ImageOps is not None:
                adicionar("autocontraste", ImageOps.autocontrast(ampliada_l))
            if ImageEnhance is not None:
                adicionar("contraste", ImageEnhance.Contrast(ampliada_l).enhance(1.9))
                adicionar("brilho_contraste", ImageEnhance.Contrast(ImageEnhance.Brightness(ampliada_l).enhance(1.25)).enhance(2.2))
            try:
                histograma = ampliada_l.histogram()
                total_pixels = max(1, sum(histograma))
                media = int(sum(indice * valor for indice, valor in enumerate(histograma)) / total_pixels)
            except Exception:
                media = 128
            limiares = [118, 135, 150, 170] if media >= 120 else [92, 110, 128, 145]
            invertida = ampliada_l.point(lambda valor: 255 - valor)
            adicionar("invertida", invertida)
            for limiar in limiares:
                adicionar(f"limiar_claro_{limiar}", ampliada_l.point(lambda valor, t=limiar: 255 if valor >= t else 0))
                adicionar(f"limiar_escuro_{limiar}", ampliada_l.point(lambda valor, t=limiar: 0 if valor >= t else 255))
                adicionar(f"invertida_limiar_{limiar}", invertida.point(lambda valor, t=limiar: 255 if valor >= t else 0))
        except Exception:
            return variantes
        return variantes

    def extrair_textos_ocr_rapido(self, imagem, texto_esperado=""):
        if not OCR_DISPONIVEL or imagem is None or Image is None:
            return []

        esperado = self.normalizar_texto_ocr(texto_esperado)
        possui_digitos = bool(re.search(r"\d", esperado))
        try:
            base = imagem.convert("L")
            maior_lado = max(int(base.width or 0), int(base.height or 0))
            fator = 4 if maior_lado < 130 else (3 if maior_lado < 260 else 2)
            ampliada = self.ampliar_imagem_para_ocr(base, fator=fator)
            variantes = [ampliada]
            try:
                histograma = ampliada.histogram()
                total = max(1, sum(histograma))
                media = int(sum(indice * valor for indice, valor in enumerate(histograma)) / total)
            except Exception:
                media = 128
            limiar = 128 if media < 120 else 150
            variantes.append(ampliada.point(lambda valor: 255 if valor >= limiar else 0))
            if media < 105:
                variantes.append(ampliada.point(lambda valor: 255 - valor))
        except Exception:
            return []

        whitelist = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ "
        if possui_digitos and not re.search(r"[a-zA-Z]", esperado):
            whitelist = "0123456789/.-:"
        configs = [
            f"--psm 7 -c tessedit_char_whitelist={whitelist}",
            f"--psm 6 -c tessedit_char_whitelist={whitelist}",
        ]
        if len(esperado.split()) <= 1:
            configs.insert(0, f"--psm 8 -c tessedit_char_whitelist={whitelist}")

        textos = []
        for variante in variantes:
            for config in configs:
                texto = self.extrair_texto_ocr_imagem_pil(variante, config=config)
                if texto and texto not in textos:
                    textos.append(texto)
                if esperado and texto:
                    comparacao = self.comparar_texto_reconhecimento_preciso(esperado, texto)
                    if (
                        comparacao.get("reference") == comparacao.get("captured")
                        or (
                            comparacao.get("has_digits")
                            and float(comparacao.get("digit_ratio", 0.0) or 0.0) >= 0.999
                        )
                        or float(comparacao.get("ratio", 0.0) or 0.0) >= 0.96
                    ):
                        return textos
        return textos

    def extrair_variantes_texto_ocr_imagem(self, imagem):
        if not OCR_DISPONIVEL or imagem is None or Image is None:
            return []

        try:
            base = imagem.convert("L")
            fator = 5 if max(int(base.width or 0), int(base.height or 0)) < 220 else 3
            ampliada = self.ampliar_imagem_para_ocr(base, fator=fator)
            invertida = ampliada.point(lambda valor: 255 - valor)
            media = 128
            try:
                histograma = ampliada.histogram()
                total = max(1, sum(histograma))
                media = int(sum(indice * valor for indice, valor in enumerate(histograma)) / total)
            except Exception:
                media = 128
            limiar_base = 132 if media < 120 else 150
            limiar_claro = ampliada.point(lambda valor: 255 if valor >= limiar_base else 0)
            limiar_escuro = ampliada.point(lambda valor: 0 if valor >= limiar_base else 255)
            limiar_invertido = invertida.point(lambda valor: 255 if valor >= 118 else 0)
        except Exception:
            return []

        configs = [
            "--psm 6",
            "--psm 7",
            "--psm 8",
            "--psm 13",
            "--psm 6 -c preserve_interword_spaces=1",
            "--psm 7 -c preserve_interword_spaces=1",
            "--psm 6 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ ",
            "--psm 7 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ ",
            "--psm 7 -c tessedit_char_whitelist=0123456789/.-:",
            "--psm 8 -c tessedit_char_whitelist=0123456789/.-:",
        ]
        variantes = [ampliada, invertida, limiar_claro, limiar_escuro, limiar_invertido]
        textos = []

        for variante in variantes:
            for config in configs:
                texto = self.extrair_texto_ocr_imagem_pil(variante, config=config)
                if texto and texto not in textos:
                    textos.append(texto)
        return textos

    def texto_referencia_critico(self, texto):
        normalizado = self.normalizar_texto_ocr(texto)
        if len(normalizado.replace(" ", "")) < 4:
            return False
        digitos = re.sub(r"\D", "", normalizado)
        return len(digitos) >= 4

    def comparar_texto_reconhecimento_preciso(self, texto_referencia, texto_capturado):
        referencia = self.normalizar_texto_ocr(texto_referencia)
        capturado = self.normalizar_texto_ocr(texto_capturado)
        if not referencia or not capturado:
            return {
                "ratio": 0.0,
                "digit_ratio": 0.0,
                "reference": referencia,
                "captured": capturado,
                "has_digits": False,
            }

        ratio = float(difflib.SequenceMatcher(None, referencia, capturado).ratio())
        digitos_referencia = re.sub(r"\D", "", referencia)
        digitos_capturados = re.sub(r"\D", "", capturado)
        possui_digitos = len(digitos_referencia) >= 4 and len(digitos_capturados) >= 4
        digit_ratio = (
            float(difflib.SequenceMatcher(None, digitos_referencia, digitos_capturados).ratio())
            if possui_digitos
            else 0.0
        )
        return {
            "ratio": ratio,
            "digit_ratio": digit_ratio,
            "reference": referencia,
            "captured": capturado,
            "has_digits": possui_digitos,
        }

    def obter_dimensoes_imagem(self, caminho_imagem):
        if not PIL_DISPONIVEL:
            return None
        try:
            with Image.open(caminho_imagem) as imagem:
                return int(imagem.width), int(imagem.height)
        except Exception:
            return None

    def carregar_template_imagem(self, caminho_imagem, incluir_ocr=False):
        if not OPENCV_DISPONIVEL or np is None:
            return None

        caminho = str(caminho_imagem)
        try:
            mtime = Path(caminho_imagem).stat().st_mtime
        except Exception:
            mtime = None

        cache = self._cache_template_imagem.get(caminho)
        if cache and cache.get("mtime") == mtime:
            if incluir_ocr and not cache.get("ocr_carregado") and OCR_DISPONIVEL and PIL_DISPONIVEL:
                try:
                    imagem_rgb = Image.fromarray(cv2.cvtColor(cache["bgr"], cv2.COLOR_BGR2RGB))
                    cache["ocr_variants"] = self.extrair_variantes_texto_ocr_imagem(imagem_rgb)
                    cache["ocr_text"] = (cache.get("ocr_variants") or [""])[0]
                except Exception:
                    cache["ocr_variants"] = []
                    cache["ocr_text"] = ""
                cache["ocr_carregado"] = True
            return cache

        try:
            referencia_bgr = cv2.imread(caminho, cv2.IMREAD_COLOR)
            if referencia_bgr is None:
                return None
            referencia_gray = cv2.cvtColor(referencia_bgr, cv2.COLOR_BGR2GRAY)
            referencia_edges = cv2.Canny(referencia_gray, 60, 160)
            referencia_binary = cv2.threshold(
                referencia_gray,
                0,
                255,
                cv2.THRESH_BINARY + cv2.THRESH_OTSU,
            )[1]
            ocr_variants = []
            if incluir_ocr and OCR_DISPONIVEL and PIL_DISPONIVEL:
                try:
                    imagem_rgb = Image.fromarray(cv2.cvtColor(referencia_bgr, cv2.COLOR_BGR2RGB))
                    ocr_variants = self.extrair_variantes_texto_ocr_imagem(imagem_rgb)
                except Exception:
                    ocr_variants = []
            cache = {
                "mtime": mtime,
                "bgr": referencia_bgr,
                "gray": referencia_gray,
                "edges": referencia_edges,
                "binary": referencia_binary,
                "ocr_text": ocr_variants[0] if ocr_variants else "",
                "ocr_variants": ocr_variants,
                "ocr_carregado": bool(incluir_ocr),
                "width": int(referencia_gray.shape[1]),
                "height": int(referencia_gray.shape[0]),
            }
            self._cache_template_imagem[caminho] = cache
            return cache
        except Exception:
            return None

    def avaliar_patch_template(self, patch_bgr, patch_gray, template):
        if not OPENCV_DISPONIVEL or np is None or not template:
            return None

        referencia_bgr = template["bgr"]
        referencia_gray = template["gray"]
        referencia_edges = template["edges"]
        referencia_binary = template.get("binary")

        if patch_bgr is None or patch_gray is None:
            return None

        try:
            if patch_bgr.shape[:2] != referencia_bgr.shape[:2]:
                patch_bgr = cv2.resize(
                    patch_bgr,
                    (referencia_bgr.shape[1], referencia_bgr.shape[0]),
                    interpolation=cv2.INTER_LINEAR,
                )
                patch_gray = cv2.resize(
                    patch_gray,
                    (referencia_gray.shape[1], referencia_gray.shape[0]),
                    interpolation=cv2.INTER_LINEAR,
                )

            diferenca = cv2.absdiff(patch_bgr, referencia_bgr)
            similaridade_cor = 1.0 - (float(diferenca.mean()) / 255.0)

            diferenca_gray = cv2.absdiff(patch_gray, referencia_gray)
            similaridade_tons = 1.0 - (float(diferenca_gray.mean()) / 255.0)

            edges_patch = cv2.Canny(patch_gray, 60, 160)
            intersecao_edges = cv2.bitwise_and(edges_patch, referencia_edges)
            uniao_edges = cv2.bitwise_or(edges_patch, referencia_edges)
            total_uniao = max(1, int(np.count_nonzero(uniao_edges)))
            similaridade_bordas = float(np.count_nonzero(intersecao_edges)) / float(total_uniao)

            patch_binary = cv2.threshold(
                patch_gray,
                0,
                255,
                cv2.THRESH_BINARY + cv2.THRESH_OTSU,
            )[1]
            if referencia_binary is None or patch_binary.shape[:2] != referencia_gray.shape[:2]:
                referencia_binary = cv2.threshold(
                    referencia_gray,
                    0,
                    255,
                    cv2.THRESH_BINARY + cv2.THRESH_OTSU,
                )[1]

            diferenca_binary = cv2.absdiff(patch_binary, referencia_binary)
            similaridade_binaria = 1.0 - (float(diferenca_binary.mean()) / 255.0)

            tinta_patch = patch_binary < 128
            tinta_ref = referencia_binary < 128
            uniao_tinta = np.logical_or(tinta_patch, tinta_ref)
            total_uniao_tinta = max(1, int(np.count_nonzero(uniao_tinta)))
            intersecao_tinta = np.logical_and(tinta_patch, tinta_ref)
            similaridade_tinta = float(np.count_nonzero(intersecao_tinta)) / float(total_uniao_tinta)

            correlacao_exata = float(
                cv2.matchTemplate(patch_gray, referencia_gray, cv2.TM_CCOEFF_NORMED)[0][0]
            )

            hist_patch = cv2.calcHist([patch_gray], [0], None, [32], [0, 256])
            hist_ref = cv2.calcHist([referencia_gray], [0], None, [32], [0, 256])
            cv2.normalize(hist_patch, hist_patch)
            cv2.normalize(hist_ref, hist_ref)
            similaridade_hist = float(cv2.compareHist(hist_patch, hist_ref, cv2.HISTCMP_CORREL))
        except Exception:
            return None

        score_composto = (
            (similaridade_cor * 0.20)
            + (similaridade_tons * 0.20)
            + (similaridade_bordas * 0.18)
            + (correlacao_exata * 0.24)
            + (similaridade_hist * 0.08)
            + (similaridade_binaria * 0.06)
            + (similaridade_tinta * 0.04)
        )

        return {
            "color": float(similaridade_cor),
            "gray": float(similaridade_tons),
            "edges": float(similaridade_bordas),
            "binary": float(similaridade_binaria),
            "ink_iou": float(similaridade_tinta),
            "corr": float(correlacao_exata),
            "hist": float(similaridade_hist),
            "composite": float(score_composto),
        }

    def obter_metricas_tela_virtual(self):
        if not hasattr(ctypes, "windll"):
            return 0, 0, 0, 0
        user32 = ctypes.windll.user32
        SM_XVIRTUALSCREEN = 76
        SM_YVIRTUALSCREEN = 77
        SM_CXVIRTUALSCREEN = 78
        SM_CYVIRTUALSCREEN = 79
        esquerda = int(user32.GetSystemMetrics(SM_XVIRTUALSCREEN))
        topo = int(user32.GetSystemMetrics(SM_YVIRTUALSCREEN))
        largura = int(user32.GetSystemMetrics(SM_CXVIRTUALSCREEN))
        altura = int(user32.GetSystemMetrics(SM_CYVIRTUALSCREEN))
        return esquerda, topo, largura, altura

    def obter_metricas_monitor_ponto(self, x, y):
        if not hasattr(ctypes, "windll"):
            return self.obter_metricas_tela_virtual()

        user32 = ctypes.windll.user32
        MONITOR_DEFAULTTONEAREST = 2

        class MONITORINFO(ctypes.Structure):
            _fields_ = [
                ("cbSize", wintypes.DWORD),
                ("rcMonitor", wintypes.RECT),
                ("rcWork", wintypes.RECT),
                ("dwFlags", wintypes.DWORD),
            ]

        try:
            ponto = wintypes.POINT(int(x), int(y))
            monitor = user32.MonitorFromPoint(ponto, MONITOR_DEFAULTTONEAREST)
            if not monitor:
                return self.obter_metricas_tela_virtual()

            info = MONITORINFO()
            info.cbSize = ctypes.sizeof(MONITORINFO)
            if not user32.GetMonitorInfoW(monitor, ctypes.byref(info)):
                return self.obter_metricas_tela_virtual()

            esquerda = int(info.rcMonitor.left)
            topo = int(info.rcMonitor.top)
            largura = int(info.rcMonitor.right - info.rcMonitor.left)
            altura = int(info.rcMonitor.bottom - info.rcMonitor.top)
            return esquerda, topo, largura, altura
        except Exception:
            return self.obter_metricas_tela_virtual()

    def listar_regioes_monitores(self):
        if not hasattr(ctypes, "windll"):
            return []

        monitores = []
        user32 = ctypes.windll.user32

        try:
            callback_type = ctypes.WINFUNCTYPE(
                ctypes.c_int,
                ctypes.c_void_p,
                ctypes.c_void_p,
                ctypes.POINTER(wintypes.RECT),
                ctypes.c_long,
            )

            def callback(_monitor, _hdc, rect, _param):
                try:
                    esquerda = int(rect.contents.left)
                    topo = int(rect.contents.top)
                    direita = int(rect.contents.right)
                    base = int(rect.contents.bottom)
                    largura = max(0, direita - esquerda)
                    altura = max(0, base - topo)
                    if largura > 0 and altura > 0:
                        monitores.append((esquerda, topo, largura, altura))
                except Exception:
                    pass
                return 1

            user32.EnumDisplayMonitors(0, 0, callback_type(callback), 0)
        except Exception:
            monitores = []

        if not monitores:
            esquerda, topo, largura, altura = self.obter_metricas_tela_virtual()
            if largura > 0 and altura > 0:
                monitores.append((esquerda, topo, largura, altura))
        return monitores

    def capturar_tela_pil(self, region=None):
        if not PIL_DISPONIVEL:
            raise RuntimeError("Pillow indisponivel para captura visual.")

        inicio_captura = time.perf_counter()
        origem_captura = "tela_inteira"

        def finalizar_captura(captura, base_left, base_top):
            duracao_ms = (time.perf_counter() - inicio_captura) * 1000.0
            self._ultima_captura_visual_info = {
                "capture_ms": duracao_ms,
                "origin": origem_captura,
                "region": list(region) if region else None,
                "base_left": int(base_left),
                "base_top": int(base_top),
                "width": int(getattr(captura, "width", 0) or 0),
                "height": int(getattr(captura, "height", 0) or 0),
            }
            return captura, base_left, base_top

        if self.alvo_execucao_ativo():
            janela_execucao = self.obter_janela_alvo_execucao_atual(obrigatoria=True)
            if region is None:
                region = (
                    int(janela_execucao.get("left", 0) or 0),
                    int(janela_execucao.get("top", 0) or 0),
                    int(janela_execucao.get("width", 0) or 0),
                    int(janela_execucao.get("height", 0) or 0),
                )
            self.registrar_origem_captura("janela_alvo", janela=janela_execucao)
            origem_captura = "janela_alvo"
        else:
            if region:
                self.registrar_origem_captura("janela_alvo", janela={"title": self._janela_contexto or self.janela_alvo or "regiao definida"})
                origem_captura = "janela_alvo"
            else:
                self.registrar_origem_captura("tela_inteira")

        if ImageGrab is None:
            captura = pyautogui.screenshot(region=region) if region else pyautogui.screenshot()
            base_left = int(region[0]) if region else 0
            base_top = int(region[1]) if region else 0
            return finalizar_captura(captura, base_left, base_top)

        try:
            if region:
                left = int(region[0])
                top = int(region[1])
                right = int(region[0] + region[2])
                bottom = int(region[1] + region[3])
                try:
                    captura = ImageGrab.grab(bbox=(left, top, right, bottom), all_screens=True)
                    if int(getattr(captura, "width", 0) or 0) > 0 and int(getattr(captura, "height", 0) or 0) > 0:
                        return finalizar_captura(captura, left, top)
                except Exception:
                    pass
                esquerda_virtual, topo_virtual, _largura_virtual, _altura_virtual = self.obter_metricas_tela_virtual()
                captura_total = ImageGrab.grab(all_screens=True)
                largura_virtual = max(1, int(_largura_virtual or getattr(captura_total, "width", 1) or 1))
                altura_virtual = max(1, int(_altura_virtual or getattr(captura_total, "height", 1) or 1))
                escala_x = float(getattr(captura_total, "width", 1) or 1) / float(largura_virtual)
                escala_y = float(getattr(captura_total, "height", 1) or 1) / float(altura_virtual)
                rel_left = int(round((left - esquerda_virtual) * escala_x))
                rel_top = int(round((top - topo_virtual) * escala_y))
                rel_right = int(round((right - esquerda_virtual) * escala_x))
                rel_bottom = int(round((bottom - topo_virtual) * escala_y))
                rel_left = max(0, min(captura_total.width, rel_left))
                rel_top = max(0, min(captura_total.height, rel_top))
                rel_right = max(rel_left, min(captura_total.width, rel_right))
                rel_bottom = max(rel_top, min(captura_total.height, rel_bottom))
                captura = captura_total.crop((rel_left, rel_top, rel_right, rel_bottom))
                base_left = int(round(esquerda_virtual + (rel_left / escala_x)))
                base_top = int(round(topo_virtual + (rel_top / escala_y)))
                return finalizar_captura(captura, base_left, base_top)

            esquerda_virtual, topo_virtual, largura_virtual, altura_virtual = self.obter_metricas_tela_virtual()
            captura = ImageGrab.grab(all_screens=True)
        except Exception:
            captura = pyautogui.screenshot(region=region) if region else pyautogui.screenshot()
            base_left = int(region[0]) if region else 0
            base_top = int(region[1]) if region else 0
            return finalizar_captura(captura, base_left, base_top)

        if not region:
            return finalizar_captura(captura, esquerda_virtual, topo_virtual)

        rel_left = int(region[0] - esquerda_virtual)
        rel_top = int(region[1] - topo_virtual)
        rel_right = int(rel_left + region[2])
        rel_bottom = int(rel_top + region[3])

        rel_left_clamped = max(0, min(captura.width, rel_left))
        rel_top_clamped = max(0, min(captura.height, rel_top))
        rel_right_clamped = max(rel_left_clamped, min(captura.width, rel_right))
        rel_bottom_clamped = max(rel_top_clamped, min(captura.height, rel_bottom))

        cortada = captura.crop((rel_left_clamped, rel_top_clamped, rel_right_clamped, rel_bottom_clamped))
        base_left = int(esquerda_virtual + rel_left_clamped)
        base_top = int(topo_virtual + rel_top_clamped)
        return finalizar_captura(cortada, base_left, base_top)

    def validar_correspondencia_visual(self, caminho_imagem, local, confianca=0.8, perfil_busca="balanced"):
        if not OPENCV_DISPONIVEL or not PIL_DISPONIVEL or np is None:
            return True

        perfil_busca = str(perfil_busca or "balanced").strip().lower()
        if perfil_busca not in {"fast", "balanced", "precise"}:
            perfil_busca = "balanced"
        # A leitura OCR da propria referencia e cara. Ela so reforca a
        # validacao precisa; no fluxo rapido/hibrido os OCRs da tela ja
        # validam o contexto antes desta etapa visual.
        template = self.carregar_template_imagem(
            caminho_imagem,
            incluir_ocr=(perfil_busca == "precise"),
        )
        if not template:
            return True

        largura_template = int(template.get("width", 0) or 0)
        altura_template = int(template.get("height", 0) or 0)
        area_template = max(1, largura_template * altura_template)
        template_pequeno = area_template <= 3600 or min(largura_template, altura_template) <= 28
        template_muito_pequeno = area_template <= 1600 or min(largura_template, altura_template) <= 20
        template_texto_critico = any(
            self.texto_referencia_critico(texto)
            for texto in (template.get("ocr_variants") or [template.get("ocr_text") or ""])
            if str(texto or "").strip()
        )

        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        try:
            captura, _base_left, _base_top = self.capturar_tela_pil(region=(esquerda, topo, largura, altura))
            captura_bgr = cv2.cvtColor(np.array(captura), cv2.COLOR_RGB2BGR)
            captura_gray = cv2.cvtColor(captura_bgr, cv2.COLOR_BGR2GRAY)
        except Exception:
            return True

        metricas = self.avaliar_patch_template(captura_bgr, captura_gray, template)
        if not metricas:
            return True

        similaridade_cor = metricas["color"]
        similaridade_tons = metricas["gray"]
        similaridade_bordas = metricas["edges"]
        similaridade_binaria = float(metricas.get("binary", 0.0) or 0.0)
        similaridade_tinta = float(metricas.get("ink_iou", 0.0) or 0.0)
        correlacao_exata = metricas["corr"]
        similaridade_hist = metricas["hist"]
        score_composto = metricas["composite"]

        confianca = float(confianca or 0.8)
        if perfil_busca == "precise":
            if confianca >= 0.99:
                limite_cor = 0.96
                limite_tons = 0.97
                limite_bordas = 0.54
                limite_correlacao = 0.90
                limite_composto = 0.88
                limite_binario = 0.94
                limite_tinta = 0.46
            elif confianca >= 0.9:
                limite_cor = 0.92
                limite_tons = 0.94
                limite_bordas = 0.38
                limite_correlacao = 0.80
                limite_composto = 0.83
                limite_binario = 0.90
                limite_tinta = 0.36
            else:
                limite_cor = 0.89
                limite_tons = 0.91
                limite_bordas = 0.28
                limite_correlacao = 0.72
                limite_composto = 0.79
                limite_binario = 0.86
                limite_tinta = 0.26
        elif confianca >= 0.99:
            limite_cor = 0.96
            limite_tons = 0.96
            limite_bordas = 0.58
            limite_correlacao = 0.92
            limite_composto = 0.89
            limite_binario = 0.0
            limite_tinta = 0.0
        elif confianca >= 0.9:
            limite_cor = 0.90
            limite_tons = 0.92
            limite_bordas = 0.46
            limite_correlacao = 0.85
            limite_composto = 0.81
            limite_binario = 0.0
            limite_tinta = 0.0
        else:
            limite_cor = 0.84
            limite_tons = 0.86
            limite_bordas = 0.34
            limite_correlacao = 0.78
            limite_composto = 0.74
            limite_binario = 0.0
            limite_tinta = 0.0

        if perfil_busca == "precise":
            if template_pequeno:
                limite_bordas = max(0.20, limite_bordas - 0.06)
                limite_correlacao = max(0.68, limite_correlacao - 0.04)
                limite_composto = max(0.76, limite_composto - 0.02)
                limite_binario = max(0.82, limite_binario - 0.03)
                limite_tinta = max(0.20, limite_tinta - 0.04)
            if template_muito_pequeno:
                limite_bordas = max(0.14, limite_bordas - 0.04)
                limite_correlacao = max(0.62, limite_correlacao - 0.04)
                limite_composto = max(0.74, limite_composto - 0.02)
                limite_binario = max(0.78, limite_binario - 0.04)
                limite_tinta = max(0.16, limite_tinta - 0.04)
        else:
            if template_pequeno:
                limite_bordas = max(0.12, limite_bordas - 0.16)
                limite_correlacao = max(0.64, limite_correlacao - 0.08)
                limite_composto = max(0.68, limite_composto - 0.05)
            if template_muito_pequeno:
                limite_bordas = max(0.06, limite_bordas - 0.10)
                limite_correlacao = max(0.58, limite_correlacao - 0.06)
                limite_composto = max(0.64, limite_composto - 0.04)

        passou_regra_principal = (
            similaridade_cor >= limite_cor
            and similaridade_tons >= limite_tons
            and similaridade_bordas >= limite_bordas
            and correlacao_exata >= limite_correlacao
        )

        passou_regra_precisa = (
            perfil_busca == "precise"
            and similaridade_cor >= limite_cor
            and similaridade_tons >= limite_tons
            and similaridade_bordas >= limite_bordas
            and correlacao_exata >= limite_correlacao
            and score_composto >= limite_composto
            and similaridade_hist >= 0.88
            and similaridade_binaria >= limite_binario
            and similaridade_tinta >= limite_tinta
        )

        passou_regra_precisa_visual = (
            perfil_busca == "precise"
            and similaridade_cor >= max(limite_cor, 0.92 if not template_pequeno else 0.90)
            and similaridade_tons >= max(limite_tons, 0.93 if not template_pequeno else 0.91)
            and similaridade_hist >= 0.90
            and similaridade_binaria >= max(limite_binario, 0.84 if template_muito_pequeno else 0.87)
            and (
                correlacao_exata >= max(limite_correlacao, 0.78 if not template_muito_pequeno else 0.70)
                or similaridade_bordas >= max(limite_bordas, 0.34 if not template_muito_pequeno else 0.24)
            )
        )

        passou_regra_texto_compacto = (
            similaridade_cor >= max(limite_cor, 0.94)
            and similaridade_tons >= max(limite_tons, 0.94)
            and similaridade_hist >= 0.92
            and correlacao_exata >= 0.62
        )

        passou_regra_visual_forte = (
            similaridade_cor >= 0.97
            and similaridade_tons >= 0.97
            and similaridade_hist >= 0.94
            and (similaridade_bordas >= 0.22 or correlacao_exata >= 0.60)
        )

        passou_regra_composta = (
            score_composto >= limite_composto
            and similaridade_hist >= 0.88
            and similaridade_cor >= max(0.90, limite_cor)
            and similaridade_tons >= max(0.90, limite_tons)
            and correlacao_exata >= max(0.58, limite_correlacao - 0.12)
        )

        passou_regra_template_pequeno = (
            template_pequeno
            and similaridade_cor >= max(0.88, limite_cor - 0.02)
            and similaridade_tons >= max(0.89, limite_tons - 0.01)
            and similaridade_hist >= 0.86
            and correlacao_exata >= max(0.60, limite_correlacao)
        )

        passou_regra_micro = (
            template_muito_pequeno
            and similaridade_cor >= 0.90
            and similaridade_tons >= 0.91
            and similaridade_hist >= 0.88
            and correlacao_exata >= 0.66
        )

        regras_aceitas = (
            (passou_regra_precisa or passou_regra_precisa_visual)
            if perfil_busca == "precise"
            else (
                passou_regra_principal
                or passou_regra_texto_compacto
                or passou_regra_visual_forte
                or passou_regra_composta
                or passou_regra_template_pequeno
                or passou_regra_micro
            )
        )

        if perfil_busca == "precise":
            textos_referencia = [
                texto
                for texto in (template.get("ocr_variants") or [template.get("ocr_text") or ""])
                if str(texto or "").strip()
            ]
            texto_referencia_critico = next(
                (texto for texto in textos_referencia if self.texto_referencia_critico(texto)),
                "",
            )
            texto_referencia_padrao = texto_referencia_critico or next(iter(textos_referencia), "")
            if texto_referencia_padrao and OCR_DISPONIVEL and Image is not None:
                try:
                    captura_rgb = cv2.cvtColor(captura_bgr, cv2.COLOR_BGR2RGB)
                    textos_patch = self.extrair_variantes_texto_ocr_imagem(Image.fromarray(captura_rgb))
                except Exception:
                    textos_patch = []

                comparacoes_texto = [
                    self.comparar_texto_reconhecimento_preciso(texto_referencia_padrao, texto_patch)
                    for texto_patch in textos_patch
                ]
                melhor_texto = max(
                    comparacoes_texto,
                    key=lambda item: (
                        float(item.get("digit_ratio", 0.0) or 0.0),
                        float(item.get("ratio", 0.0) or 0.0),
                    ),
                    default={
                        "ratio": 0.0,
                        "digit_ratio": 0.0,
                        "reference": self.normalizar_texto_ocr(texto_referencia_padrao),
                        "captured": "",
                        "has_digits": self.texto_referencia_critico(texto_referencia_padrao),
                    },
                )
                ratio_texto = float(melhor_texto.get("ratio", 0.0) or 0.0)
                ratio_digitos = float(melhor_texto.get("digit_ratio", 0.0) or 0.0)
                possui_digitos = bool(melhor_texto.get("has_digits"))
                texto_igual = (
                    str(melhor_texto.get("reference", "")).strip()
                    and str(melhor_texto.get("reference", "")).strip()
                    == str(melhor_texto.get("captured", "")).strip()
                )
                digitos_exatos = possui_digitos and ratio_digitos >= 0.999

                if texto_referencia_critico:
                    if not textos_patch and confianca >= 0.99:
                        safe_print(
                            "Correspondencia visual rejeitada no modo preciso: OCR nao confirmou a data/codigo da referencia.",
                            flush=True,
                        )
                        return False
                    if textos_patch and not (texto_igual or digitos_exatos):
                        safe_print(
                            "Correspondencia visual rejeitada pelo OCR do modo preciso "
                            f"(ref='{melhor_texto.get('reference', '')}' | capturado='{melhor_texto.get('captured', '')}' | "
                            f"texto={ratio_texto:.2f} | digitos={ratio_digitos:.2f}).",
                            flush=True,
                        )
                        return False
                elif textos_patch:
                    rejeitar_por_ocr = ratio_texto < 0.90
                    if rejeitar_por_ocr:
                        safe_print(
                            "Correspondencia visual rejeitada pelo OCR do modo preciso "
                            f"(ref='{melhor_texto.get('reference', '')}' | capturado='{melhor_texto.get('captured', '')}' | "
                            f"texto={ratio_texto:.2f}).",
                            flush=True,
                        )
                        return False

        if not regras_aceitas:
            safe_print(
                "Correspondencia visual rejeitada "
                f"(cor={similaridade_cor:.2f}, tons={similaridade_tons:.2f}, "
                f"bordas={similaridade_bordas:.2f}, bin={similaridade_binaria:.2f}, "
                f"tinta={similaridade_tinta:.2f}, corr={correlacao_exata:.2f}, "
                f"hist={similaridade_hist:.2f}, comp={score_composto:.2f}).",
                flush=True,
            )
            return False
        return True

    def procurar_texto_por_ocr(self, texto_referencia, region=None):
        if not OCR_DISPONIVEL:
            if not self._ocr_indisponivel_logado:
                safe_print("OCR indisponivel para reforco de reconhecimento.", flush=True)
                self._ocr_indisponivel_logado = True
            return None

        texto_alvo = self.normalizar_texto_ocr(texto_referencia)
        if len(texto_alvo.replace(" ", "")) < 3:
            return None
        if region is None:
            safe_print("OCR ignorado por seguranca: nenhuma area pequena foi definida para a leitura.", flush=True)
            return None

        inicio_ocr = time.perf_counter()
        regiao_chave = tuple(int(valor) for valor in region) if region else None
        chave_cache = (texto_alvo, regiao_chave)
        cache = self._cache_ocr_curto.get(chave_cache) if isinstance(self._cache_ocr_curto, dict) else None
        if cache and time.perf_counter() - float(cache.get("em", 0.0) or 0.0) <= 0.65:
            self._ultima_busca_ocr_info = dict(cache.get("info") or {})
            resultado_cache = cache.get("resultado")
            return dict(resultado_cache) if isinstance(resultado_cache, dict) else None

        try:
            captura, base_left, base_top = self.capturar_tela_pil(region=region)
        except OSError as erro:
            raise RuntimeError(f"PyAutoGUI nao conseguiu capturar a tela para OCR: {erro}") from erro

        try:
            area_ocr = max(1, int(captura.width or 0) * int(captura.height or 0))
            whitelist = "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ "
            config_ocr = (
                f"--psm {'11' if area_ocr > 420000 else '6'} "
                f"-c tessedit_char_whitelist={whitelist}"
            )
            dados = pytesseract.image_to_data(
                captura,
                output_type=pytesseract.Output.DICT,
                config=config_ocr,
            )
        except Exception:
            return None

        palavras = []
        total_itens = len(dados.get("text", []))
        for indice in range(total_itens):
            texto = self.normalizar_texto_ocr(dados["text"][indice])
            if not texto:
                continue
            palavras.append(
                {
                    "texto": texto,
                    "left": int(dados["left"][indice]),
                    "top": int(dados["top"][indice]),
                    "width": int(dados["width"][indice]),
                    "height": int(dados["height"][indice]),
                }
            )

        if not palavras:
            self._ultima_busca_ocr_info = {
                "capture_ms": float((self._ultima_captura_visual_info or {}).get("capture_ms", 0.0) or 0.0),
                "ocr_ms": (time.perf_counter() - inicio_ocr) * 1000.0,
                "region": list(region) if region else None,
                "score": 0.0,
                "used_fallback": False,
            }
            self._cache_ocr_curto[chave_cache] = {
                "em": time.perf_counter(),
                "resultado": None,
                "info": dict(self._ultima_busca_ocr_info),
            }
            return None

        alvo_palavras = texto_alvo.split()
        minimo_palavras = max(1, len(alvo_palavras) - 2)
        maximo_palavras = min(len(palavras), len(alvo_palavras) + 2)
        melhor = None
        melhor_pontuacao = 0.0

        for tamanho in range(minimo_palavras, maximo_palavras + 1):
            for inicio in range(0, len(palavras) - tamanho + 1):
                trecho = palavras[inicio : inicio + tamanho]
                candidato = " ".join(item["texto"] for item in trecho)
                pontuacao = difflib.SequenceMatcher(None, texto_alvo, candidato).ratio()
                if pontuacao <= melhor_pontuacao:
                    continue

                esquerda = min(item["left"] for item in trecho)
                topo = min(item["top"] for item in trecho)
                direita = max(item["left"] + item["width"] for item in trecho)
                base = max(item["top"] + item["height"] for item in trecho)
                centro_x = esquerda + ((direita - esquerda) / 2)
                centro_y = topo + ((base - topo) / 2)

                centro_x += base_left
                centro_y += base_top

                melhor = {
                    "x": int(centro_x),
                    "y": int(centro_y),
                    "left": int(esquerda + base_left),
                    "top": int(topo + base_top),
                    "width": int(direita - esquerda),
                    "height": int(base - topo),
                    "score": pontuacao,
                    "texto": candidato,
                }
                melhor_pontuacao = pontuacao

        limite = 0.68 if len(alvo_palavras) > 1 else 0.8
        self._ultima_busca_ocr_info = {
            "capture_ms": float((self._ultima_captura_visual_info or {}).get("capture_ms", 0.0) or 0.0),
            "ocr_ms": (time.perf_counter() - inicio_ocr) * 1000.0,
            "region": list(region) if region else None,
            "score": float(melhor_pontuacao or 0.0),
            "used_fallback": False,
        }
        if melhor and melhor_pontuacao >= limite:
            safe_print(f"Texto da referencia encontrado por OCR (score={melhor_pontuacao:.2f})", flush=True)
            self._cache_ocr_curto[chave_cache] = {
                "em": time.perf_counter(),
                "resultado": dict(melhor),
                "info": dict(self._ultima_busca_ocr_info),
            }
            return melhor
        self._cache_ocr_curto[chave_cache] = {
            "em": time.perf_counter(),
            "resultado": None,
            "info": dict(self._ultima_busca_ocr_info),
        }
        return None

    def procurar_imagem_por_template(
        self,
        caminho_imagem,
        confianca=0.8,
        region=None,
        perfil_busca="balanced",
        escala_referencia=None,
        incluir_ocr_referencia=None,
    ):
        if not OPENCV_DISPONIVEL or not PIL_DISPONIVEL or np is None:
            return None

        perfil_busca = str(perfil_busca or "balanced").strip().lower()
        if perfil_busca not in {"fast", "balanced", "precise"}:
            perfil_busca = "balanced"
        modo_estrito = float(confianca or 0.8) >= 0.99
        if incluir_ocr_referencia is None:
            incluir_ocr_referencia = perfil_busca == "precise" or modo_estrito
        else:
            incluir_ocr_referencia = bool(incluir_ocr_referencia)
        self._ultima_busca_template_info = None

        try:
            captura, base_left, base_top = self.capturar_tela_pil(region=region)
        except OSError as erro:
            raise RuntimeError(f"PyAutoGUI nao conseguiu capturar a tela para template matching: {erro}") from erro

        info_captura = self._ultima_captura_visual_info or {}
        inicio_analise = time.perf_counter()
        try:
            tela_bgr = cv2.cvtColor(np.array(captura), cv2.COLOR_RGB2BGR)
            tela_gray = cv2.cvtColor(tela_bgr, cv2.COLOR_BGR2GRAY)
            template = self.carregar_template_imagem(
                caminho_imagem,
                incluir_ocr=incluir_ocr_referencia,
            )
            if not template:
                return None
            referencia_gray = template["gray"]
        except Exception:
            return None

        largura_template = int(template.get("width", 0) or 0)
        altura_template = int(template.get("height", 0) or 0)
        area_template = max(1, largura_template * altura_template)
        template_pequeno = area_template <= 3600 or min(largura_template, altura_template) <= 28
        template_muito_pequeno = area_template <= 1600 or min(largura_template, altura_template) <= 20
        template_texto_critico = any(
            self.texto_referencia_critico(texto)
            for texto in (template.get("ocr_variants") or [template.get("ocr_text") or ""])
            if str(texto or "").strip()
        )

        altura_tela_original, largura_tela_original = tela_gray.shape[:2]
        candidatos = []
        melhor = None
        melhor_score = -1.0
        segundo_melhor_score = -1.0
        if modo_estrito:
            metodos = (
                ("sqdiff", cv2.TM_SQDIFF_NORMED),
                ("ccoeff", cv2.TM_CCOEFF_NORMED),
            )
        elif perfil_busca == "fast":
            metodos = (
                ("ccoeff", cv2.TM_CCOEFF_NORMED),
                ("sqdiff", cv2.TM_SQDIFF_NORMED),
            )
        elif perfil_busca == "precise":
            metodos = (
                ("ccoeff", cv2.TM_CCOEFF_NORMED),
                ("ccorr", cv2.TM_CCORR_NORMED),
                ("sqdiff", cv2.TM_SQDIFF_NORMED),
            )
        else:
            metodos = (
                ("ccorr", cv2.TM_CCORR_NORMED),
                ("sqdiff", cv2.TM_SQDIFF_NORMED),
                ("ccoeff", cv2.TM_CCOEFF_NORMED),
            )

        escalas_fallback = ()
        if template_muito_pequeno:
            escalas_fallback = (0.98, 1.02, 0.95, 1.05, 0.90, 1.10, 0.85, 1.15, 0.80, 1.25)
        elif template_pequeno:
            escalas_fallback = (0.98, 1.02, 0.95, 1.05, 0.90, 1.10, 0.85, 1.15, 0.80, 1.25)
        elif perfil_busca == "precise":
            escalas_fallback = (0.97, 1.03, 0.94, 1.06, 0.90, 1.10, 0.85, 1.15, 0.80, 1.25, 0.75, 1.35)
        elif perfil_busca == "balanced":
            escalas_fallback = (0.97, 1.03, 0.94, 1.06, 0.90, 1.10, 0.85, 1.15, 0.80, 1.25)
        if modo_estrito or perfil_busca == "fast":
            escalas_fallback = ()
        try:
            escala_referencia = float(escala_referencia) if escala_referencia is not None else None
        except (TypeError, ValueError):
            escala_referencia = None
        if escala_referencia is not None and not 0.50 <= escala_referencia <= 2.00:
            escala_referencia = None

        escalas_normais = (1.0,)
        if escala_referencia is not None:
            # Primeiro verifica a escala calculada pelos alvos OCR e, em seguida,
            # o tamanho original. Normalmente uma dessas duas confirma a imagem
            # sem precisar abrir a varredura multiescala.
            escalas_normais = tuple(
                dict.fromkeys(
                    round(valor, 4)
                    for valor in (escala_referencia, 1.0)
                    if 0.50 <= float(valor) <= 2.00
                )
            ) or (1.0,)
            ajustes_guiados = (
                (0.90, 0.96, 1.04, 1.10)
                if perfil_busca == "fast"
                else (1.0, 0.96, 1.04, 0.91, 1.09)
            )
            escalas_guiadas_por_ocr = tuple(
                round(escala_referencia * ajuste, 4)
                for ajuste in ajustes_guiados
            )
            # A escala vem da geometria dos OCRs confirmados. Mesmo no perfil
            # rapido testamos quatro ajustes proximos: isso cobre DPI/zoom de
            # outra maquina sem esperar o timeout mudar para "balanced".
            escalas_fallback = escalas_guiadas_por_ocr

        limite_escala_minima = 0.50 if escala_referencia is not None else 0.75
        limite_escala_maxima = 2.00 if escala_referencia is not None else 1.35
        escalas_fallback = tuple(
            dict.fromkeys(
                escala
                for escala in escalas_fallback
                if limite_escala_minima <= float(escala) <= limite_escala_maxima
                and abs(float(escala) - 1.0) > 0.0001
            )
        )

        planos_busca = []
        permite_busca_reduzida = (
            not modo_estrito
            and perfil_busca in {"fast", "balanced"}
            and min(largura_tela_original, altura_tela_original) >= 480
            and min(largura_template, altura_template) >= 14
            and (largura_tela_original * altura_tela_original) >= 240000
        )
        if permite_busca_reduzida:
            proporcao_reduzida = 0.62 if perfil_busca == "fast" else 0.82
            try:
                largura_reduzida = max(80, int(round(largura_tela_original * proporcao_reduzida)))
                altura_reduzida = max(80, int(round(altura_tela_original * proporcao_reduzida)))
                tela_gray_reduzida = cv2.resize(tela_gray, (largura_reduzida, altura_reduzida), interpolation=cv2.INTER_AREA)
                planos_busca.append(
                    {
                        "name": "reduced",
                        "ratio": proporcao_reduzida,
                        "gray": tela_gray_reduzida,
                        "scales": escalas_normais,
                        "methods": metodos[:1] if perfil_busca == "fast" else metodos[:2],
                    }
                )
            except Exception:
                pass
        planos_busca.append(
            {
                "name": "normal",
                "ratio": 1.0,
                "gray": tela_gray,
                "scales": escalas_normais,
                "methods": metodos[:1] if perfil_busca == "fast" else metodos,
            }
        )
        if escalas_fallback:
            planos_busca.append(
                {
                    "name": "multiscale",
                    "ratio": 1.0,
                    "gray": tela_gray,
                    "scales": escalas_fallback,
                    "methods": metodos,
                }
            )

        executou_multiscale = False
        executou_busca_reduzida = False

        for plano in planos_busca:
            nome_plano = str(plano.get("name") or "normal")
            proporcao_plano = float(plano.get("ratio", 1.0) or 1.0)
            tela_gray_plano = plano.get("gray")
            if tela_gray_plano is None:
                continue
            altura_tela, largura_tela = tela_gray_plano.shape[:2]
            if nome_plano == "multiscale":
                executou_multiscale = True
            if nome_plano == "reduced":
                executou_busca_reduzida = True

            for escala in tuple(plano.get("scales") or (1.0,)):
                try:
                    if escala == 1.0:
                        ref_gray_original = referencia_gray
                    else:
                        largura_ref_original = max(8, int(round(referencia_gray.shape[1] * escala)))
                        altura_ref_original = max(8, int(round(referencia_gray.shape[0] * escala)))
                        ref_gray_original = cv2.resize(
                            referencia_gray,
                            (largura_ref_original, altura_ref_original),
                            interpolation=cv2.INTER_AREA if escala < 1 else cv2.INTER_LINEAR,
                        )
                except Exception:
                    continue

                altura_ref_original, largura_ref_original = ref_gray_original.shape[:2]
                if proporcao_plano != 1.0:
                    try:
                        largura_ref = max(8, int(round(largura_ref_original * proporcao_plano)))
                        altura_ref = max(8, int(round(altura_ref_original * proporcao_plano)))
                        ref_gray = cv2.resize(
                            ref_gray_original,
                            (largura_ref, altura_ref),
                            interpolation=cv2.INTER_AREA,
                        )
                    except Exception:
                        continue
                else:
                    ref_gray = ref_gray_original
                    largura_ref = largura_ref_original
                    altura_ref = altura_ref_original

                if largura_ref >= largura_tela or altura_ref >= altura_tela:
                    continue

                template_redimensionado = {
                    **template,
                    "bgr": cv2.resize(
                        template["bgr"],
                        (largura_ref_original, altura_ref_original),
                        interpolation=cv2.INTER_LINEAR,
                    )
                    if (template["width"] != largura_ref_original or template["height"] != altura_ref_original)
                    else template["bgr"],
                    "gray": ref_gray_original,
                    "edges": cv2.Canny(ref_gray_original, 60, 160),
                }

                for nome_metodo, metodo_cv in tuple(plano.get("methods") or metodos):
                    try:
                        resultado = cv2.matchTemplate(tela_gray_plano, ref_gray, metodo_cv)
                    except Exception:
                        continue

                    try:
                        resultado_norm = (1.0 - resultado) if metodo_cv == cv2.TM_SQDIFF_NORMED else resultado
                        total_valores = int(resultado_norm.size or 0)
                        if total_valores <= 0:
                            continue
                        if perfil_busca == "fast":
                            top_k = 3 if template_muito_pequeno else (2 if template_pequeno else 1)
                        else:
                            top_k = 2 if modo_estrito else (6 if perfil_busca == "precise" else 4)
                            if template_muito_pequeno:
                                top_k += 4
                            elif template_pequeno:
                                top_k += 2
                        if nome_plano == "reduced":
                            top_k = min(top_k, 3)
                        top_k = min(top_k, total_valores)
                        flat = resultado_norm.reshape(-1)
                        indices = np.argpartition(flat, -top_k)[-top_k:]
                        indices = indices[np.argsort(flat[indices])[::-1]]
                    except Exception:
                        continue

                    for indice_flat in indices:
                        top_plano, left_plano = np.unravel_index(int(indice_flat), resultado_norm.shape)
                        left_plano = int(left_plano)
                        top_plano = int(top_plano)
                        score_atual = float(resultado_norm[top_plano, left_plano])
                        score_minimo = max(0.62, float(confianca) - 0.18)
                        if template_muito_pequeno:
                            score_minimo = max(0.52, float(confianca) - 0.26)
                        elif template_pequeno:
                            score_minimo = max(0.57, float(confianca) - 0.22)
                        if template_texto_critico and (perfil_busca == "precise" or modo_estrito):
                            if template_muito_pequeno:
                                score_minimo = max(score_minimo, 0.78 if modo_estrito else 0.74)
                            elif template_pequeno:
                                score_minimo = max(score_minimo, 0.82 if modo_estrito else 0.78)
                            else:
                                score_minimo = max(score_minimo, 0.86 if modo_estrito else 0.82)
                        if score_atual < score_minimo:
                            continue

                        if proporcao_plano != 1.0:
                            left = int(round(left_plano / proporcao_plano))
                            top = int(round(top_plano / proporcao_plano))
                        else:
                            left = left_plano
                            top = top_plano

                        left = max(0, left)
                        top = max(0, top)
                        if left + largura_ref_original > largura_tela_original or top + altura_ref_original > altura_tela_original:
                            continue

                        patch_bgr = tela_bgr[top : top + altura_ref_original, left : left + largura_ref_original]
                        patch_gray = tela_gray[top : top + altura_ref_original, left : left + largura_ref_original]
                        if (
                            patch_bgr is None
                            or patch_gray is None
                            or patch_bgr.shape[1] != largura_ref_original
                            or patch_bgr.shape[0] != altura_ref_original
                        ):
                            continue

                        metricas = self.avaliar_patch_template(
                            patch_bgr,
                            patch_gray,
                            template_redimensionado,
                        )
                        score_composto = float(metricas["composite"]) if metricas else float(score_atual)
                        similaridade_binaria = float((metricas or {}).get("binary", score_composto) or score_composto)
                        similaridade_tinta = float((metricas or {}).get("ink_iou", 0.0) or 0.0)
                        if perfil_busca == "precise":
                            score_ranqueado = (
                                (score_atual * 0.36)
                                + (score_composto * 0.34)
                                + (similaridade_binaria * 0.20)
                                + (max(similaridade_tinta, float((metricas or {}).get("edges", 0.0) or 0.0)) * 0.10)
                            )
                        else:
                            score_ranqueado = (score_atual * 0.58) + (score_composto * 0.42)

                        candidato = {
                            "left": int(left + base_left),
                            "top": int(top + base_top),
                            "width": int(largura_ref_original),
                            "height": int(altura_ref_original),
                            "x": int(left + (largura_ref_original / 2) + base_left),
                            "y": int(top + (altura_ref_original / 2) + base_top),
                            "score": float(score_atual),
                            "rank_score": float(score_ranqueado),
                            "quality_score": float(score_composto),
                            "scale": float(escala),
                            "method": nome_metodo,
                            "metrics": metricas or {},
                            "search_stage": nome_plano,
                            "search_ratio": float(proporcao_plano),
                        }
                        # Os quatro OCRs ja confirmaram o contexto e deram a
                        # escala esperada. Uma correspondencia visual muito forte
                        # na passagem normal e conclusiva; continuar testando todas
                        # as escalas apenas acrescentava varios segundos ao clique.
                        if (
                            escala_referencia is not None
                            and nome_plano == "normal"
                            and float(candidato.get("score", 0.0) or 0.0) >= max(0.94, float(confianca) + 0.12)
                            and float(candidato.get("quality_score", 0.0) or 0.0) >= 0.92
                        ):
                            duracao_analise_ms = (time.perf_counter() - inicio_analise) * 1000.0
                            self._ultimo_template_debug = candidato
                            self._ultima_busca_template_info = {
                                "capture_ms": float(info_captura.get("capture_ms", 0.0) or 0.0),
                                "recognition_ms": duracao_analise_ms,
                                "region": list(region) if region else (info_captura.get("region") or None),
                                "score": float(candidato.get("rank_score", 0.0) or 0.0),
                                "raw_score": float(candidato.get("score", 0.0) or 0.0),
                                "quality_score": float(candidato.get("quality_score", 0.0) or 0.0),
                                "search_stage": "ocr_guided",
                                "used_multiscale": False,
                                "used_reduced": bool(executou_busca_reduzida),
                                "used_fallback": False,
                            }
                            safe_print(
                                "Imagem encontrada por template guiado pelos OCRs "
                                f"(score={float(candidato.get('score', 0.0)):.2f}, qualidade={float(candidato.get('quality_score', 0.0)):.2f}, "
                                f"escala={candidato['scale']:.2f}, captura={float(info_captura.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                                f"analise={duracao_analise_ms:.0f}ms).",
                                flush=True,
                            )
                            return candidato
                        if perfil_busca == "fast" and nome_plano == "reduced":
                            if (
                                float(candidato.get("rank_score", 0.0)) >= max(0.76, float(confianca) - 0.05)
                                and float(candidato.get("quality_score", 0.0)) >= 0.78
                                and float((metricas or {}).get("corr", 0.0) or 0.0) >= 0.60
                                and float((metricas or {}).get("gray", 0.0) or 0.0) >= 0.88
                            ):
                                duracao_analise_ms = (time.perf_counter() - inicio_analise) * 1000.0
                                self._ultimo_template_debug = candidato
                                self._ultima_busca_template_info = {
                                    "capture_ms": float(info_captura.get("capture_ms", 0.0) or 0.0),
                                    "recognition_ms": duracao_analise_ms,
                                    "region": list(region) if region else (info_captura.get("region") or None),
                                    "score": float(candidato.get("rank_score", 0.0) or 0.0),
                                    "raw_score": float(candidato.get("score", 0.0) or 0.0),
                                    "quality_score": float(candidato.get("quality_score", 0.0) or 0.0),
                                    "search_stage": "reduced",
                                    "used_multiscale": False,
                                    "used_reduced": True,
                                    "used_fallback": False,
                                }
                                safe_print(
                                    "Imagem encontrada por template rapido "
                                    f"(score={float(candidato.get('score', 0.0)):.2f}, rank={float(candidato.get('rank_score', 0.0)):.2f}, "
                                    f"qualidade={float(candidato.get('quality_score', 0.0)):.2f}, "
                                    f"captura={float(info_captura.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                                    f"analise={duracao_analise_ms:.0f}ms, area={self.formatar_regiao_log(region)}, etapa=reduced)",
                                    flush=True,
                                )
                                return candidato
                        candidatos.append(candidato)

        if candidatos:
            candidatos.sort(key=lambda item: float(item.get("rank_score", 0.0)), reverse=True)
            filtrados = []
            for candidato in candidatos:
                if any(
                    abs(int(candidato["left"]) - int(existente["left"])) < 10
                    and abs(int(candidato["top"]) - int(existente["top"])) < 10
                    and abs(int(candidato["width"]) - int(existente["width"])) < 12
                    and abs(int(candidato["height"]) - int(existente["height"])) < 12
                    for existente in filtrados
                ):
                    continue
                filtrados.append(candidato)
                if len(filtrados) >= 12:
                    break
            candidatos = filtrados
            melhor = candidatos[0]
            melhor_score = float(melhor.get("rank_score", 0.0))
            segundo_melhor_score = float(candidatos[1].get("rank_score", 0.0)) if len(candidatos) > 1 else -1.0

        self._ultimo_template_debug = melhor

        if modo_estrito:
            limite = 0.88
        elif perfil_busca == "precise":
            limite = max(0.82, float(confianca) - 0.01)
        elif perfil_busca == "balanced":
            limite = max(0.75, float(confianca) - 0.05)
        else:
            limite = max(0.73, float(confianca) - 0.06)

        if perfil_busca == "precise":
            if template_muito_pequeno:
                limite = max(0.78, limite - 0.03)
            elif template_pequeno:
                limite = max(0.80, limite - 0.02)
            if template_texto_critico:
                limite = max(
                    limite,
                    0.86 if template_muito_pequeno else (0.88 if template_pequeno else 0.90),
                )
        else:
            if template_muito_pequeno:
                limite = max(0.66, limite - 0.08)
            elif template_pequeno:
                limite = max(0.70, limite - 0.05)

        ambiguidade = (
            melhor
            and segundo_melhor_score > 0
            and (
                (modo_estrito and abs(float(melhor.get("rank_score", 0.0)) - float(segundo_melhor_score)) < 0.004)
                or (
                    perfil_busca == "precise"
                    and abs(float(melhor.get("rank_score", 0.0)) - float(segundo_melhor_score))
                    < (0.015 if template_texto_critico else 0.008)
                )
            )
        )

        metricas_melhor = (melhor or {}).get("metrics") or {}
        aceite_por_score_bruto = (
            melhor
            and float(melhor.get("score", 0.0)) >= max(0.84, float(confianca))
            and float(melhor.get("quality_score", 0.0)) >= 0.72
            and float(metricas_melhor.get("corr", 0.0)) >= 0.58
            and float(metricas_melhor.get("gray", 0.0)) >= 0.88
        )
        if melhor and template_pequeno:
            aceite_por_score_bruto = aceite_por_score_bruto or (
                float(melhor.get("score", 0.0)) >= max(0.74, float(confianca) - 0.08)
                and float(melhor.get("quality_score", 0.0)) >= 0.68
                and float(metricas_melhor.get("corr", 0.0)) >= 0.54
                and float(metricas_melhor.get("gray", 0.0)) >= 0.86
            )
        if melhor and template_muito_pequeno:
            aceite_por_score_bruto = aceite_por_score_bruto or (
                float(melhor.get("score", 0.0)) >= max(0.68, float(confianca) - 0.12)
                and float(melhor.get("quality_score", 0.0)) >= 0.64
                and float(metricas_melhor.get("corr", 0.0)) >= 0.50
                and float(metricas_melhor.get("gray", 0.0)) >= 0.84
            )

        if candidatos:
            for candidato in candidatos:
                metricas_candidato = candidato.get("metrics") or {}
                if perfil_busca == "precise":
                    aceite_por_score_bruto_candidato = (
                        float(candidato.get("score", 0.0)) >= max(0.86, float(confianca))
                        and float(candidato.get("quality_score", 0.0)) >= 0.80
                        and float(metricas_candidato.get("corr", 0.0)) >= 0.72
                        and float(metricas_candidato.get("gray", 0.0)) >= 0.90
                        and float(metricas_candidato.get("binary", 0.0)) >= 0.86
                        and float(metricas_candidato.get("ink_iou", 0.0)) >= 0.24
                    )
                    if template_pequeno:
                        aceite_por_score_bruto_candidato = aceite_por_score_bruto_candidato or (
                            float(candidato.get("score", 0.0)) >= max(0.80, float(confianca) - 0.02)
                            and float(candidato.get("quality_score", 0.0)) >= 0.78
                            and float(metricas_candidato.get("corr", 0.0)) >= 0.68
                            and float(metricas_candidato.get("gray", 0.0)) >= 0.89
                            and float(metricas_candidato.get("binary", 0.0)) >= 0.82
                            and float(metricas_candidato.get("ink_iou", 0.0)) >= 0.18
                        )
                    if template_muito_pequeno:
                        aceite_por_score_bruto_candidato = aceite_por_score_bruto_candidato or (
                            float(candidato.get("score", 0.0)) >= max(0.76, float(confianca) - 0.04)
                            and float(candidato.get("quality_score", 0.0)) >= 0.76
                            and float(metricas_candidato.get("corr", 0.0)) >= 0.64
                            and float(metricas_candidato.get("gray", 0.0)) >= 0.88
                            and float(metricas_candidato.get("binary", 0.0)) >= 0.78
                            and float(metricas_candidato.get("ink_iou", 0.0)) >= 0.14
                        )
                else:
                    if perfil_busca == "fast" and str(candidato.get("search_stage") or "") == "reduced":
                        if (
                            float(candidato.get("quality_score", 0.0)) < 0.78
                            or float(metricas_candidato.get("corr", 0.0)) < 0.60
                            or float(metricas_candidato.get("gray", 0.0)) < 0.88
                        ):
                            continue
                    aceite_por_score_bruto_candidato = (
                        float(candidato.get("score", 0.0)) >= max(0.84, float(confianca))
                        and float(candidato.get("quality_score", 0.0)) >= 0.72
                        and float(metricas_candidato.get("corr", 0.0)) >= 0.58
                        and float(metricas_candidato.get("gray", 0.0)) >= 0.88
                    )
                    if perfil_busca == "fast" and str(candidato.get("search_stage") or "") == "reduced":
                        aceite_por_score_bruto_candidato = aceite_por_score_bruto_candidato or (
                            float(candidato.get("score", 0.0)) >= max(0.72, float(confianca) - 0.10)
                            and float(candidato.get("quality_score", 0.0)) >= 0.70
                            and float(metricas_candidato.get("corr", 0.0)) >= 0.52
                            and float(metricas_candidato.get("gray", 0.0)) >= 0.86
                        )
                    if template_pequeno:
                        aceite_por_score_bruto_candidato = aceite_por_score_bruto_candidato or (
                            float(candidato.get("score", 0.0)) >= max(0.74, float(confianca) - 0.08)
                            and float(candidato.get("quality_score", 0.0)) >= 0.68
                            and float(metricas_candidato.get("corr", 0.0)) >= 0.54
                            and float(metricas_candidato.get("gray", 0.0)) >= 0.86
                        )
                    if template_muito_pequeno:
                        aceite_por_score_bruto_candidato = aceite_por_score_bruto_candidato or (
                            float(candidato.get("score", 0.0)) >= max(0.68, float(confianca) - 0.12)
                            and float(candidato.get("quality_score", 0.0)) >= 0.64
                            and float(metricas_candidato.get("corr", 0.0)) >= 0.50
                            and float(metricas_candidato.get("gray", 0.0)) >= 0.84
                        )
                if float(candidato.get("rank_score", 0.0)) < limite and not aceite_por_score_bruto_candidato:
                    continue
                if perfil_busca == "precise" and ambiguidade:
                    safe_print(
                        "Correspondencia visual rejeitada por ambiguidade no modo preciso.",
                        flush=True,
                    )
                    return None
                duracao_analise_ms = (time.perf_counter() - inicio_analise) * 1000.0
                self._ultima_busca_template_info = {
                    "capture_ms": float(info_captura.get("capture_ms", 0.0) or 0.0),
                    "recognition_ms": duracao_analise_ms,
                    "region": list(region) if region else (info_captura.get("region") or None),
                    "score": float(candidato.get("rank_score", 0.0) or 0.0),
                    "raw_score": float(candidato.get("score", 0.0) or 0.0),
                    "quality_score": float(candidato.get("quality_score", 0.0) or 0.0),
                    "search_stage": str(candidato.get("search_stage") or "normal"),
                    "used_multiscale": bool(candidato.get("search_stage") == "multiscale"),
                    "used_reduced": bool(candidato.get("search_stage") == "reduced"),
                    "used_fallback": bool(candidato.get("search_stage") in {"multiscale"}),
                }
                safe_print(
                    "Imagem encontrada por template "
                    f"(score={float(candidato.get('score', 0.0)):.2f}, rank={float(candidato.get('rank_score', 0.0)):.2f}, "
                    f"qualidade={float(candidato.get('quality_score', 0.0)):.2f}, escala={candidato['scale']:.2f}, "
                    f"metodo={candidato.get('method', 'desconhecido')}, corr={float((metricas_candidato or {}).get('corr', 0.0)):.2f}, "
                    f"bin={float((metricas_candidato or {}).get('binary', 0.0)):.2f}, "
                    f"captura={float(info_captura.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                    f"analise={duracao_analise_ms:.0f}ms, area={self.formatar_regiao_log(region)}, "
                    f"etapa={candidato.get('search_stage', 'normal')})",
                    flush=True,
                )
                return candidato
        if ambiguidade:
            safe_print(
                "Correspondencia visual rejeitada por ambiguidade entre candidatos proximos.",
                flush=True,
            )
        elif melhor:
            metricas = melhor.get("metrics") or {}
            safe_print(
                "Template rejeitado por rigidez "
                f"(score={float(melhor.get('score', 0.0)):.2f}, rank={float(melhor.get('rank_score', 0.0)):.2f}, "
                f"qualidade={float(melhor.get('quality_score', 0.0)):.2f}, cor={float(metricas.get('color', 0.0)):.2f}, "
                f"tons={float(metricas.get('gray', 0.0)):.2f}, bordas={float(metricas.get('edges', 0.0)):.2f}, "
                f"corr={float(metricas.get('corr', 0.0)):.2f}).",
                flush=True,
            )
        self._ultima_busca_template_info = {
            "capture_ms": float(info_captura.get("capture_ms", 0.0) or 0.0),
            "recognition_ms": (time.perf_counter() - inicio_analise) * 1000.0,
            "region": list(region) if region else (info_captura.get("region") or None),
            "score": float((melhor or {}).get("rank_score", 0.0) or 0.0),
            "raw_score": float((melhor or {}).get("score", 0.0) or 0.0),
            "quality_score": float((melhor or {}).get("quality_score", 0.0) or 0.0),
            "search_stage": str((melhor or {}).get("search_stage") or ("multiscale" if executou_multiscale else ("reduced" if executou_busca_reduzida else "normal"))),
            "used_multiscale": bool(executou_multiscale),
            "used_reduced": bool(executou_busca_reduzida),
            "used_fallback": bool(executou_multiscale),
        }
        return None

    def salvar_diagnostico_reconhecimento(self, caminho_imagem, region=None, titulo_janela=""):
        if not PIL_DISPONIVEL:
            return None

        try:
            captura, base_left, base_top = self.capturar_tela_pil(region=region)
        except Exception:
            return None

    def salvar_captura_bruta_reconhecimento(self, region=None, titulo_janela=""):
        if not PIL_DISPONIVEL:
            return None

        try:
            captura, base_left, base_top = self.capturar_tela_pil(region=region)
            PASTA_DIAGNOSTICOS_RECONHECIMENTO.mkdir(parents=True, exist_ok=True)
            timestamp = time.strftime("%Y%m%d_%H%M%S")
            imagem_saida = PASTA_DIAGNOSTICOS_RECONHECIMENTO / f"captura_bruta_{timestamp}.png"
            captura.save(imagem_saida)
            metadata_saida = PASTA_DIAGNOSTICOS_RECONHECIMENTO / f"captura_bruta_{timestamp}.json"
            metadata_saida.write_text(
                json.dumps(
                    {
                        "janela_alvo": str(titulo_janela or self.janela_alvo or ""),
                        "regiao_busca": list(region) if region else None,
                        "base_left": int(base_left),
                        "base_top": int(base_top),
                        "width": int(captura.width),
                        "height": int(captura.height),
                    },
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            return str(imagem_saida)
        except Exception:
            return None

        try:
            PASTA_DIAGNOSTICOS_RECONHECIMENTO.mkdir(parents=True, exist_ok=True)
            draw = ImageDraw.Draw(captura) if ImageDraw is not None else None
            melhor = self._ultimo_template_debug or {}
            score = float(melhor.get("score", 0.0) or 0.0)
            if draw and melhor:
                left = int(melhor.get("left", 0) - base_left)
                top = int(melhor.get("top", 0) - base_top)
                width = int(melhor.get("width", 0))
                height = int(melhor.get("height", 0))
                if width > 0 and height > 0:
                    draw.rectangle((left, top, left + width, top + height), outline="#ff9f43", width=3)
                    legenda = f"Melhor score: {score:.2f}"
                    draw.rectangle((max(0, left), max(0, top - 24), max(120, left + 180), max(20, top)), fill="#111827")
                    draw.text((max(4, left + 4), max(2, top - 20)), legenda, fill="#ff9f43")

            timestamp = time.strftime("%Y%m%d_%H%M%S")
            arquivo_base = f"diagnostico_{timestamp}"
            imagem_saida = PASTA_DIAGNOSTICOS_RECONHECIMENTO / f"{arquivo_base}.png"
            captura.save(imagem_saida)

            metadados = {
                "imagem_referencia": str(caminho_imagem),
                "janela_alvo": str(titulo_janela or self.janela_alvo or ""),
                "regiao_busca": list(region) if region else None,
                "base_left": int(base_left),
                "base_top": int(base_top),
                "melhor_template": melhor or None,
            }
            metadata_saida = PASTA_DIAGNOSTICOS_RECONHECIMENTO / f"{arquivo_base}.json"
            metadata_saida.write_text(json.dumps(metadados, ensure_ascii=False, indent=2), encoding="utf-8")
            self._ultimo_diagnostico_reconhecimento = str(imagem_saida)
            return str(imagem_saida)
        except Exception:
            return None

    def emitir_progresso(
        self,
        atual,
        total,
        descricao,
        empresa="",
        sequencia="",
        ordem=0,
        linha_atual=None,
        iteracao_atual=0,
        iteracao_total=0,
    ):
        total_seguro = max(1, int(total or 1))
        percentual = round((max(0, atual) / total_seguro) * 100, 1)
        payload = {
            "current": int(max(0, atual)),
            "total": total_seguro,
            "percent": percentual,
            "description": self.limpar_saida_status(descricao),
            "company": self.limpar_saida_status(empresa),
            "sequence": self.normalizar_nome_sequencia(sequencia) if str(sequencia or "").strip() else "",
            "order": int(ordem or 0),
            "line": int(linha_atual) if linha_atual not in (None, "") else 0,
            "iteration_current": int(iteracao_atual or 0),
            "iteration_total": int(iteracao_total or 0),
        }
        safe_print(
            "PROGRESS_JSON|" + json.dumps(payload, ensure_ascii=False),
            flush=True,
        )

    def iniciar_relatorio_loop(self, evento_planilha, indice_inicial, linha_inicial_real, total_iteracoes, linhas_loop):
        detalhes = evento_planilha.get("detalhes", {}) or {}
        self._relatorio_loop = {
            "iniciado_em": datetime.now(),
            "status": "em_andamento",
            "erro_fatal": "",
            "arquivo": self.normalizar_id_planilha(detalhes.get("arquivo")),
            "aba": str(detalhes.get("aba") or ""),
            "coluna": str(detalhes.get("coluna") or ""),
            "iteracao_inicial": int(indice_inicial) + 1,
            "linha_inicial": int(linha_inicial_real),
            "total_iteracoes": int(total_iteracoes),
            "linhas_loop": [int(linha) for linha in (linhas_loop or []) if linha is not None],
            "resultados": [],
        }

    def registrar_resultado_relatorio_loop(
        self,
        *,
        indice,
        linha_referencia,
        empresa,
        sucesso,
        erro="",
        duracao_segundos=0.0,
    ):
        if not self._relatorio_loop:
            return
        self._relatorio_loop["resultados"].append(
            {
                "indice": int(indice) + 1,
                "linha": int(linha_referencia) if linha_referencia not in (None, "") else None,
                "empresa": str(empresa or "").strip(),
                "sucesso": bool(sucesso),
                "erro": str(erro or "").strip(),
                "duracao_segundos": round(max(0.0, float(duracao_segundos or 0.0)), 2),
                "finalizado_em": datetime.now(),
            }
        )

    def gerar_relatorio_loop_pdf(self, relatorio):
        if not relatorio:
            return ""
        if not PDF_RELATORIO_DISPONIVEL:
            safe_print(
                "Relatorio do loop nao foi gerado em PDF porque a biblioteca reportlab nao esta instalada.",
                flush=True,
            )
            return ""

        PASTA_RELATORIOS_LOOP.mkdir(parents=True, exist_ok=True)
        iniciado_em = relatorio.get("iniciado_em") or datetime.now()
        arquivo_base = re.sub(r"[^A-Za-z0-9._-]+", "_", str(relatorio.get("arquivo") or "planilha")) or "planilha"
        nome_arquivo = f"relatorio_loop_{arquivo_base}_{iniciado_em.strftime('%Y%m%d_%H%M%S')}.pdf"
        caminho_pdf = PASTA_RELATORIOS_LOOP / nome_arquivo

        largura, altura = A4
        margem_x = 42
        topo = altura - 42
        rodape = 42
        largura_texto = 94
        y = topo

        def nova_pagina(pdf):
            pdf.showPage()
            pdf.setTitle("Relatorio de execucao do loop")
            return topo

        def escrever(pdf, texto, *, fonte="Helvetica", tamanho=10, espacamento=14):
            nonlocal y
            for linha in wrap(str(texto or ""), width=largura_texto) or [""]:
                if y <= rodape:
                    y = nova_pagina(pdf)
                pdf.setFont(fonte, tamanho)
                pdf.drawString(margem_x, y, linha)
                y -= espacamento

        pdf = canvas.Canvas(str(caminho_pdf), pagesize=A4)
        pdf.setTitle("Relatorio de execucao do loop")

        resultados = relatorio.get("resultados", [])
        total_linhas = len(resultados)
        total_ok = sum(1 for item in resultados if item.get("sucesso"))
        total_erro = total_linhas - total_ok
        status_final = str(relatorio.get("status") or "desconhecido").upper()

        escrever(pdf, "Relatorio de execucao do loop", fonte="Helvetica-Bold", tamanho=16, espacamento=18)
        escrever(pdf, f"Status final: {status_final}", fonte="Helvetica-Bold")
        escrever(pdf, f"Iniciado em: {iniciado_em.strftime('%d/%m/%Y %H:%M:%S')}")
        finalizado_em = relatorio.get("finalizado_em")
        if isinstance(finalizado_em, datetime):
            escrever(pdf, f"Finalizado em: {finalizado_em.strftime('%d/%m/%Y %H:%M:%S')}")
        escrever(pdf, f"Arquivo da planilha: {relatorio.get('arquivo') or '-'}")
        escrever(pdf, f"Aba: {relatorio.get('aba') or '-'} | Coluna: {relatorio.get('coluna') or '-'}")
        escrever(
            pdf,
            f"Inicio do loop: iteracao {relatorio.get('iteracao_inicial') or 1} | linha real {relatorio.get('linha_inicial') or '-'}",
        )
        escrever(
            pdf,
            f"Resumo: {total_ok} linha(s) com sucesso | {total_erro} linha(s) com erro/interrupcao | {total_linhas} linha(s) processada(s)",
            fonte="Helvetica-Bold",
        )

        erro_fatal = str(relatorio.get("erro_fatal") or "").strip()
        if erro_fatal:
            escrever(pdf, f"Erro fatal: {erro_fatal}", fonte="Helvetica-Bold")

        y -= 6
        escrever(pdf, "Detalhamento por linha", fonte="Helvetica-Bold", tamanho=13, espacamento=16)

        if not resultados:
            escrever(pdf, "Nenhuma linha da planilha foi processada.")
        else:
            for item in resultados:
                linha = item.get("linha")
                empresa = item.get("empresa") or "-"
                status_linha = "OK" if item.get("sucesso") else "ERRO"
                duracao = item.get("duracao_segundos", 0)
                cabecalho = (
                    f"Linha {linha if linha is not None else '-'} | "
                    f"Iteracao {item.get('indice')} | "
                    f"Status: {status_linha} | "
                    f"Duracao: {duracao:.2f}s"
                )
                escrever(pdf, cabecalho, fonte="Helvetica-Bold")
                escrever(pdf, f"Empresa/valor: {empresa}")
                if item.get("sucesso"):
                    escrever(pdf, "Resultado: todas as etapas do loop foram executadas com sucesso.")
                else:
                    escrever(pdf, f"Resultado: {item.get('erro') or 'Falha durante a execucao desta linha.'}")
                y -= 4

        pdf.save()
        return str(caminho_pdf)

    def finalizar_relatorio_loop(self, status="concluido", erro_fatal=""):
        if not self._relatorio_loop:
            return ""
        self._relatorio_loop["status"] = str(status or "concluido")
        self._relatorio_loop["erro_fatal"] = str(erro_fatal or "").strip()
        self._relatorio_loop["finalizado_em"] = datetime.now()
        try:
            return self.gerar_relatorio_loop_pdf(self._relatorio_loop)
        finally:
            self._relatorio_loop = None

    def identificar_coluna_empresa(self, colunas):
        candidatos_exatos = {
            "empresa",
            "razao social",
            "nome empresa",
            "cliente",
            "nome",
        }
        candidatos_parciais = ("empresa", "razao", "cliente")
        for coluna in colunas:
            nome = str(coluna).strip().lower()
            if nome in candidatos_exatos:
                return coluna
        for coluna in colunas:
            nome = str(coluna).strip().lower()
            if any(parte in nome for parte in candidatos_parciais):
                return coluna
        return None

    def obter_registros_planilha(self, arquivo, aba):
        return self.planilha_cache.get(self.normalizar_id_planilha(arquivo), {}).get(aba, {}).get("linhas", [])

    def obter_valores_evento_planilha(self, evento):
        detalhes = evento.get("detalhes", {})
        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        aba = detalhes.get("aba")
        coluna = detalhes.get("coluna")
        if not aba or not coluna:
            return [valor for valor in detalhes.get("valores", [])]

        registros = self.planilha_cache.get(arquivo, {}).get(aba, {}).get("linhas", [])
        if not registros:
            return [valor for valor in detalhes.get("valores", [])]

        linhas_selecionadas = [int(linha) for linha in (detalhes.get("linhas_selecionadas") or [])]
        registros_por_linha = {int(registro.get("__linha__", 0)): registro for registro in registros}

        if linhas_selecionadas:
            valores = []
            for numero_linha in linhas_selecionadas:
                registro = registros_por_linha.get(numero_linha)
                if not registro:
                    continue
                valor = registro.get(coluna)
                if valor is None or str(valor).strip() == "":
                    continue
                valores.append(valor)
            return valores

        valores = []
        for registro in registros:
            valor = registro.get(coluna)
            if valor is None or str(valor).strip() == "":
                continue
            valores.append(valor)
        return valores if valores else [valor for valor in detalhes.get("valores", [])]

    def obter_linhas_evento_planilha(self, detalhes):
        linhas_selecionadas = detalhes.get("linhas_selecionadas") or []
        if linhas_selecionadas:
            return [int(linha) for linha in linhas_selecionadas]

        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        aba = detalhes.get("aba")
        coluna = detalhes.get("coluna")
        if not aba or not coluna:
            return []

        linhas = []
        for registro in self.obter_registros_planilha(arquivo, aba):
            valor = registro.get(coluna)
            if valor is None or str(valor).strip() == "":
                continue
            linhas.append(int(registro.get("__linha__", 0)))
        return linhas

    def resolver_inicio_loop(self, evento_planilha, loop_start):
        detalhes = evento_planilha.get("detalhes", {})
        valores = self.obter_valores_evento_planilha(evento_planilha)
        total_iteracoes = len(valores)
        if total_iteracoes == 0:
            raise ValueError("A sequencia de planilha nao possui valores para o loop.")

        linha_solicitada = max(1, int(loop_start or 1))
        linhas = self.obter_linhas_evento_planilha(detalhes)

        if not linhas:
            indice_inicial = linha_solicitada - 1
            if indice_inicial >= total_iteracoes:
                raise ValueError(
                    f"O item inicial {linha_solicitada} esta fora do intervalo do loop (1 a {total_iteracoes})."
                )
            return indice_inicial, linha_solicitada, total_iteracoes, linhas

        for indice, numero_linha in enumerate(linhas):
            if numero_linha >= linha_solicitada:
                return indice, numero_linha, total_iteracoes, linhas

        raise ValueError(
            f"A linha inicial {linha_solicitada} esta fora do intervalo do loop. "
            f"Ultima linha disponivel: {linhas[-1]}."
        )

    def obter_valor_evento_planilha_no_loop(self, evento, linha_referencia, indice_fallback):
        detalhes = evento.get("detalhes", {})
        coluna = detalhes.get("coluna")
        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        aba = detalhes.get("aba")

        if linha_referencia is not None and aba and coluna:
            for registro in self.obter_registros_planilha(arquivo, aba):
                if int(registro.get("__linha__", 0)) != int(linha_referencia):
                    continue
                valor = registro.get(coluna)
                if valor is None or str(valor).strip() == "":
                    break
                return valor

        valores = self.obter_valores_evento_planilha(evento)
        if 0 <= indice_fallback < len(valores):
            return valores[indice_fallback]
        return None

    def obter_contexto_linha_loop(self, contexto_loop=None):
        if not contexto_loop:
            return {}
        escopo_principal = contexto_loop.get("loop_principal", {}) if isinstance(contexto_loop, dict) else {}
        linha_referencia = contexto_loop.get("linha_referencia")
        workbook_preferido = contexto_loop.get("workbook_id") or escopo_principal.get("arquivo")
        aba_preferida = contexto_loop.get("aba") or escopo_principal.get("aba")
        if linha_referencia in (None, ""):
            return {}

        cache_iteravel = self.planilha_cache.items()
        if workbook_preferido and workbook_preferido in self.planilha_cache:
            cache_iteravel = [(workbook_preferido, self.planilha_cache.get(workbook_preferido, {}))]

        for workbook_id, abas in cache_iteravel:
            for nome_aba, cache_aba in (abas or {}).items():
                if aba_preferida and str(nome_aba) != str(aba_preferida):
                    continue
                for registro in (cache_aba or {}).get("linhas", []):
                    if int(registro.get("__linha__", 0) or 0) != int(linha_referencia):
                        continue
                    contexto_linha = {
                        chave: valor
                        for chave, valor in registro.items()
                        if chave != "__linha__"
                    }
                    contexto_linha["__arquivo__"] = workbook_id
                    contexto_linha["__aba__"] = nome_aba
                    contexto_linha["__linha__"] = int(linha_referencia)
                    return contexto_linha
        return {}

    def obter_linha_atual_contexto(self, contexto_loop=None):
        if not contexto_loop:
            return None
        linha_atual = contexto_loop.get("linha_atual")
        if linha_atual not in (None, ""):
            return int(linha_atual)
        linha_referencia = contexto_loop.get("linha_referencia")
        if linha_referencia not in (None, ""):
            return int(linha_referencia)
        escopo_principal = contexto_loop.get("loop_principal", {}) if isinstance(contexto_loop, dict) else {}
        linha_principal = escopo_principal.get("linha_atual")
        if linha_principal not in (None, ""):
            return int(linha_principal)
        return None

    def salvar_variavel_contexto(self, nome, valor):
        chave = str(nome or "").strip()
        if chave.startswith("{{") and chave.endswith("}}"):
            chave = chave[2:-2].strip()
        if not chave:
            return
        self.contexto_variaveis[chave] = value = "" if valor is None else str(valor)
        safe_print(f"Variavel de contexto atualizada: {chave}={value}", flush=True)

    def montar_contexto_substituicao(self, contexto_loop=None):
        contexto = {}
        for chave, valor in (self.contexto_variaveis or {}).items():
            contexto[str(chave)] = valor

        linha = self.obter_contexto_linha_loop(contexto_loop)
        contexto["linha"] = linha
        contexto["linha_atual"] = (
            self.obter_linha_atual_contexto(contexto_loop)
            if contexto_loop
            else linha.get("__linha__") if isinstance(linha, dict) else ""
        )
        contexto["planilha_atual"] = linha.get("__arquivo__", "") if isinstance(linha, dict) else ""
        contexto["aba_atual"] = linha.get("__aba__", "") if isinstance(linha, dict) else ""

        for chave, valor in linha.items():
            if str(chave).startswith("__"):
                continue
            contexto.setdefault(str(chave), valor)

        if contexto_loop:
            contexto["loop_principal"] = contexto_loop.get("loop_principal", {})
            contexto["loop_local"] = contexto_loop.get("loop_local", {})
            for chave, valor in contexto_loop.items():
                if isinstance(valor, (str, int, float, bool)) or valor is None:
                    contexto[str(chave)] = valor

        return contexto

    def resolver_variavel_contexto(self, token, contexto):
        chave = str(token or "").strip()
        if not chave:
            return ""

        if chave in contexto:
            return contexto[chave]

        atual = contexto
        for parte in chave.split("."):
            if isinstance(atual, dict) and parte in atual:
                atual = atual[parte]
                continue
            atual = None
            break
        if atual is not None:
            return atual

        linha = contexto.get("linha") if isinstance(contexto, dict) else None
        if isinstance(linha, dict):
            if chave in linha:
                return linha[chave]
            if chave.startswith("linha."):
                ultimo = chave.split(".", 1)[1]
                return linha.get(ultimo, "")

        return ""

    def substituir_variaveis_texto(self, texto, contexto_loop=None):
        if texto is None:
            return ""

        contexto = self.montar_contexto_substituicao(contexto_loop)
        bruto = str(texto)
        if "{{" not in bruto or "}}" not in bruto:
            return bruto

        def substituir(match):
            token = str(match.group(1) or "").strip()
            valor = self.resolver_variavel_contexto(token, contexto)
            return "" if valor is None else str(valor)

        return re.sub(r"\{\{\s*([^{}]+?)\s*\}\}", substituir, bruto)

    def resumir_html(self, html, limite=140):
        texto = re.sub(r"\s+", " ", str(html or "")).strip()
        if len(texto) <= limite:
            return texto
        return texto[: max(0, limite - 3)] + "..."

    def slug_texto(self, texto):
        texto = self.normalizar_texto_busca(texto).replace(" ", "_")
        return texto or "bloco"

    def garantir_playwright(self):
        if not PLAYWRIGHT_DISPONIVEL or sync_playwright is None:
            raise RuntimeError(
                "Playwright nao esta disponivel neste ambiente. "
                "Instale com: pip install playwright"
            )
        if self._playwright_runtime is None:
            self._playwright_runtime = sync_playwright().start()
        return self._playwright_runtime

    def selecionar_pagina_chrome(self, browser):
        contextos = list(getattr(browser, "contexts", []) or [])
        if contextos:
            for contexto in contextos:
                paginas = list(getattr(contexto, "pages", []) or [])
                for pagina in reversed(paginas):
                    try:
                        if not pagina.is_closed():
                            return contexto, pagina
                    except Exception:
                        continue
            try:
                return contextos[0], contextos[0].new_page()
            except Exception:
                pass

        contexto = browser.new_context(no_viewport=True, accept_downloads=True)
        return contexto, contexto.new_page()

    def garantir_sessao_chrome(self):
        if self._chrome_page is not None:
            try:
                if not self._chrome_page.is_closed():
                    return self._chrome_page
            except Exception:
                self._chrome_page = None

        runtime = self.garantir_playwright()
        chromium = runtime.chromium
        ultimo_erro = None

        for cdp_url in ("http://127.0.0.1:9222", "http://localhost:9222"):
            try:
                browser = chromium.connect_over_cdp(cdp_url)
                contexto, pagina = self.selecionar_pagina_chrome(browser)
                self._chrome_browser = browser
                self._chrome_context = contexto
                self._chrome_page = pagina
                self._chrome_source = f"cdp:{cdp_url}"
                self._chrome_launched_by_executor = False
                safe_print(f"Sessao Chrome conectada via CDP: {cdp_url}", flush=True)
                return pagina
            except Exception as erro:
                ultimo_erro = erro

        candidatos_launch = [
            {
                "channel": "chrome",
                "headless": False,
                "args": ["--start-maximized"],
            },
            {
                "headless": False,
                "args": ["--start-maximized"],
            },
        ]
        for parametros in candidatos_launch:
            try:
                browser = chromium.launch(**parametros)
                contexto = browser.new_context(no_viewport=True, accept_downloads=True)
                pagina = contexto.new_page()
                self._chrome_browser = browser
                self._chrome_context = contexto
                self._chrome_page = pagina
                self._chrome_source = "launch:chrome" if parametros.get("channel") == "chrome" else "launch:chromium"
                self._chrome_launched_by_executor = True
                safe_print(
                    "Sessao Chrome iniciada em modo controlado. "
                    "Se voce queria usar uma aba ja aberta, abra o Chrome com --remote-debugging-port=9222.",
                    flush=True,
                )
                return pagina
            except Exception as erro:
                ultimo_erro = erro

        raise RuntimeError(
            "Nao foi possivel conectar ou iniciar uma sessao Playwright/Chrome. "
            f"Ultimo erro: {ultimo_erro}"
        )

    def obter_pagina_chrome(self):
        pagina = self.garantir_sessao_chrome()
        if self._chrome_context is None and self._chrome_browser is not None:
            try:
                self._chrome_context, self._chrome_page = self.selecionar_pagina_chrome(self._chrome_browser)
                pagina = self._chrome_page
            except Exception:
                pass
        return pagina

    def criar_locator_playwright(self, page, alvo):
        alvo = alvo or {}
        metodo = str(alvo.get("metodo", "css") or "css").strip().lower()
        valor = str(alvo.get("valor", "") or "").strip()
        seletor_css = str(alvo.get("seletor_css", "") or "").strip()
        role = str(alvo.get("role", "") or "").strip()

        if not valor and not seletor_css:
            raise ValueError("O bloco Chrome nao possui alvo suficiente para localizar o elemento.")

        if metodo == "placeholder":
            return page.get_by_placeholder(valor)
        if metodo == "name":
            seletor = seletor_css or f'[name="{valor.replace("\\", "\\\\").replace("\"", "\\\"")}"]'
            return page.locator(seletor)
        if metodo == "aria-label":
            return page.get_by_label(valor)
        if metodo == "labelledby":
            seletor = seletor_css or f'[aria-labelledby="{valor.replace("\\", "\\\\").replace("\"", "\\\"")}"]'
            return page.locator(seletor)
        if metodo == "text":
            return page.get_by_text(valor, exact=True)
        if metodo == "role":
            if not role:
                raise ValueError("O metodo Role exige que o bloco tenha um role salvo.")
            return page.get_by_role(role, name=valor, exact=True)
        if metodo == "xpath":
            xpath = valor if valor.startswith("xpath=") else f"xpath={valor}"
            return page.locator(xpath)
        if metodo == "id":
            seletor = seletor_css or f'[id="{valor.replace("\\", "\\\\").replace("\"", "\\\"")}"]'
            return page.locator(seletor)
        if metodo == "css":
            return page.locator(valor or seletor_css)

        if seletor_css:
            return page.locator(seletor_css)
        return page.locator(valor)

    def salvar_captura_erro_chrome(self, nome_bloco, page=None):
        PASTA_DIAGNOSTICOS_CHROME.mkdir(parents=True, exist_ok=True)
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        arquivo = PASTA_DIAGNOSTICOS_CHROME / f"{self.slug_texto(nome_bloco)}_{timestamp}.png"
        pagina = page or self._chrome_page
        try:
            if pagina is not None:
                pagina.screenshot(path=str(arquivo), full_page=True)
            else:
                pyautogui.screenshot(str(arquivo))
            return str(arquivo)
        except Exception:
            try:
                pyautogui.screenshot(str(arquivo))
                return str(arquivo)
            except Exception:
                return ""

    def executar_fallback_visual_chrome(self, evento, contexto_loop=None):
        detalhes = evento.get("detalhes", {}) or {}
        if not detalhes.get("fallback_visual"):
            return False

        imagem_fallback = str(detalhes.get("fallback_imagem", "") or "").strip()
        if not imagem_fallback:
            safe_print("Fallback visual habilitado, mas nenhuma imagem foi vinculada ao bloco Chrome.", flush=True)
            return False

        nome_bloco = str(detalhes.get("nome") or "Chrome / Elemento HTML")
        acao = dict(detalhes.get("acao") or {})
        tipo_acao = str(acao.get("tipo") or "clicar").strip().lower()
        timeout = max(0.5, float(detalhes.get("timeout", 10000) or 10000) / 1000.0)

        safe_print(f"[Chrome: {nome_bloco}] Tentando fallback visual com {imagem_fallback}.", flush=True)

        if tipo_acao in {"esperar", "verificar_existe"}:
            return bool(
                self.clicar_imagem(
                    imagem=imagem_fallback,
                    timeout=timeout,
                    confianca=0.8,
                    clicar=False,
                    perfil_busca="balanced",
                )
            )

        local = self.clicar_imagem(
            imagem=imagem_fallback,
            timeout=timeout,
            confianca=0.8,
            clicar=True,
            clicks=1,
            perfil_busca="balanced",
        )
        if not local:
            return False

        if tipo_acao == "limpar":
            pyautogui.hotkey("ctrl", "a")
            pyautogui.press("backspace")
            return True

        if tipo_acao == "preencher":
            valor = self.substituir_variaveis_texto(acao.get("valor", ""), contexto_loop)
            if acao.get("limpar_antes", True):
                pyautogui.hotkey("ctrl", "a")
                time.sleep(0.05)
                pyautogui.press("backspace")
            if valor:
                self.escrever_texto(valor)
            tecla_depois = str(acao.get("pressionar_depois", "") or "").strip()
            if tecla_depois:
                self.executar_acao(tecla_depois.lower())
            return True

        if tipo_acao in {"selecionar_texto", "selecionar_valor", "selecionar_indice"}:
            valor = self.substituir_variaveis_texto(acao.get("valor", ""), contexto_loop)
            if valor:
                self.escrever_texto(valor)
                pyautogui.press("enter")
            return True

        if tipo_acao in {"clicar", "clicar_aguardar_carregamento", "clicar_aguardar_download", "pegar_texto"}:
            return True

        return False

    def executar_bloco_chrome_elemento_html(self, evento, contexto_loop=None):
        detalhes = evento.get("detalhes", {}) or {}
        nome_bloco = str(detalhes.get("nome") or "Chrome / Elemento HTML").strip() or "Chrome / Elemento HTML"
        html_resumido = self.resumir_html(detalhes.get("html_original", ""))
        alvo = dict(detalhes.get("alvo") or {})
        acao = dict(detalhes.get("acao") or {})
        metodo = str(alvo.get("metodo") or "css").strip()
        valor_alvo = str(alvo.get("valor") or "").strip()
        tipo_acao = str(acao.get("tipo") or "clicar").strip().lower()
        timeout = max(500, int(detalhes.get("timeout", 10000) or 10000))
        tentativas = max(1, int(detalhes.get("quantidade_tentativas", 1) or 1))
        intervalo_tentativas = max(0.0, float(detalhes.get("intervalo_tentativas", 0.5) or 0.0))
        ao_sucesso = str(detalhes.get("ao_sucesso", "continuar") or "continuar").strip().lower()
        bloco_sucesso = detalhes.get("bloco_sucesso", "")
        sequencia_sucesso = str(detalhes.get("sequencia_sucesso", "") or "").strip()
        atraso_sequencia_sucesso = max(0.0, float(detalhes.get("atraso_antes_sequencia_sucesso", 0) or 0))
        ao_falhar = str(detalhes.get("ao_falhar", "parar") or "parar").strip().lower()
        bloco_falha = detalhes.get("bloco_falha", "")
        sequencia_falha = str(detalhes.get("sequencia_falha", "") or "").strip()
        atraso_sequencia_falha = max(0.0, float(detalhes.get("atraso_antes_sequencia_falha", 0) or 0))
        ultima_falha = None

        def finalizar_sucesso():
            if ao_sucesso == "ir_para_sequencia" and sequencia_sucesso:
                safe_print(f"[Chrome: {nome_bloco}] Sucesso. Redirecionando para pagina {sequencia_sucesso}.", flush=True)
                return self.resultado_fluxo("ir_para_sequencia", sequencia=sequencia_sucesso, atraso=atraso_sequencia_sucesso)
            if ao_sucesso == "ir_para_bloco" and str(bloco_sucesso or "").strip():
                safe_print(f"[Chrome: {nome_bloco}] Sucesso. Redirecionando para bloco {bloco_sucesso}.", flush=True)
                return self.resultado_fluxo("ir_para_bloco", bloco=bloco_sucesso, atraso=atraso_sequencia_sucesso)
            return self.resultado_fluxo("continuar")

        for tentativa in range(1, tentativas + 1):
            if self.cancelamento_solicitado():
                return self.resultado_fluxo("parar")

            safe_print(f"[Chrome: {nome_bloco}] Reconhecendo... tentativa {tentativa}/{tentativas}", flush=True)
            safe_print(f"[Chrome: {nome_bloco}] HTML: {html_resumido}", flush=True)
            safe_print(
                f"[Chrome: {nome_bloco}] Metodo: {metodo} | Alvo: {valor_alvo} | Acao: {tipo_acao}",
                flush=True,
            )

            try:
                page = self.obter_pagina_chrome()
                try:
                    page.bring_to_front()
                except Exception:
                    pass
                locator = self.criar_locator_playwright(page, alvo).first

                valor_acao = self.substituir_variaveis_texto(acao.get("valor", ""), contexto_loop)
                tecla_depois = str(acao.get("pressionar_depois", "") or "").strip()
                salvar_em = str(acao.get("salvar_em", "") or "").strip()

                if tipo_acao == "preencher":
                    locator.wait_for(state="visible", timeout=timeout)
                    if acao.get("limpar_antes", True):
                        locator.fill("", timeout=timeout)
                    locator.fill(valor_acao, timeout=timeout)
                    if tecla_depois:
                        locator.press(tecla_depois, timeout=timeout)
                    safe_print(f"[Chrome: {nome_bloco}] Valor usado: {valor_acao}", flush=True)

                elif tipo_acao == "limpar":
                    locator.wait_for(state="visible", timeout=timeout)
                    locator.fill("", timeout=timeout)

                elif tipo_acao == "clicar":
                    locator.wait_for(state="visible", timeout=timeout)
                    locator.click(timeout=timeout)

                elif tipo_acao == "esperar":
                    locator.wait_for(state="visible", timeout=timeout)

                elif tipo_acao == "verificar_existe":
                    locator.wait_for(state="attached", timeout=timeout)

                elif tipo_acao == "pegar_texto":
                    locator.wait_for(state="visible", timeout=timeout)
                    texto_capturado = locator.inner_text(timeout=timeout)
                    if salvar_em:
                        self.salvar_variavel_contexto(salvar_em, texto_capturado)
                    safe_print(f"[Chrome: {nome_bloco}] Texto capturado: {texto_capturado}", flush=True)

                elif tipo_acao == "clicar_aguardar_carregamento":
                    locator.wait_for(state="visible", timeout=timeout)
                    locator.click(timeout=timeout)
                    page.wait_for_load_state("load", timeout=timeout)

                elif tipo_acao == "clicar_aguardar_download":
                    locator.wait_for(state="visible", timeout=timeout)
                    with page.expect_download(timeout=timeout) as download_info:
                        locator.click(timeout=timeout)
                    download = download_info.value
                    if salvar_em:
                        self.salvar_variavel_contexto(salvar_em, download.suggested_filename)
                    safe_print(
                        f"[Chrome: {nome_bloco}] Download detectado: {download.suggested_filename}",
                        flush=True,
                    )

                elif tipo_acao == "selecionar_texto":
                    locator.wait_for(state="visible", timeout=timeout)
                    locator.select_option(label=valor_acao, timeout=timeout)

                elif tipo_acao == "selecionar_valor":
                    locator.wait_for(state="visible", timeout=timeout)
                    locator.select_option(value=valor_acao, timeout=timeout)

                elif tipo_acao == "selecionar_indice":
                    locator.wait_for(state="visible", timeout=timeout)
                    indice = max(0, int(valor_acao or 0))
                    locator.select_option(index=indice, timeout=timeout)

                else:
                    raise ValueError(f"A acao Chrome nao eh suportada: {tipo_acao}")

                safe_print(f"[Chrome: {nome_bloco}] Resultado: sucesso", flush=True)
                return finalizar_sucesso()

            except Exception as erro:
                ultima_falha = erro
                caminho_print = self.salvar_captura_erro_chrome(nome_bloco)
                safe_print(f"[Chrome: {nome_bloco}] Resultado: erro | {erro}", flush=True)
                if caminho_print:
                    safe_print(f"[Chrome: {nome_bloco}] Print salvo em: {caminho_print}", flush=True)

                if detalhes.get("fallback_visual") and detalhes.get("fallback_imagem"):
                    try:
                        if self.executar_fallback_visual_chrome(evento, contexto_loop=contexto_loop):
                            safe_print(f"[Chrome: {nome_bloco}] Fallback visual executado com sucesso.", flush=True)
                            return finalizar_sucesso()
                    except Exception as erro_fallback:
                        safe_print(f"[Chrome: {nome_bloco}] Falha no fallback visual: {erro_fallback}", flush=True)

                if tentativa < tentativas and intervalo_tentativas > 0:
                    if not self.esperar(intervalo_tentativas):
                        return self.resultado_fluxo("parar")

        if ao_falhar == "continuar":
            safe_print(f"[Chrome: {nome_bloco}] Falha final. Acao configurada: continuar.", flush=True)
            return self.resultado_fluxo("continuar")
        if ao_falhar == "ir_para_sequencia" and sequencia_falha:
            safe_print(f"[Chrome: {nome_bloco}] Falha final. Redirecionando para pagina {sequencia_falha}.", flush=True)
            return self.resultado_fluxo("ir_para_sequencia", sequencia=sequencia_falha, atraso=atraso_sequencia_falha)
        if ao_falhar == "ir_para_bloco" and str(bloco_falha or "").strip():
            safe_print(f"[Chrome: {nome_bloco}] Falha final. Redirecionando para bloco {bloco_falha}.", flush=True)
            return self.resultado_fluxo("ir_para_bloco", bloco=bloco_falha, atraso=atraso_sequencia_falha)

        if ao_falhar == "tentar_novamente":
            raise RuntimeError(
                f"Falha ao executar bloco Chrome apos {tentativas} tentativa(s): {nome_bloco} | erro: {ultima_falha}"
            ) from ultima_falha

        raise RuntimeError(
            f"Falha ao executar bloco Chrome: {nome_bloco} | metodo={metodo} | alvo={valor_alvo} | erro: {ultima_falha}"
        ) from ultima_falha

    def carregar_uia_runtime(self):
        try:
            import uiautomation as auto

            return auto
        except Exception as erro:
            raise RuntimeError(
                "A biblioteca uiautomation nao esta instalada neste computador. "
                "Atualize as dependencias do FlowDetect/Agent antes de executar este bloco."
            ) from erro

    def controle_uia_existe(self, control, timeout=0.2):
        try:
            return bool(control.Exists(timeout))
        except Exception:
            return control is not None

    def propriedades_controle_uia(self, control):
        try:
            parent = control.GetParentControl()
            parent_name = str(getattr(parent, "Name", "") or "") if parent else ""
        except Exception:
            parent_name = ""
        return {
            "automation_id": str(getattr(control, "AutomationId", "") or ""),
            "name": str(getattr(control, "Name", "") or ""),
            "control_type": str(getattr(control, "ControlTypeName", "") or ""),
            "class_name": str(getattr(control, "ClassName", "") or ""),
            "parent": parent_name,
        }

    def candidatos_raiz_uia(self, auto, elemento_base=None):
        elemento_base = elemento_base or {}
        source_window = elemento_base.get("source_window") or {}
        candidatos = []
        vistos = set()

        def adicionar(control):
            if not control:
                return
            try:
                chave = int(getattr(control, "NativeWindowHandle", 0) or 0)
            except Exception:
                chave = id(control)
            if chave in vistos:
                return
            vistos.add(chave)
            candidatos.append(control)

        handle = int(source_window.get("handle") or elemento_base.get("handle") or 0)
        if handle:
            try:
                control = auto.ControlFromHandle(handle)
                if self.controle_uia_existe(control, 0.2):
                    adicionar(control)
            except Exception:
                pass

        title = str(source_window.get("title") or "").strip()
        class_name = str(source_window.get("class_name") or "").strip()
        if title:
            try:
                control = auto.WindowControl(searchDepth=1, Name=title)
                if self.controle_uia_existe(control, 0.4):
                    adicionar(control)
            except Exception:
                pass

        try:
            adicionar(auto.GetForegroundControl())
        except Exception:
            pass

        try:
            root = auto.GetRootControl()
            for child in root.GetChildren():
                props = self.propriedades_controle_uia(child)
                if title and title.lower() not in props.get("name", "").lower():
                    if class_name and class_name != props.get("class_name", ""):
                        continue
                    if not class_name:
                        continue
                elif class_name and class_name != props.get("class_name", ""):
                    continue
                adicionar(child)
        except Exception:
            pass

        return candidatos

    def consultas_selector_uia(self, selector, elemento_base=None):
        elemento_base = elemento_base or {}
        selector = selector or {}
        query_original = dict(selector.get("query") or {})
        permitidas = {"automation_id", "name", "control_type", "class_name", "parent"}
        consultas = []

        def adicionar(query):
            filtrada = {k: str(v) for k, v in (query or {}).items() if k in permitidas and str(v).strip()}
            if filtrada and filtrada not in consultas:
                consultas.append(filtrada)

        adicionar(query_original)
        automation_id = elemento_base.get("automation_id")
        name = elemento_base.get("name")
        control_type = elemento_base.get("control_type")
        class_name = elemento_base.get("class_name")
        parent = elemento_base.get("parent")
        adicionar({"automation_id": automation_id, "control_type": control_type})
        adicionar({"automation_id": automation_id})
        adicionar({"name": name, "control_type": control_type})
        adicionar({"class_name": class_name, "name": name, "parent": parent})
        adicionar({"class_name": class_name, "name": name})
        adicionar({"class_name": class_name, "control_type": control_type, "parent": parent})
        adicionar({"class_name": class_name, "control_type": control_type})
        adicionar({"control_type": control_type, "parent": parent})
        adicionar({"name": name})
        return consultas

    def controle_uia_compativel_com_elemento(self, control, elemento_base=None):
        elemento_base = elemento_base or {}
        props = self.propriedades_controle_uia(control)
        control_type = str(elemento_base.get("control_type") or "").strip()
        class_name = str(elemento_base.get("class_name") or "").strip()
        parent = str(elemento_base.get("parent") or "").strip()
        if control_type and props.get("control_type") != control_type:
            return False
        if class_name and props.get("class_name") != class_name:
            return False
        if parent and props.get("parent") and props.get("parent") != parent:
            return False
        return True

    def retangulo_controle_uia(self, control):
        try:
            rect = getattr(control, "BoundingRectangle", None)
            if not rect:
                return None
            return {
                "left": int(rect.left),
                "top": int(rect.top),
                "width": int(rect.right) - int(rect.left),
                "height": int(rect.bottom) - int(rect.top),
            }
        except Exception:
            return None

    def referencia_relativa_uia(self, elemento_base=None, root=None):
        elemento_base = elemento_base or {}
        rect = elemento_base.get("bounding_rectangle") or {}
        source_rect = (elemento_base.get("source_window") or {}).get("bounding_rectangle") or {}
        try:
            left = int(rect.get("left", 0) or 0)
            top = int(rect.get("top", 0) or 0)
            width = int(rect.get("width", 0) or 0)
            height = int(rect.get("height", 0) or 0)
        except Exception:
            return None
        if width <= 0 or height <= 0:
            return None
        if root:
            root_rect = self.retangulo_controle_uia(root)
            try:
                source_left = int(source_rect.get("left", 0) or 0)
                source_top = int(source_rect.get("top", 0) or 0)
            except Exception:
                source_left = 0
                source_top = 0
            if root_rect and source_rect:
                left = int(root_rect["left"]) + (left - source_left)
                top = int(root_rect["top"]) + (top - source_top)
        return {"left": left, "top": top, "width": width, "height": height}

    def localizar_elemento_uia_por_area_salva(self, auto, elemento_base=None, roots=None):
        ref = self.referencia_relativa_uia(elemento_base)
        if not ref:
            return None
        if roots:
            for root in roots:
                ajustada = self.referencia_relativa_uia(elemento_base, root)
                if ajustada:
                    ref = ajustada
                    break
        left = ref["left"]
        top = ref["top"]
        width = ref["width"]
        height = ref["height"]
        pontos = [
            (left + width // 2, top + height // 2),
            (left + min(width - 1, max(1, width // 4)), top + height // 2),
            (left + max(1, (width * 3) // 4), top + height // 2),
        ]
        for x, y in pontos:
            try:
                control = auto.ControlFromPoint(x, y)
                if control and self.controle_uia_compativel_com_elemento(control, elemento_base):
                    return control
            except Exception:
                continue
        return None

    def caminho_uia_parece(self, path_atual="", path_salvo=""):
        if not path_atual or not path_salvo:
            return False
        atual = str(path_atual).strip()
        salvo = str(path_salvo).strip()
        return atual == salvo or atual.endswith(salvo.split("/", 1)[-1])

    def pontuar_controle_uia(self, control, elemento_base=None, root=None, path_atual=""):
        elemento_base = elemento_base or {}
        score = 0
        try:
            if not bool(getattr(control, "IsOffscreen", False)):
                score += 5
        except Exception:
            pass
        try:
            if bool(getattr(control, "IsEnabled", False)):
                score += 3
        except Exception:
            pass
        try:
            rect = self.retangulo_controle_uia(control)
            ref_rect = self.referencia_relativa_uia(elemento_base, root)
            ref_left = int((ref_rect or {}).get("left", 0) or 0)
            ref_top = int((ref_rect or {}).get("top", 0) or 0)
            if rect and ref_left and ref_top:
                distancia = abs(int(rect["left"]) - ref_left) + abs(int(rect["top"]) - ref_top)
                score += max(0, 60 - min(60, distancia // 8))
        except Exception:
            pass
        if self.caminho_uia_parece(path_atual, elemento_base.get("path") or ""):
            score += 40
        return score

    def localizar_elemento_uia_por_selector(self, selector, elemento_base=None, timeout=5):
        auto = self.carregar_uia_runtime()
        consultas = self.consultas_selector_uia(selector, elemento_base)
        strategy = str((selector or {}).get("strategy") or "").strip()
        deadline = time.time() + max(0.5, float(timeout or 5))
        ultimo_erro = None
        while time.time() < deadline:
            roots = self.candidatos_raiz_uia(auto, elemento_base)
            for query in consultas:
                encontrados = []
                for root in roots:
                    root_type = self.propriedades_controle_uia(root).get("control_type") or "WindowControl"
                    fila = [(root, f"/{root_type}[0]")]
                    visitados = 0
                    while fila and visitados < 1200:
                        visitados += 1
                        control, path_atual = fila.pop(0)
                        try:
                            props = self.propriedades_controle_uia(control)
                            if all(str(props.get(k, "")) == str(v) for k, v in query.items()):
                                encontrados.append((control, root, path_atual))
                            filhos = control.GetChildren()
                            for index, filho in enumerate(filhos):
                                child_type = self.propriedades_controle_uia(filho).get("control_type") or "Control"
                                fila.append((filho, f"{path_atual}/{child_type}[{index}]"))
                        except Exception as erro:
                            ultimo_erro = erro
                if encontrados:
                    escolhidos = sorted(
                        encontrados,
                        key=lambda item: self.pontuar_controle_uia(item[0], elemento_base, root=item[1], path_atual=item[2]),
                        reverse=True,
                    )
                    return escolhidos[0][0]
            por_area = self.localizar_elemento_uia_por_area_salva(auto, elemento_base, roots=roots)
            if por_area:
                return por_area
            if not consultas:
                for root in roots:
                    try:
                        if self.propriedades_controle_uia(root).get("name"):
                            return root
                    except Exception as erro:
                        ultimo_erro = erro
            time.sleep(0.12)
        alvo = (elemento_base or {}).get("name") or strategy or "elemento"
        raise RuntimeError(f"Elemento UIA nao encontrado: {alvo}. Ultimo erro: {ultimo_erro}")

    def resolver_valor_acao_uia(self, detalhes, contexto_loop=None):
        acao_configurada = detalhes.get("acao_configurada") or {}
        origem = str(acao_configurada.get("origem_valor") or "texto_fixo").strip().lower()
        if origem == "planilha":
            arquivo = self.normalizar_id_planilha(acao_configurada.get("arquivo"))
            aba = str(acao_configurada.get("aba", "") or "").strip()
            coluna = str(acao_configurada.get("coluna", "") or "").strip()
            valor = self.obter_valor_coluna_linha_atual(arquivo, aba, coluna, contexto_loop)
            return "" if valor is None else str(valor)
        return self.substituir_variaveis_texto(str(acao_configurada.get("valor", "") or detalhes.get("valor", "") or ""), contexto_loop)

    def pressionar_tecla_depois_uia(self, tecla):
        tecla_normalizada = str(tecla or "").strip().lower()
        if not tecla_normalizada:
            return
        mapa = {
            "enter": "enter",
            "tab": "tab",
        }
        pyautogui.press(mapa.get(tecla_normalizada, tecla_normalizada))

    def invocar_controle_uia(self, control):
        tentativas = [
            lambda: control.GetInvokePattern().Invoke(),
            lambda: control.Invoke(),
            lambda: control.GetSelectionItemPattern().Select(),
        ]
        ultimo_erro = None
        for tentativa in tentativas:
            try:
                tentativa()
                return "interno"
            except Exception as erro:
                ultimo_erro = erro
        try:
            control.SetFocus()
            pyautogui.press("space")
            return "teclado"
        except Exception as erro:
            ultimo_erro = erro
        raise RuntimeError(f"Nao foi possivel acionar o elemento internamente: {ultimo_erro}")

    def ler_valor_uia(self, control):
        try:
            value_pattern = control.GetValuePattern()
            return str(getattr(value_pattern, "Value", "") or "")
        except Exception:
            return ""

    def valor_uia_confere(self, valor_atual, valor_esperado):
        atual = str(valor_atual or "").strip()
        esperado = str(valor_esperado or "").strip()
        if atual == esperado:
            return True
        normalizar = getattr(self, "normalizar_texto_ocr", None)
        if callable(normalizar):
            return normalizar(atual) == normalizar(esperado)
        return atual.lower() == esperado.lower()

    def colar_texto_rapido(self, texto):
        try:
            import pyperclip

            pyperclip.copy(str(texto))
            pyautogui.hotkey("ctrl", "v")
            return "area_transferencia"
        except Exception:
            pyautogui.write(str(texto), interval=0)
            return "teclado"

    def definir_texto_uia(self, control, valor, limpar_antes=True):
        valor_texto = str(valor or "")
        try:
            control.GetValuePattern().SetValue(valor_texto)
            time.sleep(0.05)
            if self.valor_uia_confere(self.ler_valor_uia(control), valor_texto):
                return "interno"
        except Exception:
            pass
        try:
            control.SetFocus()
        except Exception as erro:
            raise RuntimeError(
                "Nao foi possivel focar o campo sem mover o mouse. "
                "Capture o elemento editavel interno ou use um bloco visual para clicar antes de preencher."
            ) from erro
        time.sleep(0.08)
        if limpar_antes:
            pyautogui.hotkey("ctrl", "a")
            time.sleep(0.03)
            pyautogui.press("backspace")
            time.sleep(0.03)
        modo = self.colar_texto_rapido(valor_texto)
        time.sleep(0.05)
        if self.valor_uia_confere(self.ler_valor_uia(control), valor_texto):
            return modo
        return modo

    def estado_checkbox_uia(self, control):
        try:
            return int(control.GetTogglePattern().ToggleState)
        except Exception:
            return None

    def selecionar_opcao_uia(self, control, valor):
        valor_normalizado = self.normalizar_texto_ocr(valor)
        try:
            control.GetExpandCollapsePattern().Expand()
            time.sleep(0.15)
        except Exception:
            pass
        fila = []
        try:
            fila.extend(control.GetChildren())
        except Exception:
            fila = []
        while fila:
            filho = fila.pop(0)
            try:
                nome_filho = str(getattr(filho, "Name", "") or "")
                if self.normalizar_texto_ocr(nome_filho) == valor_normalizado:
                    try:
                        filho.GetSelectionItemPattern().Select()
                    except Exception:
                        self.invocar_controle_uia(filho)
                    return
                fila.extend(filho.GetChildren())
            except Exception:
                continue
        self.definir_texto_uia(control, valor, True)

    def executar_bloco_uia_elemento(self, evento, contexto_loop=None):
        detalhes = evento.get("detalhes", {}) or {}
        nome = str(detalhes.get("block_name") or "Mapear elemento").strip()
        selector = detalhes.get("selector") or {}
        elemento = detalhes.get("elemento") or {}
        acao_configurada = detalhes.get("acao_configurada") or {}
        acao = str(acao_configurada.get("tipo") or detalhes.get("configured_action") or detalhes.get("recommended_action") or "reference_only").strip().lower()
        ao_sucesso = str(detalhes.get("ao_sucesso", "continuar") or "continuar").strip().lower()
        bloco_sucesso = detalhes.get("bloco_sucesso", "")
        sequencia_sucesso = str(detalhes.get("sequencia_sucesso", "") or "").strip()
        atraso_sucesso = max(0.0, float(detalhes.get("atraso_antes_sequencia_sucesso", 0) or 0))
        ao_falhar = str(detalhes.get("ao_falhar", "parar") or "parar").strip().lower()
        bloco_falha = detalhes.get("bloco_falha", "")
        sequencia_falha = str(detalhes.get("sequencia_falha", "") or "").strip()
        atraso_falha = max(0.0, float(detalhes.get("atraso_antes_sequencia_falha", 0) or 0))

        def finalizar_sucesso():
            if ao_sucesso == "ir_para_sequencia" and sequencia_sucesso:
                return self.resultado_fluxo("ir_para_sequencia", sequencia=sequencia_sucesso, atraso=atraso_sucesso)
            if ao_sucesso == "ir_para_bloco" and str(bloco_sucesso or "").strip():
                return self.resultado_fluxo("ir_para_bloco", bloco=bloco_sucesso, atraso=atraso_sucesso)
            return self.resultado_fluxo("continuar")

        try:
            safe_print(f"[UIA: {nome}] Procurando elemento por seletor {selector.get('quality', 'FRAGIL')}", flush=True)
            control = self.localizar_elemento_uia_por_selector(selector, elemento, timeout=5)
            if acao in {"click", "select_tab", "select"}:
                modo_acionamento = self.invocar_controle_uia(control)
                safe_print(f"[UIA: {nome}] Clique executado por acionamento {modo_acionamento}, sem mover o mouse.", flush=True)
            elif acao in {"set_text", "set_or_select", "select_option"}:
                valor = self.resolver_valor_acao_uia(detalhes, contexto_loop)
                if acao == "select_option":
                    self.selecionar_opcao_uia(control, valor)
                else:
                    modo_preenchimento = self.definir_texto_uia(control, valor, bool(acao_configurada.get("limpar_antes", True)))
                    safe_print(f"[UIA: {nome}] Preenchimento aplicado por {modo_preenchimento}.", flush=True)
                self.pressionar_tecla_depois_uia(acao_configurada.get("pressionar_depois"))
                safe_print(f"[UIA: {nome}] Valor aplicado: {self.limpar_saida_status(valor)}", flush=True)
            elif acao == "toggle":
                control.GetTogglePattern().Toggle()
            elif acao in {"check", "uncheck"}:
                estado_atual = self.estado_checkbox_uia(control)
                deve_marcar = acao == "check"
                if estado_atual is None or bool(estado_atual) != deve_marcar:
                    control.GetTogglePattern().Toggle()
            elif acao == "expand":
                control.GetExpandCollapsePattern().Expand()
            elif acao in {"read_text", "read_grid", "reference_only"}:
                texto = str(getattr(control, "Name", "") or "")
                salvar_em = str(acao_configurada.get("salvar_em") or "").strip()
                if salvar_em:
                    self.contexto_variaveis[salvar_em] = texto
                safe_print(f"[UIA: {nome}] Referencia lida: {texto}", flush=True)
            else:
                safe_print(f"[UIA: {nome}] Elemento encontrado. Nenhuma acao executada.", flush=True)
            safe_print(f"[UIA: {nome}] Resultado: sucesso", flush=True)
            return finalizar_sucesso()
        except Exception as erro:
            safe_print(f"[UIA: {nome}] Resultado: erro | {erro}", flush=True)
            if ao_falhar == "continuar":
                return self.resultado_fluxo("continuar")
            if ao_falhar == "ir_para_sequencia" and sequencia_falha:
                return self.resultado_fluxo("ir_para_sequencia", sequencia=sequencia_falha, atraso=atraso_falha)
            if ao_falhar == "ir_para_bloco" and str(bloco_falha or "").strip():
                return self.resultado_fluxo("ir_para_bloco", bloco=bloco_falha, atraso=atraso_falha)
            return self.resultado_fluxo("parar")

    def obter_empresa_evento_planilha(self, evento, indice_item, linha_referencia=None):
        detalhes = evento.get("detalhes", {})
        aba = detalhes.get("aba")
        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        if not aba:
            return ""

        cache_aba = self.planilha_cache.get(arquivo, {}).get(aba, {})
        coluna_empresa = self.identificar_coluna_empresa(cache_aba.get("colunas", []))
        if not coluna_empresa:
            return ""

        registros = self.planilha_cache.get(arquivo, {}).get(aba, {}).get("linhas", [])
        registros_por_linha = {int(registro.get("__linha__", 0)): registro for registro in registros}

        try:
            linha_referencia_int = int(linha_referencia or 0)
        except (TypeError, ValueError):
            linha_referencia_int = 0
        if linha_referencia_int > 0:
            registro = registros_por_linha.get(linha_referencia_int)
            if registro:
                empresa = registro.get(coluna_empresa)
                if empresa is not None and str(empresa).strip():
                    return str(empresa).strip()

        linhas_evento = self.obter_linhas_evento_planilha(detalhes)

        if indice_item < len(linhas_evento):
            registro = registros_por_linha.get(int(linhas_evento[indice_item]))
            if registro:
                empresa = registro.get(coluna_empresa)
                if empresa is not None and str(empresa).strip():
                    return str(empresa).strip()

        valores = detalhes.get("valores", [])
        if detalhes.get("coluna") == coluna_empresa and indice_item < len(valores):
            return str(valores[indice_item]).strip()
        return ""

    def descrever_evento(self, evento, indice_item=None, total_itens=None):
        tipo = evento.get("tipo")
        ordem = evento.get("ordem", "?")

        if tipo == "inicio_loop":
            return f"Evento #{ordem}: inicio do loop"

        if tipo == "fim_loop":
            return f"Evento #{ordem}: fim do loop"

        if tipo == "texto_livre":
            texto = self.limpar_saida_status(evento.get("texto", ""))
            return f"Evento #{ordem}: texto {texto[:40]}"

        if tipo in TIPOS_SEQUENCIA_PLANILHA:
            detalhes = evento.get("detalhes", {})
            coluna = detalhes.get("coluna", "coluna")
            valores = self.obter_valores_evento_planilha(evento)
            total = total_itens if total_itens is not None else len(valores)
            if indice_item is not None and indice_item < len(valores):
                valor = self.limpar_saida_status(valores[indice_item])
                return f"Evento #{ordem}: {coluna} {indice_item + 1}/{total} ({valor})"
            return f"Evento #{ordem}: sequencia da coluna {coluna}"

        if tipo in TIPOS_ESCRITA_PLANILHA:
            detalhes = evento.get("detalhes", {})
            return f"Evento #{ordem}: escrever em {detalhes.get('coluna', 'coluna')}"

        if tipo in TIPOS_COMPARACAO_PLANILHA:
            detalhes = evento.get("detalhes", {})
            return f"Evento #{ordem}: conferir {detalhes.get('coluna_esperada', 'valor')} com OCR"

        if tipo in TIPOS_MOUSE:
            detalhes = evento.get("detalhes", {})
            cliques = max(1, int(detalhes.get("cliques", 1) or 1))
            return (
                f"Evento #{ordem}: mouse "
                f"{str(detalhes.get('botao', 'left')).upper()} "
                f"{'duplo ' if cliques == 2 else ''}"
                f"X={int(detalhes.get('x', 0))} Y={int(detalhes.get('y', 0))}"
            )

        if tipo in TIPOS_IMAGEM:
            detalhes = evento.get("detalhes", {})
            nome_imagem = Path(str(detalhes.get("imagem", ""))).name or "imagem"
            return f"Evento #{ordem}: imagem {nome_imagem}"

        if tipo in TIPOS_OCR:
            detalhes = evento.get("detalhes", {})
            texto = self.limpar_saida_status(detalhes.get("texto", "texto"))
            acao = self.limpar_saida_status(detalhes.get("acao", "clicar"))
            return f"Evento #{ordem}: procurar texto {texto[:42]} [{acao}]"

        if tipo in TIPOS_VERIFICACAO:
            detalhes = evento.get("detalhes", {})
            return f"Evento #{ordem}: {detalhes.get('nome', 'verificar estado do botao')}"

        if tipo in TIPOS_CHROME:
            detalhes = evento.get("detalhes", {})
            alvo = detalhes.get("alvo", {}) or {}
            acao = detalhes.get("acao", {}) or {}
            return (
                f"Evento #{ordem}: chrome "
                f"{self.limpar_saida_status(detalhes.get('nome', 'elemento html'))} "
                f"[{self.limpar_saida_status(alvo.get('metodo', 'css'))} -> "
                f"{self.limpar_saida_status(acao.get('tipo', 'clicar'))}]"
            )

        if tipo in TIPOS_UIA:
            detalhes = evento.get("detalhes", {})
            return (
                f"Evento #{ordem}: janela "
                f"{self.limpar_saida_status(detalhes.get('block_name', 'elemento'))} "
                f"[{self.limpar_saida_status(detalhes.get('recommended_action', 'reference_only'))}]"
            )

        return f"Evento #{ordem}: {self.limpar_saida_status(evento.get('texto', 'acao'))}"

    def descrever_iteracao_loop(self, indice, total_iteracoes, linha_referencia=None, empresa="", contexto="loop"):
        partes = [f"Iteracao {int(indice) + 1}/{int(total_iteracoes)}"]
        if linha_referencia not in (None, ""):
            partes.append(f"linha {int(linha_referencia)}")
        if str(empresa or "").strip():
            partes.append(f"empresa {self.limpar_saida_status(empresa)}")
        if contexto:
            partes.append(str(contexto))
        return " | ".join(partes)

    def emitir_preparacao_proxima_iteracao_loop(
        self,
        sequencia_loop,
        proximo_indice,
        total_iteracoes,
        linhas_loop,
        evento_planilha,
        contexto="loop",
    ):
        if int(proximo_indice) >= int(total_iteracoes):
            return

        linha_referencia = linhas_loop[proximo_indice] if proximo_indice < len(linhas_loop) else None
        empresa_atual = self.obter_empresa_evento_planilha(
            evento_planilha,
            proximo_indice,
            linha_referencia=linha_referencia,
        )
        descricao = self.descrever_iteracao_loop(
            proximo_indice,
            total_iteracoes,
            linha_referencia=linha_referencia,
            empresa=empresa_atual,
            contexto=contexto,
        )
        self.emitir_progresso(
            proximo_indice + 1,
            total_iteracoes,
            descricao,
            empresa_atual,
            sequencia=sequencia_loop,
            ordem=0,
            linha_atual=linha_referencia,
            iteracao_atual=proximo_indice + 1,
            iteracao_total=total_iteracoes,
        )

    def calcular_total_passos_normal(self):
        total = 0
        for evento in self.eventos:
            if evento.get("tipo") in TIPOS_MARCADORES_LOOP:
                continue
            if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA:
                total += len(self.obter_valores_evento_planilha(evento))
            else:
                total += 1
        return max(1, total)

    def resultado_fluxo(self, acao="continuar", sequencia="", bloco="", atraso=0):
        return {
            "acao": str(acao or "continuar"),
            "sequencia": self.normalizar_nome_sequencia(sequencia) if sequencia else "",
            "bloco": int(bloco) if str(bloco or "").strip() else "",
            "atraso": max(0.0, float(atraso or 0)),
        }

    def interpretar_config_acao(self, config, padrao="continuar"):
        if isinstance(config, str):
            return self.resultado_fluxo(config)
        if isinstance(config, dict):
            return self.resultado_fluxo(
                config.get("acao", padrao),
                config.get("sequencia", ""),
                config.get("bloco", ""),
                config.get("atraso", 0),
            )
        return self.resultado_fluxo(padrao)

    def resultado_sucesso_evento(self, evento, prefixo_log=""):
        if self.cancelamento_solicitado():
            return self.resultado_fluxo("parar")
        detalhes = evento.get("detalhes") or {}
        acao = str(detalhes.get("ao_sucesso", "continuar") or "continuar").strip().lower()
        bloco = detalhes.get("bloco_sucesso", "")
        sequencia = str(detalhes.get("sequencia_sucesso", "") or "").strip()
        atraso_sequencia = max(0.0, float(detalhes.get("atraso_antes_sequencia_sucesso", 0) or 0))
        if acao == "ir_para_sequencia" and sequencia:
            if prefixo_log:
                safe_print(f"{prefixo_log} Sucesso. Redirecionando para pagina {sequencia}.", flush=True)
            else:
                ordem = int(evento.get("ordem", 0) or 0)
                safe_print(f"Bloco {ordem}: concluido com sucesso. Redirecionando para pagina {sequencia}.", flush=True)
            return self.resultado_fluxo("ir_para_sequencia", sequencia=sequencia, atraso=atraso_sequencia)
        if acao == "ir_para_bloco" and str(bloco or "").strip():
            if prefixo_log:
                safe_print(f"{prefixo_log} Sucesso. Redirecionando para bloco {bloco}.", flush=True)
            else:
                ordem = int(evento.get("ordem", 0) or 0)
                safe_print(f"Bloco {ordem}: concluido com sucesso. Redirecionando para bloco {bloco}.", flush=True)
            return self.resultado_fluxo("ir_para_bloco", bloco=bloco, atraso=atraso_sequencia)
        return self.resultado_fluxo("continuar")

    def localizar_bloco(self, ordem_destino):
        for nome_sequencia, eventos in self.sequencias.items():
            for indice, evento in enumerate(eventos):
                if int(evento.get("ordem", 0) or 0) == int(ordem_destino):
                    return nome_sequencia, indice
        raise ValueError(f"Bloco nao encontrado: {ordem_destino}")

    def obter_valor_coluna_linha_atual(self, arquivo, aba, coluna, contexto_loop=None):
        linha_atual = self.obter_linha_atual_contexto(contexto_loop)
        registros = self.obter_registros_planilha(arquivo, aba)
        if linha_atual is not None:
            for registro in registros:
                if int(registro.get("__linha__", 0) or 0) == int(linha_atual):
                    return registro.get(coluna)
        indice = int((contexto_loop or {}).get("indice", 0) or 0) if contexto_loop else 0
        if 0 <= indice < len(registros):
            return registros[indice].get(coluna)
        return None

    def escrever_celula_planilha(self, arquivo, aba, coluna, linha, valor):
        caminho_planilha, _workbook_id = self.resolver_planilha(arquivo)
        if not caminho_planilha.exists():
            raise FileNotFoundError(f"Planilha nao encontrada: {caminho_planilha}")
        workbook = load_workbook(caminho_planilha)
        try:
            if aba not in workbook.sheetnames:
                raise ValueError(f"Aba nao encontrada na planilha: {aba}")
            worksheet = workbook[aba]
            cabecalhos = [
                str(celula.value).strip() if celula.value is not None else ""
                for celula in worksheet[1]
            ]
            if coluna not in cabecalhos:
                cabecalhos.append(coluna)
                worksheet.cell(row=1, column=len(cabecalhos), value=coluna)
            indice_coluna = cabecalhos.index(coluna) + 1
            worksheet.cell(row=int(linha), column=indice_coluna, value=valor)
            workbook.save(caminho_planilha)
        finally:
            workbook.close()

    def resolver_texto_ocr_comparacao(self, source="ultimo_ocr"):
        origem = str(source or "ultimo_ocr").strip().lower()
        if origem.startswith("ocr_bloco_"):
            valor = self.contexto_variaveis.get(origem, "")
            if valor:
                return valor
            numero = origem.replace("ocr_bloco_", "", 1)
            return self.contexto_variaveis.get(f"bloco_{numero}_ocr", "")
        if origem.isdigit():
            return self.contexto_variaveis.get(f"ocr_bloco_{origem}", "") or self.contexto_variaveis.get(f"bloco_{origem}_ocr", "")
        return self.contexto_variaveis.get("ultimo_ocr", "") or self.contexto_variaveis.get("ocr_tela", "")

    def normalizar_valor_comparacao(self, valor, modo="texto"):
        texto = str(valor if valor is not None else "").strip()
        modo = str(modo or "texto").strip().lower()
        if modo == "numero":
            texto = texto.replace(".", "").replace(",", ".")
            match = re.search(r"-?\d+(?:\.\d+)?", texto)
            return str(float(match.group(0))) if match else ""
        if modo == "data":
            return re.sub(r"\D", "", texto)
        return self.normalizar_texto_ocr(texto)

    def executar_bloco_escrever_planilha(self, evento, contexto_loop=None):
        detalhes = evento.get("detalhes") or {}
        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        aba = str(detalhes.get("aba", "") or "").strip()
        coluna = str(detalhes.get("coluna", "") or "").strip()
        valor = self.substituir_variaveis_texto(str(detalhes.get("valor", "") or ""), contexto_loop)
        linha = self.obter_linha_atual_contexto(contexto_loop)
        if not aba or not coluna:
            raise RuntimeError("Bloco Escrever celula sem aba ou coluna configurada.")
        if linha is None:
            raise RuntimeError("Bloco Escrever celula precisa estar dentro de um loop de planilha.")
        self.escrever_celula_planilha(arquivo, aba, coluna, linha, valor)
        safe_print(f"Planilha atualizada: {Path(arquivo).name} | {aba}!{coluna}{linha} = {valor}", flush=True)
        return self.resultado_sucesso_evento(evento)

    def executar_bloco_comparar_planilha(self, evento, contexto_loop=None):
        detalhes = evento.get("detalhes") or {}
        arquivo = self.normalizar_id_planilha(detalhes.get("arquivo"))
        aba = str(detalhes.get("aba", "") or "").strip()
        coluna = str(detalhes.get("coluna_esperada", detalhes.get("coluna", "")) or "").strip()
        modo = str(detalhes.get("modo_comparacao", "texto") or "texto")
        esperado = self.obter_valor_coluna_linha_atual(arquivo, aba, coluna, contexto_loop)
        capturado = self.resolver_texto_ocr_comparacao(detalhes.get("origem_tela", "ultimo_ocr"))
        esperado_norm = self.normalizar_valor_comparacao(esperado, modo)
        capturado_norm = self.normalizar_valor_comparacao(capturado, modo)
        bateu = bool(esperado_norm) and esperado_norm == capturado_norm
        safe_print(
            f"Conferir valor: esperado='{esperado}' | OCR='{capturado}' | resultado={'OK' if bateu else 'DIVERGENTE'}",
            flush=True,
        )
        if bateu:
            return self.resultado_sucesso_evento(evento)
        return self.resultado_fluxo(
            detalhes.get("ao_divergir", "parar"),
            sequencia=detalhes.get("sequencia_divergir", ""),
            bloco=detalhes.get("bloco_divergir", ""),
            atraso=detalhes.get("atraso_antes_sequencia_divergir", 0),
        )

    def similaridade_texto(self, esperado, capturado):
        esperado_norm = self.normalizar_texto_ocr(esperado)
        capturado_norm = self.normalizar_texto_ocr(capturado)
        if not esperado_norm or not capturado_norm:
            return 0.0
        if esperado_norm == capturado_norm:
            return 1.0
        if esperado_norm in capturado_norm:
            return min(0.98, 0.82 + (len(esperado_norm) / max(len(capturado_norm), 1)) * 0.16)
        esperado_compacto = re.sub(r"[^a-z0-9]+", "", esperado_norm)
        capturado_compacto = re.sub(r"[^a-z0-9]+", "", capturado_norm)
        scores = [float(difflib.SequenceMatcher(None, esperado_norm, capturado_norm).ratio())]
        if esperado_compacto and capturado_compacto:
            if esperado_compacto == capturado_compacto:
                return 0.995
            if esperado_compacto in capturado_compacto:
                return min(0.985, 0.84 + (len(esperado_compacto) / max(len(capturado_compacto), 1)) * 0.14)
            scores.append(float(difflib.SequenceMatcher(None, esperado_compacto, capturado_compacto).ratio()))
        palavras_esperadas = esperado_norm.split()
        palavras_capturadas = capturado_norm.split()
        if len(palavras_esperadas) >= 2 and len(palavras_capturadas) >= 2:
            tamanho = len(palavras_esperadas)
            melhor_janela = 0.0
            for inicio in range(0, max(1, len(palavras_capturadas) - tamanho + 1)):
                janela = palavras_capturadas[inicio : inicio + tamanho]
                if len(janela) != tamanho:
                    continue
                palavra_scores = [
                    float(difflib.SequenceMatcher(None, ref, cap).ratio())
                    for ref, cap in zip(palavras_esperadas, janela)
                ]
                media_palavras = sum(palavra_scores) / max(1, len(palavra_scores))
                primeira_ref = palavras_esperadas[0]
                primeira_cap = janela[0]
                primeira_score = palavra_scores[0]
                if (
                    primeira_score < 0.67
                    or (
                        len(primeira_ref) <= 4
                        and primeira_ref[:1] != primeira_cap[:1]
                        and primeira_score < 0.82
                    )
                ):
                    media_palavras *= 0.55
                if any(score < 0.62 for score in palavra_scores):
                    media_palavras *= 0.65
                melhor_janela = max(melhor_janela, media_palavras)
            if melhor_janela < 0.70:
                return min(max(scores), 0.64)
            return max(max(scores), min(0.98, melhor_janela))
        return max(scores)

    def interpretar_texto_ocr_composto(self, texto):
        bruto = str(texto or "").strip()
        if not bruto:
            return None
        for separador in ("=>", "->", "→", "⇢"):
            if separador in bruto:
                partes = [parte.strip() for parte in bruto.split(separador) if parte.strip()]
                if len(partes) >= 2:
                    return {"partes": partes, "relacional": True, "separador": separador}

        linhas = [linha.strip() for linha in re.split(r"[\r\n]+", bruto) if linha.strip()]
        if len(linhas) >= 2:
            segunda_tem_numero = any(re.search(r"\d", linha) for linha in linhas[1:])
            return {"partes": linhas, "relacional": bool(segunda_tem_numero), "separador": "linha"}

        partes_multiplas = [parte.strip() for parte in re.split(r"\s*(?:\||;|\+)\s*", bruto) if parte.strip()]
        if len(partes_multiplas) >= 2:
            return {"partes": partes_multiplas, "relacional": False, "separador": "multi"}
        return None

    def _retangulo_texto_ocr(self, candidato):
        if not isinstance(candidato, dict):
            return None
        try:
            left = int(candidato.get("left", 0) or 0)
            top = int(candidato.get("top", 0) or 0)
            width = max(1, int(candidato.get("width", 0) or 0))
            height = max(1, int(candidato.get("height", 0) or 0))
            return left, top, width, height
        except Exception:
            return None

    def _unir_retangulos_ocr(self, candidatos):
        retangulos = [self._retangulo_texto_ocr(candidato) for candidato in candidatos]
        retangulos = [retangulo for retangulo in retangulos if retangulo]
        if not retangulos:
            return None
        left = min(retangulo[0] for retangulo in retangulos)
        top = min(retangulo[1] for retangulo in retangulos)
        right = max(retangulo[0] + retangulo[2] for retangulo in retangulos)
        bottom = max(retangulo[1] + retangulo[3] for retangulo in retangulos)
        return left, top, max(1, right - left), max(1, bottom - top)

    def pontuar_relacao_ocr_vertical(self, ancora, valor):
        rect_ancora = self._retangulo_texto_ocr(ancora)
        rect_valor = self._retangulo_texto_ocr(valor)
        if not rect_ancora or not rect_valor:
            return 0.0
        ax, ay, aw, ah = rect_ancora
        vx, vy, vw, vh = rect_valor
        acx = ax + (aw / 2.0)
        acy = ay + (ah / 2.0)
        vcx = vx + (vw / 2.0)
        vcy = vy + (vh / 2.0)
        if vcy <= acy:
            return 0.0
        distancia_y = max(0.0, vy - (ay + ah))
        limite_y = max(42.0, min(280.0, ah * 12.0 + 90.0))
        if distancia_y > limite_y:
            return 0.0
        sobreposicao_x = max(0.0, min(ax + aw, vx + vw) - max(ax, vx))
        menor_largura = max(1.0, float(min(aw, vw)))
        taxa_sobreposicao = sobreposicao_x / menor_largura
        centro_diferenca = abs(vcx - acx)
        tolerancia_centro = max(70.0, aw * 0.62)
        dentro_coluna = (ax - 24) <= vcx <= (ax + aw + 24)
        score_x = 1.0 if dentro_coluna else max(0.0, 1.0 - (centro_diferenca / tolerancia_centro))
        score_x = max(score_x, min(1.0, taxa_sobreposicao))
        if score_x < 0.18:
            return 0.0
        score_y = max(0.0, 1.0 - (distancia_y / limite_y))
        return max(0.0, min(1.0, (score_x * 0.66) + (score_y * 0.34)))

    def montar_layouts_alvos_ocr(self, leitura_relativa=None, dimensoes_referencia=None, quantidade=0):
        config = leitura_relativa if isinstance(leitura_relativa, dict) else {}
        layouts = config.get("ocr_target_layouts") if isinstance(config.get("ocr_target_layouts"), list) else []
        ref_w, ref_h = dimensoes_referencia or (0, 0)
        alvos = []
        total = max(int(quantidade or 0), len(layouts))
        for index in range(total):
            layout = layouts[index] if index < len(layouts) and isinstance(layouts[index], dict) else {}
            if index == 0:
                width = config.get("largura", layout.get("width", 140))
                height = config.get("altura", layout.get("height", 40))
                offset_x = config.get("offset_x", layout.get("offsetX", 0))
                offset_y = config.get("offset_y", layout.get("offsetY", 0))
            else:
                width = layout.get("width", 140)
                height = layout.get("height", 40)
                offset_x = layout.get("offsetX")
                offset_y = layout.get("offsetY")
                if offset_x is None:
                    offset_x = config.get("offset_x", 0) + index * 0.032 * max(1, ref_w)
                if offset_y is None:
                    offset_y = config.get("offset_y", 0) + index * 0.072 * max(1, ref_h)
            alvos.append({
                "offset_x": float(offset_x or 0),
                "offset_y": float(offset_y or 0),
                "width": float(width or 140),
                "height": float(height or 40),
            })
        return alvos

    def pontuar_relacao_ocr_por_layout(self, ancora, valor, layout_ancora=None, layout_valor=None):
        rect_ancora = self._retangulo_texto_ocr(ancora)
        rect_valor = self._retangulo_texto_ocr(valor)
        if not rect_ancora or not rect_valor or not layout_ancora or not layout_valor:
            return 0.0
        ax, ay, aw, ah = rect_ancora
        vx, vy, vw, vh = rect_valor
        dx_detectado = (vx + vw / 2.0) - (ax + aw / 2.0)
        dy_detectado = (vy + vh / 2.0) - (ay + ah / 2.0)
        dx_esperado = float(layout_valor.get("offset_x", 0.0) or 0.0) - float(layout_ancora.get("offset_x", 0.0) or 0.0)
        dy_esperado = float(layout_valor.get("offset_y", 0.0) or 0.0) - float(layout_ancora.get("offset_y", 0.0) or 0.0)
        if abs(dx_esperado) < 1 and abs(dy_esperado) < 1:
            return 0.0

        escalas = [1.0]
        if abs(dx_esperado) >= 6:
            escalas.append(max(0.45, min(2.4, dx_detectado / dx_esperado)))
        if abs(dy_esperado) >= 6:
            escalas.append(max(0.45, min(2.4, dy_detectado / dy_esperado)))

        melhor = 0.0
        tamanho_ref = max(
            float(layout_ancora.get("width", 0.0) or 0.0),
            float(layout_ancora.get("height", 0.0) or 0.0),
            float(layout_valor.get("width", 0.0) or 0.0),
            float(layout_valor.get("height", 0.0) or 0.0),
            24.0,
        )
        for escala in escalas:
            erro_x = abs(dx_detectado - dx_esperado * escala)
            erro_y = abs(dy_detectado - dy_esperado * escala)
            tolerancia_x = max(46.0, tamanho_ref * max(0.72, escala) * 1.35)
            tolerancia_y = max(32.0, tamanho_ref * max(0.72, escala) * 1.05)
            score_x = max(0.0, 1.0 - (erro_x / tolerancia_x))
            score_y = max(0.0, 1.0 - (erro_y / tolerancia_y))
            melhor = max(melhor, (score_x * 0.58) + (score_y * 0.42))
        return max(0.0, min(1.0, melhor))

    def localizar_texto_por_layout_relativo_ocr(self, ancora, texto, layout_ancora=None, layout_valor=None, region=None):
        rect_ancora = self._retangulo_texto_ocr(ancora)
        if not rect_ancora or not layout_ancora or not layout_valor:
            return None
        texto_esperado = str(texto or "").strip()
        if not texto_esperado:
            return None
        ax, ay, aw, ah = rect_ancora
        layout_anchor_w = max(1.0, float(layout_ancora.get("width", 0) or 1))
        layout_anchor_h = max(1.0, float(layout_ancora.get("height", 0) or 1))
        escala_x = max(0.45, min(2.5, aw / layout_anchor_w))
        escala_y = max(0.45, min(2.5, ah / layout_anchor_h))
        dx = (float(layout_valor.get("offset_x", 0) or 0) - float(layout_ancora.get("offset_x", 0) or 0)) * escala_x
        dy = (float(layout_valor.get("offset_y", 0) or 0) - float(layout_ancora.get("offset_y", 0) or 0)) * escala_y
        largura = max(18, int(round(float(layout_valor.get("width", 40) or 40) * escala_x)))
        altura = max(14, int(round(float(layout_valor.get("height", 20) or 20) * escala_y)))
        centro_x = ax + aw / 2.0 + dx
        centro_y = ay + ah / 2.0 + dy
        margem_x = max(8, int(largura * 0.35))
        margem_y = max(6, int(altura * 0.45))
        regiao = self.limitar_regiao(
            (
                int(round(centro_x - largura / 2.0 - margem_x)),
                int(round(centro_y - altura / 2.0 - margem_y)),
                int(round(largura + margem_x * 2)),
                int(round(altura + margem_y * 2)),
            ),
            limite=region,
        )
        if not regiao:
            return None
        try:
            captura, _base_left, _base_top = self.capturar_tela_pil(region=regiao)
        except Exception:
            return None
        textos = self.extrair_textos_ocr_rapido(captura, texto_esperado=texto_esperado)
        if not textos:
            textos = [
                self.extrair_texto_ocr_imagem_pil(captura, config="--psm 6"),
                self.extrair_texto_ocr_imagem_pil(captura, config="--psm 7"),
            ]
        capturado = " ".join(str(item or "").strip() for item in textos if str(item or "").strip()).strip()
        comparacao = self.comparar_texto_reconhecimento_preciso(texto_esperado, capturado)
        confirmado = self.texto_ocr_confirma_alvo(texto_esperado, capturado, limite_ratio=0.86)
        if not confirmado:
            return None
        return {
            "left": int(regiao[0]),
            "top": int(regiao[1]),
            "width": int(regiao[2]),
            "height": int(regiao[3]),
            "text": capturado or texto_esperado,
            "texto": capturado or texto_esperado,
            "score": max(0.72, float(comparacao.get("ratio", 0.0) or 0.0)),
            "source": "ocr_layout_relativo",
        }

    def montar_candidato_ocr_composto(self, partes, candidatos_escolhidos, relacional=False, score_layout=1.0):
        if not candidatos_escolhidos:
            return None
        alvo_clique = candidatos_escolhidos[1] if relacional and len(candidatos_escolhidos) >= 2 else candidatos_escolhidos[0]
        rect = self._retangulo_texto_ocr(alvo_clique) or self._unir_retangulos_ocr(candidatos_escolhidos)
        uniao = self._unir_retangulos_ocr(candidatos_escolhidos)
        if not rect:
            return None
        left, top, width, height = rect
        scores = [float(candidato.get("score", 0.0) or 0.0) for candidato in candidatos_escolhidos]
        score_texto = sum(scores) / max(1, len(scores))
        score_total = (score_texto * 0.78) + (float(score_layout or 0.0) * 0.22)
        textos_lidos = [str(candidato.get("text") or candidato.get("texto") or "").strip() for candidato in candidatos_escolhidos]
        bounds = []
        for indice, candidato in enumerate(candidatos_escolhidos):
            r = self._retangulo_texto_ocr(candidato)
            if not r:
                continue
            bounds.append({
                "parte": str(partes[indice] if indice < len(partes) else "").strip(),
                "texto": textos_lidos[indice] if indice < len(textos_lidos) else "",
                "left": int(r[0]),
                "top": int(r[1]),
                "width": int(r[2]),
                "height": int(r[3]),
            })
        candidato = {
            "left": int(left),
            "top": int(top),
            "width": int(width),
            "height": int(height),
            "text": " | ".join(texto for texto in textos_lidos if texto) or " | ".join(partes),
            "score": max(0.0, min(1.0, score_total)),
            "source": "ocr_relacional" if relacional else "ocr_multiplo",
            "ocr_composto": True,
            "ocr_relacional": bool(relacional),
            "ocr_partes": list(partes),
            "ocr_textos": textos_lidos,
            "ocr_related_bounds": bounds,
        }
        if uniao:
            candidato["ocr_union_bounds"] = {
                "left": int(uniao[0]),
                "top": int(uniao[1]),
                "width": int(uniao[2]),
                "height": int(uniao[3]),
            }
        if relacional and len(bounds) >= 2:
            candidato["ocr_anchor_bounds"] = bounds[0]
            candidato["ocr_value_bounds"] = bounds[1]
            candidato["ocr_relation_score"] = float(score_layout or 0.0)
        return candidato

    def localizar_texto_composto_por_ocr_screenshot(
        self,
        texto_original,
        spec,
        confianca=0.72,
        janela_alvo=None,
        region=None,
        retornar_candidatos=False,
        marcar_candidatos=False,
        limite_candidatos=8,
        layout_alvos=None,
        _escala_ocr=None,
        _tentativa_ampliada=False,
    ):
        partes = [str(parte or "").strip() for parte in (spec or {}).get("partes", []) if str(parte or "").strip()]
        if len(partes) < 2:
            return [] if retornar_candidatos else None
        relacional = bool((spec or {}).get("relacional"))
        conjuntos = []
        limite_por_parte = 5 if relacional else 3
        for indice_parte, parte in enumerate(partes):
            encontrados = self.localizar_texto_por_ocr_screenshot(
                parte,
                confianca=max(0.50, min(float(confianca or 0.72), 0.80)),
                janela_alvo=janela_alvo,
                region=region,
                retornar_candidatos=True,
                marcar_candidatos=False,
                limite_candidatos=limite_por_parte,
                _permitir_composto=False,
                _escala_ocr=_escala_ocr,
                _tentativa_ampliada=_tentativa_ampliada,
            )
            encontrados = [
                candidato
                for candidato in (encontrados or [])
                if self.texto_ocr_confirma_alvo(
                    parte,
                    candidato.get("texto") or candidato.get("text") or "",
                    limite_ratio=0.86,
                )
            ]
            if not encontrados:
                if not (layout_alvos and indice_parte > 0):
                    safe_print(f"OCR composto: parte obrigatoria nao encontrada: '{parte}'.", flush=True)
                    return [] if retornar_candidatos else None
                safe_print(
                    f"OCR composto: parte '{parte}' nao apareceu na busca geral; tentando area roxa do layout.",
                    flush=True,
                )
                encontrados = []
            conjuntos.append(encontrados)

        candidatos_compostos = []
        if relacional:
            for ancora in conjuntos[0][:5]:
                valores = conjuntos[1][:5]
                if not valores and layout_alvos and len(layout_alvos) >= 2:
                    valor_layout = self.localizar_texto_por_layout_relativo_ocr(
                        ancora,
                        partes[1],
                        layout_alvos[0],
                        layout_alvos[1],
                        region=region,
                    )
                    valores = [valor_layout] if valor_layout else []
                for valor in valores:
                    score_vertical = self.pontuar_relacao_ocr_vertical(ancora, valor)
                    score_por_layout = 0.0
                    if layout_alvos and len(layout_alvos) >= 2:
                        score_por_layout = self.pontuar_relacao_ocr_por_layout(ancora, valor, layout_alvos[0], layout_alvos[1])
                    score_layout = max(score_vertical, score_por_layout)
                    if score_layout <= 0.0:
                        continue
                    escolhidos = [ancora, valor]
                    valido = True
                    for extra_indice, extras in enumerate(conjuntos[2:], start=2):
                        melhor_extra = None
                        melhor_extra_score = -1.0
                        if not extras and layout_alvos and len(layout_alvos) > extra_indice:
                            extra_layout = self.localizar_texto_por_layout_relativo_ocr(
                                ancora,
                                partes[extra_indice],
                                layout_alvos[0],
                                layout_alvos[extra_indice],
                                region=region,
                            )
                            extras = [extra_layout] if extra_layout else []
                        for extra in extras[:5]:
                            score_extra = 1.0
                            if layout_alvos and len(layout_alvos) > extra_indice:
                                score_extra = self.pontuar_relacao_ocr_por_layout(ancora, extra, layout_alvos[0], layout_alvos[extra_indice])
                            if score_extra > melhor_extra_score:
                                melhor_extra_score = score_extra
                                melhor_extra = extra
                        if melhor_extra is None or melhor_extra_score <= 0.0:
                            valido = False
                            break
                        escolhidos.append(melhor_extra)
                    if not valido:
                        continue
                    candidato = self.montar_candidato_ocr_composto(partes, escolhidos, relacional=True, score_layout=score_layout)
                    if candidato:
                        candidatos_compostos.append(candidato)
        else:
            if layout_alvos and len(layout_alvos) >= len(partes):
                # Mesmo quando os textos nao formam uma relacao "texto => valor",
                # eles pertencem ao mesmo recorte. Escolher o primeiro candidato de
                # cada lista podia misturar textos de lugares diferentes da tela.
                for ancora in conjuntos[0][:5]:
                    escolhidos = [ancora]
                    pontuacoes_layout = []
                    valido = True
                    for extra_indice, extras_originais in enumerate(conjuntos[1:], start=1):
                        extras = [
                            extra
                            for extra in extras_originais[:5]
                            if self.texto_ocr_confirma_alvo(
                                partes[extra_indice],
                                extra.get("texto") or extra.get("text") or "",
                                limite_ratio=0.86,
                            )
                        ]
                        if not extras:
                            extra_layout = self.localizar_texto_por_layout_relativo_ocr(
                                ancora,
                                partes[extra_indice],
                                layout_alvos[0],
                                layout_alvos[extra_indice],
                                region=region,
                            )
                            extras = [extra_layout] if extra_layout else []
                        melhor_extra = None
                        melhor_score = -1.0
                        for extra in extras:
                            score_extra = self.pontuar_relacao_ocr_por_layout(
                                ancora,
                                extra,
                                layout_alvos[0],
                                layout_alvos[extra_indice],
                            )
                            if score_extra > melhor_score:
                                melhor_score = score_extra
                                melhor_extra = extra
                        if melhor_extra is None or melhor_score < 0.42:
                            valido = False
                            break
                        escolhidos.append(melhor_extra)
                        pontuacoes_layout.append(melhor_score)
                    if not valido or len(escolhidos) != len(partes):
                        continue
                    score_layout = sum(pontuacoes_layout) / max(1, len(pontuacoes_layout))
                    candidato = self.montar_candidato_ocr_composto(
                        partes,
                        escolhidos,
                        relacional=False,
                        score_layout=score_layout,
                    )
                    if candidato:
                        candidatos_compostos.append(candidato)
            else:
                escolhidos = [grupo[0] for grupo in conjuntos if grupo]
                if len(escolhidos) == len(partes):
                    candidato = self.montar_candidato_ocr_composto(partes, escolhidos, relacional=False, score_layout=1.0)
                    if candidato:
                        candidatos_compostos.append(candidato)

        if not candidatos_compostos:
            modo = "relacao" if relacional else "multiplos textos"
            safe_print(f"OCR composto: encontrou as partes de '{texto_original}', mas nao validou {modo}.", flush=True)
            return [] if retornar_candidatos else None

        candidatos_compostos.sort(key=lambda item: float(item.get("score", 0.0) or 0.0), reverse=True)
        candidatos_compostos = candidatos_compostos[: max(1, int(limite_candidatos or 8))]
        if marcar_candidatos:
            for indice, candidato in enumerate(candidatos_compostos[:4], start=1):
                bounds = candidato.get("ocr_related_bounds") if isinstance(candidato.get("ocr_related_bounds"), list) else []
                if bounds:
                    for parte_indice, bound in enumerate(bounds[:4], start=1):
                        self.marcar_area_reconhecida(
                            bound,
                            duracao=0.18,
                            bloqueante=False,
                            tema=self.tema_alvo_ocr(parte_indice - 1),
                            etiqueta=f"OCR {indice}.{parte_indice}",
                            margem=3,
                        )
                else:
                    self.marcar_area_reconhecida(
                        candidato,
                        duracao=0.18,
                        bloqueante=False,
                        tema="ocr_busca",
                        etiqueta=f"OCR {indice}",
                        margem=3,
                    )
        safe_print(
            f"OCR composto confirmado ({'relacional' if relacional else 'multiplo'}): "
            f"{', '.join(partes)}.",
            flush=True,
        )
        if retornar_candidatos:
            return candidatos_compostos
        return candidatos_compostos[0]

    def localizar_texto_por_uia(self, texto, confianca=0.72, janela_alvo=None, limite=650):
        texto_alvo = self.normalizar_texto_ocr(texto)
        if not texto_alvo:
            return None
        try:
            auto = self.carregar_uia_runtime()
        except Exception:
            return None

        try:
            self.preparar_janela_para_acao(janela_alvo)
        except Exception:
            pass

        roots = []
        try:
            if self.alvo_execucao_ativo():
                hwnd = int((self.contexto_execucao_alvo or {}).get("target_hwnd", 0) or 0)
                if hwnd:
                    roots.append(auto.ControlFromHandle(hwnd))
            elif janela_alvo:
                janela = self.localizar_janela_alvo(janela_alvo)
                if janela and int(janela.get("hwnd", 0) or 0):
                    roots.append(auto.ControlFromHandle(int(janela.get("hwnd"))))
        except Exception:
            pass
        if not roots:
            try:
                roots.append(auto.GetRootControl())
            except Exception:
                return None

        melhor = None
        visitados = 0
        fila = list(roots)
        while fila and visitados < max(80, int(limite or 650)):
            if self.cancelamento_solicitado():
                return None
            control = fila.pop(0)
            visitados += 1
            try:
                props = self.propriedades_controle_uia(control)
                nome = str(props.get("name") or "").strip()
                score = self.similaridade_texto(texto_alvo, nome)
                rect = self.retangulo_controle_uia(control)
                if nome and rect and score >= float(confianca):
                    candidato = {
                        "left": int(rect["left"]),
                        "top": int(rect["top"]),
                        "width": max(12, int(rect["right"] - rect["left"])),
                        "height": max(10, int(rect["bottom"] - rect["top"])),
                        "text": nome,
                        "score": score,
                        "source": "uia",
                        "control": control,
                    }
                    if melhor is None or candidato["score"] > melhor["score"]:
                        melhor = candidato
                        if score >= 0.98:
                            break
                for child in control.GetChildren():
                    fila.append(child)
            except Exception:
                continue
        return melhor

    def capturar_itens_ocr_screenshot(
        self,
        janela_alvo=None,
        region=None,
        retornar_candidatos=False,
        escala_ocr=None,
    ):
        """Lê uma captura uma vez e devolve as palavras reconhecidas com posição.

        Um bloco hibrido com vários OCRs chama a busca de cada texto em sequência.
        Reexecutar o Tesseract para todos eles era a maior fonte de demora. A
        captura fica válida somente por um intervalo curto, apenas para os outros
        alvos da mesma tentativa.
        """
        try:
            regiao_ocr = region or self.obter_regiao_busca(janela_alvo)
            chave_regiao = (
                tuple(int(round(float(valor))) for valor in regiao_ocr)
                if regiao_ocr is not None
                else ("tela",)
            )
        except Exception:
            return None

        modo_teste = bool(getattr(self, "modo_teste_reconhecimento", False))
        try:
            escala_cache = round(float(escala_ocr), 2) if escala_ocr is not None else None
        except (TypeError, ValueError):
            escala_cache = None
        chave_cache = (chave_regiao, bool(retornar_candidatos), modo_teste, escala_cache)
        cache_itens = getattr(self, "_cache_ocr_screenshot", None)
        cache = cache_itens.get(chave_cache) if isinstance(cache_itens, dict) else None
        if cache and time.perf_counter() - float(cache.get("em", 0.0) or 0.0) <= 1.25:
            captura_cache = cache.get("captura")
            if captura_cache is not None:
                return (
                    captura_cache,
                    int(cache.get("base_left", 0) or 0),
                    int(cache.get("base_top", 0) or 0),
                    [dict(item) for item in (cache.get("itens") or []) if isinstance(item, dict)],
                    regiao_ocr,
                )

        try:
            if regiao_ocr:
                safe_print(f"OCR capturando regiao: {self.formatar_regiao_log(regiao_ocr)}", flush=True)
            captura, base_left, base_top = self.capturar_tela_pil(region=regiao_ocr)
            safe_print(
                f"OCR base da captura: x={int(base_left)} y={int(base_top)} "
                f"w={int(getattr(captura, 'width', 0) or 0)} h={int(getattr(captura, 'height', 0) or 0)}",
                flush=True,
            )
        except Exception:
            return None

        itens = []
        try:
            area_captura = max(1, int(captura.width or 0) * int(captura.height or 0))
            captura_grande = area_captura >= 900000
            if retornar_candidatos:
                # O caminho hibrido usa esta leitura unica para localizar todos
                # os textos do bloco. Ampliar um dialogo inteiro para 155% fazia
                # o Tesseract gastar 8--10s antes de devolver a primeira palavra.
                # Formularios desktop ja sao capturados com texto legivel; ler a
                # resolucao nativa preserva as coordenadas e reduz muito a area
                # que o motor precisa analisar. As rotas de OCR de recorte ainda
                # usam ampliacao quando o texto realmente for pequeno.
                fator_padrao = 1.0 if area_captura >= 180000 else 1.25
                try:
                    fator_solicitado = float(escala_ocr) if escala_ocr is not None else fator_padrao
                except (TypeError, ValueError):
                    fator_solicitado = fator_padrao
                fator_ocr = max(0.85, min(1.45, fator_solicitado))
                variantes = self.gerar_variantes_imagem_ocr(captura, fator=fator_ocr, incluir_cor=False)[:1]
            else:
                fator_ocr = 1.35 if captura_grande else 2
                variantes = self.gerar_variantes_imagem_ocr(captura, fator=fator_ocr, incluir_cor=True)
            if modo_teste or retornar_candidatos:
                variantes = variantes[:1]
            imagem_escala = variantes[0][1] if variantes else captura
            escala_x = float(getattr(imagem_escala, "width", 1) or 1) / max(1.0, float(getattr(captura, "width", 1) or 1))
            escala_y = float(getattr(imagem_escala, "height", 1) or 1) / max(1.0, float(getattr(captura, "height", 1) or 1))
        except Exception:
            return None

        if modo_teste or retornar_candidatos:
            configs = ["--psm 11 -c preserve_interword_spaces=1"]
        else:
            configs = [
                "--psm 6 -c preserve_interword_spaces=1",
                "--psm 11 -c preserve_interword_spaces=1",
                "--psm 12 -c preserve_interword_spaces=1",
                "--psm 7 -c preserve_interword_spaces=1",
                "--psm 6 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ ",
                "--psm 11 -c tessedit_char_whitelist=ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789/.-:_ ",
            ]
        for nome_variante, imagem in variantes:
            for config in configs:
                inicio_leitura_ocr = time.perf_counter()
                try:
                    dados = pytesseract.image_to_data(imagem, config=config, output_type=pytesseract.Output.DICT)
                except Exception as erro:
                    nome_erro = erro.__class__.__name__.lower()
                    mensagem_erro = str(erro).lower()
                    if (
                        "tesseractnotfound" in nome_erro
                        or "tesseract is not installed" in mensagem_erro
                        or "tesseract.exe" in mensagem_erro
                    ) and not self._ocr_indisponivel_logado:
                        safe_print(
                            "OCR por screenshot indisponivel: tesseract.exe nao foi encontrado. "
                            "Inclua o Tesseract-OCR junto ao FlowDetect ou instale o motor OCR no sistema.",
                            flush=True,
                        )
                        self._ocr_indisponivel_logado = True
                    continue

                safe_print(
                    f"OCR screenshot concluido em {(time.perf_counter() - inicio_leitura_ocr) * 1000.0:.0f}ms "
                    f"(escala={fator_ocr:.2f}, psm={config.split()[1] if len(config.split()) > 1 else 'auto'}).",
                    flush=True,
                )
                total = len(dados.get("text", []))
                for indice in range(total):
                    bruto = str(dados.get("text", [""])[indice] or "").strip()
                    if not bruto:
                        continue
                    try:
                        conf_ocr = float(dados.get("conf", ["0"])[indice])
                    except Exception:
                        conf_ocr = 0.0
                    if conf_ocr < 18:
                        continue
                    left = int(float(dados.get("left", [0])[indice]) / escala_x) + int(base_left)
                    top = int(float(dados.get("top", [0])[indice]) / escala_y) + int(base_top)
                    width = max(8, int(float(dados.get("width", [0])[indice]) / escala_x))
                    height = max(8, int(float(dados.get("height", [0])[indice]) / escala_y))
                    itens.append({
                        "text": bruto,
                        "left": left,
                        "top": top,
                        "width": width,
                        "height": height,
                        "variant": nome_variante,
                        "conf": conf_ocr,
                    })

        if not isinstance(cache_itens, dict):
            cache_itens = {}
            self._cache_ocr_screenshot = cache_itens
        cache_itens[chave_cache] = {
            "em": time.perf_counter(),
            "captura": captura,
            "base_left": int(base_left),
            "base_top": int(base_top),
            "itens": [dict(item) for item in itens],
        }
        return captura, int(base_left), int(base_top), itens, regiao_ocr

    def localizar_texto_por_ocr_screenshot(
        self,
        texto,
        confianca=0.72,
        janela_alvo=None,
        region=None,
        retornar_candidatos=False,
        marcar_candidatos=False,
        limite_candidatos=8,
        _permitir_composto=True,
        _escala_ocr=None,
        _tentativa_ampliada=False,
    ):
        if not OCR_DISPONIVEL or pytesseract is None:
            if not self._ocr_indisponivel_logado:
                safe_print(
                    "OCR por screenshot indisponivel: Pillow/PyTesseract nao esta carregado.",
                    flush=True,
                )
                self._ocr_indisponivel_logado = True
            return None
        if _permitir_composto:
            spec_composto = self.interpretar_texto_ocr_composto(texto)
            if spec_composto:
                return self.localizar_texto_composto_por_ocr_screenshot(
                    texto,
                    spec_composto,
                    confianca=confianca,
                    janela_alvo=janela_alvo,
                    region=region,
                    retornar_candidatos=retornar_candidatos,
                    marcar_candidatos=marcar_candidatos,
                    limite_candidatos=limite_candidatos,
                    _escala_ocr=_escala_ocr,
                    _tentativa_ampliada=_tentativa_ampliada,
                )
        texto_alvo = self.normalizar_texto_ocr(texto)
        if not texto_alvo:
            return [] if retornar_candidatos else None
        leitura_ocr = self.capturar_itens_ocr_screenshot(
            janela_alvo=janela_alvo,
            region=region,
            retornar_candidatos=retornar_candidatos,
            escala_ocr=_escala_ocr,
        )
        if not leitura_ocr:
            return [] if retornar_candidatos else None
        captura, base_left, base_top, itens, regiao_ocr = leitura_ocr

        candidatos = []
        if itens:
            itens_unicos = {}
            for item in itens:
                texto_norm = self.normalizar_texto_ocr(item.get("text", ""))
                if not texto_norm:
                    continue
                chave = (
                    texto_norm,
                    int(int(item.get("left", 0) or 0) / 5),
                    int(int(item.get("top", 0) or 0) / 5),
                    int(int(item.get("width", 0) or 0) / 5),
                    int(int(item.get("height", 0) or 0) / 5),
                )
                atual = itens_unicos.get(chave)
                if atual is None or float(item.get("conf", 0.0) or 0.0) > float(atual.get("conf", 0.0) or 0.0):
                    itens_unicos[chave] = item
            itens = list(itens_unicos.values())
            itens.sort(key=lambda item: (int(item["top"] / 14), item["left"]))
            largura_maxima_candidato = max(90, min(420, int(len(texto_alvo) * 15)))
            altura_maxima_candidato = 96

            linhas = []
            for item in itens:
                centro_y = int(item["top"] + item["height"] / 2)
                linha = next((grupo for grupo in linhas if abs(grupo["centro_y"] - centro_y) <= max(10, int(item["height"] * 0.7))), None)
                if linha is None:
                    linha = {"centro_y": centro_y, "itens": []}
                    linhas.append(linha)
                linha["itens"].append(item)

            palavras_alvo = texto_alvo.split()
            menor_janela = max(1, len(palavras_alvo) - 1)
            maior_janela = max(menor_janela, min(max(4, len(palavras_alvo) + 2), 7))
            for linha in linhas:
                grupo_linha = sorted(linha["itens"], key=lambda item: item["left"])
                for inicio in range(len(grupo_linha)):
                    max_tamanho = min(maior_janela, len(grupo_linha) - inicio)
                    for tamanho in range(menor_janela, max_tamanho + 1):
                        grupo = grupo_linha[inicio : inicio + tamanho]
                        texto_grupo = " ".join(item["text"] for item in grupo)
                        score = self.similaridade_texto(texto_alvo, texto_grupo)
                        if score < float(confianca):
                            continue
                        left = min(item["left"] for item in grupo)
                        top = min(item["top"] for item in grupo)
                        right = max(item["left"] + item["width"] for item in grupo)
                        bottom = max(item["top"] + item["height"] for item in grupo)
                        largura = max(12, right - left)
                        altura = max(10, bottom - top)
                        if largura > largura_maxima_candidato or altura > altura_maxima_candidato:
                            continue
                        candidatos.append({
                            "left": left,
                            "top": top,
                            "width": largura,
                            "height": altura,
                            "text": texto_grupo,
                            "score": score,
                            "source": "ocr",
                        })

        if not candidatos and itens:
            amostras = " | ".join(str(item.get("text", "")) for item in itens[:24])
            if amostras:
                safe_print(f"OCR leu, mas nao bateu com o texto esperado. Amostra: {amostras}", flush=True)
        if not candidatos:
            # R?tulos pequenos como "Selecionar" podem sumir na leitura nativa.
            # Antes disso disparava uma fila de ate dez recortes, cada um com
            # varias variantes/PSM, que prendia a execucao por muitos segundos.
            # Uma unica leitura ampliada do mesmo dialogo e previsivel e mantem
            # a busca hibrida responsiva.
            if retornar_candidatos and not _tentativa_ampliada:
                safe_print(
                    "OCR rapido nao encontrou o texto; conferindo uma vez o dialogo ampliado.",
                    flush=True,
                )
                resultado_ampliado = self.localizar_texto_por_ocr_screenshot(
                    texto,
                    confianca=confianca,
                    janela_alvo=janela_alvo,
                    region=region,
                    retornar_candidatos=True,
                    marcar_candidatos=marcar_candidatos,
                    limite_candidatos=limite_candidatos,
                    _permitir_composto=False,
                    _escala_ocr=1.35,
                    _tentativa_ampliada=True,
                )
                if resultado_ampliado:
                    return resultado_ampliado
            if not retornar_candidatos and not bool(getattr(self, "modo_teste_reconhecimento", False)):
                resultado_varredura = self.localizar_texto_por_ocr_em_regioes(
                    captura,
                    texto_alvo,
                    confianca=confianca,
                    base_left=base_left,
                    base_top=base_top,
                )
                if resultado_varredura:
                    return resultado_varredura
            chave_diag = f"{texto_alvo}|{int(base_left)}|{int(base_top)}|{getattr(captura, 'width', 0)}|{getattr(captura, 'height', 0)}"
            if chave_diag not in self._diagnosticos_ocr_salvos:
                self._diagnosticos_ocr_salvos.add(chave_diag)
                diagnostico = self.salvar_captura_bruta_reconhecimento(
                    region=regiao_ocr,
                    titulo_janela=janela_alvo or self._janela_contexto,
                )
                if diagnostico:
                    safe_print(f"OCR diagnostico salvo em: {diagnostico}", flush=True)
            return [] if retornar_candidatos else None

        candidatos_ordenados = sorted(
            candidatos,
            key=lambda item: float(item.get("score", 0.0) or 0.0),
            reverse=True,
        )
        candidatos_filtrados = [
            item for item in candidatos_ordenados
            if float(item.get("score", 0.0) or 0.0) >= max(0.45, float(confianca) - 0.18)
        ][: max(1, int(limite_candidatos or 8))]
        if marcar_candidatos:
            for indice, candidato in enumerate(candidatos_filtrados[:6], start=1):
                self.marcar_area_reconhecida(
                    candidato,
                    duracao=0.18,
                    bloqueante=False,
                    tema="ocr_busca",
                    etiqueta=f"OCR {indice}",
                    margem=3,
                )
        if retornar_candidatos:
            return candidatos_filtrados
        return candidatos_ordenados[0]

    def extrair_textos_ocr_leve(self, imagem):
        if not OCR_DISPONIVEL or pytesseract is None or imagem is None or Image is None:
            return []
        try:
            area = max(1, int(imagem.width or 0) * int(imagem.height or 0))
            fator = 1.4 if area >= 450000 else (2 if area < 180000 else 1.6)
            variantes = self.gerar_variantes_imagem_ocr(imagem, fator=fator, incluir_cor=False)
        except Exception:
            return []
        configs = [
            "--psm 6 -c preserve_interword_spaces=1",
            "--psm 11 -c preserve_interword_spaces=1",
            "--psm 12 -c preserve_interword_spaces=1",
            "--psm 7 -c preserve_interword_spaces=1",
        ]
        textos = []
        for _nome, variante in variantes:
            for config in configs:
                try:
                    texto = pytesseract.image_to_string(variante, config=config)
                except Exception:
                    continue
                texto = self.normalizar_texto_ocr(texto)
                if texto and texto not in textos:
                    textos.append(texto)
        return textos

    def texto_ocr_imagem_corresponde(self, imagem, texto_alvo, confianca=0.72, pesado=False):
        textos = self.extrair_textos_ocr_leve(imagem)
        if pesado:
            for texto in self.extrair_variantes_texto_ocr_imagem(imagem):
                if texto and texto not in textos:
                    textos.append(texto)
        melhor_texto = ""
        melhor_score = 0.0
        for texto_lido in textos:
            score = self.similaridade_texto(texto_alvo, texto_lido)
            if score > melhor_score:
                melhor_score = score
                melhor_texto = texto_lido
            if score >= float(confianca):
                return True, texto_lido, score
        return False, melhor_texto, melhor_score

    def subdividir_regiao_ocr(self, largura, altura):
        regioes = []
        metade_largura = max(1, int(largura / 2))
        metade_altura = max(1, int(altura / 2))
        # Prioriza o lado direito, onde botões de login/acesso aparecem com frequência.
        regioes.extend([
            (metade_largura, 0, largura - metade_largura, altura),
            (0, 0, metade_largura, altura),
            (metade_largura, 0, largura - metade_largura, metade_altura),
            (metade_largura, metade_altura, largura - metade_largura, altura - metade_altura),
            (0, 0, metade_largura, metade_altura),
            (0, metade_altura, metade_largura, altura - metade_altura),
        ])
        filtradas = []
        vistos = set()
        for x, y, w, h in regioes:
            if w < 80 or h < 40:
                continue
            chave = (x, y, w, h)
            if chave in vistos:
                continue
            vistos.add(chave)
            filtradas.append(chave)
        return filtradas

    def localizar_texto_por_ocr_em_regioes(self, captura, texto_alvo, confianca=0.72, base_left=0, base_top=0):
        if captura is None or Image is None:
            return None

        fila = [(0, 0, int(captura.width or 0), int(captura.height or 0), 0)]
        melhor = None
        melhor_score = 0.0
        max_niveis = 3
        analisadas = 0
        while fila and analisadas < 10:
            x, y, largura, altura, nivel = fila.pop(0)
            if largura < 80 or altura < 40:
                continue
            analisadas += 1
            try:
                recorte = captura.crop((x, y, x + largura, y + altura))
            except Exception:
                continue
            area_recorte = max(1, largura * altura)
            encontrou, texto_lido, score = self.texto_ocr_imagem_corresponde(
                recorte,
                texto_alvo,
                confianca=max(0.58, float(confianca) - 0.05),
                pesado=area_recorte <= 260000,
            )
            if score > melhor_score:
                melhor_score = score
                melhor = (x, y, largura, altura, texto_lido, score)
            if not encontrou:
                continue
            if nivel >= max_niveis or largura <= 260 or altura <= 140:
                safe_print(
                    f"OCR encontrou texto por varredura regional: '{texto_lido}' "
                    f"(similaridade={score:.2f}).",
                    flush=True,
                )
                return {
                    "left": int(base_left + x),
                    "top": int(base_top + y),
                    "width": max(24, int(largura)),
                    "height": max(24, int(altura)),
                    "text": texto_lido,
                    "score": score,
                    "source": "ocr-regiao",
                }
            for sub_x, sub_y, sub_w, sub_h in self.subdividir_regiao_ocr(largura, altura):
                fila.append((x + sub_x, y + sub_y, sub_w, sub_h, nivel + 1))

        if melhor and melhor_score >= max(0.70, float(confianca) + 0.08):
            x, y, largura, altura, texto_lido, score = melhor
            safe_print(
                f"OCR encontrou texto por leitura regional ampla: '{texto_lido}' "
                f"(similaridade={score:.2f}).",
                flush=True,
            )
            return {
                "left": int(base_left + x),
                "top": int(base_top + y),
                "width": max(24, int(largura)),
                "height": max(24, int(altura)),
                "text": texto_lido,
                "score": score,
                "source": "ocr-regiao",
            }
        return None

    def clicar_resultado_texto(self, resultado, clicar=True, clicks=1, offset_x=0, offset_y=0):
        if not resultado:
            return None
        self.marcar_area_reconhecida(resultado, duracao=0.45, bloqueante=True)
        destino_x = int(resultado["left"] + resultado["width"] / 2 + int(offset_x or 0))
        destino_y = int(resultado["top"] + resultado["height"] / 2 + int(offset_y or 0))
        if clicar:
            time.sleep(0.08)
            pyautogui.click(destino_x, destino_y, clicks=max(1, int(clicks or 1)))
            self._ultimo_clique_visual_em = time.perf_counter()
            self._ultima_posicao_clique_visual = {
                "x": destino_x,
                "y": destino_y,
                "button": "left",
                "clicks": max(1, int(clicks or 1)),
            }
        return {"x": destino_x, "y": destino_y}

    def executar_bloco_procurar_texto_ocr(self, evento):
        detalhes = evento.get("detalhes") or {}
        ordem = int(evento.get("ordem", 0) or 0)
        texto = str(detalhes.get("texto", "") or "").strip()
        if not texto:
            raise RuntimeError("Bloco OCR inteligente sem texto configurado.")
        teste_ocr_direto = bool(detalhes.get("teste_reconhecimento_ocr")) or bool(getattr(self, "modo_teste_reconhecimento", False))
        acao = str(detalhes.get("acao", "clicar") or "clicar").strip().lower()
        deve_clicar = bool(detalhes.get("clicar", True)) and acao in {"clicar", "click"}
        janela_evento = "" if teste_ocr_direto else self.resolver_janela_evento(
            detalhes.get("janela_alvo"),
            registrar_conflito=not bool(str(detalhes.get("janela_alvo") or "").strip()),
            priorizar_evento=bool(str(detalhes.get("janela_alvo") or "").strip()),
        )
        if not janela_evento:
            janela_auto = None if teste_ocr_direto else self.escolher_janela_ocr_automatica(texto)
            if janela_auto:
                janela_evento = str(janela_auto.get("title") or "").strip()
        tentativas = max(1, int(detalhes.get("quantidade_tentativas", 3) or 3))
        intervalo = max(0.0, float(detalhes.get("intervalo_tentativas", 0.7) or 0.7))
        confianca = max(0.35, min(1.0, float(detalhes.get("confianca", 0.72) or 0.72)))
        timeout = max(0.1, float(detalhes.get("timeout", 8) or 8))
        permitir_scroll = str(detalhes.get("varredura_scroll", True)).strip().lower() not in {"0", "false", "nao", "não", "off", "desativado"}
        if teste_ocr_direto:
            tentativas = min(tentativas, 2)
            timeout = min(timeout, 4.0)
            permitir_scroll = False
        scrolls = max(0, min(8, int(detalhes.get("varredura_scroll_tentativas", 4) or 4)))
        pixels_scroll = max(120, int(detalhes.get("varredura_scroll_pixels", 520) or 520))
        direcao_scroll = str(detalhes.get("varredura_scroll_direcao", "baixo") or "baixo")
        fim = time.perf_counter() + timeout
        ultimo_resultado = None

        for tentativa in range(tentativas):
            if self.cancelamento_solicitado():
                return self.resultado_fluxo("parar")
            safe_print(f"Bloco {ordem}: procurando texto '{texto}' | tentativa {tentativa + 1}/{tentativas}", flush=True)
            etapas_scroll = scrolls + 1 if permitir_scroll else 1
            for etapa in range(etapas_scroll):
                if self.cancelamento_solicitado():
                    return self.resultado_fluxo("parar")
                if not janela_evento and not teste_ocr_direto:
                    janela_auto = self.escolher_janela_ocr_automatica(texto) or self.adotar_janela_ativa_para_ocr()
                    if janela_auto:
                        janela_evento = str(janela_auto.get("title") or "").strip()
                if teste_ocr_direto:
                    safe_print("Teste OCR direto: lendo a tela visivel inteira, sem fixar janela.", flush=True)
                    candidatos = self.localizar_texto_por_ocr_screenshot(
                        texto,
                        confianca=confianca,
                        janela_alvo=None,
                        region=None,
                        retornar_candidatos=True,
                        marcar_candidatos=True,
                        limite_candidatos=1,
                    )
                    resultado = candidatos[0] if isinstance(candidatos, list) and candidatos else None
                else:
                    resultado = self.localizar_texto_por_uia(texto, confianca=confianca, janela_alvo=janela_evento)
                    if not resultado:
                        candidatos = self.localizar_texto_por_ocr_screenshot(
                            texto,
                            confianca=confianca,
                            janela_alvo=janela_evento,
                            retornar_candidatos=True,
                            marcar_candidatos=True,
                            limite_candidatos=1,
                        )
                        resultado = candidatos[0] if isinstance(candidatos, list) and candidatos else None
                if resultado:
                    ultimo_resultado = resultado
                    capturado = str(resultado.get("text") or texto).strip()
                    self.salvar_variavel_contexto("ultimo_ocr", capturado)
                    self.salvar_variavel_contexto("ocr_tela", capturado)
                    self.salvar_variavel_contexto(f"ocr_bloco_{ordem}", capturado)
                    salvar_em = str(detalhes.get("salvar_em", "") or "").strip()
                    if salvar_em:
                        self.salvar_variavel_contexto(salvar_em, capturado)
                    safe_print(
                        f"Texto encontrado por {resultado.get('source', 'texto')}: '{capturado}' "
                        f"(similaridade={float(resultado.get('score', 0.0) or 0.0):.2f}).",
                        flush=True,
                    )
                    if acao == "verificar":
                        return self.resultado_sucesso_evento(evento)
                    self.clicar_resultado_texto(
                        resultado,
                        clicar=deve_clicar,
                        clicks=detalhes.get("cliques", 1),
                        offset_x=detalhes.get("offset_x", 0),
                        offset_y=detalhes.get("offset_y", 0),
                    )
                    return self.resultado_sucesso_evento(evento)
                if etapa < etapas_scroll - 1:
                    self.rolar_para_varredura_visual(
                        titulo_janela=janela_evento,
                        direcao=direcao_scroll,
                        pixels=pixels_scroll,
                    )
                if time.perf_counter() >= fim:
                    break
            if tentativa < tentativas - 1 and intervalo > 0:
                if not self.esperar(intervalo):
                    return self.resultado_fluxo("parar")

        acao_falha = str(detalhes.get("ao_nao_encontrar", "parar") or "parar")
        if acao_falha in {"continuar", "pular"}:
            safe_print(f"Bloco {ordem}: texto '{texto}' nao encontrado. Continuando.", flush=True)
            return self.resultado_fluxo("continuar")
        if acao_falha == "ir_para_sequencia":
            sequencia = str(detalhes.get("sequencia_nao_encontrar", "") or "").strip()
            return self.resultado_fluxo(
                "ir_para_sequencia",
                sequencia=sequencia,
                atraso=detalhes.get("atraso_antes_sequencia_nao_encontrar", 0),
            )
        if acao_falha == "ir_para_bloco":
            return self.resultado_fluxo(
                "ir_para_bloco",
                bloco=detalhes.get("bloco_nao_encontrar", ""),
                atraso=detalhes.get("atraso_antes_sequencia_nao_encontrar", 0),
            )
        extra = ""
        if ultimo_resultado:
            extra = f" | melhor texto: {ultimo_resultado.get('text', '')}"
        raise RuntimeError(f"Texto nao encontrado na tela: {texto}{extra}")

    def cancelamento_solicitado(self):
        if self.cancelar.is_set():
            return True
        if self.arquivo_cancelamento is None:
            return False
        return self.arquivo_cancelamento.exists()

    def carregar_eventos(self):
        return list(self.payload.get("eventos", []))

    def carregar_planilha(self):
        arquivos = {self.normalizar_id_planilha(None)}
        for evento in self.eventos:
            tipo = str(evento.get("tipo") or "").strip()
            detalhes = evento.get("detalhes") or {}
            if tipo in {"sequencia_planilha", "escrever_planilha", "comparar_planilha_tela"}:
                arquivos.add(self.normalizar_id_planilha(detalhes.get("arquivo")))
        cache = {}
        for arquivo in arquivos:
            caminho_planilha, workbook_id = self.resolver_planilha(arquivo)
            if not caminho_planilha.exists():
                continue

            workbook = load_workbook(caminho_planilha, read_only=True, data_only=True)
            cache_workbook = {}
            for nome_aba in workbook.sheetnames:
                worksheet = workbook[nome_aba]
                linhas = list(worksheet.iter_rows(values_only=True))
                if not linhas:
                    cache_workbook[nome_aba] = {"colunas": [], "linhas": []}
                    continue

                cabecalhos = [
                    str(valor).strip() if valor is not None else f"Coluna {indice + 1}"
                    for indice, valor in enumerate(linhas[0])
                ]
                registros = []
                for numero_linha, linha in enumerate(linhas[1:], start=2):
                    registro = {"__linha__": numero_linha}
                    for indice, coluna in enumerate(cabecalhos):
                        registro[coluna] = linha[indice] if indice < len(linha) else None
                    registros.append(registro)
                cache_workbook[nome_aba] = {"colunas": cabecalhos, "linhas": registros}

            workbook.close()
            cache[workbook_id] = cache_workbook

        return cache

    def normalizar_acao(self, acao):
        if not acao:
            return []
        if isinstance(acao, str):
            texto = acao.strip().lower()
            if not texto or texto == "nenhuma":
                return []
            return [parte.strip() for parte in texto.replace(" ", "").split("+") if parte.strip()]
        return [str(parte).strip() for parte in acao if str(parte).strip()]

    def tecla_para_pyautogui(self, nome):
        if len(nome) == 1:
            return nome
        return MAPA_PYAUTOGUI.get(nome, nome)

    def eh_texto_simples(self, nome):
        return len(nome) == 1 and nome not in {"\r", "\n", "\t"}

    def escrever_texto(self, texto):
        if not self.preparar_foco_para_digitacao():
            return False
        for caractere in str(texto):
            if self.cancelamento_solicitado():
                return False
            if caractere == "\r":
                continue
            if caractere == "\n":
                pyautogui.press("enter")
                continue
            if caractere == "\t":
                pyautogui.press("tab")
                continue
            pyautogui.write(caractere, interval=INTERVALO_TEXTO)
        return True

    def esperar_imagem(
        self,
        imagem,
        timeout=10,
        confianca=0.8,
        janela_alvo=None,
        perfil_busca="balanced",
        bloco_id=None,
        modo_area_busca="target_window",
        area_personalizada=None,
        texto_ocr_esperado="",
        leitura_relativa=None,
        preservar_janelas=True,
    ):
        caminho_imagem, imagem_id = self.resolver_imagem(imagem)
        if not imagem_id:
            raise FileNotFoundError("Nenhuma imagem foi informada para o bloco clicar por imagem.")
        if not caminho_imagem.exists():
            raise FileNotFoundError(f"Imagem nao encontrada no caminho informado: {caminho_imagem}")

        # Cada bloco precisa iniciar sobre uma tela nova. O cache abaixo so e
        # compartilhado entre os varios textos OCR da mesma tentativa hibrida.
        cache_ocr_screenshot = getattr(self, "_cache_ocr_screenshot", None)
        if isinstance(cache_ocr_screenshot, dict):
            cache_ocr_screenshot.clear()

        janela_alvo_explicita = str(janela_alvo or "").strip()
        # Os blocos recebem o titulo do alvo selecionado como contexto herdado.
        # Esse titulo nao pode fazer a execucao abandonar o HWND fixado pelo
        # usuario: ele descreve a mesma janela, apenas sem carregar o handle.
        # Um titulo diferente ainda funciona como substituicao explicita para
        # aquele bloco.
        alvo_execucao_ativo = self.alvo_execucao_ativo()
        titulo_contexto_execucao = str(
            (getattr(self, "contexto_execucao_alvo", None) or {}).get("target_title")
            or self.janela_alvo
            or ""
        ).strip()
        titulo_herdado_do_alvo = (
            not janela_alvo_explicita
            or (
                bool(titulo_contexto_execucao)
                and self.normalizar_texto_busca(janela_alvo_explicita)
                == self.normalizar_texto_busca(titulo_contexto_execucao)
            )
        )
        usando_alvo_execucao = alvo_execucao_ativo and titulo_herdado_do_alvo
        janela_alvo_resolvida = str(
            self.resolver_janela_evento(
                janela_alvo_explicita,
                registrar_conflito=not bool(janela_alvo_explicita),
                priorizar_evento=bool(janela_alvo_explicita),
            ) or ""
        ).strip()
        titulo_busca = str(janela_alvo_resolvida or self._janela_contexto or self.janela_alvo or "").strip()
        janela_info_cache = None
        usando_contexto_herdado = False
        if usando_alvo_execucao:
            janela_resolvida = self.obter_janela_alvo_execucao_atual(obrigatoria=True)
            if janela_resolvida and str(janela_resolvida.get("title") or "").strip():
                titulo_busca = str(janela_resolvida.get("title") or "").strip()
                self._janela_contexto = titulo_busca
                janela_info_cache = janela_resolvida
        elif not titulo_busca:
            janela_fixada = self.obter_janela_fixada()
            if janela_fixada and str(janela_fixada.get("title") or "").strip():
                titulo_busca = str(janela_fixada.get("title") or "").strip()
                self._janela_contexto = titulo_busca
                janela_info_cache = janela_fixada
                usando_contexto_herdado = True
                safe_print(f"Usando janela de contexto para reconhecimento: {titulo_busca}", flush=True)
        if titulo_busca and not usando_alvo_execucao:
            if usando_contexto_herdado:
                janela_resolvida = self.obter_janela_fixada()
            else:
                janela_resolvida = self.localizar_janela_alvo(titulo_busca)
            if janela_resolvida and str(janela_resolvida.get("title") or "").strip():
                titulo_busca = str(janela_resolvida.get("title") or "").strip()
                self._janela_contexto = titulo_busca
                janela_info_cache = janela_resolvida
        janela_alvo = titulo_busca or None
        preservar_janelas = bool(preservar_janelas)
        # Reconhecimento deve apenas observar os pixels que ja estao visiveis.
        # Trazer uma janela para frente durante a busca pode esconder o dialogo
        # que contem o alvo e altera a ordem escolhida pelo usuario.
        reativar_janela_durante_busca = (
            not preservar_janelas
            and not usando_alvo_execucao
            and not usando_contexto_herdado
            and bool(janela_alvo or self.janela_alvo)
        )

        confianca = float(confianca)
        perfil_busca = str(perfil_busca or "balanced").strip().lower()
        if perfil_busca not in {"fast", "balanced", "precise"}:
            perfil_busca = "balanced"
        bloco_id = str(bloco_id or imagem_id or "").strip() or None
        modo_area_busca = self.normalizar_modo_area_busca(modo_area_busca)
        modo_estrito = confianca >= 0.99

        safe_print(f"Procurando imagem: {imagem_id}", flush=True)
        safe_print(
            f"Reconhecendo... perfil={perfil_busca} | timeout={max(0.1, float(timeout))}s | confianca={confianca:.2f}",
            flush=True,
        )
        safe_print(f"Area configurada para o bloco: {modo_area_busca}", flush=True)
        timeout_total = max(0.1, float(timeout))
        permitir_ocr_geral = (
            perfil_busca == "precise"
            and not modo_estrito
            and OCR_DISPONIVEL
            and timeout_total >= 4.0
        )
        inicio_busca = time.perf_counter()
        fim = inicio_busca + timeout_total
        ultima_falha = None
        ultima_tentativa_ocr = 0.0
        ocr_hibrido_pendente = {}
        ultima_ativacao_janela = 0.0
        self._ultima_regiao_reconhecimento = None
        self._ultimo_template_debug = None
        self._ultimo_diagnostico_reconhecimento = None
        texto_referencia = self.extrair_texto_referencia(caminho_imagem) if permitir_ocr_geral else ""
        dimensoes_referencia = self.obter_dimensoes_imagem(caminho_imagem)
        texto_ocr_ancora = str(texto_ocr_esperado or "").strip()
        hibrido_ocr_visual_obrigatorio = bool(texto_ocr_ancora and dimensoes_referencia and not modo_estrito)
        janela_ativa_hibrida = None
        janela_companheira_hibrida = None
        chave_janela_companheira_logada = ""

        # Sem uma janela alvo definida, o OCR da tela inteira torna textos de
        # abas e dialogos pequenos demais (por exemplo, "SPED Reinf"). Lemos
        # somente a regiao que ja esta em primeiro plano, sem fixar, ativar,
        # maximizar ou alterar a ordem de nenhuma janela.
        if (
            hibrido_ocr_visual_obrigatorio
            and modo_area_busca == "target_window"
            and not usando_alvo_execucao
            and not janela_alvo
            and not self.janela_alvo
        ):
            candidata_ativa = self.obter_janela_ativa()
            if (
                not isinstance(candidata_ativa, dict)
                or self.janela_parece_interface_flowdetect(candidata_ativa)
                or int(candidata_ativa.get("width", 0) or 0) < 260
                or int(candidata_ativa.get("height", 0) or 0) < 180
            ):
                candidata_ativa = self.adotar_janela_visivel_passiva_para_ocr()
            if (
                isinstance(candidata_ativa, dict)
                and not self.janela_parece_interface_flowdetect(candidata_ativa)
                and int(candidata_ativa.get("width", 0) or 0) >= 260
                and int(candidata_ativa.get("height", 0) or 0) >= 180
            ):
                janela_ativa_hibrida = candidata_ativa
                self.memorizar_contexto_visual_passivo(candidata_ativa)
                safe_print(
                    "Hibrido OCR + imagem: lendo somente a janela visivel do sistema, sem altera-la.",
                    flush=True,
                )

        def adicionar_regiao_busca(destino, regiao, etapa, fallback=False):
            regiao_normalizada = self.limitar_regiao(regiao) if regiao is not None else None
            if regiao is not None and regiao_normalizada is None:
                return
            for existente in destino:
                if existente.get("region") == regiao_normalizada:
                    return
            destino.append(
                {
                    "region": regiao_normalizada,
                    "stage": str(etapa or "area"),
                    "fallback": bool(fallback),
                }
            )

        usa_confianca = OPENCV_DISPONIVEL
        if confianca != 1.0 and not OPENCV_DISPONIVEL:
            safe_print(
                "OpenCV indisponivel. Continuando o reconhecimento sem confidence e com reforco por OCR quando possivel.",
                flush=True,
            )

        confiancas_busca = [1.0 if modo_estrito else (confianca if usa_confianca else 1.0)]
        if usa_confianca and perfil_busca == "precise" and not modo_estrito and confianca > 0.72:
            confiancas_busca.append(round(max(0.55, confianca - 0.12), 2))

        while time.perf_counter() <= fim:
            if self.cancelamento_solicitado():
                return None
            agora = time.perf_counter()
            tempo_decorrido = max(0.0, agora - inicio_busca)
            if perfil_busca == "fast":
                perfil_iteracao = "fast"
            elif perfil_busca == "balanced":
                # Depois que os OCRs delimitaram uma area pequena, nao faz
                # sentido aguardar 68% do timeout para ampliar a validacao
                # visual. Esse atraso era o motivo dos blocos apos SPED
                # demorarem 8--10 segundos apesar de o texto ja estar lido.
                limite_fast_balanceado = timeout_total * 0.68
                if hibrido_ocr_visual_obrigatorio:
                    limite_fast_balanceado = min(1.10, max(0.35, timeout_total * 0.18))
                perfil_iteracao = "fast" if tempo_decorrido < limite_fast_balanceado else "balanced"
            else:
                limite_fast_preciso = timeout_total * 0.40
                if hibrido_ocr_visual_obrigatorio:
                    limite_fast_preciso = min(1.30, max(0.40, timeout_total * 0.22))
                perfil_iteracao = "fast" if tempo_decorrido < limite_fast_preciso else "precise"
            intervalo_reativacao = 2.4 if perfil_busca == "fast" else (1.8 if perfil_busca == "balanced" else 1.0)
            if usando_alvo_execucao:
                janela_info_cache = self.obter_janela_alvo_execucao_atual(obrigatoria=True)
            elif reativar_janela_durante_busca and time.perf_counter() - ultima_ativacao_janela >= intervalo_reativacao:
                janela_ativada = self.ativar_janela_alvo(janela_alvo)
                if janela_ativada:
                    janela_info_cache = janela_ativada
                ultima_ativacao_janela = time.perf_counter()
            elif reativar_janela_durante_busca and janela_info_cache is None:
                janela_info_cache = self.localizar_janela_alvo(janela_alvo)
            elif usando_contexto_herdado:
                janela_info_cache = self.obter_janela_fixada() or janela_info_cache
            if hibrido_ocr_visual_obrigatorio and not usando_alvo_execucao and (janela_alvo or self.janela_alvo):
                janela_atualizada = self.localizar_janela_alvo(janela_alvo or self.janela_alvo)
                if janela_atualizada:
                    janela_info_cache = janela_atualizada
            if hibrido_ocr_visual_obrigatorio and usando_alvo_execucao and janela_info_cache:
                janela_companheira_hibrida = self.obter_janela_companheira_execucao(janela_info_cache)
                if janela_companheira_hibrida:
                    chave_companheira = (
                        f"{int(janela_companheira_hibrida.get('hwnd', 0) or 0)}|"
                        f"{janela_companheira_hibrida.get('title', '')}"
                    )
                    if chave_companheira != chave_janela_companheira_logada:
                        safe_print(
                            "Hibrido OCR + imagem: lendo passivamente o dialogo visivel do mesmo programa "
                            f"({janela_companheira_hibrida.get('title', 'dialogo atual')}).",
                            flush=True,
                        )
                        chave_janela_companheira_logada = chave_companheira
            margens_busca = [0]
            if usando_alvo_execucao or janela_alvo or self.janela_alvo:
                if modo_estrito:
                    margens_busca = [0] if tempo_decorrido < timeout_total * 0.92 else [0, 6]
                elif perfil_busca == "fast":
                    margens_busca = [0] if tempo_decorrido < timeout_total * 0.9 else [0, 12]
                elif perfil_busca == "precise":
                    margens_busca = [0, 16] if tempo_decorrido < timeout_total * 0.45 else [0, 16, 32]
                else:
                    margens_busca = [0] if tempo_decorrido < timeout_total * 0.65 else [0, 24]

            regiao_historica = self.obter_regiao_historica_reconhecimento(
                imagem_id,
                janela_alvo=janela_alvo,
                bloco_id=bloco_id,
                janela_info=janela_info_cache,
            )
            regiao_personalizada = self.construir_regiao_personalizada(
                area_personalizada,
                janela_info=janela_info_cache,
            )
            regioes_busca = []
            if regiao_historica and not hibrido_ocr_visual_obrigatorio:
                adicionar_regiao_busca(regioes_busca, regiao_historica, "cache", fallback=False)

            if modo_area_busca == "custom_area" and regiao_personalizada:
                adicionar_regiao_busca(regioes_busca, regiao_personalizada, "custom_area", fallback=bool(regiao_historica))

            if usando_alvo_execucao or janela_alvo or self.janela_alvo:
                for margem in margens_busca:
                    if janela_info_cache:
                        margem = max(0, int(margem or 0))
                        regiao_candidata = (
                            int(janela_info_cache["left"]) - margem,
                            int(janela_info_cache["top"]) - margem,
                            int(janela_info_cache["width"]) + margem * 2,
                            int(janela_info_cache["height"]) + margem * 2,
                        )
                    else:
                        regiao_candidata = self.obter_regiao_busca(janela_alvo, margem=margem)
                    adicionar_regiao_busca(
                        regioes_busca,
                        regiao_candidata,
                        "target_window",
                        fallback=(modo_area_busca != "target_window") or bool(regiao_historica),
                    )
            if janela_companheira_hibrida:
                adicionar_regiao_busca(
                    regioes_busca,
                    (
                        int(janela_companheira_hibrida.get("left", 0) or 0),
                        int(janela_companheira_hibrida.get("top", 0) or 0),
                        int(janela_companheira_hibrida.get("width", 0) or 0),
                        int(janela_companheira_hibrida.get("height", 0) or 0),
                    ),
                    "target_companion",
                    fallback=False,
                )
                # Em fluxos fiscais o alvo normalmente aparece em um dialogo
                # modal do mesmo programa (por exemplo, SPED Reinf). A busca
                # anterior examinava primeiro a janela principal inteira e so
                # depois o dialogo menor que ja esta visivel. Priorizar esse
                # dialogo reduz muito o OCR/template sem ativar ou mover janela;
                # a janela principal permanece como alternativa na mesma rodada.
                if hibrido_ocr_visual_obrigatorio:
                    regioes_busca.sort(
                        key=lambda entrada: 0 if entrada.get("stage") == "target_companion" else 1
                    )
            elif modo_area_busca == "custom_area" and regiao_personalizada:
                adicionar_regiao_busca(regioes_busca, regiao_personalizada, "custom_area", fallback=bool(regiao_historica))

            if janela_ativa_hibrida:
                adicionar_regiao_busca(
                    regioes_busca,
                    (
                        int(janela_ativa_hibrida.get("left", 0) or 0),
                        int(janela_ativa_hibrida.get("top", 0) or 0),
                        int(janela_ativa_hibrida.get("width", 0) or 0),
                        int(janela_ativa_hibrida.get("height", 0) or 0),
                    ),
                    "active_window",
                    fallback=False,
                )

            if hibrido_ocr_visual_obrigatorio and not regioes_busca:
                monitores = self.listar_regioes_monitores()
                if len(monitores) > 1:
                    for indice_monitor, regiao_monitor in enumerate(monitores, start=1):
                        adicionar_regiao_busca(
                            regioes_busca,
                            regiao_monitor,
                            f"monitor_{indice_monitor}",
                            fallback=False,
                        )
                else:
                    adicionar_regiao_busca(regioes_busca, None, "screen", fallback=False)

            # A janela ativa resolve quase todos os dialogos pequenos com muito
            # mais rapidez. Se ela ainda nao tiver o elemento depois de alguns
            # segundos, tentamos a tela inteira uma unica vez como reserva.
            if (
                hibrido_ocr_visual_obrigatorio
                and janela_ativa_hibrida
                and tempo_decorrido >= min(2.5, max(0.9, timeout_total * 0.55))
            ):
                adicionar_regiao_busca(regioes_busca, None, "screen", fallback=True)

            self._ultima_regiao_reconhecimento = regioes_busca[-1]["region"] if regioes_busca else None
            if (usando_alvo_execucao or janela_alvo or self.janela_alvo) and not regioes_busca:
                if usando_alvo_execucao:
                    raise RuntimeError("A janela alvo selecionada nao gerou uma regiao valida de captura.")
                time.sleep(0.12)
                continue

            # O intervalo e entre ciclos de captura, e nao entre as regioes do
            # mesmo ciclo. Assim, quando o sistema abre um dialogo do mesmo
            # processo, a janela selecionada e o dialogo visivel sao avaliados
            # na mesma tentativa passiva.
            executar_hibrido_neste_ciclo = True
            if hibrido_ocr_visual_obrigatorio and OCR_DISPONIVEL:
                # Cada OCR de dialogo leva centenas de milissegundos. Repetir
                # imediatamente ao terminar mantinha um nucleo em 100% e deixava
                # o cursor lento. A proxima rodada usa o cache ou uma captura
                # nova depois deste intervalo curto e previsivel.
                intervalo_hibrido = 0.38 if (
                    perfil_busca == "fast" or bool(getattr(self, "modo_teste_reconhecimento", False))
                ) else 0.58
                agora_hibrido = time.perf_counter()
                executar_hibrido_neste_ciclo = agora_hibrido - ultima_tentativa_ocr >= intervalo_hibrido
                if executar_hibrido_neste_ciclo:
                    ultima_tentativa_ocr = agora_hibrido
                    self._ultimo_ocr_global_em = agora_hibrido

            for entrada_regiao in (regioes_busca or [{"region": None, "stage": "screen", "fallback": False}]):
                regiao = entrada_regiao.get("region")
                etapa_busca = str(entrada_regiao.get("stage") or "screen")
                if regiao and dimensoes_referencia:
                    ref_largura, ref_altura = dimensoes_referencia
                    if ref_largura > regiao[2] or ref_altura > regiao[3]:
                        if not self._needle_maior_logada:
                            safe_print(
                                "A referencia visual eh maior que a janela alvo selecionada. "
                                "Ajuste a imagem ou escolha uma janela/regiao maior.",
                                flush=True,
                            )
                            self._needle_maior_logada = True
                        continue

                perfil_busca_tentativa = "fast" if etapa_busca == "cache" else perfil_iteracao
                if hibrido_ocr_visual_obrigatorio:
                    if not OCR_DISPONIVEL:
                        if not self._ocr_indisponivel_logado:
                            safe_print("Hibrido OCR + imagem requer OCR disponivel, mas o motor OCR nao esta carregado.", flush=True)
                            self._ocr_indisponivel_logado = True
                        time.sleep(0.12)
                        continue
                    if not executar_hibrido_neste_ciclo:
                        continue
                    chave_hibrido = tuple(int(valor) for valor in regiao) if regiao else ("tela",)
                    cache_hibrido = ocr_hibrido_pendente.get(chave_hibrido)
                    ocr_cacheado = None
                    ocr_cacheado_em = 0.0
                    if cache_hibrido and agora_hibrido - float(cache_hibrido.get("em", 0.0) or 0.0) <= 1.4:
                        ocr_cacheado = cache_hibrido.get("local")
                        ocr_cacheado_em = float(cache_hibrido.get("em", 0.0) or 0.0)
                    if ocr_cacheado:
                        safe_print(
                            "Hibrido OCR + imagem: reutilizando os OCRs ja confirmados para a etapa 2/2.",
                            flush=True,
                        )
                    else:
                        safe_print(
                            f"Hibrido OCR + imagem: etapa 1/2 procurando os OCRs de '{texto_ocr_ancora}'.",
                            flush=True,
                        )
                    local_hibrido = self.localizar_referencia_hibrida_ocr_visual(
                        caminho_imagem,
                        texto_ocr_esperado=texto_ocr_ancora,
                        region=regiao,
                        dimensoes_referencia=dimensoes_referencia,
                        confianca_ocr=0.72,
                        confianca_visual=confianca,
                        perfil_busca=perfil_busca_tentativa,
                        leitura_relativa=leitura_relativa,
                        ocr_confirmado=ocr_cacheado,
                    )
                    if isinstance(local_hibrido, dict) and "ocr_validation" in local_hibrido:
                        # O teste completo ja tem o resultado de cada alvo,
                        # mesmo quando algum falhou. Nao reinicie noutra area.
                        return local_hibrido
                    fluxo_ocr_individual = (
                        local_hibrido.get("ocr_target_flow")
                        if isinstance(local_hibrido, dict)
                        else None
                    )
                    if (
                        isinstance(fluxo_ocr_individual, dict)
                        and str(fluxo_ocr_individual.get("acao") or "") in {"continuar", "parar"}
                    ):
                        # Uma falha individual nao e uma tentativa pendente de
                        # OCR: ela ja decidiu o fluxo. Antes esse resultado era
                        # descartado pelo cache hibrido e o comportamento geral
                        # do card acabava substituindo a regra do OCR.
                        ocr_hibrido_pendente.pop(chave_hibrido, None)
                        return local_hibrido
                    if local_hibrido and local_hibrido.get("visual_confirmed"):
                        ocr_hibrido_pendente.pop(chave_hibrido, None)
                        self.registrar_reconhecimento_bem_sucedido(
                            imagem_id,
                            local_hibrido,
                            janela_alvo=janela_alvo,
                            bloco_id=bloco_id,
                            janela_info=janela_info_cache,
                        )
                        return local_hibrido
                    if local_hibrido and self.ocr_hibrido_tem_todos_alvos(
                        local_hibrido, texto_ocr_ancora, leitura_relativa
                    ):
                        ocr_hibrido_pendente[chave_hibrido] = {
                            "local": local_hibrido,
                            # A leitura pode levar mais de um segundo. Registrar o
                            # horario anterior ao OCR fazia o cache expirar antes da
                            # proxima tentativa visual e repetia todos os alvos.
                            "em": time.perf_counter(),
                        }
                    else:
                        ocr_hibrido_pendente.pop(chave_hibrido, None)
                    continue
                local_template_preferencial = self.procurar_imagem_por_template(
                    caminho_imagem,
                    confianca=confianca,
                    region=regiao,
                    perfil_busca=perfil_busca_tentativa,
                )
                if local_template_preferencial and perfil_busca == "fast" and etapa_busca == "cache":
                    info_template = self._ultima_busca_template_info or {}
                    score_cache = float(info_template.get("score", local_template_preferencial.get("rank_score", 0.0)) or 0.0)
                    qualidade_cache = float(info_template.get("quality_score", local_template_preferencial.get("quality_score", 0.0)) or 0.0)
                    score_bruto_cache = float(info_template.get("raw_score", local_template_preferencial.get("score", 0.0)) or 0.0)
                    if (
                        score_cache >= max(0.86, float(confianca) - 0.02)
                        and qualidade_cache >= 0.74
                        and score_bruto_cache >= max(0.82, float(confianca) - 0.04)
                    ):
                        self.registrar_reconhecimento_bem_sucedido(
                            imagem_id,
                            local_template_preferencial,
                            janela_alvo=janela_alvo,
                            bloco_id=bloco_id,
                            janela_info=janela_info_cache,
                        )
                        safe_print(
                            "Imagem encontrada pelo cache rapido "
                            f"(captura={float(info_template.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                            f"reconhecimento={float(info_template.get('recognition_ms', 0.0) or 0.0):.0f}ms, "
                            f"area={self.formatar_regiao_log(regiao)}, confianca={score_cache:.2f}).",
                            flush=True,
                        )
                        return local_template_preferencial
                if local_template_preferencial and self.validar_correspondencia_visual(
                    caminho_imagem,
                    local_template_preferencial,
                    confianca,
                    perfil_busca=perfil_busca_tentativa,
                ):
                    local_confirmado = self.confirmar_localizacao_imagem(
                        caminho_imagem,
                        local_template_preferencial,
                        confianca,
                        perfil_busca=perfil_busca_tentativa,
                    )
                    if local_confirmado:
                        self.registrar_reconhecimento_bem_sucedido(
                            imagem_id,
                            local_confirmado,
                            janela_alvo=janela_alvo,
                            bloco_id=bloco_id,
                            janela_info=janela_info_cache,
                        )
                        info_template = self._ultima_busca_template_info or {}
                        safe_print(
                            "Imagem encontrada por confirmacao visual "
                            f"(etapa={etapa_busca}, captura={float(info_template.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                            f"reconhecimento={float(info_template.get('recognition_ms', 0.0) or 0.0):.0f}ms, "
                            f"area={self.formatar_regiao_log(regiao)}, confianca={float(info_template.get('score', 0.0) or 0.0):.2f}, "
                            f"fallback={'sim' if info_template.get('used_fallback') or entrada_regiao.get('fallback') else 'nao'}).",
                            flush=True,
                        )
                        return local_confirmado
                    safe_print("Correspondencia visual descartada apos confirmacao.", flush=True)

                if modo_estrito:
                    grayscale_modes = (False,)
                elif perfil_iteracao == "fast":
                    grayscale_modes = (False,)
                elif perfil_iteracao == "precise":
                    grayscale_modes = (False,) if tempo_decorrido < timeout_total * 0.72 else (False, True)
                else:
                    grayscale_modes = (False,) if tempo_decorrido < timeout_total * 0.4 else (False, True)
                permitir_locate = (
                    perfil_iteracao == "precise"
                    and not modo_estrito
                    and tempo_decorrido >= timeout_total * 0.72
                )
                if permitir_locate:
                    for confianca_atual in confiancas_busca:
                        for grayscale in grayscale_modes:
                            try:
                                if usa_confianca:
                                    local = pyautogui.locateOnScreen(
                                        str(caminho_imagem),
                                        confidence=confianca_atual,
                                        grayscale=grayscale,
                                        region=regiao,
                                    )
                                else:
                                    local = pyautogui.locateOnScreen(
                                        str(caminho_imagem),
                                        grayscale=grayscale,
                                        region=regiao,
                                    )
                            except TypeError as erro:
                                raise RuntimeError(
                                    "Falha ao usar confidence na busca da imagem. Verifique se opencv-python esta instalado corretamente."
                                ) from erro
                            except OSError as erro:
                                raise RuntimeError(f"PyAutoGUI nao conseguiu capturar a tela para localizar a imagem: {erro}") from erro
                            except Exception as erro:
                                if "needle dimension" in str(erro).lower() or "haystack" in str(erro).lower():
                                    if not self._needle_maior_logada:
                                        safe_print(
                                            "A referencia visual eh maior que a area de busca atual. "
                                            "Verifique a janela alvo ou reduza o recorte da imagem.",
                                            flush=True,
                                        )
                                        self._needle_maior_logada = True
                                    ultima_falha = None
                                    local = None
                                    continue
                                ultima_falha = erro
                                local = None

                            if local and self.validar_correspondencia_visual(
                                caminho_imagem,
                                local,
                                confianca_atual,
                                perfil_busca=perfil_iteracao,
                            ):
                                local_confirmado = self.confirmar_localizacao_imagem(
                                    caminho_imagem,
                                    local,
                                    confianca_atual,
                                    perfil_busca=perfil_busca,
                                )
                                if local_confirmado:
                                    self.registrar_reconhecimento_bem_sucedido(
                                        imagem_id,
                                        local_confirmado,
                                        janela_alvo=janela_alvo,
                                        bloco_id=bloco_id,
                                        janela_info=janela_info_cache,
                                    )
                                    if grayscale:
                                        safe_print(
                                            f"Imagem encontrada em modo tons de cinza (conf={confianca_atual:g})",
                                            flush=True,
                                        )
                                    else:
                                        safe_print(f"Imagem encontrada (conf={confianca_atual:g})", flush=True)
                                    return local_confirmado
                                safe_print("Correspondencia visual descartada apos confirmacao.", flush=True)

                        if (
                            perfil_iteracao == "precise"
                            or confianca_atual != confianca
                            or (perfil_iteracao == "balanced" and tempo_decorrido >= timeout_total * 0.82)
                        ):
                            local_template = self.procurar_imagem_por_template(
                                caminho_imagem,
                                confianca=confianca_atual,
                                region=regiao,
                                perfil_busca=perfil_iteracao,
                            )
                            if local_template and self.validar_correspondencia_visual(
                                caminho_imagem,
                                local_template,
                                confianca_atual,
                                perfil_busca=perfil_iteracao,
                            ):
                                local_confirmado = self.confirmar_localizacao_imagem(
                                    caminho_imagem,
                                    local_template,
                                    confianca_atual,
                                    perfil_busca=perfil_iteracao,
                                )
                                if local_confirmado:
                                    self.registrar_reconhecimento_bem_sucedido(
                                        imagem_id,
                                        local_confirmado,
                                        janela_alvo=janela_alvo,
                                        bloco_id=bloco_id,
                                        janela_info=janela_info_cache,
                                    )
                                    info_template = self._ultima_busca_template_info or {}
                                    safe_print(
                                        "Imagem encontrada por template confirmado "
                                        f"(etapa={etapa_busca}, captura={float(info_template.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                                        f"reconhecimento={float(info_template.get('recognition_ms', 0.0) or 0.0):.0f}ms, "
                                        f"area={self.formatar_regiao_log(regiao)}, confianca={float(info_template.get('score', 0.0) or 0.0):.2f}, "
                                        f"fallback={'sim' if info_template.get('used_fallback') or entrada_regiao.get('fallback') else 'nao'}).",
                                        flush=True,
                                    )
                                    return local_confirmado
                                safe_print("Correspondencia por template descartada apos confirmacao.", flush=True)

            intervalo_ocr_ancora = 0.9 if perfil_busca == "fast" else 1.25
            permitir_ocr_ancora = (
                texto_ocr_ancora
                and OCR_DISPONIVEL
                and dimensoes_referencia
                and not hibrido_ocr_visual_obrigatorio
                and not modo_estrito
                and time.perf_counter() - ultima_tentativa_ocr >= intervalo_ocr_ancora
                and time.perf_counter() - self._ultimo_ocr_global_em >= intervalo_ocr_ancora
            )
            if permitir_ocr_ancora:
                ultima_tentativa_ocr = time.perf_counter()
                self._ultimo_ocr_global_em = ultima_tentativa_ocr
                regioes_ocr = regioes_busca or [{"region": None, "stage": "screen", "fallback": False}]
                for entrada_ocr in regioes_ocr:
                    regiao_ocr = entrada_ocr.get("region")
                    local_ancora_visual = self.localizar_referencia_por_recorte_ocr_visual(
                        caminho_imagem,
                        region=regiao_ocr,
                        leitura_relativa=leitura_relativa,
                        dimensoes_referencia=dimensoes_referencia,
                        confianca=0.76 if perfil_busca == "fast" else 0.80,
                    )
                    if local_ancora_visual:
                        self.registrar_reconhecimento_bem_sucedido(
                            imagem_id,
                            local_ancora_visual,
                            janela_alvo=janela_alvo,
                            bloco_id=bloco_id,
                            janela_info=janela_info_cache,
                        )
                        return local_ancora_visual
                    safe_print(
                        f"Tentando localizar pela ancora OCR: {texto_ocr_ancora} "
                        f"(area={self.formatar_regiao_log(regiao_ocr)}).",
                        flush=True,
                    )
                    local_ancora = self.localizar_referencia_por_ancora_ocr(
                        caminho_imagem,
                        texto_ocr_esperado=texto_ocr_ancora,
                        region=regiao_ocr,
                        leitura_relativa=leitura_relativa,
                        dimensoes_referencia=dimensoes_referencia,
                    )
                    if local_ancora:
                        self.registrar_reconhecimento_bem_sucedido(
                            imagem_id,
                            local_ancora,
                            janela_alvo=janela_alvo,
                            bloco_id=bloco_id,
                            janela_info=janela_info_cache,
                        )
                        return local_ancora

            intervalo_ocr = max(1.8, min(3.2, timeout_total * 0.45))
            permitir_ocr = (
                permitir_ocr_geral
                and not hibrido_ocr_visual_obrigatorio
                and not modo_estrito
                and len(str(texto_referencia or "").strip()) >= 6
            )
            limite_inicio_ocr = timeout_total * 0.82
            if (
                permitir_ocr
                and texto_referencia
                and regiao is not None
                and tempo_decorrido >= limite_inicio_ocr
                and time.perf_counter() - ultima_tentativa_ocr >= intervalo_ocr
                and time.perf_counter() - self._ultimo_ocr_global_em >= intervalo_ocr
            ):
                ultima_tentativa_ocr = time.perf_counter()
                self._ultimo_ocr_global_em = ultima_tentativa_ocr
                safe_print(f"Tentando reconhecimento por texto (OCR): {texto_referencia}", flush=True)
                local_ocr = self.procurar_texto_por_ocr(texto_referencia, region=regiao)
                if local_ocr:
                    self.registrar_reconhecimento_bem_sucedido(
                        imagem_id,
                        local_ocr,
                        janela_alvo=janela_alvo,
                        bloco_id=bloco_id,
                        janela_info=janela_info_cache,
                    )
                    info_ocr = self._ultima_busca_ocr_info or {}
                    safe_print(
                        "OCR confirmou a referencia "
                        f"(captura={float(info_ocr.get('capture_ms', 0.0) or 0.0):.0f}ms, "
                        f"ocr={float(info_ocr.get('ocr_ms', 0.0) or 0.0):.0f}ms, "
                        f"area={self.formatar_regiao_log(regiao)}, confianca={float(info_ocr.get('score', 0.0) or 0.0):.2f}).",
                        flush=True,
                    )
                    return local_ocr

            if perfil_busca == "fast":
                time.sleep(0.02)
            elif perfil_busca == "balanced":
                time.sleep(0.03)
            else:
                time.sleep(0.025)

        safe_print("Imagem nao encontrada", flush=True)
        safe_print("Reconhecimento finalizado sem encontrar correspondencia.", flush=True)
        if self._ultimo_template_debug:
            safe_print(
                "Melhor aproximacao por template: "
                f"score={float(self._ultimo_template_debug.get('score', 0.0)):.2f} "
                f"escala={float(self._ultimo_template_debug.get('scale', 1.0)):.2f}",
                flush=True,
            )
        diagnostico = self.salvar_diagnostico_reconhecimento(
            caminho_imagem,
            region=self._ultima_regiao_reconhecimento,
            titulo_janela=str(janela_alvo or self.janela_alvo or ""),
        )
        if diagnostico:
            safe_print(f"Diagnostico salvo em: {diagnostico}", flush=True)
        if ultima_falha is not None:
            raise RuntimeError(f"Falha ao procurar imagem na tela: {ultima_falha}") from ultima_falha
        return None

    def clicar_imagem(
        self,
        imagem,
        timeout=10,
        confianca=0.8,
        offset_x=0,
        offset_y=0,
        clicar=True,
        clicks=1,
        janela_alvo=None,
        perfil_busca="balanced",
        bloco_id=None,
        modo_area_busca="target_window",
        area_personalizada=None,
        texto_ocr_esperado="",
        leitura_relativa=None,
        preservar_janelas=True,
    ):
        caminho_imagem, _imagem_id = self.resolver_imagem(imagem)
        dimensoes_referencia = self.obter_dimensoes_imagem(caminho_imagem) if caminho_imagem.exists() else None
        local = self.esperar_imagem(
            imagem=imagem,
            timeout=timeout,
            confianca=confianca,
            janela_alvo=janela_alvo,
            perfil_busca=perfil_busca,
            bloco_id=bloco_id,
            modo_area_busca=modo_area_busca,
            area_personalizada=area_personalizada,
            texto_ocr_esperado=texto_ocr_esperado,
            leitura_relativa=leitura_relativa,
            preservar_janelas=preservar_janelas,
        )
        if not local:
            return None
        fluxo_individual = local.get("ocr_target_flow") if isinstance(local, dict) else None
        if isinstance(fluxo_individual, dict) and str(fluxo_individual.get("acao") or "") in {"continuar", "parar"}:
            # A regra de falha de um OCR individual encerra a busca do bloco
            # antes da confirmacao visual/click padrao.
            return local
        hibrido_ocr_visual = bool(
            texto_ocr_esperado
            and dimensoes_referencia
            and float(confianca or 0.8) < 0.99
        )
        if hibrido_ocr_visual and not bool(local.get("visual_confirmed")):
            safe_print(
                "Hibrido OCR + imagem bloqueado: os OCRs foram lidos, mas a imagem ainda nao confirmou a referencia.",
                flush=True,
            )
            return None
        ordem_bloco = None
        if str(bloco_id or "").startswith("bloco:"):
            ordem_bloco = str(bloco_id).split(":", 1)[1].strip()
        leitura = {
            **(leitura_relativa if isinstance(leitura_relativa, dict) else {}),
            # Mantem o mesmo ponto "Clique" configurado na aba Busca para
            # uma acao individual de OCR.
            "_click_offset_x": offset_x,
            "_click_offset_y": offset_y,
            "_dimensoes_referencia": dimensoes_referencia,
        }
        salvar_em = str(leitura.get("salvar_em", "") or "").strip()
        ocr_ja_confirmado = (
            isinstance(local, dict)
            and (
                self.ocr_hibrido_tem_todos_alvos(local, texto_ocr_esperado, leitura_relativa)
                or (
                    local.get("ocr_target_individual")
                    and local.get("visual_confirmed")
                    and local.get("ocr_anchor_confirmed", True)
                )
                or (
                    local.get("ocr_visual_anchor")
                    and texto_ocr_esperado
                    and not salvar_em
                    and self.quantidade_alvos_ocr_esperados(texto_ocr_esperado, leitura_relativa) <= 1
                )
            )
        )
        if texto_ocr_esperado and ocr_ja_confirmado:
            capturado_ocr = str(local.get("ocr_text") or texto_ocr_esperado).strip()
            if capturado_ocr:
                self.salvar_variavel_contexto("ultimo_ocr", capturado_ocr)
                if ordem_bloco:
                    self.salvar_variavel_contexto(f"ocr_bloco_{ordem_bloco}", capturado_ocr)
                    self.salvar_variavel_contexto(f"bloco_{ordem_bloco}_ocr", capturado_ocr)
                if salvar_em:
                    self.salvar_variavel_contexto(salvar_em, capturado_ocr)
            safe_print(
                f"OCR confirmou pela ancora: esperado='{texto_ocr_esperado}' | capturado='{capturado_ocr}'",
                flush=True,
            )
            ocr_bounds = local.get("ocr_bounds") if isinstance(local.get("ocr_bounds"), dict) else None
            if ocr_bounds and local.get("search_stage") not in {"ocr_visual_hybrid", "ocr_text_fallback"}:
                self.marcar_regiao_ocr(
                    (
                        int(ocr_bounds.get("left", 0) or 0),
                        int(ocr_bounds.get("top", 0) or 0),
                        int(ocr_bounds.get("width", 0) or 0),
                        int(ocr_bounds.get("height", 0) or 0),
                    ),
                    confirmado=True,
                    duracao=0.8,
                )
        elif not self.confirmar_ocr_relativo(
            local,
            texto_esperado=texto_ocr_esperado,
            leitura_relativa=leitura_relativa,
            ordem_bloco=ordem_bloco,
            dimensoes_referencia=dimensoes_referencia,
            caminho_referencia=caminho_imagem,
        ):
            safe_print(
                "Imagem encontrada, mas o OCR obrigatorio nao confirmou. "
                "O bloco nao vai clicar nem avancar nesta tentativa.",
                flush=True,
            )
            return None
        acao_ocr_individual = self.aplicar_acao_individual_ocr(local, leitura_relativa=leitura)
        if acao_ocr_individual == "cancelado":
            return None
        if acao_ocr_individual in {"clicar", "ir_para_bloco"}:
            return local

        # A borda verde continua informando a confirmacao, so que com uma
        # exibicao curta para nao atrasar cada proximo bloco da automacao.
        self.marcar_area_reconhecida(local, duracao=0.28, bloqueante=True, etiqueta="Reconhecimento")
        offset_click_x, offset_click_y = self.escalar_offset_referencia(
            local,
            offset_x=offset_x,
            offset_y=offset_y,
            dimensoes_referencia=dimensoes_referencia,
        )
        if int(offset_x or 0) or int(offset_y or 0):
            safe_print(
                "Ponto do clique ajustado pela referencia: "
                f"offset {int(offset_x or 0)}/{int(offset_y or 0)} -> "
                f"{offset_click_x}/{offset_click_y}",
                flush=True,
            )
        if clicar:
            time.sleep(0.04)
            destino_click_x, destino_click_y = self.centro_area_reconhecida(
                local,
                offset_x=offset_click_x,
                offset_y=offset_click_y,
            )
            self.marcar_regiao_absoluta(
                int(destino_click_x) - 9,
                int(destino_click_y) - 9,
                18,
                18,
                duracao=0.12,
                bloqueante=True,
                tema="visual",
                etiqueta="Clique",
                margem=2,
            )
            return self.clicar_area_reconhecida(
                local,
                offset_x=offset_click_x,
                offset_y=offset_click_y,
                button="left",
                clicks=max(1, int(clicks or 1)),
                janela_alvo=janela_alvo,
                preservar_janelas=preservar_janelas,
            )
        destino_x, destino_y = self.centro_area_reconhecida(local, offset_x=offset_click_x, offset_y=offset_click_y)
        return {"x": destino_x, "y": destino_y}

    def testar_reconhecimento(
        self,
        imagem,
        timeout=10,
        confianca=0.8,
        janela_alvo=None,
        clicar=False,
        perfil_busca="balanced",
        bloco_id=None,
        modo_area_busca="target_window",
        area_personalizada=None,
        texto_ocr_esperado="",
        leitura_relativa=None,
        mostrar_marcador=True,
        marcador_bloqueante=True,
        duracao_marcador=0.24,
    ):
        leitura_relativa = dict(leitura_relativa or {})
        leitura_relativa["_mostrar_marcadores_ocr"] = bool(mostrar_marcador)
        leitura_relativa["_testar_todos_alvos_ocr"] = bool(
            texto_ocr_esperado and self.modo_acao_alvos_ocr(leitura_relativa) == "individual"
        )
        titulo_janela = str(janela_alvo or self.janela_alvo or "").strip()
        if titulo_janela:
            safe_print(f"Janela alvo configurada para teste: {titulo_janela}", flush=True)
        # No teste isolado, nao forcamos foco da janela alvo.
        # Isso evita fechar popups, menus ou estados transitivos
        # que a sequencia real deixa abertos antes do reconhecimento.
        time.sleep(0.06)

        janela_teste = self.localizar_janela_alvo(titulo_janela) if titulo_janela else None
        modo_area_teste = modo_area_busca
        area_teste = area_personalizada
        janela_ativa_teste = None
        regiao_teste = None
        if self.normalizar_modo_area_busca(modo_area_busca) == "custom_area":
            regiao_teste = self.construir_regiao_personalizada(area_personalizada, janela_info=janela_teste)
        if regiao_teste is None and titulo_janela:
            regiao_teste = self.obter_regiao_busca(titulo_janela)

        # Ao usar o botao Procurar, o Studio e minimizado antes do worker
        # iniciar. O elemento pode estar em um modal cujo titulo nao coincide
        # com a janela gravada no bloco. Para o teste isolado, os pixels visiveis
        # sao a fonte de verdade: priorizamos a janela externa em primeiro plano
        # sem ativa-la, move-la ou depender do titulo antigo.
        if (
            bool(getattr(self, "modo_teste_reconhecimento", False))
            and self.normalizar_modo_area_busca(modo_area_busca) == "target_window"
        ):
            janela_ativa_teste = self.obter_janela_ativa()
            if (
                isinstance(janela_ativa_teste, dict)
                and not self.janela_parece_interface_flowdetect(janela_ativa_teste)
                and int(janela_ativa_teste.get("width", 0) or 0) >= 260
                and int(janela_ativa_teste.get("height", 0) or 0) >= 180
            ):
                titulo_janela = str(janela_ativa_teste.get("title") or titulo_janela or "").strip()
                self.janela_alvo = titulo_janela
                janela_teste = janela_ativa_teste
                area_teste = {
                    "x": int(janela_ativa_teste.get("left", 0) or 0),
                    "y": int(janela_ativa_teste.get("top", 0) or 0),
                    "width": int(janela_ativa_teste.get("width", 0) or 0),
                    "height": int(janela_ativa_teste.get("height", 0) or 0),
                }
                modo_area_teste = "custom_area"
                regiao_teste = self.construir_regiao_personalizada(area_teste)
                safe_print(
                    "Teste rapido: usando a janela externa em primeiro plano "
                    f"({janela_ativa_teste.get('title', 'janela atual')}).",
                    flush=True,
                )
        if os.environ.get("FLOWDETECT_DEBUG_RECONHECIMENTO", "").strip() == "1":
            captura_bruta = self.salvar_captura_bruta_reconhecimento(region=regiao_teste, titulo_janela=titulo_janela)
            if captura_bruta:
                safe_print(f"Captura bruta do teste salva em: {captura_bruta}", flush=True)

        inicio_teste = time.perf_counter()
        timeout_total = max(0.1, float(timeout or 0.1))
        timeout_primeira_tentativa = timeout_total
        if janela_ativa_teste and modo_area_teste == "custom_area" and not leitura_relativa["_testar_todos_alvos_ocr"]:
            timeout_primeira_tentativa = min(timeout_total, 2.5)
        local = self.esperar_imagem(
            imagem=imagem,
            timeout=timeout_primeira_tentativa,
            confianca=confianca,
            janela_alvo=titulo_janela,
            perfil_busca=perfil_busca,
            bloco_id=bloco_id,
            modo_area_busca=modo_area_teste,
            area_personalizada=area_teste,
            texto_ocr_esperado=texto_ocr_esperado,
            leitura_relativa=leitura_relativa,
        )
        if not local and janela_ativa_teste and modo_area_teste == "custom_area":
            restante = timeout_total - (time.perf_counter() - inicio_teste)
            if restante > 0.12:
                safe_print(
                    "Teste rapido nao encontrou na janela em primeiro plano; ampliando a busca para a tela.",
                    flush=True,
                )
                local = self.esperar_imagem(
                    imagem=imagem,
                    timeout=restante,
                    confianca=confianca,
                    janela_alvo="",
                    perfil_busca=perfil_busca,
                    bloco_id=bloco_id,
                    modo_area_busca=modo_area_busca,
                    area_personalizada=area_personalizada,
                    texto_ocr_esperado=texto_ocr_esperado,
                    leitura_relativa=leitura_relativa,
                )
        if not local:
            safe_print("Teste de reconhecimento: nenhuma correspondencia encontrada.", flush=True)
            return False

        diagnostico = local.get("ocr_validation") if isinstance(local, dict) else None
        if isinstance(diagnostico, dict):
            # O modo individual ja leu os alvos na ordem e confirmou a imagem.
            # Revalidar aqui como OCR composto repetia 1/2 e alterava o resultado.
            safe_print("OCR_VALIDATION_JSON|" + json.dumps(diagnostico, ensure_ascii=True), flush=True)
            if mostrar_marcador and diagnostico.get("visual_confirmed"):
                self.marcar_area_reconhecida(
                    local, duracao=self.duracao_confirmacao_visual(duracao_marcador),
                    bloqueante=bool(marcador_bloqueante), etiqueta="Reconhecimento",
                )
            sucesso = bool(diagnostico.get("success"))
            if sucesso:
                safe_print("Teste concluido: todos os OCRs e a imagem confirmaram.", flush=True)
            else:
                falhas = ", ".join(
                    f"OCR {item['index'] + 1}" for item in diagnostico["targets"] if item["status"] != "confirmed"
                )
                safe_print(f"Teste concluido com falha em: {falhas}.", flush=True)
            return sucesso

        esquerda, topo, largura, altura = self.bounds_reconhecimento(local)
        safe_print(
            f"Teste de reconhecimento localizou a referencia em x={esquerda}, y={topo}, largura={largura}, altura={altura}.",
            flush=True,
        )
        caminho_imagem, _imagem_id = self.resolver_imagem(imagem)
        dimensoes_referencia = self.obter_dimensoes_imagem(caminho_imagem) if caminho_imagem.exists() else None
        if texto_ocr_esperado and isinstance(local, dict) and self.ocr_hibrido_tem_todos_alvos(local, texto_ocr_esperado, leitura_relativa):
            capturado_ocr = str(local.get("ocr_text") or texto_ocr_esperado).strip()
            safe_print(
                f"Teste de reconhecimento: OCR confirmado pela ancora '{capturado_ocr}'.",
                flush=True,
            )
            ocr_bounds = local.get("ocr_bounds") if isinstance(local.get("ocr_bounds"), dict) else None
            if mostrar_marcador and ocr_bounds and local.get("search_stage") not in {"ocr_visual_hybrid", "ocr_text_fallback"}:
                self.marcar_regiao_ocr(
                    (
                        int(ocr_bounds.get("left", 0) or 0),
                        int(ocr_bounds.get("top", 0) or 0),
                        int(ocr_bounds.get("width", 0) or 0),
                        int(ocr_bounds.get("height", 0) or 0),
                    ),
                    confirmado=True,
                    duracao=max(0.12, float(duracao_marcador or 0.24)),
                )
        elif not self.confirmar_ocr_relativo(
            local,
            texto_esperado=texto_ocr_esperado,
            leitura_relativa=leitura_relativa,
            ordem_bloco=None,
            dimensoes_referencia=dimensoes_referencia,
            marcar_ocr=mostrar_marcador,
            caminho_referencia=caminho_imagem,
        ):
            safe_print("Teste de reconhecimento: imagem encontrada, mas o OCR nao confirmou o texto esperado.", flush=True)
            return False
        if mostrar_marcador:
            self.marcar_area_reconhecida(
                local,
                duracao=self.duracao_confirmacao_visual(duracao_marcador),
                bloqueante=bool(marcador_bloqueante),
            )
        if clicar:
            self.clicar_area_reconhecida(local, button="left")

        safe_print("Teste de reconhecimento concluido com sucesso.", flush=True)
        return True

    def verificar_estado_botao(
        self,
        imagem_ativo,
        imagem_desativado,
        timeout=5,
        confianca=0.85,
        se_ativo=None,
        se_desativado=None,
        se_nao_encontrar=None,
    ):
        safe_print(f"Procurando botao ativo: {imagem_ativo}", flush=True)
        ativo = self.esperar_imagem(imagem_ativo, timeout=timeout, confianca=confianca)
        if ativo:
            safe_print("Botao ativo encontrado", flush=True)
            self.marcar_area_reconhecida(ativo, duracao=1.35)
            time.sleep(0.18)
            return self.interpretar_config_acao(se_ativo, "continuar")

        safe_print(f"Procurando botao desativado: {imagem_desativado}", flush=True)
        desativado = self.esperar_imagem(imagem_desativado, timeout=timeout, confianca=confianca)
        if desativado:
            safe_print("Botao desativado encontrado", flush=True)
            self.marcar_area_reconhecida(desativado, duracao=1.35)
            time.sleep(0.18)
            return self.interpretar_config_acao(se_desativado, "continuar")

        safe_print("Nenhum estado encontrado", flush=True)
        return self.interpretar_config_acao(se_nao_encontrar, "parar")

    def esperar(self, segundos):
        fim = time.perf_counter() + max(0, float(segundos))
        while time.perf_counter() < fim:
            if self.cancelamento_solicitado():
                return False
            time.sleep(0.05)
        return True

    def executar_acao(self, acao):
        if self.cancelamento_solicitado():
            return False
        partes = self.normalizar_acao(acao)
        if not partes:
            return True
        if len(partes) == 1:
            pyautogui.press(partes[0])
            return not self.cancelamento_solicitado()
        pyautogui.hotkey(*partes)
        return not self.cancelamento_solicitado()

    def normalizar_url_para_abertura(self, url):
        valor = str(url or "").strip()
        if not valor:
            raise RuntimeError("Bloco Abrir URL sem URL configurada.")
        if not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://", valor):
            valor = f"https://{valor}"
        return valor

    def localizar_executavel_navegador(self, navegador):
        navegador = str(navegador or "padrao").strip().lower()
        if navegador == "chrome":
            candidatos = [
                shutil.which("chrome"),
                shutil.which("chrome.exe"),
                r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
            ]
        elif navegador == "edge":
            candidatos = [
                shutil.which("msedge"),
                shutil.which("msedge.exe"),
                r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            ]
        elif navegador in {"padrao", "default", ""}:
            candidatos = [
                shutil.which("chrome"),
                shutil.which("chrome.exe"),
                r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
                shutil.which("msedge"),
                shutil.which("msedge.exe"),
                r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
            ]
        else:
            return None
        for candidato in candidatos:
            if candidato and Path(str(candidato)).exists():
                return str(candidato)
        return None

    def aguardar_e_fixar_janela_navegador(self, url="", navegador="padrao", timeout=6.0, tela_cheia=False):
        fim = time.perf_counter() + max(0.5, float(timeout or 0))
        texto_url = self.normalizar_texto_busca(url)
        compacto_url = re.sub(r"[^a-z0-9]+", "", texto_url)
        termos_url = [
            termo
            for termo in re.split(r"[^a-zA-Z0-9]+", str(url or "").lower())
            if len(termo) >= 3 and termo not in {"http", "https", "www", "com", "br", "gov"}
        ]
        navegador = str(navegador or "padrao").strip().lower()
        melhor = None
        melhor_score = -1

        while time.perf_counter() <= fim:
            if self.cancelamento_solicitado():
                return None
            for janela in self.listar_janelas_visiveis():
                if self.janela_parece_interface_flowdetect(janela):
                    continue
                if int(janela.get("width", 0) or 0) < 360 or int(janela.get("height", 0) or 0) < 240:
                    continue
                titulo = self.normalizar_texto_busca(janela.get("title", ""))
                classe = self.normalizar_texto_busca(janela.get("class_name", ""))
                compacto_titulo = re.sub(r"[^a-z0-9]+", "", titulo)
                score = 0
                if self.janela_parece_navegador(janela):
                    score += 50
                if navegador == "chrome" and "chrome_widgetwin" in classe:
                    score += 18
                if navegador == "edge" and ("edge" in titulo or "applicationframe" in classe):
                    score += 14
                if compacto_url and compacto_url[:18] and compacto_url[:18] in compacto_titulo:
                    score += 35
                for termo in termos_url:
                    termo_norm = self.normalizar_texto_busca(termo)
                    termo_compacto = re.sub(r"[^a-z0-9]+", "", termo_norm)
                    if termo_norm and termo_norm in titulo:
                        score += 12
                    elif termo_compacto and termo_compacto in compacto_titulo:
                        score += 10
                if "nova guia" in titulo or "new tab" in titulo:
                    score -= 8
                if score > melhor_score:
                    melhor = janela
                    melhor_score = score

            if melhor and melhor_score >= 50:
                titulo_original = str(melhor.get("title") or "").strip()
                self._janela_contexto = titulo_original
                self._janela_fixada = melhor
                try:
                    self.preparar_janela_para_acao(titulo_original)
                    if not self.maximizar_janela_por_hwnd(melhor.get("hwnd")):
                        pyautogui.hotkey("win", "up")
                except Exception:
                    pass
                safe_print(
                    f"Janela do navegador fixada para proximos reconhecimentos: {titulo_original}",
                    flush=True,
                )
                return melhor
            time.sleep(0.25)

        if melhor and melhor_score >= 35:
            titulo_original = str(melhor.get("title") or "").strip()
            self._janela_contexto = titulo_original
            self._janela_fixada = melhor
            try:
                self.preparar_janela_para_acao(titulo_original)
                if not self.maximizar_janela_por_hwnd(melhor.get("hwnd")):
                    pyautogui.hotkey("win", "up")
            except Exception:
                pass
            safe_print(
                f"Janela provavel do navegador fixada: {titulo_original}",
                flush=True,
            )
            return melhor
        return None

    def executar_bloco_abrir_url(self, evento):
        detalhes = evento.get("detalhes") or {}
        url = self.normalizar_url_para_abertura(detalhes.get("url"))
        navegador = str(detalhes.get("navegador", "padrao") or "padrao").strip().lower()
        modo_abertura = str(detalhes.get("modo_abertura", "nova_janela") or "nova_janela").strip().lower()
        aguardar = max(0.0, float(detalhes.get("aguardar_segundos", 3) or 0))
        tela_cheia = bool(detalhes.get("tela_cheia", False))
        nova_janela = modo_abertura != "nova_aba"
        safe_print(
            f"Abrindo URL: {url} | navegador={navegador} | modo={modo_abertura}",
            flush=True,
        )
        executavel = self.localizar_executavel_navegador(navegador)
        if executavel:
            argumentos = [executavel]
            nome_executavel = Path(executavel).name.lower()
            navegador_automatizavel = navegador in {"chrome", "edge", "padrao", "default", ""} or nome_executavel in {"chrome.exe", "msedge.exe"}
            if navegador_automatizavel:
                argumentos.extend([
                    "--start-fullscreen" if tela_cheia else "--start-maximized",
                    "--force-renderer-accessibility",
                    "--enable-accessibility-object-model",
                ])
            if nova_janela:
                argumentos.append("--new-window")
            argumentos.append(url)
            popen_silencioso(argumentos)
        else:
            webbrowser.open(url, new=1 if nova_janela else 2)
        janela = self.aguardar_e_fixar_janela_navegador(
            url=url,
            navegador=navegador,
            timeout=max(aguardar, 5.0),
            tela_cheia=tela_cheia,
        )
        if janela is None:
            safe_print(
                "Aviso: a URL foi solicitada, mas nenhuma janela de navegador ficou visivel para fixar. "
                "Os proximos reconhecimentos podem capturar a area de trabalho.",
                flush=True,
            )
        if janela is None and aguardar > 0 and not self.esperar(aguardar):
            return self.resultado_fluxo("parar")
        return self.resultado_sucesso_evento(evento)

    def executar_bloco_abrir_programa(self, evento):
        detalhes = evento.get("detalhes") or {}
        caminho = str(detalhes.get("caminho", "") or "").strip()
        if not caminho:
            raise RuntimeError("Bloco Abrir Programa sem executavel configurado.")
        parametros = str(detalhes.get("parametros", "") or "").strip()
        aguardar = max(0.0, float(detalhes.get("aguardar_segundos", 2) or 0))
        try:
            argumentos_extra = shlex.split(parametros, posix=False) if parametros else []
        except ValueError:
            argumentos_extra = [parametros] if parametros else []
        comando = [caminho, *argumentos_extra]
        cwd = None
        caminho_path = Path(caminho)
        if caminho_path.exists() and caminho_path.parent.exists():
            cwd = str(caminho_path.parent)
        safe_print(
            f"Abrindo programa: {caminho}" + (f" | parametros={parametros}" if parametros else ""),
            flush=True,
        )
        popen_silencioso(comando, cwd=cwd)
        if aguardar > 0 and not self.esperar(aguardar):
            return self.resultado_fluxo("parar")
        return self.resultado_sucesso_evento(evento)

    def executar_evento(self, evento, contexto_loop=None):
        tipo = evento.get("tipo")

        if tipo in TIPOS_MARCADORES_LOOP:
            return self.resultado_fluxo("continuar")

        if tipo == "abrir_url":
            return self.executar_bloco_abrir_url(evento)

        if tipo == "abrir_programa":
            return self.executar_bloco_abrir_programa(evento)

        if tipo == "texto_livre":
            self.preparar_janela_para_acao()
            return self.resultado_sucesso_evento(evento) if self.escrever_texto(evento.get("texto", "")) else self.resultado_fluxo("parar")

        if tipo in TIPOS_SEQUENCIA_PLANILHA:
            detalhes = evento.get("detalhes", {})
            self.preparar_janela_para_acao()
            acao_apos_item = detalhes.get("acao_apos_item")
            if acao_apos_item is None and detalhes.get("pressionar_enter"):
                acao_apos_item = ["enter"]

            if contexto_loop is not None:
                valor_loop = self.obter_valor_evento_planilha_no_loop(
                    evento,
                    contexto_loop.get("linha_referencia"),
                    contexto_loop.get("indice", 0),
                )
                if valor_loop is not None:
                    if not self.escrever_texto(str(valor_loop)):
                        return self.resultado_fluxo("parar")
                    if acao_apos_item and not self.executar_acao(acao_apos_item):
                        return self.resultado_fluxo("parar")
                return self.resultado_sucesso_evento(evento)

            valores = self.obter_valores_evento_planilha(evento)
            intervalo_entre_itens = detalhes.get(
                "intervalo_entre_itens",
                INTERVALO_ENTRE_ITENS_PLANILHA,
            )
            for indice, valor in enumerate(valores):
                if self.cancelamento_solicitado():
                    return self.resultado_fluxo("parar")
                if not self.escrever_texto(valor):
                    return self.resultado_fluxo("parar")
                if indice < len(valores) - 1 and acao_apos_item:
                    if not self.executar_acao(acao_apos_item):
                        return self.resultado_fluxo("parar")
                if indice < len(valores) - 1 and intervalo_entre_itens > 0:
                    if not self.esperar(intervalo_entre_itens):
                        return self.resultado_fluxo("parar")
            return self.resultado_sucesso_evento(evento)

        if tipo in TIPOS_ESCRITA_PLANILHA:
            return self.executar_bloco_escrever_planilha(evento, contexto_loop=contexto_loop)

        if tipo in TIPOS_COMPARACAO_PLANILHA:
            return self.executar_bloco_comparar_planilha(evento, contexto_loop=contexto_loop)

        if tipo in TIPOS_MOUSE:
            if self.cancelamento_solicitado():
                return self.resultado_fluxo("parar")
            self.preparar_janela_para_acao()
            detalhes = evento.get("detalhes", {})
            destino_x, destino_y = self.ajustar_coordenadas_para_janela_alvo(
                detalhes.get("x", 0),
                detalhes.get("y", 0),
            )
            pyautogui.click(
                x=destino_x,
                y=destino_y,
                button=str(detalhes.get("botao", "left")),
                clicks=max(1, int(detalhes.get("cliques", 1) or 1)),
            )
            return self.resultado_sucesso_evento(evento)

        if tipo in TIPOS_IMAGEM:
            detalhes = evento.get("detalhes", {})
            modo_reconhecimento = str(detalhes.get("modo_reconhecimento", "imagem") or "imagem").strip().lower()
            texto_ocr_hibrido = str(detalhes.get("texto_ocr_busca", "") or "").strip() if modo_reconhecimento == "texto_ocr" else ""
            ordem_evento = int(evento.get("ordem", 0) or 0)
            nome_imagem = str(detalhes.get("imagem") or "")
            janela_evento = self.resolver_janela_evento(
                detalhes.get("janela_alvo"),
                registrar_conflito=not bool(str(detalhes.get("janela_alvo") or "").strip()),
                priorizar_evento=bool(str(detalhes.get("janela_alvo") or "").strip()),
            )
            area_personalizada = {
                "x": detalhes.get("area_busca_x", 0),
                "y": detalhes.get("area_busca_y", 0),
                "width": detalhes.get("area_busca_largura", 0),
                "height": detalhes.get("area_busca_altura", 0),
            }
            tentativas = max(1, int(detalhes.get("quantidade_tentativas", 1) or 1))
            intervalo_tentativas = max(0.0, float(detalhes.get("intervalo_tentativas", 1) or 1))
            timeout_reconhecimento = max(0.1, float(detalhes.get("timeout", 10) or 10))
            ao_nao_encontrar = str(detalhes.get("ao_nao_encontrar", "parar") or "parar")
            ultima_falha_reconhecimento = None
            for tentativa in range(tentativas):
                if self.cancelamento_solicitado():
                    return False
                inicio_tentativa = time.perf_counter()
                janela_reconhecimento = str(
                    janela_evento
                    or self._janela_contexto
                    or self.janela_alvo
                    or ""
                ).strip()
                safe_print(
                    f"Bloco {ordem_evento}: reconhecendo {nome_imagem} | tentativa {tentativa + 1}/{tentativas}"
                    + (f" | janela={janela_reconhecimento}" if janela_reconhecimento else ""),
                    flush=True,
                )
                try:
                    modo_area_visual = self.normalizar_modo_area_busca(detalhes.get("modo_area_busca", "target_window"))
                    auto_scroll = str(detalhes.get("varredura_scroll", detalhes.get("rolagem_ao_nao_encontrar", "auto"))).strip().lower()
                    permitir_varredura_scroll = (
                        auto_scroll not in {"0", "false", "nao", "não", "off", "desativado"}
                        and modo_area_visual != "custom_area"
                    )
                    quantidade_scrolls = max(
                        0,
                        min(
                            8,
                            int(
                                detalhes.get(
                                    "varredura_scroll_tentativas",
                                    detalhes.get("rolagem_tentativas", 4 if permitir_varredura_scroll else 0),
                                )
                                or 0
                            ),
                        ),
                    )
                    pixels_scroll = max(120, int(detalhes.get("varredura_scroll_pixels", detalhes.get("rolagem_pixels", 560)) or 560))
                    timeout_curto_scroll = max(1.2, min(3.0, timeout_reconhecimento * 0.35))
                    resultado = None
                    for etapa_scroll in range(quantidade_scrolls + 1):
                        if self.cancelamento_solicitado():
                            return self.resultado_fluxo("parar")
                        resultado = self.clicar_imagem(
                            imagem=nome_imagem,
                            timeout=timeout_reconhecimento if etapa_scroll == 0 else timeout_curto_scroll,
                            confianca=detalhes.get("confianca", 0.8),
                            perfil_busca=detalhes.get("perfil_busca", "balanced"),
                            offset_x=detalhes.get("offset_x", 0),
                            offset_y=detalhes.get("offset_y", 0),
                            clicar=detalhes.get("clicar", True),
                            clicks=detalhes.get("cliques", 1 if detalhes.get("clicar", True) else 0),
                            janela_alvo=janela_evento,
                            bloco_id=f"bloco:{ordem_evento}",
                            modo_area_busca=detalhes.get("modo_area_busca", "target_window"),
                            area_personalizada=area_personalizada,
                            texto_ocr_esperado=texto_ocr_hibrido or detalhes.get("texto_ocr_esperado", ""),
                            leitura_relativa={**(detalhes.get("leitura_relativa") or {}), "ocr_target_layouts": detalhes.get("ocr_target_layouts") or []},
                        )
                        if resultado or etapa_scroll >= quantidade_scrolls:
                            break
                        if not permitir_varredura_scroll:
                            break
                        self.rolar_para_varredura_visual(
                            titulo_janela=janela_evento,
                            direcao=detalhes.get("varredura_scroll_direcao", detalhes.get("rolagem_direcao", "baixo")),
                            pixels=pixels_scroll,
                        )
                    ultima_falha_reconhecimento = None
                except Exception as erro:
                    ultima_falha_reconhecimento = erro
                    resultado = None
                    duracao_tentativa = time.perf_counter() - inicio_tentativa
                    restante = max(0.0, timeout_reconhecimento - duracao_tentativa)
                    if restante > 0 and not self.esperar(restante):
                        return self.resultado_fluxo("parar")
                    safe_print(
                        f"Falha no reconhecimento visual: {erro}",
                        flush=True,
                    )

                if resultado:
                    fluxo_ocr = resultado.get("ocr_target_flow") if isinstance(resultado, dict) else None
                    acao_fluxo_ocr = str(fluxo_ocr.get("acao") or "") if isinstance(fluxo_ocr, dict) else ""
                    if acao_fluxo_ocr == "parar":
                        safe_print(
                            f"Bloco {ordem_evento}: OCR individual acionou parar automacao.",
                            flush=True,
                        )
                        return self.resultado_fluxo("parar")
                    if acao_fluxo_ocr == "continuar":
                        safe_print(
                            f"Bloco {ordem_evento}: OCR individual nao confirmou; continuando a automacao.",
                            flush=True,
                        )
                        return self.resultado_fluxo("continuar")
                    if acao_fluxo_ocr == "ir_para_bloco":
                        bloco_destino_ocr = str(fluxo_ocr.get("bloco") or "").strip()
                        if bloco_destino_ocr:
                            safe_print(
                                f"Bloco {ordem_evento}: OCR individual confirmou {nome_imagem}. Indo para bloco {bloco_destino_ocr}.",
                                flush=True,
                            )
                            return self.resultado_fluxo("ir_para_bloco", bloco=bloco_destino_ocr)
                    safe_print(
                        f"Bloco {ordem_evento}: reconhecimento confirmou {nome_imagem}.",
                        flush=True,
                    )
                    return self.resultado_sucesso_evento(evento)

                if tentativa < tentativas - 1:
                    safe_print(
                        f"Bloco {ordem_evento}: reconhecimento nao encontrou {nome_imagem}. Tentativa {tentativa + 1}/{tentativas}.",
                        flush=True,
                    )
                    if intervalo_tentativas > 0 and not self.esperar(intervalo_tentativas):
                        return self.resultado_fluxo("parar")
                    continue

            if ao_nao_encontrar in {"pular", "continuar"}:
                safe_print(
                    f"Bloco {ordem_evento}: reconhecimento nao encontrou {nome_imagem} apos {tentativas} tentativa(s). Acao: continuar.",
                    flush=True,
                )
                # A busca que falhou nao deve reativar nem reorganizar janelas.
                # O proximo bloco prepara o proprio foco apenas se precisar digitar
                # ou executar uma acao nele.
                safe_print("Reconhecimento nao encontrado. Continuando a sequencia.", flush=True)
                return self.resultado_fluxo("continuar")

            if ao_nao_encontrar == "ir_para_sequencia":
                sequencia_destino = str(detalhes.get("sequencia_nao_encontrar", "") or "").strip()
                atraso_sequencia_destino = max(0.0, float(detalhes.get("atraso_antes_sequencia_nao_encontrar", 0) or 0))
                if sequencia_destino:
                    safe_print(
                        f"Bloco {ordem_evento}: reconhecimento nao encontrou {nome_imagem} apos {tentativas} tentativa(s). Acao: ir para sequencia {sequencia_destino}.",
                        flush=True,
                    )
                    safe_print(f"Redirecionando para sequencia {sequencia_destino}", flush=True)
                    return self.resultado_fluxo("ir_para_sequencia", sequencia_destino, atraso=atraso_sequencia_destino)
                return self.resultado_fluxo("continuar")

            if ao_nao_encontrar == "ir_para_bloco":
                bloco_destino = detalhes.get("bloco_nao_encontrar", "")
                atraso_bloco_destino = max(0.0, float(detalhes.get("atraso_antes_sequencia_nao_encontrar", 0) or 0))
                if str(bloco_destino or "").strip():
                    safe_print(
                        f"Bloco {ordem_evento}: reconhecimento nao encontrou {nome_imagem} apos {tentativas} tentativa(s). Acao: ir para bloco {bloco_destino}.",
                        flush=True,
                    )
                    safe_print(f"Redirecionando para bloco {bloco_destino}", flush=True)
                    return self.resultado_fluxo("ir_para_bloco", bloco=bloco_destino, atraso=atraso_bloco_destino)
                return self.resultado_fluxo("continuar")

            if ao_nao_encontrar == "tentar_novamente":
                mensagem = (
                    f"Bloco {ordem_evento}: reconhecimento/OCR nao confirmou {nome_imagem} "
                    f"apos {tentativas} tentativa(s). Parando para evitar repetir o mesmo bloco sem fim."
                )
                safe_print(mensagem, flush=True)
                raise RuntimeError(mensagem)

            complemento = ""
            if self._ultimo_diagnostico_reconhecimento:
                complemento = f" | diagnostico: {self._ultimo_diagnostico_reconhecimento}"
            if ultima_falha_reconhecimento is not None:
                raise RuntimeError(
                    f"Falha ao reconhecer imagem: {detalhes.get('imagem')} | erro: {ultima_falha_reconhecimento}{complemento}"
                ) from ultima_falha_reconhecimento
            raise RuntimeError(f"Imagem nao encontrada: {detalhes.get('imagem')}{complemento}")

            return self.resultado_fluxo("parar")

        if tipo in TIPOS_OCR:
            return self.executar_bloco_procurar_texto_ocr(evento)

        if tipo in TIPOS_VERIFICACAO:
            detalhes = evento.get("detalhes", {})
            resultado = self.verificar_estado_botao(
                imagem_ativo=detalhes.get("imagem_ativo"),
                imagem_desativado=detalhes.get("imagem_desativado"),
                timeout=detalhes.get("timeout", 5),
                confianca=detalhes.get("confianca", 0.85),
                se_ativo=detalhes.get("se_ativo"),
                se_desativado=detalhes.get("se_desativado"),
                se_nao_encontrar=detalhes.get("se_nao_encontrar"),
            )
            if resultado.get("acao") == "ir_para_sequencia" and resultado.get("sequencia"):
                safe_print(f"Redirecionando para sequencia {resultado['sequencia']}", flush=True)
            return resultado

        if tipo in TIPOS_CHROME:
            return self.executar_bloco_chrome_elemento_html(evento, contexto_loop=contexto_loop)

        if tipo in TIPOS_UIA:
            return self.executar_bloco_uia_elemento(evento, contexto_loop=contexto_loop)

        teclas = evento.get("teclas", [])
        if not teclas:
            return self.resultado_fluxo("continuar")

        self.preparar_janela_para_acao()

        if len(teclas) > 1:
            repeticoes = max(1, int((evento.get("detalhes") or {}).get("repeticoes", 1) or 1))
            for _ in range(repeticoes):
                if self.cancelamento_solicitado():
                    return self.resultado_fluxo("parar")
                pyautogui.hotkey(*[self.tecla_para_pyautogui(tecla) for tecla in teclas])
            return self.resultado_sucesso_evento(evento)

        tecla = teclas[0]
        repeticoes = max(1, int((evento.get("detalhes") or {}).get("repeticoes", 1) or 1))
        if self.eh_texto_simples(tecla):
            pyautogui.write(tecla * repeticoes, interval=INTERVALO_TEXTO)
            return self.resultado_sucesso_evento(evento)

        for _ in range(repeticoes):
            if self.cancelamento_solicitado():
                return self.resultado_fluxo("parar")
            pyautogui.press(self.tecla_para_pyautogui(tecla))
        return self.resultado_sucesso_evento(evento)

    def executar_fluxo_controlado(
        self,
        contexto_loop=None,
        sequencia_inicial=None,
        indice_inicial=0,
        parar_antes_da_sequencia="",
        limitar_a_sequencia="",
        parar_ao_entrar_em_loop=False,
        sequencias_loop_ativas=None,
        limitar_as_sequencias_loop=False,
        encerrar_iteracao_ao_voltar_para_sequencia="",
        sequencia_limite_blocos="",
        bloco_inicial_permitido=0,
        parar_antes_do_bloco=0,
    ):
        total_passos = self.calcular_total_passos_normal()
        passo_atual = 0
        sequencia_atual = (
            self.normalizar_nome_sequencia(sequencia_inicial)
            if str(sequencia_inicial or "").strip()
            else self.sequencia_inicial
        )
        indice = max(0, int(indice_inicial or 0))
        alvo_parada = (
            self.normalizar_nome_sequencia(parar_antes_da_sequencia)
            if str(parar_antes_da_sequencia or "").strip()
            else ""
        )
        sequencia_limitada = (
            self.normalizar_nome_sequencia(limitar_a_sequencia)
            if str(limitar_a_sequencia or "").strip()
            else ""
        )
        sequencias_loop = {
            self.normalizar_nome_sequencia(nome)
            for nome in (sequencias_loop_ativas or set())
            if str(nome or "").strip()
        }
        sequencia_reinicio_loop = (
            self.normalizar_nome_sequencia(encerrar_iteracao_ao_voltar_para_sequencia)
            if str(encerrar_iteracao_ao_voltar_para_sequencia or "").strip()
            else ""
        )
        sequencia_blocos = (
            self.normalizar_nome_sequencia(sequencia_limite_blocos)
            if str(sequencia_limite_blocos or "").strip()
            else ""
        )
        bloco_minimo = int(bloco_inicial_permitido or 0)
        bloco_maximo_exclusivo = int(parar_antes_do_bloco or 0)

        if alvo_parada and sequencia_atual == alvo_parada and indice == 0:
            return {
                "success": True,
                "status": "aguardando_sequencia",
                "sequencia": sequencia_atual,
                "indice": indice,
                "resultado": self.resultado_fluxo("continuar"),
            }

        if parar_ao_entrar_em_loop and sequencia_atual in sequencias_loop and indice == 0:
            return {
                "success": True,
                "status": "aguardando_loop",
                "sequencia": sequencia_atual,
                "indice": indice,
                "resultado": self.resultado_fluxo("continuar"),
            }

        self.emitir_progresso(
            0,
            total_passos,
            f"Iniciando sequencia {sequencia_atual}",
            sequencia=sequencia_atual,
            linha_atual=self.obter_linha_atual_contexto(contexto_loop),
            iteracao_atual=(contexto_loop or {}).get("indice", 0) + 1 if contexto_loop else 0,
            iteracao_total=(contexto_loop or {}).get("total", 0) if contexto_loop else 0,
        )

        limite_iteracoes = 0
        while True:
            limite_iteracoes += 1
            if limite_iteracoes > 10000:
                raise ValueError("O fluxo entrou em um ciclo inesperado durante a execucao.")

            eventos_sequencia = self.sequencias.get(sequencia_atual, [])
            if indice >= len(eventos_sequencia):
                if limitar_as_sequencias_loop and sequencia_atual in sequencias_loop:
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("continuar"),
                    }
                return {
                    "success": True,
                    "status": "concluido",
                    "sequencia": sequencia_atual,
                    "indice": indice,
                    "resultado": self.resultado_fluxo("continuar"),
                }

            evento = eventos_sequencia[indice]
            ordem_evento = int(evento.get("ordem", 0) or 0)
            if sequencia_blocos and sequencia_atual == sequencia_blocos:
                if bloco_minimo and ordem_evento < bloco_minimo:
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("ir_para_bloco", bloco=ordem_evento),
                    }
                if bloco_maximo_exclusivo and ordem_evento >= bloco_maximo_exclusivo:
                    return {
                        "success": True,
                        "status": "aguardando_bloco",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("continuar"),
                    }
                pass
            if self.cancelamento_solicitado():
                return {
                    "success": False,
                    "status": "cancelado",
                    "sequencia": sequencia_atual,
                    "indice": indice,
                    "resultado": self.resultado_fluxo("parar"),
                }

            intervalo_antes = evento.get("intervalo_antes", 0)
            if intervalo_antes > 0 and not self.esperar(intervalo_antes):
                return {
                    "success": False,
                    "status": "cancelado",
                    "sequencia": sequencia_atual,
                    "indice": indice,
                    "resultado": self.resultado_fluxo("parar"),
                }

            contabiliza_passo = evento.get("tipo") not in TIPOS_MARCADORES_LOOP
            if contabiliza_passo:
                passo_atual += 1
                descricao = self.descrever_evento(evento)
                empresa = ""
                if contexto_loop is not None:
                    empresa = contexto_loop.get("empresa", "")
                    descricao = (
                        f"Iteracao {contexto_loop.get('indice', 0) + 1}/{contexto_loop.get('total', 1)} - "
                        f"{descricao}"
                    )
                self.emitir_progresso(
                    passo_atual,
                    total_passos,
                    descricao,
                    empresa,
                    sequencia=sequencia_atual,
                    ordem=ordem_evento,
                    linha_atual=self.obter_linha_atual_contexto(contexto_loop),
                    iteracao_atual=(contexto_loop or {}).get("indice", 0) + 1 if contexto_loop else 0,
                    iteracao_total=(contexto_loop or {}).get("total", 0) if contexto_loop else 0,
                )

            resultado = self.executar_evento(evento, contexto_loop=contexto_loop)
            acao = str((resultado or {}).get("acao", "continuar"))
            destino = str((resultado or {}).get("sequencia", "")).strip()
            destino_bloco = (resultado or {}).get("bloco", "")
            atraso_destino = max(0.0, float((resultado or {}).get("atraso", 0) or 0))

            if acao == "parar":
                return {
                    "success": False,
                    "status": "parado",
                    "sequencia": sequencia_atual,
                    "indice": indice,
                    "resultado": resultado or self.resultado_fluxo("parar"),
                }
            if acao == "pular_card":
                indice += 2
                continue
            if acao == "ir_para_sequencia":
                if not destino:
                    raise ValueError("A acao ir_para_sequencia exige um nome de sequencia.")
                destino = self.normalizar_nome_sequencia(destino)
                if destino not in self.sequencias:
                    raise ValueError(f"Sequencia nao encontrada: {destino}")
                if (
                    sequencia_reinicio_loop
                    and limitar_as_sequencias_loop
                    and sequencia_atual in sequencias_loop
                    and sequencia_atual != sequencia_reinicio_loop
                    and destino == sequencia_reinicio_loop
                ):
                    if atraso_destino > 0:
                        safe_print(
                            f"Aguardando {atraso_destino:g}s antes de voltar para a pagina {destino}.",
                            flush=True,
                        )
                        if not self.esperar(atraso_destino):
                            return {
                                "success": False,
                                "status": "cancelado",
                                "sequencia": sequencia_atual,
                                "indice": indice,
                                "resultado": self.resultado_fluxo("parar"),
                            }
                    safe_print(
                        f"Retorno para a pagina inicial do loop ({destino}) detectado. "
                        "Avancando para a proxima linha da planilha.",
                        flush=True,
                    )
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": resultado,
                    }
                if sequencia_limitada and sequencia_atual == sequencia_limitada and destino != sequencia_limitada:
                    return {
                        "success": True,
                        "status": "saida_sequencia",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": resultado,
                    }
                if limitar_as_sequencias_loop and sequencia_atual in sequencias_loop and destino not in sequencias_loop:
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": resultado,
                    }
                if atraso_destino > 0:
                    safe_print(
                        f"Aguardando {atraso_destino:g}s antes de abrir a pagina {destino}.",
                        flush=True,
                    )
                    if not self.esperar(atraso_destino):
                        return {
                            "success": False,
                            "status": "cancelado",
                            "sequencia": sequencia_atual,
                            "indice": indice,
                            "resultado": self.resultado_fluxo("parar"),
                        }
                sequencia_atual = destino
                indice = 0
                if parar_ao_entrar_em_loop and sequencia_atual in sequencias_loop:
                    return {
                        "success": True,
                        "status": "aguardando_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": resultado,
                    }
                if alvo_parada and sequencia_atual == alvo_parada:
                    return {
                        "success": True,
                        "status": "aguardando_sequencia",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("continuar"),
                    }
                continue

            if acao == "ir_para_bloco":
                if not str(destino_bloco or "").strip():
                    raise ValueError("A acao ir_para_bloco exige um numero de bloco.")
                sequencia_destino, indice_destino = self.localizar_bloco(destino_bloco)
                ordem_destino = int(destino_bloco)
                if (
                    sequencia_reinicio_loop
                    and limitar_as_sequencias_loop
                    and sequencia_atual in sequencias_loop
                    and sequencia_atual != sequencia_reinicio_loop
                    and sequencia_destino == sequencia_reinicio_loop
                ):
                    safe_print(
                        f"Retorno para o bloco {ordem_destino} da pagina inicial do loop "
                        f"({sequencia_destino}) detectado. Avancando para a proxima linha da planilha.",
                        flush=True,
                    )
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("ir_para_bloco", bloco=destino_bloco),
                    }
                if sequencia_limitada and sequencia_atual == sequencia_limitada and sequencia_destino != sequencia_limitada:
                    return {
                        "success": True,
                        "status": "saida_bloco",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("ir_para_bloco", bloco=destino_bloco),
                    }
                if sequencia_blocos and sequencia_destino == sequencia_blocos:
                    if bloco_minimo and ordem_destino < bloco_minimo:
                        return {
                            "success": True,
                            "status": "saida_loop",
                            "sequencia": sequencia_atual,
                            "indice": indice,
                            "resultado": self.resultado_fluxo("ir_para_bloco", bloco=ordem_destino),
                        }
                    if bloco_maximo_exclusivo and ordem_destino >= bloco_maximo_exclusivo:
                        return {
                            "success": True,
                            "status": "saida_loop",
                            "sequencia": sequencia_atual,
                            "indice": indice,
                            "resultado": self.resultado_fluxo("ir_para_bloco", bloco=ordem_destino),
                        }
                if limitar_as_sequencias_loop and sequencia_atual in sequencias_loop and sequencia_destino not in sequencias_loop:
                    return {
                        "success": True,
                        "status": "saida_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("ir_para_bloco", bloco=destino_bloco),
                    }
                if atraso_destino > 0:
                    safe_print(
                        f"Aguardando {atraso_destino:g}s antes de abrir o bloco {ordem_destino}.",
                        flush=True,
                    )
                    if not self.esperar(atraso_destino):
                        return {
                            "success": False,
                            "status": "cancelado",
                            "sequencia": sequencia_atual,
                            "indice": indice,
                            "resultado": self.resultado_fluxo("parar"),
                        }
                sequencia_atual, indice = sequencia_destino, indice_destino
                if parar_ao_entrar_em_loop and sequencia_atual in sequencias_loop:
                    return {
                        "success": True,
                        "status": "aguardando_loop",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("continuar"),
                    }
                if alvo_parada and sequencia_atual == alvo_parada and indice == 0:
                    return {
                        "success": True,
                        "status": "aguardando_sequencia",
                        "sequencia": sequencia_atual,
                        "indice": indice,
                        "resultado": self.resultado_fluxo("continuar"),
                    }
                continue

            indice += 1

    def executar_fluxo(self, contexto_loop=None, sequencia_inicial=None, indice_inicial=0):
        resultado = self.executar_fluxo_controlado(
            contexto_loop=contexto_loop,
            sequencia_inicial=sequencia_inicial,
            indice_inicial=indice_inicial,
        )
        return bool(resultado.get("success"))

    def executar_normal(self):
        safe_print("Executando roteiro completo...", flush=True)
        self._janela_contexto = ""
        return self.executar_fluxo(sequencia_inicial=self.sequencia_inicial)

    def executar_loop_legado(self, loop_start=1):
        evento_planilha = next(
            (evento for evento in self.eventos if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA),
            None,
        )
        if evento_planilha is None:
            raise ValueError("Nao existe nenhuma sequencia de planilha no roteiro gravado.")
        indice_inicial, linha_inicial_real, total_iteracoes, linhas_loop = self.resolver_inicio_loop(
            evento_planilha,
            loop_start,
        )
        detalhes_planilha = evento_planilha.get("detalhes", {}) or {}
        workbook_loop = self.normalizar_id_planilha(detalhes_planilha.get("arquivo"))
        aba_loop = str(detalhes_planilha.get("aba") or "")

        iteracoes_restantes = total_iteracoes - indice_inicial
        safe_print(
            f"Executando loop: {iteracoes_restantes} iteracoes a partir da linha {linha_inicial_real}.",
            flush=True,
        )
        self.iniciar_relatorio_loop(
            evento_planilha,
            indice_inicial,
            linha_inicial_real,
            total_iteracoes,
            linhas_loop,
        )

        status_relatorio = "concluido"
        erro_fatal = ""
        caminho_relatorio = ""

        try:
            for indice in range(indice_inicial, total_iteracoes):
                if self.cancelamento_solicitado():
                    status_relatorio = "interrompido"
                    return False

                self._janela_contexto = ""

                linha_referencia = linhas_loop[indice] if indice < len(linhas_loop) else None
                empresa_atual = self.obter_empresa_evento_planilha(
                    evento_planilha,
                    indice,
                    linha_referencia=linha_referencia,
                )
                safe_print(
                    self.descrever_iteracao_loop(
                        indice,
                        total_iteracoes,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                    ),
                    flush=True,
                )

                inicio_iteracao = time.perf_counter()
                try:
                    contexto_principal = {
                        "indice": indice,
                        "total": total_iteracoes,
                        "linha_atual": linha_referencia,
                        "linha_referencia": linha_referencia,
                        "empresa": empresa_atual,
                        "arquivo": workbook_loop,
                        "aba": aba_loop,
                    }
                    sucesso_fluxo = self.executar_fluxo(
                        contexto_loop={
                            "indice": indice,
                            "total": total_iteracoes,
                            "linha_atual": linha_referencia,
                            "linha_referencia": linha_referencia,
                            "empresa": empresa_atual,
                            "workbook_id": workbook_loop,
                            "aba": aba_loop,
                            "loop_principal": contexto_principal,
                            "loop_local": {},
                        }
                    )
                except Exception as erro:
                    duracao_iteracao = time.perf_counter() - inicio_iteracao
                    mensagem_erro = str(erro or "Falha durante a execucao da linha do loop.").strip()
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "erro"
                    erro_fatal = mensagem_erro
                    raise

                duracao_iteracao = time.perf_counter() - inicio_iteracao
                if not sucesso_fluxo:
                    execucao_interrompida = self.cancelamento_solicitado()
                    mensagem_erro = (
                        "Execucao interrompida pelo usuario."
                        if execucao_interrompida
                        else "Falha em uma das etapas do loop."
                    )
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "interrompido" if execucao_interrompida else "erro"
                    if not execucao_interrompida:
                        erro_fatal = mensagem_erro
                    return False

                self.registrar_resultado_relatorio_loop(
                    indice=indice,
                    linha_referencia=linha_referencia,
                    empresa=empresa_atual,
                    sucesso=True,
                    duracao_segundos=duracao_iteracao,
                )

            return True
        finally:
            caminho_relatorio = self.finalizar_relatorio_loop(
                status=status_relatorio,
                erro_fatal=erro_fatal,
            )
            if caminho_relatorio:
                safe_print(f"Relatorio PDF do loop salvo em: {caminho_relatorio}", flush=True)

    def resolver_destino_pos_loop(self, sequencia_loop, sequencia_apos_loop="", resultado_saida=None):
        sequencia_loop = self.normalizar_nome_sequencia(sequencia_loop)
        destino_configurado = str(sequencia_apos_loop or "").strip()
        destino_configurado = (
            self.normalizar_nome_sequencia(destino_configurado) if destino_configurado else ""
        )
        if destino_configurado and destino_configurado != sequencia_loop:
            atraso = 0.0
            if resultado_saida and str(resultado_saida.get("acao") or "") == "ir_para_sequencia":
                if self.normalizar_nome_sequencia(resultado_saida.get("sequencia")) == destino_configurado:
                    atraso = max(0.0, float(resultado_saida.get("atraso", 0) or 0))
            return {"tipo": "sequencia", "sequencia": destino_configurado, "atraso": atraso}

        if not resultado_saida:
            return {}

        acao = str(resultado_saida.get("acao") or "continuar").strip().lower()
        if acao == "ir_para_sequencia":
            destino = self.normalizar_nome_sequencia(resultado_saida.get("sequencia"))
            if destino and destino != sequencia_loop:
                return {
                    "tipo": "sequencia",
                    "sequencia": destino,
                    "atraso": max(0.0, float(resultado_saida.get("atraso", 0) or 0)),
                }
        if acao == "ir_para_bloco" and str(resultado_saida.get("bloco") or "").strip():
            return {"tipo": "bloco", "bloco": int(resultado_saida.get("bloco"))}
        return {}

    def executar_finalizacao_pos_loop(self, destino):
        if not destino:
            return True

        tipo_destino = str(destino.get("tipo") or "").strip().lower()
        if tipo_destino == "sequencia":
            nome_sequencia = self.normalizar_nome_sequencia(destino.get("sequencia"))
            atraso = max(0.0, float(destino.get("atraso", 0) or 0))
            if atraso > 0:
                safe_print(
                    f"Aguardando {atraso:g}s antes de abrir a pagina {nome_sequencia} apos o loop.",
                    flush=True,
                )
                if not self.esperar(atraso):
                    return False
            safe_print(f"Loop concluido. Abrindo a pagina {nome_sequencia}.", flush=True)
            resultado = self.executar_fluxo_controlado(sequencia_inicial=nome_sequencia)
            return bool(resultado.get("success"))

        if tipo_destino == "bloco":
            sequencia_destino, indice_destino = self.localizar_bloco(destino.get("bloco"))
            safe_print(
                f"Loop concluido. Continuando a partir do bloco {destino.get('bloco')}.",
                flush=True,
            )
            resultado = self.executar_fluxo_controlado(
                sequencia_inicial=sequencia_destino,
                indice_inicial=indice_destino,
            )
            return bool(resultado.get("success"))

        return True

    def executar_loop_por_pagina(self, loop_start, configuracao_loop):
        sequencias_loop = {
            self.normalizar_nome_sequencia(nome)
            for nome in (configuracao_loop.get("sequencias_loop") or set())
            if str(nome or "").strip()
        }
        evento_planilha = configuracao_loop.get("evento_planilha")
        if evento_planilha is None:
            raise ValueError("Nao foi possivel localizar o bloco de planilha do grupo em loop.")

        if not sequencias_loop:
            raise ValueError("Nenhuma pagina foi marcada como loop no roteiro atual.")

        sequencia_loop = self.sequencia_inicial if self.sequencia_inicial in sequencias_loop else ""
        if not sequencia_loop:
            safe_print(
                "Modo loop por pagina ativo. Executando as etapas iniciais ate o primeiro grupo em loop.",
                flush=True,
            )
            preparo = self.executar_fluxo_controlado(
                sequencia_inicial=self.sequencia_inicial,
                parar_ao_entrar_em_loop=True,
                sequencias_loop_ativas=sequencias_loop,
            )
            if not preparo.get("success"):
                return False
            if preparo.get("status") != "aguardando_loop":
                raise ValueError(
                    "O grupo de paginas marcado como loop nao foi alcancado a partir da sequencia inicial."
                )
            sequencia_loop = self.normalizar_nome_sequencia(preparo.get("sequencia"))

        indice_inicial, linha_inicial_real, total_iteracoes, linhas_loop = self.resolver_inicio_loop(
            evento_planilha,
            loop_start,
        )
        detalhes_planilha = evento_planilha.get("detalhes", {}) or {}
        workbook_loop = self.normalizar_id_planilha(detalhes_planilha.get("arquivo"))
        aba_loop = str(detalhes_planilha.get("aba") or "")
        iteracoes_restantes = total_iteracoes - indice_inicial
        paginas_loop_label = ", ".join(sorted(sequencias_loop))
        safe_print(
            f"Executando grupo em loop ({paginas_loop_label}): {iteracoes_restantes} iteracoes a partir da linha {linha_inicial_real}.",
            flush=True,
        )

        self.iniciar_relatorio_loop(
            evento_planilha,
            indice_inicial,
            linha_inicial_real,
            total_iteracoes,
            linhas_loop,
        )

        status_relatorio = "concluido"
        erro_fatal = ""
        resultado_saida_loop = None
        sequencia_saida_loop = sequencia_loop

        try:
            for indice in range(indice_inicial, total_iteracoes):
                if self.cancelamento_solicitado():
                    status_relatorio = "interrompido"
                    return False

                self._janela_contexto = ""
                linha_referencia = linhas_loop[indice] if indice < len(linhas_loop) else None
                empresa_atual = self.obter_empresa_evento_planilha(
                    evento_planilha,
                    indice,
                    linha_referencia=linha_referencia,
                )
                safe_print(
                    self.descrever_iteracao_loop(
                        indice,
                        total_iteracoes,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        contexto=f"grupo em loop: {paginas_loop_label}",
                    ),
                    flush=True,
                )

                inicio_iteracao = time.perf_counter()
                try:
                    contexto_principal = {
                        "indice": indice,
                        "total": total_iteracoes,
                        "linha_atual": linha_referencia,
                        "linha_referencia": linha_referencia,
                        "empresa": empresa_atual,
                        "arquivo": workbook_loop,
                        "aba": aba_loop,
                    }
                    resultado_iteracao = self.executar_fluxo_controlado(
                        contexto_loop={
                            "indice": indice,
                            "total": total_iteracoes,
                            "linha_atual": linha_referencia,
                            "linha_referencia": linha_referencia,
                            "empresa": empresa_atual,
                            "workbook_id": workbook_loop,
                            "aba": aba_loop,
                            "loop_principal": contexto_principal,
                            "loop_local": {},
                        },
                        sequencia_inicial=sequencia_loop,
                        sequencias_loop_ativas=sequencias_loop,
                        limitar_as_sequencias_loop=True,
                        encerrar_iteracao_ao_voltar_para_sequencia=sequencia_loop,
                    )
                except Exception as erro:
                    duracao_iteracao = time.perf_counter() - inicio_iteracao
                    mensagem_erro = str(erro or "Falha durante a execucao da linha do loop.").strip()
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "erro"
                    erro_fatal = mensagem_erro
                    raise

                duracao_iteracao = time.perf_counter() - inicio_iteracao
                if not resultado_iteracao.get("success"):
                    execucao_interrompida = self.cancelamento_solicitado()
                    mensagem_erro = (
                        "Execucao interrompida pelo usuario."
                        if execucao_interrompida
                        else "Falha em uma das etapas do loop."
                    )
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "interrompido" if execucao_interrompida else "erro"
                    if not execucao_interrompida:
                        erro_fatal = mensagem_erro
                    return False

                if resultado_iteracao.get("status") == "saida_loop":
                    resultado_saida_loop = resultado_iteracao.get("resultado") or resultado_saida_loop
                    sequencia_saida_loop = self.normalizar_nome_sequencia(
                        resultado_iteracao.get("sequencia") or sequencia_saida_loop
                    )
                    if resultado_saida_loop:
                        safe_print(
                            "A saida do grupo em loop foi registrada e sera aplicada somente depois que todas as linhas terminarem.",
                            flush=True,
                        )
                    self.emitir_preparacao_proxima_iteracao_loop(
                        sequencia_loop,
                        indice + 1,
                        total_iteracoes,
                        linhas_loop,
                        evento_planilha,
                        contexto=f"preparando grupo em loop: {paginas_loop_label}",
                    )

                self.registrar_resultado_relatorio_loop(
                    indice=indice,
                    linha_referencia=linha_referencia,
                    empresa=empresa_atual,
                    sucesso=True,
                    duracao_segundos=duracao_iteracao,
                )

            configuracao_saida = self.obter_configuracao_sequencia(sequencia_saida_loop)
            destino_pos_loop = self.resolver_destino_pos_loop(
                sequencia_saida_loop,
                configuracao_saida.get("sequencia_apos_loop", ""),
                resultado_saida_loop,
            )
            if destino_pos_loop and not self.executar_finalizacao_pos_loop(destino_pos_loop):
                status_relatorio = "erro"
                erro_fatal = "Falha ao executar a finalizacao apos o loop."
                return False
            return True
        finally:
            caminho_relatorio = self.finalizar_relatorio_loop(
                status=status_relatorio,
                erro_fatal=erro_fatal,
            )
            if caminho_relatorio:
                safe_print(f"Relatorio PDF do loop salvo em: {caminho_relatorio}", flush=True)

    def executar_loop_por_blocos(self, loop_start, configuracao_loop_blocos):
        sequencia_loop = self.normalizar_nome_sequencia(configuracao_loop_blocos.get("sequencia_loop"))
        inicio_indice = int(configuracao_loop_blocos.get("inicio_indice", 0) or 0)
        fim_indice = int(configuracao_loop_blocos.get("fim_indice", 0) or 0)
        inicio_ordem = int(configuracao_loop_blocos.get("inicio_ordem", 0) or 0)
        fim_ordem = int(configuracao_loop_blocos.get("fim_ordem", 0) or 0)
        evento_planilha = configuracao_loop_blocos.get("evento_planilha")
        if evento_planilha is None:
            raise ValueError("Nao foi possivel localizar o bloco de planilha dentro do trecho em loop.")

        safe_print(
            "Modo loop por blocos ativo. Executando as etapas iniciais ate o marcador Inicio do loop.",
            flush=True,
        )
        preparo = self.executar_fluxo_controlado(
            sequencia_inicial=self.sequencia_inicial,
            sequencia_limite_blocos=sequencia_loop,
            parar_antes_do_bloco=inicio_ordem,
        )
        if not preparo.get("success"):
            return False
        if preparo.get("status") != "aguardando_bloco":
            raise ValueError(
                "O marcador Inicio do loop nao foi alcancado a partir da sequencia inicial."
            )

        indice_inicial, linha_inicial_real, total_iteracoes, linhas_loop = self.resolver_inicio_loop(
            evento_planilha,
            loop_start,
        )
        detalhes_planilha = evento_planilha.get("detalhes", {}) or {}
        workbook_loop = self.normalizar_id_planilha(detalhes_planilha.get("arquivo"))
        aba_loop = str(detalhes_planilha.get("aba") or "")
        iteracoes_restantes = total_iteracoes - indice_inicial
        safe_print(
            f"Executando trecho em loop entre os blocos {inicio_ordem} e {fim_ordem}: "
            f"{iteracoes_restantes} iteracoes a partir da linha {linha_inicial_real}.",
            flush=True,
        )

        self.iniciar_relatorio_loop(
            evento_planilha,
            indice_inicial,
            linha_inicial_real,
            total_iteracoes,
            linhas_loop,
        )

        status_relatorio = "concluido"
        erro_fatal = ""
        resultado_saida_loop = None

        try:
            for indice in range(indice_inicial, total_iteracoes):
                if self.cancelamento_solicitado():
                    status_relatorio = "interrompido"
                    return False

                self._janela_contexto = ""
                linha_referencia = linhas_loop[indice] if indice < len(linhas_loop) else None
                empresa_atual = self.obter_empresa_evento_planilha(
                    evento_planilha,
                    indice,
                    linha_referencia=linha_referencia,
                )
                safe_print(
                    self.descrever_iteracao_loop(
                        indice,
                        total_iteracoes,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        contexto="trecho em loop",
                    ),
                    flush=True,
                )

                inicio_iteracao = time.perf_counter()
                try:
                    contexto_principal = {
                        "indice": indice,
                        "total": total_iteracoes,
                        "linha_atual": linha_referencia,
                        "linha_referencia": linha_referencia,
                        "empresa": empresa_atual,
                        "arquivo": workbook_loop,
                        "aba": aba_loop,
                    }
                    resultado_iteracao = self.executar_fluxo_controlado(
                        contexto_loop={
                            "indice": indice,
                            "total": total_iteracoes,
                            "linha_atual": linha_referencia,
                            "linha_referencia": linha_referencia,
                            "empresa": empresa_atual,
                            "workbook_id": workbook_loop,
                            "aba": aba_loop,
                            "loop_principal": contexto_principal,
                            "loop_local": {},
                        },
                        sequencia_inicial=sequencia_loop,
                        indice_inicial=inicio_indice + 1,
                        limitar_a_sequencia=sequencia_loop,
                        sequencia_limite_blocos=sequencia_loop,
                        bloco_inicial_permitido=inicio_ordem + 1,
                        parar_antes_do_bloco=fim_ordem,
                    )
                except Exception as erro:
                    duracao_iteracao = time.perf_counter() - inicio_iteracao
                    mensagem_erro = str(erro or "Falha durante a execucao da linha do loop.").strip()
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "erro"
                    erro_fatal = mensagem_erro
                    raise

                duracao_iteracao = time.perf_counter() - inicio_iteracao
                if not resultado_iteracao.get("success"):
                    execucao_interrompida = self.cancelamento_solicitado()
                    mensagem_erro = (
                        "Execucao interrompida pelo usuario."
                        if execucao_interrompida
                        else "Falha em uma das etapas do loop."
                    )
                    self.registrar_resultado_relatorio_loop(
                        indice=indice,
                        linha_referencia=linha_referencia,
                        empresa=empresa_atual,
                        sucesso=False,
                        erro=mensagem_erro,
                        duracao_segundos=duracao_iteracao,
                    )
                    status_relatorio = "interrompido" if execucao_interrompida else "erro"
                    if not execucao_interrompida:
                        erro_fatal = mensagem_erro
                    return False

                if resultado_iteracao.get("status") in {"saida_loop", "saida_sequencia", "saida_bloco"}:
                    resultado_saida_loop = resultado_iteracao.get("resultado") or resultado_saida_loop
                    if resultado_saida_loop:
                        safe_print(
                            "A saida do trecho em loop foi registrada e sera aplicada somente depois que todas as linhas terminarem.",
                            flush=True,
                        )
                    self.emitir_preparacao_proxima_iteracao_loop(
                        sequencia_loop,
                        indice + 1,
                        total_iteracoes,
                        linhas_loop,
                        evento_planilha,
                        contexto="preparando trecho em loop",
                    )

                self.registrar_resultado_relatorio_loop(
                    indice=indice,
                    linha_referencia=linha_referencia,
                    empresa=empresa_atual,
                    sucesso=True,
                    duracao_segundos=duracao_iteracao,
                )

            if resultado_saida_loop:
                destino_pos_loop = self.resolver_destino_pos_loop(
                    sequencia_loop,
                    "",
                    resultado_saida_loop,
                )
                if destino_pos_loop and not self.executar_finalizacao_pos_loop(destino_pos_loop):
                    status_relatorio = "erro"
                    erro_fatal = "Falha ao executar a finalizacao apos o loop."
                    return False
                return True

            safe_print("Loop concluido. Continuando apos o bloco Fim do loop.", flush=True)
            resultado_final = self.executar_fluxo_controlado(
                sequencia_inicial=sequencia_loop,
                indice_inicial=fim_indice + 1,
            )
            return bool(resultado_final.get("success"))
        finally:
            caminho_relatorio = self.finalizar_relatorio_loop(
                status=status_relatorio,
                erro_fatal=erro_fatal,
            )
            if caminho_relatorio:
                safe_print(f"Relatorio PDF do loop salvo em: {caminho_relatorio}", flush=True)

    def executar_loop(self, loop_start=1):
        configuracao_loop_blocos = self.localizar_faixa_loop_por_blocos()
        if configuracao_loop_blocos:
            return self.executar_loop_por_blocos(loop_start, configuracao_loop_blocos)
        configuracao_loop = self.localizar_configuracao_loop_planilha()
        if configuracao_loop:
            return self.executar_loop_por_pagina(loop_start, configuracao_loop)
        return self.executar_loop_legado(loop_start)

    def fechar_sessao_chrome(self):
        try:
            if self._playwright_runtime is not None:
                self._playwright_runtime.stop()
        except Exception:
            pass
        finally:
            self._playwright_runtime = None
            self._chrome_browser = None
            self._chrome_context = None
            self._chrome_page = None
            self._chrome_source = ""
            self._chrome_launched_by_executor = False


def criar_parser():
    parser = argparse.ArgumentParser(description="Executor separado da automacao gravada.")
    parser.add_argument("--modo", choices=("normal", "loop", "reconhecimento"), default="normal")
    parser.add_argument("--atraso", type=float, default=3.0)
    parser.add_argument("--loop-start", type=int, default=1)
    parser.add_argument("--arquivo", default=str(ARQUIVO_MAPEAMENTO))
    parser.add_argument("--planilha", default=str(ARQUIVO_PLANILHA))
    parser.add_argument("--cancelamento-arquivo", default=str(ARQUIVO_CANCELAMENTO))
    parser.add_argument("--janela-alvo", default="")
    parser.add_argument("--target-context", default="")
    parser.add_argument("--imagem-teste", default="")
    parser.add_argument("--timeout-imagem", type=float, default=10.0)
    parser.add_argument("--confianca-imagem", type=float, default=0.8)
    parser.add_argument("--perfil-busca-imagem", default="balanced")
    parser.add_argument("--modo-area-busca-imagem", default="target_window")
    parser.add_argument("--area-busca-x-imagem", type=int, default=0)
    parser.add_argument("--area-busca-y-imagem", type=int, default=0)
    parser.add_argument("--area-busca-largura-imagem", type=int, default=0)
    parser.add_argument("--area-busca-altura-imagem", type=int, default=0)
    parser.add_argument("--texto-ocr-esperado", default="")
    parser.add_argument("--modo-reconhecimento", choices=("imagem", "texto_ocr"), default="imagem")
    parser.add_argument("--texto-ocr-busca", default="")
    parser.add_argument("--acao-ocr", choices=("clicar", "localizar", "verificar"), default="localizar")
    parser.add_argument("--confianca-ocr", type=float, default=0.72)
    parser.add_argument("--timeout-ocr", type=float, default=8.0)
    parser.add_argument("--varredura-scroll", default="true")
    parser.add_argument("--varredura-scroll-tentativas", type=int, default=4)
    parser.add_argument("--varredura-scroll-pixels", type=int, default=520)
    parser.add_argument("--varredura-scroll-direcao", choices=("baixo", "cima"), default="baixo")
    parser.add_argument("--leitura-relativa-ativa", default="false")
    parser.add_argument("--leitura-relativa-lado", default="livre")
    parser.add_argument("--leitura-relativa-largura", type=int, default=140)
    parser.add_argument("--leitura-relativa-altura", type=int, default=40)
    parser.add_argument("--leitura-relativa-offset-x", type=int, default=0)
    parser.add_argument("--leitura-relativa-offset-y", type=int, default=0)
    parser.add_argument("--leitura-relativa-salvar-em", default="")
    parser.add_argument("--ocr-target-layouts", default="[]")
    parser.add_argument("--text-overrides", default="")
    return parser


def main():
    parser = criar_parser()
    args = parser.parse_args()

    if args.modo == "reconhecimento" and args.cancelamento_arquivo:
        try:
            Path(args.cancelamento_arquivo).unlink(missing_ok=True)
        except Exception:
            pass

    executor = ExecutorAutomacao(
        arquivo_mapeamento=args.arquivo,
        arquivo_planilha=args.planilha,
        arquivo_cancelamento=args.cancelamento_arquivo,
        janela_alvo=args.janela_alvo,
        target_context=args.target_context,
        modo_teste_reconhecimento=args.modo == "reconhecimento",
    )
    executor.aplicar_text_overrides(args.text_overrides)

    atraso = max(0, round(float(args.atraso), 2))
    safe_print(f"Automacao pronta. Atraso inicial: {atraso}s", flush=True)
    if executor.alvo_execucao_ativo():
        contexto = executor.contexto_execucao_alvo or {}
        safe_print(
            f"Janela alvo selecionada: {contexto.get('target_title', '')} / hwnd {int(contexto.get('target_hwnd', 0) or 0)}",
            flush=True,
        )
        if contexto.get("target_class"):
            safe_print(f"Classe da janela alvo: {contexto.get('target_class')}", flush=True)
    elif executor.janela_alvo:
        safe_print(f"Janela alvo configurada: {executor.janela_alvo}", flush=True)
    if atraso > 0:
        safe_print(f"Voce tem {atraso:g}s para focar a janela onde a automacao deve executar.", flush=True)
        time.sleep(atraso)

    if args.modo != "reconhecimento" and not executor.janela_alvo and not executor.alvo_execucao_ativo():
        janela_passiva = executor.adotar_janela_ativa_para_ocr()
        if janela_passiva is None:
            janela_passiva = executor.adotar_janela_visivel_passiva_para_ocr()
        if janela_passiva:
            safe_print(
                f"Execucao sem alvo fixo: OCR limitado passivamente a {janela_passiva.get('title', 'janela visivel')}.",
                flush=True,
            )
        else:
            safe_print(
                "Nenhuma janela alvo selecionada. Captura e reconhecimento vao usar a tela inteira.",
                flush=True,
            )
    elif args.modo == "reconhecimento" and not executor.janela_alvo and not executor.alvo_execucao_ativo():
        safe_print(
            "Teste de reconhecimento sem janela alvo fixa. Procurando a referencia na tela inteira.",
            flush=True,
        )

    try:
        try:
            if args.modo == "reconhecimento":
                if args.modo_reconhecimento == "texto_ocr" and not str(args.imagem_teste or "").strip():
                    evento_teste = {
                        "ordem": 0,
                        "tipo": "clicar_imagem",
                        "detalhes": {
                            "modo_reconhecimento": "texto_ocr",
                            "texto": args.texto_ocr_busca,
                            "texto_ocr_busca": args.texto_ocr_busca,
                            "acao": "localizar",
                            "acao_ocr": args.acao_ocr,
                            "timeout": args.timeout_ocr,
                            "confianca": args.confianca_ocr,
                            "clicar": False,
                            "cliques": 0,
                            "janela_alvo": args.janela_alvo,
                            "teste_reconhecimento_ocr": True,
                            "quantidade_tentativas": 1,
                            "intervalo_tentativas": 0,
                            "varredura_scroll": str(args.varredura_scroll or "").strip().lower() not in {"0", "false", "nao", "não", "off"},
                            "varredura_scroll_tentativas": args.varredura_scroll_tentativas,
                            "varredura_scroll_pixels": args.varredura_scroll_pixels,
                            "varredura_scroll_direcao": args.varredura_scroll_direcao,
                            "ao_nao_encontrar": "parar",
                            "ao_sucesso": "continuar",
                        },
                    }
                    resultado = executor.executar_bloco_procurar_texto_ocr(evento_teste)
                    sucesso = resultado.get("acao") != "parar"
                    safe_print("Teste de OCR inteligente concluido com sucesso.", flush=True)
                else:
                    leitura_relativa_teste = {
                        "ativo": str(args.leitura_relativa_ativa or "").strip().lower() in {"1", "true", "sim", "yes"},
                        "lado": args.leitura_relativa_lado,
                        "largura": args.leitura_relativa_largura,
                        "altura": args.leitura_relativa_altura,
                        "offset_x": args.leitura_relativa_offset_x,
                        "offset_y": args.leitura_relativa_offset_y,
                        "salvar_em": args.leitura_relativa_salvar_em,
                        "ocr_target_layouts": json.loads(args.ocr_target_layouts or "[]"),
                    }
                    sucesso = executor.testar_reconhecimento(
                        imagem=args.imagem_teste,
                        timeout=args.timeout_imagem,
                        confianca=args.confianca_imagem,
                        janela_alvo=args.janela_alvo,
                        clicar=False,
                        perfil_busca=args.perfil_busca_imagem,
                        modo_area_busca=args.modo_area_busca_imagem,
                        area_personalizada={
                            "x": args.area_busca_x_imagem,
                            "y": args.area_busca_y_imagem,
                            "width": args.area_busca_largura_imagem,
                            "height": args.area_busca_altura_imagem,
                        },
                        texto_ocr_esperado=args.texto_ocr_busca or args.texto_ocr_esperado,
                        leitura_relativa=leitura_relativa_teste,
                    )
            elif args.modo == "loop":
                sucesso = executor.executar_loop(loop_start=args.loop_start)
            else:
                sucesso = executor.executar_normal()
        except KeyboardInterrupt:
            executor.cancelar.set()
            safe_print("Execucao interrompida pelo teclado.", flush=True)
            return 1
        except Exception as erro:
            executor.cancelar.set()
            safe_print(str(erro or "Falha na execucao da automacao."), flush=True)
            return 1
    finally:
        executor.fechar_sessao_chrome()

    if sucesso:
        safe_print("Automacao concluida com sucesso.", flush=True)
        return 0

    safe_print("Automacao interrompida.", flush=True)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
