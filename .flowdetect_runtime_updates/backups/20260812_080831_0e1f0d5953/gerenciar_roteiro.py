import argparse
import hashlib
import json
import re
import shutil
import sys
import tempfile
import zipfile
from html import unescape
from html.parser import HTMLParser
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ARQUIVO_SAIDA = Path("mapeamento_teclado.json")
ARQUIVO_RESUMO = Path("resumo_teclado.txt")
ARQUIVO_SCRIPT = Path("comandos_teclado.py")
ARQUIVO_PLANILHA = Path("db.xlsx")
DIRETORIO_PLANILHAS = Path("planilhas")
PASTA_IMAGENS = Path("imagens")
DIRETORIO_AUTOMACOES = Path("automacoes")
DIRETORIO_EXPORTACOES_AUTOMACAO = Path("automacoes_exportadas")
DIRETORIO_ATUALIZACOES_RUNTIME = Path(".flowdetect_runtime_updates")
ARQUIVO_BIBLIOTECA = DIRETORIO_AUTOMACOES / "biblioteca.json"
ARQUIVO_AUTOMACAO_ATUAL = DIRETORIO_AUTOMACOES / "automacao_atual.json"
ARQUIVO_BUNDLE_AUTOMACAO = "flowdetect_automation.json"
ARQUIVO_BUNDLE_RUNTIME = "flowdetect_runtime_update.json"
BUNDLE_AUTOMACAO_TIPO = "flowdetect_automation_bundle"
BUNDLE_AUTOMACAO_VERSAO = 1
BUNDLE_RUNTIME_TIPO = "flowdetect_runtime_bundle"
BUNDLE_RUNTIME_VERSAO = 1
EXTENSOES_PLANILHA = {".xlsx", ".xlsm", ".xltx", ".xltm"}
INTERVALO_TEXTO = 0.01
INTERVALO_ENTRE_ITENS_PLANILHA = 0.5
INSTRUCOES_PADRAO = [
    "Teclas simples e atalhos sao gravados individualmente.",
    "F8 abre uma caixa de texto grande e salva o bloco como um unico comando.",
    "O botao Capturar mouse salva a coordenada exata do proximo clique.",
    "Cada evento tambem salva o intervalo em segundos desde a acao anterior.",
    "A aba Planilha pode adicionar uma sequencia inteira ou so linhas selecionadas ao roteiro.",
    "Backspace remove a ultima entrada salva e F9 ou ESC encerram a captura.",
]
TIPOS_SEQUENCIA_PLANILHA = {"sequencia_planilha"}
TIPOS_ESCRITA_PLANILHA = {"escrever_planilha"}
TIPOS_COMPARACAO_PLANILHA = {"comparar_planilha_tela"}
TIPOS_MARCADORES_LOOP = {"inicio_loop", "fim_loop"}
TIPOS_MOUSE = {"mouse_click"}
TIPOS_IMAGEM = {"clicar_imagem"}
TIPOS_VERIFICACAO = {"verificar_estado_botao"}
TIPOS_CHROME = {"chrome_elemento_html"}
TIPOS_PREPARACAO_TELA = {"abrir_url", "abrir_programa"}
TIPOS_DOWNLOAD = {"salvar_download"}
TIPOS_SUPORTADOS_RUNTIME = {
    "texto_livre",
    "tecla",
    "atalho",
    *TIPOS_SEQUENCIA_PLANILHA,
    *TIPOS_ESCRITA_PLANILHA,
    *TIPOS_COMPARACAO_PLANILHA,
    *TIPOS_MARCADORES_LOOP,
    *TIPOS_MOUSE,
    *TIPOS_IMAGEM,
    *TIPOS_VERIFICACAO,
    *TIPOS_CHROME,
    *TIPOS_PREPARACAO_TELA,
    *TIPOS_DOWNLOAD,
}
MODOS_EXECUCAO_SEQUENCIA = {"normal", "loop_planilha"}
ARQUIVOS_RUNTIME_ATUALIZACAO = [
    "executar_automacao.py",
    "executar_automacao.bat",
    "gerenciar_roteiro.py",
    "gerenciar_dados_planilha.py",
    "capturar_imagem_tela.py",
    "capturar_cliques.py",
    "capturar_mouse_global.py",
    "comandos_teclado.py",
    "overlay_janela_alvo.py",
    "listar_janelas_abertas.py",
    "validar_dependencias_instaladas.py",
    "diagnostico_ambiente.py",
    "instalar_dependencias_silencioso.bat",
    "requirements.txt",
]
TECLAS_MODIFICADORAS = {
    "shift",
    "shift_r",
    "ctrl",
    "ctrl_r",
    "alt",
    "alt_r",
    "cmd",
    "cmd_r",
}
MAPA_EXIBICAO = {
    "space": "SPACE",
    "enter": "ENTER",
    "tab": "TAB",
    "backspace": "BACKSPACE",
    "esc": "ESC",
    "shift": "SHIFT",
    "shift_r": "SHIFT_R",
    "ctrl": "CTRL",
    "ctrl_r": "CTRL_R",
    "alt": "ALT",
    "alt_r": "ALT_R",
    "cmd": "WIN",
    "cmd_r": "WIN_R",
    "up": "UP",
    "down": "DOWN",
    "left": "LEFT",
    "right": "RIGHT",
    "delete": "DELETE",
    "insert": "INSERT",
    "home": "HOME",
    "end": "END",
    "page_up": "PAGE_UP",
    "page_down": "PAGE_DOWN",
    "f8": "F8",
    "f9": "F9",
}
MAPA_PYAUTOGUI = {
    "space": "space",
    "enter": "enter",
    "tab": "tab",
    "backspace": "backspace",
    "esc": "esc",
    "shift": "shift",
    "shift_r": "shift",
    "ctrl": "ctrl",
    "ctrl_r": "ctrl",
    "alt": "alt",
    "alt_r": "alt",
    "cmd": "win",
    "cmd_r": "win",
    "up": "up",
    "down": "down",
    "left": "left",
    "right": "right",
    "delete": "delete",
    "insert": "insert",
    "home": "home",
    "end": "end",
    "page_up": "pageup",
    "page_down": "pagedown",
    "f8": "f8",
    "f9": "f9",
}


def agora_iso():
    return datetime.now().isoformat(timespec="seconds")


def garantir_diretorio_automacoes():
    DIRETORIO_AUTOMACOES.mkdir(parents=True, exist_ok=True)


def ler_json_arquivo(caminho):
    caminho = Path(caminho)
    return json.loads(caminho.read_text(encoding="utf-8-sig"))


def criar_payload_vazio():
    return {
        "criado_em": agora_iso(),
        "instrucoes": list(INSTRUCOES_PADRAO),
        "resumo_digitado": "",
        "eventos": [],
        "sequencia_inicial": "principal",
        "configuracoes_sequencia": {},
    }


def biblioteca_vazia():
    return {"items": []}


def gerar_id_automacao():
    return f"auto_{datetime.now().strftime('%Y%m%d%H%M%S%f')}"


def normalizar_item_biblioteca(item):
    criado = str(item.get("criado_em") or agora_iso())
    atualizado = str(item.get("atualizado_em") or criado)
    return {
        "id": str(item.get("id") or gerar_id_automacao()),
        "nome": str(item.get("nome") or "Automacao sem nome").strip() or "Automacao sem nome",
        "descricao": str(item.get("descricao") or "").strip(),
        "categoria": str(item.get("categoria") or "Outros").strip() or "Outros",
        "criado_em": criado,
        "atualizado_em": atualizado,
        "vezes_executada": max(0, int(item.get("vezes_executada") or 0)),
        "ultima_execucao": item.get("ultima_execucao"),
        "sequencia_inicial": str(item.get("sequencia_inicial") or "principal").strip() or "principal",
    }


def diretorio_automacao(automacao_id):
    return DIRETORIO_AUTOMACOES / str(automacao_id)


def arquivo_roteiro_automacao(automacao_id):
    return diretorio_automacao(automacao_id) / ARQUIVO_SAIDA.name


def carregar_biblioteca():
    garantir_diretorio_automacoes()
    if not ARQUIVO_BIBLIOTECA.exists():
        biblioteca = biblioteca_vazia()
    else:
        biblioteca = ler_json_arquivo(ARQUIVO_BIBLIOTECA)
    itens = [normalizar_item_biblioteca(item) for item in biblioteca.get("items", [])]
    biblioteca = {"items": itens}
    biblioteca = migrar_legado_se_necessario(biblioteca)
    salvar_biblioteca(biblioteca)
    return biblioteca


def salvar_biblioteca(biblioteca):
    garantir_diretorio_automacoes()
    biblioteca["items"] = [normalizar_item_biblioteca(item) for item in biblioteca.get("items", [])]
    ARQUIVO_BIBLIOTECA.write_text(
        json.dumps(biblioteca, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )


def normalizar_nome_arquivo(texto, fallback="automacao"):
    nome = re.sub(r'[<>:"/\\\\|?*]+', "_", str(texto or "").strip())
    nome = re.sub(r"\s+", "_", nome)
    nome = re.sub(r"_+", "_", nome).strip("._ ")
    return nome or fallback


def caminho_arquivo_unico(diretorio, nome_arquivo):
    diretorio = Path(diretorio)
    diretorio.mkdir(parents=True, exist_ok=True)
    nome_base = Path(nome_arquivo).name or "arquivo"
    caminho = diretorio / nome_base
    if not caminho.exists():
        return caminho
    stem = caminho.stem or "arquivo"
    suffix = caminho.suffix
    for indice in range(2, 10000):
        candidato = diretorio / f"{stem}_{indice:03}{suffix}"
        if not candidato.exists():
            return candidato
    return caminho


def listar_arquivos_runtime_existentes(base_dir=None):
    raiz = Path(base_dir or Path.cwd())
    encontrados = []
    for relativo in ARQUIVOS_RUNTIME_ATUALIZACAO:
        caminho = raiz / relativo
        if caminho.exists() and caminho.is_file():
            encontrados.append(relativo.replace("\\", "/"))
    return encontrados


def calcular_hash_runtime(base_dir=None, arquivos=None):
    raiz = Path(base_dir or Path.cwd())
    lista_arquivos = sorted(arquivos or listar_arquivos_runtime_existentes(raiz))
    digest = hashlib.sha256()
    for relativo in lista_arquivos:
        caminho = raiz / relativo
        if not caminho.exists() or not caminho.is_file():
            continue
        digest.update(relativo.replace("\\", "/").encode("utf-8"))
        digest.update(b"\0")
        digest.update(caminho.read_bytes())
        digest.update(b"\0")
    return digest.hexdigest()


def detectar_tipos_bloco_payload(payload):
    eventos = (payload or {}).get("eventos", [])
    return sorted({str((evento or {}).get("tipo") or "").strip() for evento in eventos if str((evento or {}).get("tipo") or "").strip()})


def gerar_manifesto_runtime(base_dir=None):
    raiz = Path(base_dir or Path.cwd())
    arquivos = listar_arquivos_runtime_existentes(raiz)
    return {
        "bundle_type": BUNDLE_RUNTIME_TIPO,
        "bundle_version": BUNDLE_RUNTIME_VERSAO,
        "generated_at": agora_iso(),
        "runtime_hash": calcular_hash_runtime(raiz, arquivos),
        "supported_block_types": sorted(TIPOS_SUPORTADOS_RUNTIME),
        "files": arquivos,
    }


def gerar_resumo_compatibilidade_payload(payload, manifesto_runtime=None):
    manifesto = manifesto_runtime or gerar_manifesto_runtime()
    tipos_detectados = detectar_tipos_bloco_payload(payload)
    suportados = set(manifesto.get("supported_block_types") or [])
    nao_suportados = sorted(tipo for tipo in tipos_detectados if tipo not in suportados)
    return {
        "required_block_types": tipos_detectados,
        "unsupported_block_types": nao_suportados,
        "runtime_hash": manifesto.get("runtime_hash", ""),
        "compatible": not nao_suportados,
    }


def localizar_manifesto_runtime_bundle(diretorio):
    diretorio = Path(diretorio)
    manifestos = sorted(diretorio.rglob(ARQUIVO_BUNDLE_RUNTIME), key=lambda caminho: len(caminho.parts))
    return manifestos[0] if manifestos else None


def localizar_pasta_runtime_importado(diretorio, manifesto_path=None):
    diretorio = Path(diretorio)
    if manifesto_path:
        candidata = manifesto_path.parent / "runtime_update" / "files"
        if candidata.exists():
            return candidata
    candidatos = sorted(
        (caminho for caminho in diretorio.rglob("runtime_update") if (caminho / "files").exists()),
        key=lambda caminho: len(caminho.parts),
    )
    return candidatos[0] / "files" if candidatos else None


def exportar_runtime_para_pasta(destino_base):
    destino_base = Path(destino_base)
    pasta_runtime = destino_base / "runtime_update" / "files"
    pasta_runtime.mkdir(parents=True, exist_ok=True)
    arquivos_copiados = []
    for relativo in listar_arquivos_runtime_existentes():
        origem = Path(relativo)
        destino = pasta_runtime / relativo
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(origem, destino)
        arquivos_copiados.append(relativo)
    manifesto_runtime = gerar_manifesto_runtime()
    (destino_base / ARQUIVO_BUNDLE_RUNTIME).write_text(
        json.dumps(manifesto_runtime, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    return {
        "manifest": manifesto_runtime,
        "files": arquivos_copiados,
        "runtime_dir": str(pasta_runtime),
    }


def aplicar_pacote_runtime_importado(pasta_runtime_files, manifesto_runtime=None, destino_base=None):
    pasta_runtime_files = Path(pasta_runtime_files)
    if not pasta_runtime_files.exists():
        raise ValueError("O pacote de atualizacao nao contem a pasta de arquivos do runtime.")

    destino_base = Path(destino_base or Path.cwd())
    manifesto_runtime = manifesto_runtime or {}
    arquivos_manifesto = [
        str(item).replace("\\", "/")
        for item in (manifesto_runtime.get("files") or [])
        if str(item).strip()
    ]
    arquivos_disponiveis = []
    for relativo in (arquivos_manifesto or listar_arquivos_runtime_existentes(pasta_runtime_files)):
        origem = pasta_runtime_files / relativo
        if origem.exists() and origem.is_file():
            arquivos_disponiveis.append(relativo)

    if not arquivos_disponiveis:
        raise ValueError("Nenhum arquivo valido foi encontrado no pacote de atualizacao do runtime.")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    hash_curto = str(manifesto_runtime.get("runtime_hash") or "manual")[:10] or "manual"
    pasta_backup = destino_base / DIRETORIO_ATUALIZACOES_RUNTIME / "backups" / f"{timestamp}_{hash_curto}"
    pasta_backup.mkdir(parents=True, exist_ok=True)

    arquivos_aplicados = []
    for relativo in arquivos_disponiveis:
        destino = destino_base / relativo
        if destino.exists() and destino.is_file():
            backup = pasta_backup / relativo
            backup.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(destino, backup)
        destino.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(pasta_runtime_files / relativo, destino)
        arquivos_aplicados.append(relativo)

    manifesto_local = gerar_manifesto_runtime(destino_base)
    return {
        "applied": True,
        "applied_files": arquivos_aplicados,
        "backup_dir": str(pasta_backup),
        "local_runtime": manifesto_local,
        "source_runtime_hash": str(manifesto_runtime.get("runtime_hash") or ""),
    }


def comparar_runtime_manifestos(local_manifesto, remoto_manifesto):
    local_hash = str((local_manifesto or {}).get("runtime_hash") or "")
    remoto_hash = str((remoto_manifesto or {}).get("runtime_hash") or "")
    return {
        "same_hash": bool(local_hash and remoto_hash and local_hash == remoto_hash),
        "local_hash": local_hash,
        "remote_hash": remoto_hash,
        "local_supported_block_types": sorted(local_manifesto.get("supported_block_types") or []),
        "remote_supported_block_types": sorted(remoto_manifesto.get("supported_block_types") or []),
    }


def definir_automacao_ativa(automacao_id):
    garantir_diretorio_automacoes()
    if not automacao_id:
        if ARQUIVO_AUTOMACAO_ATUAL.exists():
            ARQUIVO_AUTOMACAO_ATUAL.unlink()
        return
    ARQUIVO_AUTOMACAO_ATUAL.write_text(
        json.dumps({"id": str(automacao_id)}, ensure_ascii=False),
        encoding="utf-8",
    )


def obter_automacao_ativa_id():
    if not ARQUIVO_AUTOMACAO_ATUAL.exists():
        return ""
    try:
        return str(ler_json_arquivo(ARQUIVO_AUTOMACAO_ATUAL).get("id") or "")
    except Exception:
        return ""


def localizar_item_biblioteca(biblioteca, automacao_id):
    for indice, item in enumerate(biblioteca.get("items", [])):
        if str(item.get("id")) == str(automacao_id):
            return indice, item
    raise ValueError(f"Automacao nao encontrada: {automacao_id}")


def copiar_arquivo_se_existir(origem, destino):
    origem = Path(origem)
    destino = Path(destino)
    if not origem.exists() or origem.name.startswith("~$"):
        return False
    destino.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copy2(origem, destino)
        return True
    except OSError:
        # Arquivos abertos pelo Excel/Windows podem ficar bloqueados. Nao deixe
        # isso impedir a troca de automacao; o usuario pode fechar e recarregar.
        return False


def sincronizar_diretorio(origem, destino):
    origem = Path(origem)
    destino = Path(destino)
    if not origem.exists():
        return
    destino.mkdir(parents=True, exist_ok=True)
    for caminho in origem.rglob("*"):
        if caminho.is_dir() or caminho.name.startswith("~$"):
            continue
        try:
            relativo = caminho.relative_to(origem)
        except ValueError:
            continue
        copiar_arquivo_se_existir(caminho, destino / relativo)


def contar_blocos_payload(payload):
    return len((payload or {}).get("eventos", []))


def migrar_legado_se_necessario(biblioteca):
    if biblioteca.get("items"):
        return biblioteca
    if not ARQUIVO_SAIDA.exists():
        return biblioteca
    payload = carregar_payload()
    if not payload.get("eventos") and not payload.get("resumo_digitado"):
        return biblioteca
    automacao_id = gerar_id_automacao()
    item = normalizar_item_biblioteca(
        {
            "id": automacao_id,
            "nome": "Automacao atual",
            "descricao": "Importada automaticamente do roteiro existente.",
            "categoria": "Outros",
            "criado_em": payload.get("criado_em") or agora_iso(),
            "atualizado_em": agora_iso(),
        }
    )
    biblioteca["items"].append(item)
    salvar_payload_em_automacao(automacao_id, payload)
    definir_automacao_ativa(automacao_id)
    return biblioteca


def salvar_payload_em_automacao(automacao_id, payload):
    pasta = diretorio_automacao(automacao_id)
    pasta.mkdir(parents=True, exist_ok=True)
    normalizar_eventos(payload)
    normalizar_sequencia_inicial_payload(payload)
    normalizar_configuracoes_sequencia(payload)
    roteiro_texto = json.dumps(payload, indent=2, ensure_ascii=False)
    (pasta / ARQUIVO_SAIDA.name).write_text(roteiro_texto, encoding="utf-8")
    (pasta / ARQUIVO_RESUMO.name).write_text(montar_resumo(payload["eventos"]), encoding="utf-8")
    (pasta / ARQUIVO_SCRIPT.name).write_text(montar_script_python(payload["eventos"]), encoding="utf-8")
    copiar_arquivo_se_existir(ARQUIVO_PLANILHA, pasta / ARQUIVO_PLANILHA.name)
    sincronizar_diretorio(DIRETORIO_PLANILHAS, pasta / DIRETORIO_PLANILHAS.name)
    sincronizar_diretorio(PASTA_IMAGENS, pasta / PASTA_IMAGENS.name)


def sincronizar_automacao_ativa(payload):
    automacao_id = obter_automacao_ativa_id()
    if not automacao_id:
        return
    salvar_payload_em_automacao(automacao_id, payload)
    biblioteca = carregar_biblioteca()
    indice, item = localizar_item_biblioteca(biblioteca, automacao_id)
    item["atualizado_em"] = agora_iso()
    item["sequencia_inicial"] = normalizar_sequencia_inicial_payload(payload)
    biblioteca["items"][indice] = item
    salvar_biblioteca(biblioteca)


def carregar_payload_da_automacao(automacao_id):
    arquivo = arquivo_roteiro_automacao(automacao_id)
    if not arquivo.exists():
        return criar_payload_vazio()
    return ler_json_arquivo(arquivo)


def carregar_automacao_na_raiz(automacao_id):
    payload = carregar_payload_da_automacao(automacao_id)
    pasta = diretorio_automacao(automacao_id)
    copiar_arquivo_se_existir(pasta / ARQUIVO_PLANILHA.name, ARQUIVO_PLANILHA)
    if (pasta / DIRETORIO_PLANILHAS.name).exists():
        sincronizar_diretorio(pasta / DIRETORIO_PLANILHAS.name, DIRETORIO_PLANILHAS)
    if (pasta / PASTA_IMAGENS.name).exists():
        sincronizar_diretorio(pasta / PASTA_IMAGENS.name, PASTA_IMAGENS)
    definir_automacao_ativa(automacao_id)
    salvar_payload(payload)
    return payload


def escrever_payload_em_pasta(destino, payload):
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    payload_copia = json.loads(json.dumps(payload, ensure_ascii=False))
    payload_copia["criado_em"] = str(payload_copia.get("criado_em") or agora_iso())
    payload_copia["instrucoes"] = list(INSTRUCOES_PADRAO)
    normalizar_eventos(payload_copia)
    normalizar_sequencia_inicial_payload(payload_copia)
    normalizar_configuracoes_sequencia(payload_copia)
    payload_copia["resumo_digitado"] = montar_resumo(payload_copia.get("eventos", []))
    (destino / ARQUIVO_SAIDA.name).write_text(
        json.dumps(payload_copia, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    (destino / ARQUIVO_RESUMO.name).write_text(payload_copia["resumo_digitado"], encoding="utf-8")
    (destino / ARQUIVO_SCRIPT.name).write_text(
        montar_script_python(payload_copia.get("eventos", [])),
        encoding="utf-8",
    )
    return payload_copia


def contar_sequencias_payload(payload):
    eventos = (payload or {}).get("eventos", [])
    return len({normalizar_nome_sequencia(evento.get("sequencia")) for evento in eventos}) or 1


def exportar_automacao_para_zip(automacao_id, item, destino_zip=None):
    origem = diretorio_automacao(automacao_id)
    if not origem.exists():
        raise ValueError(f"Automacao nao encontrada para exportar: {automacao_id}")

    payload = carregar_payload_da_automacao(automacao_id)
    nome_zip = normalizar_nome_arquivo(item.get("nome") or automacao_id, "automacao")
    if destino_zip:
        destino_zip = Path(destino_zip)
        if destino_zip.suffix.lower() != ".zip":
            destino_zip = destino_zip.with_suffix(".zip")
        destino_zip.parent.mkdir(parents=True, exist_ok=True)
    else:
        destino_zip = caminho_arquivo_unico(DIRETORIO_EXPORTACOES_AUTOMACAO, f"{nome_zip}.zip")

    manifesto_runtime = gerar_manifesto_runtime()
    compatibilidade = gerar_resumo_compatibilidade_payload(payload, manifesto_runtime)
    manifesto = {
        "bundle_type": BUNDLE_AUTOMACAO_TIPO,
        "bundle_version": BUNDLE_AUTOMACAO_VERSAO,
        "exported_at": agora_iso(),
        "automation": {
            "nome": str(item.get("nome") or "Automacao importada"),
            "descricao": str(item.get("descricao") or ""),
            "categoria": str(item.get("categoria") or "Outros"),
            "criado_em": str(item.get("criado_em") or agora_iso()),
            "atualizado_em": str(item.get("atualizado_em") or agora_iso()),
            "sequencia_inicial": str(item.get("sequencia_inicial") or "principal"),
            "command_count": contar_blocos_payload(payload),
            "sequence_count": contar_sequencias_payload(payload),
        },
        "compatibility": compatibilidade,
        "runtime": manifesto_runtime,
    }

    with tempfile.TemporaryDirectory(prefix="flowdetect_export_") as temp_dir:
        raiz_bundle = Path(temp_dir) / "bundle"
        raiz_bundle.mkdir(parents=True, exist_ok=True)
        shutil.copytree(origem, raiz_bundle / "automacao", dirs_exist_ok=True)
        exportar_runtime_para_pasta(raiz_bundle)
        (raiz_bundle / ARQUIVO_BUNDLE_AUTOMACAO).write_text(
            json.dumps(manifesto, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        (raiz_bundle / "COMO_IMPORTAR.txt").write_text(
            "Abra o FlowDetect no outro computador e use a opcao Importar ZIP na biblioteca de automacoes.",
            encoding="utf-8",
        )
        with zipfile.ZipFile(destino_zip, "w", compression=zipfile.ZIP_DEFLATED) as arquivo_zip:
            for caminho in sorted(raiz_bundle.rglob("*")):
                if caminho.is_file():
                    arquivo_zip.write(caminho, caminho.relative_to(raiz_bundle))

    return destino_zip


def extrair_zip_com_seguranca(arquivo_zip, destino):
    destino = Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    destino_resolvido = destino.resolve()
    with zipfile.ZipFile(arquivo_zip, "r") as pacote:
        for item in pacote.infolist():
            caminho_item = (destino / item.filename).resolve()
            if caminho_item != destino_resolvido and destino_resolvido not in caminho_item.parents:
                raise ValueError("O arquivo ZIP contem caminhos invalidos e nao pode ser importado.")
            pacote.extract(item, destino)


def localizar_manifesto_bundle(diretorio):
    diretorio = Path(diretorio)
    manifestos = sorted(diretorio.rglob(ARQUIVO_BUNDLE_AUTOMACAO), key=lambda caminho: len(caminho.parts))
    return manifestos[0] if manifestos else None


def localizar_pasta_automacao_importada(diretorio, manifesto_path=None):
    diretorio = Path(diretorio)
    if manifesto_path:
        candidata = manifesto_path.parent / "automacao"
        if (candidata / ARQUIVO_SAIDA.name).exists():
            return candidata
    candidatos = sorted(diretorio.rglob(ARQUIVO_SAIDA.name), key=lambda caminho: len(caminho.parts))
    return candidatos[0].parent if candidatos else None


def listar_automacoes_resumidas():
    biblioteca = carregar_biblioteca()
    ativos = obter_automacao_ativa_id()
    itens = []
    for item in biblioteca.get("items", []):
        payload = carregar_payload_da_automacao(item["id"])
        itens.append(
            {
                **item,
                "active": str(item["id"]) == str(ativos),
                "command_count": contar_blocos_payload(payload),
                "sequence_count": len({normalizar_nome_sequencia(evento.get("sequencia")) for evento in payload.get("eventos", [])}) or 1,
            }
        )
    categorias = sorted({str(item.get("categoria") or "Outros") for item in itens})
    return {
        "items": itens,
        "active_id": ativos,
        "categories": categorias,
    }


def normalizar_id_planilha(workbook_file=None):
    texto = str(workbook_file or ARQUIVO_PLANILHA.name).strip().replace("\\", "/")
    if not texto:
        return ARQUIVO_PLANILHA.name
    return texto


def normalizar_origem_ocr_comparacao(source=None):
    texto = str(source or "ultimo_ocr").strip()
    if not texto or texto in {"ultimo_ocr", "ocr_tela"}:
        return texto or "ultimo_ocr"
    compacto = texto.lower().replace(" ", "")
    encontrado = re.search(r"(?:ocr_bloco_|bloco[:_-]?)(\d+)$", compacto)
    if encontrado:
        return f"ocr_bloco_{encontrado.group(1)}"
    if compacto.isdigit():
        return f"ocr_bloco_{compacto}"
    return texto


def resolver_planilha(workbook_file=None):
    workbook_id = normalizar_id_planilha(workbook_file)
    caminho = Path(workbook_id)
    if caminho.is_absolute():
        return caminho
    if caminho.exists():
        return caminho
    if "/" not in workbook_id and "\\" not in workbook_id and workbook_id != ARQUIVO_PLANILHA.name:
        alternativa = DIRETORIO_PLANILHAS / workbook_id
        if alternativa.exists():
            return alternativa
    return caminho


def id_planilha_por_caminho(caminho):
    caminho = Path(caminho)
    try:
        return caminho.relative_to(Path.cwd()).as_posix()
    except ValueError:
        return caminho.as_posix()


def normalizar_id_imagem(imagem):
    texto = str(imagem or "").strip().replace("\\", "/")
    if not texto:
        return ""
    caminho = Path(texto)
    if caminho.is_absolute():
        return caminho.as_posix()
    if caminho.exists():
        try:
            return caminho.relative_to(Path.cwd()).as_posix()
        except ValueError:
            return caminho.as_posix()
    if "/" not in texto and "\\" not in texto:
        return (PASTA_IMAGENS / texto).as_posix()
    return texto


def normalizar_nome_sequencia(nome):
    texto = str(nome or "").strip()
    return texto or "principal"


def normalizar_sequencia_inicial_payload(payload):
    if not isinstance(payload, dict):
        return "principal"
    payload["sequencia_inicial"] = normalizar_nome_sequencia(payload.get("sequencia_inicial"))
    return payload["sequencia_inicial"]


def normalizar_modo_execucao_sequencia(modo):
    texto = str(modo or "normal").strip().lower()
    return texto if texto in MODOS_EXECUCAO_SEQUENCIA else "normal"


def normalizar_configuracoes_sequencia(payload):
    bruto = payload.get("configuracoes_sequencia")
    if not isinstance(bruto, dict):
        bruto = {}

    normalizado = {}
    for nome_sequencia, configuracao in bruto.items():
        sequencia = normalizar_nome_sequencia(nome_sequencia)
        config = configuracao if isinstance(configuracao, dict) else {}
        modo_execucao = normalizar_modo_execucao_sequencia(
            config.get("modo_execucao") or config.get("execution_mode")
        )
        sequencia_apos_loop = str(
            config.get("sequencia_apos_loop") or config.get("after_loop_sequence") or ""
        ).strip()
        sequencia_apos_loop = normalizar_nome_sequencia(sequencia_apos_loop) if sequencia_apos_loop else ""
        if sequencia_apos_loop == sequencia:
            sequencia_apos_loop = ""
        if modo_execucao != "loop_planilha":
            sequencia_apos_loop = ""
        if modo_execucao != "normal" or sequencia_apos_loop:
            normalizado[sequencia] = {
                "modo_execucao": modo_execucao,
                "sequencia_apos_loop": sequencia_apos_loop,
            }

    payload["configuracoes_sequencia"] = normalizado
    return normalizado


def listar_arquivos_planilha():
    encontrados = []
    vistos = set()

    candidatos = [ARQUIVO_PLANILHA]
    if DIRETORIO_PLANILHAS.exists():
        candidatos.extend(sorted(DIRETORIO_PLANILHAS.iterdir()))

    for caminho in candidatos:
        if (
            not caminho.exists()
            or caminho.suffix.lower() not in EXTENSOES_PLANILHA
            or caminho.name.startswith("~$")
        ):
            continue
        workbook_id = id_planilha_por_caminho(caminho)
        if workbook_id in vistos:
            continue
        vistos.add(workbook_id)
        encontrados.append(
            {
                "id": workbook_id,
                "file_name": caminho.name,
                "name": caminho.stem,
                "path": str(caminho.resolve()),
                "exists": True,
            }
        )

    if ARQUIVO_PLANILHA.exists() and ARQUIVO_PLANILHA.name not in vistos:
        encontrados.insert(
            0,
            {
                "id": ARQUIVO_PLANILHA.name,
                "file_name": ARQUIVO_PLANILHA.name,
                "name": ARQUIVO_PLANILHA.stem,
                "path": str(ARQUIVO_PLANILHA.resolve()),
                "exists": ARQUIVO_PLANILHA.exists(),
            },
        )

    return encontrados


def tecla_para_exibicao(nome):
    if len(nome) == 1:
        if nome == " ":
            return "SPACE"
        return nome.upper()
    return MAPA_EXIBICAO.get(nome, nome.upper())


def tecla_para_pyautogui(nome):
    if len(nome) == 1:
        return nome
    return MAPA_PYAUTOGUI.get(nome, nome)


def eh_texto_simples(nome):
    return len(nome) == 1 and nome not in {"\r", "\n", "\t"}


def normalizar_acao_apos_item(texto):
    texto_limpo = (texto or "").strip().lower()
    if not texto_limpo or texto_limpo in {"nenhuma", "somente_contexto", "contexto", "so_contexto"}:
        return []
    return [parte.strip() for parte in texto_limpo.replace(" ", "").split("+") if parte.strip()]


def normalizar_modo_area_busca_reconhecimento(valor):
    texto = str(valor or "").strip().lower()
    aliases = {
        "janela": "target_window",
        "janela_alvo": "target_window",
        "janela_alvo_inteira": "target_window",
        "window": "target_window",
        "target_window": "target_window",
        "area": "custom_area",
        "area_personalizada": "custom_area",
        "custom": "custom_area",
        "custom_area": "custom_area",
        "ultima": "last_found",
        "ultima_posicao": "last_found",
        "proximo_ultima_posicao": "last_found",
        "cache": "last_found",
        "last_found": "last_found",
    }
    return aliases.get(texto, "target_window")


def escape_css_attr_value(valor):
    return str(valor or "").replace("\\", "\\\\").replace('"', '\\"')


def xpath_literal(valor):
    texto = str(valor or "")
    if '"' not in texto:
        return f'"{texto}"'
    if "'" not in texto:
        return f"'{texto}'"
    partes = texto.split('"')
    tokens = []
    for indice, parte in enumerate(partes):
        if parte:
            tokens.append(f'"{parte}"')
        if indice < len(partes) - 1:
            tokens.append("'\"'")
    return f"concat({', '.join(tokens)})"


def id_parece_dinamico(valor):
    texto = str(valor or "").strip()
    if not texto:
        return False
    if re.fullmatch(r"id[0-9a-f]{8,}", texto, flags=re.IGNORECASE):
        return True
    if re.fullmatch(r"[a-f0-9]{10,}", texto, flags=re.IGNORECASE):
        return True
    if re.search(r"(_ngcontent|react|ember|chakra|mui|ant|css-)", texto, flags=re.IGNORECASE):
        return True
    if len(texto) >= 14 and sum(char.isdigit() for char in texto) >= 6:
        return True
    return False


class ElementoHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root_tag = ""
        self.root_attrs = {}
        self.stack = []
        self.textos = []

    def handle_starttag(self, tag, attrs):
        atributos = {str(chave): "" if valor is None else str(valor) for chave, valor in attrs}
        if not self.root_tag:
            self.root_tag = str(tag or "").lower()
            self.root_attrs = atributos
        self.stack.append(str(tag or "").lower())

    def handle_startendtag(self, tag, attrs):
        atributos = {str(chave): "" if valor is None else str(valor) for chave, valor in attrs}
        if not self.root_tag:
            self.root_tag = str(tag or "").lower()
            self.root_attrs = atributos

    def handle_endtag(self, tag):
        nome = str(tag or "").lower()
        if self.stack and self.stack[-1] == nome:
            self.stack.pop()

    def handle_data(self, data):
        if self.root_tag and self.stack:
            texto = str(data or "").strip()
            if texto:
                self.textos.append(texto)


def classificar_tipo_elemento_html(elemento):
    tag = str((elemento or {}).get("tag", "")).lower()
    tipo_input = str((elemento or {}).get("type", "")).lower()
    role = str((elemento or {}).get("role", "")).lower()
    if tag == "select":
        return "select"
    if tag == "textarea":
        return "textarea"
    if tag == "a" or role == "link":
        return "link"
    if tag == "button" or role == "button" or tipo_input in {"button", "submit", "reset"}:
        return "button"
    if tag == "input":
        if tipo_input in {"checkbox", "radio"}:
            return "choice"
        if tipo_input in {"text", "email", "password", "search", "tel", "number", "url", ""}:
            return "text_input"
    if str((elemento or {}).get("contenteditable", "")).lower() in {"true", "plaintext-only"}:
        return "editable"
    return "generic"


def sugestao_valor_chrome(elemento):
    pistas = " ".join(
        [
            str((elemento or {}).get("placeholder", "")),
            str((elemento or {}).get("aria_label", "")),
            str((elemento or {}).get("name", "")),
            str((elemento or {}).get("text", "")),
        ]
    ).lower()
    if "cnpj" in pistas:
        return "{{CNPJ}}"
    if "cpf" in pistas:
        return "{{CPF}}"
    if "empresa" in pistas:
        return "{{EMPRESA}}"
    if "nome" in pistas:
        return "{{NOME}}"
    return ""


def sugerir_nome_bloco_chrome(elemento, acao_tipo):
    placeholder = str((elemento or {}).get("placeholder", "")).strip()
    aria_label = str((elemento or {}).get("aria_label", "")).strip()
    texto = str((elemento or {}).get("text", "")).strip()
    alvo = placeholder or aria_label or texto or str((elemento or {}).get("name", "")).strip()
    alvo = alvo or str((elemento or {}).get("tag", "elemento")).strip()
    alvo = re.sub(r"\s+", " ", alvo).strip()
    if acao_tipo == "preencher":
        return f"Preencher {alvo}"
    if acao_tipo == "clicar":
        return f"Clicar {alvo}"
    if acao_tipo.startswith("selecionar"):
        return f"Selecionar {alvo}"
    return f"Usar {alvo}"


def gerar_seletores(elemento):
    elemento = elemento or {}
    tag = str(elemento.get("tag", "")).strip().lower() or "*"
    role = str(elemento.get("role", "")).strip().lower()
    seletores = []

    def adicionar(metodo, valor, codigo, confianca="media", motivo="", extra=None):
        if not str(valor or "").strip():
            return
        seletores.append(
            {
                "metodo": metodo,
                "valor": str(valor).strip(),
                "codigo_playwright": codigo,
                "confianca": confianca,
                "motivo": motivo,
                **(extra or {}),
            }
        )

    placeholder = str(elemento.get("placeholder", "")).strip()
    name = str(elemento.get("name", "")).strip()
    aria_label = str(elemento.get("aria_label", "")).strip()
    aria_labelledby = str(elemento.get("aria_labelledby", "")).strip()
    element_id = str(elemento.get("id", "")).strip()
    texto = str(elemento.get("text", "")).strip()
    texto_proximo = str(elemento.get("texto_proximo", "") or elemento.get("nearby_text", "")).strip()
    css_capturado = str(elemento.get("css_selector", "") or elemento.get("css", "")).strip()
    xpath_capturado = str(elemento.get("xpath", "")).strip()
    tipo = str(elemento.get("type", "")).strip().lower()

    adicionar(
        "placeholder",
        placeholder,
        f"page.get_by_placeholder({placeholder!r})",
        "alta",
        extra={"seletor_css": f'{tag}[placeholder="{escape_css_attr_value(placeholder)}"]'},
    )
    adicionar(
        "name",
        name,
        f"page.locator({f'{tag}[name=\"{escape_css_attr_value(name)}\"]'!r})",
        "alta",
        extra={"seletor_css": f'{tag}[name="{escape_css_attr_value(name)}"]'},
    )
    adicionar(
        "aria-label",
        aria_label,
        f"page.get_by_label({aria_label!r})",
        "alta",
        extra={"seletor_css": f'{tag}[aria-label="{escape_css_attr_value(aria_label)}"]'},
    )
    adicionar(
        "labelledby",
        aria_labelledby,
        f"page.locator({f'{tag}[aria-labelledby=\"{escape_css_attr_value(aria_labelledby)}\"]'!r})",
        "media",
        "Referencia indireta via aria-labelledby.",
        extra={"seletor_css": f'{tag}[aria-labelledby="{escape_css_attr_value(aria_labelledby)}"]'},
    )
    if element_id:
        confianca_id = "baixa" if id_parece_dinamico(element_id) else "media"
        motivo_id = "ID parece gerado dinamicamente" if confianca_id == "baixa" else ""
        adicionar(
            "id",
            element_id,
            f"page.locator({f'[id=\"{escape_css_attr_value(element_id)}\"]'!r})",
            confianca_id,
            motivo_id,
            extra={"seletor_css": f'[id="{escape_css_attr_value(element_id)}"]'},
        )
    if texto:
        adicionar(
            "text",
            texto,
            f"page.get_by_text({texto!r})",
            "media",
            extra={"seletor_css": tag},
        )
    if texto_proximo:
        adicionar(
            "nearby_text",
            texto_proximo,
            f"page.get_by_text({texto_proximo!r})",
            "media",
            "Texto perto do elemento capturado.",
        )
    if role and texto:
        adicionar(
            "role",
            texto,
            f"page.get_by_role({role!r}, name={texto!r})",
            "alta",
            extra={"role": role},
        )

    if css_capturado:
        adicionar(
            "css",
            css_capturado,
            f"page.locator({css_capturado!r})",
            "media",
            "CSS capturado direto da pagina.",
            extra={"seletor_css": css_capturado},
        )

    atributos_css = []
    for chave in ("placeholder", "name", "aria-label", "type"):
        valor = str(elemento.get(chave.replace("-", "_"), "") or elemento.get(chave, "")).strip()
        if valor:
            atributos_css.append((chave, valor))
    if not atributos_css and tipo:
        atributos_css.append(("type", tipo))
    if atributos_css:
        seletor_css = tag + "".join(
            f'[{atributo}="{escape_css_attr_value(valor)}"]' for atributo, valor in atributos_css
        )
        adicionar(
            "css",
            seletor_css,
            f"page.locator({seletor_css!r})",
            "alta" if any(atributo in {"placeholder", "name", "aria-label"} for atributo, _ in atributos_css) else "media",
            extra={"seletor_css": seletor_css},
        )

    predicados_xpath = []
    for atributo in ("placeholder", "name", "aria-label", "type"):
        valor = str(elemento.get(atributo.replace("-", "_"), "") or elemento.get(atributo, "")).strip()
        if valor:
            predicados_xpath.append(f"@{atributo}={xpath_literal(valor)}")
    if texto:
        predicados_xpath.append(f"normalize-space(.)={xpath_literal(texto)}")
    if not predicados_xpath:
        predicados_xpath.append("1=1")
    xpath = xpath_capturado or f"//{tag}[{' and '.join(predicados_xpath)}]"
    adicionar(
        "xpath",
        xpath,
        f"page.locator({f'xpath={xpath}'!r})",
        "baixa" if not xpath_capturado else "media",
        "XPath deve ser ultimo recurso." if not xpath_capturado else "XPath capturado direto da pagina.",
    )

    prioridade = {
        "id": 0,
        "name": 1,
        "placeholder": 2,
        "aria-label": 3,
        "text": 4,
        "nearby_text": 5,
        "css": 6,
        "labelledby": 7,
        "role": 8,
        "xpath": 8,
    }
    confianca_score = {"alta": 0, "media": 1, "baixa": 2}
    seletores.sort(
        key=lambda seletor: (
            prioridade.get(seletor.get("metodo"), 99),
            confianca_score.get(seletor.get("confianca"), 9),
            len(str(seletor.get("valor", ""))),
        )
    )
    return seletores


def parse_html_elemento(html):
    trecho = unescape(str(html or "").strip())
    if not trecho:
        raise ValueError("Nenhum HTML foi informado para analise.")

    parser = ElementoHTMLParser()
    parser.feed(trecho)
    parser.close()

    if not parser.root_tag:
        raise ValueError("Nao foi possivel identificar um elemento HTML valido no trecho informado.")

    atributos = {str(chave): str(valor) for chave, valor in (parser.root_attrs or {}).items()}
    classes = [classe for classe in str(atributos.get("class", "")).split() if classe]
    elemento = {
        "tag": parser.root_tag,
        "type": str(atributos.get("type", "")).strip().lower(),
        "id": str(atributos.get("id", "")).strip(),
        "name": str(atributos.get("name", "")).strip(),
        "placeholder": str(atributos.get("placeholder", "")).strip(),
        "aria_label": str(atributos.get("aria-label", "")).strip(),
        "aria_labelledby": str(atributos.get("aria-labelledby", "")).strip(),
        "role": str(atributos.get("role", "")).strip().lower(),
        "classes": classes,
        "text": " ".join(parser.textos).strip(),
        "value": str(atributos.get("value", "")).strip(),
        "maxlength": str(atributos.get("maxlength", "")).strip(),
        "required": "required" in atributos and str(atributos.get("required", "true")).lower() not in {"false", "0", "nao"},
        "attributes": atributos,
        "contenteditable": str(atributos.get("contenteditable", "")).strip(),
    }
    tipo_provavel = classificar_tipo_elemento_html(elemento)
    seletores = gerar_seletores(elemento)
    alvo_recomendado = dict(seletores[0]) if seletores else {
        "metodo": "css",
        "valor": elemento["tag"],
        "codigo_playwright": f"page.locator({elemento['tag']!r})",
        "confianca": "baixa",
        "seletor_css": elemento["tag"],
    }
    acao_padrao = "preencher" if tipo_provavel in {"text_input", "textarea", "editable"} else (
        "selecionar_texto" if tipo_provavel == "select" else "clicar"
    )
    nome_sugerido = sugerir_nome_bloco_chrome(elemento, acao_padrao)
    return {
        "html_original": trecho,
        "elemento": elemento,
        "tipo_provavel": tipo_provavel,
        "seletores": seletores,
        "alvo_recomendado": alvo_recomendado,
        "acao_sugerida": {
            "tipo": acao_padrao,
            "origem_valor": "texto_fixo",
            "valor": sugestao_valor_chrome(elemento),
            "limpar_antes": True,
            "pressionar_depois": "enter" if acao_padrao == "preencher" else "",
            "salvar_em": "",
        },
        "nome_sugerido": nome_sugerido,
        "confianca": alvo_recomendado.get("confianca", "media"),
    }


def formatar_evento(teclas, tipo_evento, texto_livre=None, metadados=None):
    if tipo_evento == "texto_livre":
        return texto_livre or ""

    if tipo_evento == "inicio_loop":
        return "[INICIO LOOP]"

    if tipo_evento == "fim_loop":
        return "[FIM LOOP]"

    if tipo_evento in TIPOS_SEQUENCIA_PLANILHA:
        coluna = (metadados or {}).get("coluna", "Coluna")
        quantidade = len((metadados or {}).get("valores", []))
        acao = (metadados or {}).get("acao_apos_item_texto") or "nenhuma"
        if (metadados or {}).get("modo_loop") == "somente_contexto":
            return f"[LOOP DADOS:{coluna} x{quantidade} | sem digitar]"
        return f"[PLANILHA:{coluna} x{quantidade} | apos:{acao}]"

    if tipo_evento in TIPOS_ESCRITA_PLANILHA:
        detalhes = metadados or {}
        coluna = str(detalhes.get("coluna", "Coluna") or "Coluna")
        valor = str(detalhes.get("valor", "") or "")
        return f"[ESCREVER PLANILHA:{coluna}='{valor[:28]}' | linha atual]"

    if tipo_evento in TIPOS_COMPARACAO_PLANILHA:
        detalhes = metadados or {}
        coluna = str(detalhes.get("coluna_esperada", "Coluna") or "Coluna")
        status = str(detalhes.get("coluna_status", "Status") or "Status")
        return f"[CONFERIR:{coluna} x OCR | resultado:{status}]"

    if tipo_evento in TIPOS_MOUSE:
        detalhes = metadados or {}
        x = detalhes.get("x", 0)
        y = detalhes.get("y", 0)
        botao = str(detalhes.get("botao", "left")).upper()
        cliques = max(1, int(detalhes.get("cliques", 1) or 1))
        sufixo = " duplo" if cliques == 2 else ""
        return f"[MOUSE:{botao}{sufixo} x={x} y={y}]"

    if tipo_evento in TIPOS_IMAGEM:
        detalhes = metadados or {}
        imagem = Path(str(detalhes.get("imagem", ""))).name or "imagem"
        timeout = float(detalhes.get("timeout", 10) or 10)
        confianca = float(detalhes.get("confianca", 0.8) or 0.8)
        perfil_busca = str(detalhes.get("perfil_busca", "balanced") or "balanced").strip().lower()
        modo_area_busca = normalizar_modo_area_busca_reconhecimento(detalhes.get("modo_area_busca", "target_window"))
        clicar = bool(detalhes.get("clicar", True))
        cliques = max(0, int(detalhes.get("cliques", 1 if clicar else 0) or 0))
        acao = str(detalhes.get("ao_nao_encontrar", "parar"))
        sequencia_falha = str(detalhes.get("sequencia_nao_encontrar", "") or "").strip()
        bloco_falha = str(detalhes.get("bloco_nao_encontrar", "") or "").strip()
        janela_alvo = str(detalhes.get("janela_alvo", "") or "").strip()
        texto_ocr = str(detalhes.get("texto_ocr_esperado", "") or "").strip()
        sufixo = "2 cliques" if cliques >= 2 else ("clicar" if clicar else "localizar")
        alvo = ""
        if acao == "ir_para_sequencia" and sequencia_falha:
            alvo = f" -> {sequencia_falha}"
        elif acao == "ir_para_bloco" and bloco_falha:
            alvo = f" -> bloco {bloco_falha}"
        janela = f" | janela:{janela_alvo}" if janela_alvo else ""
        ocr = f" | OCR:{texto_ocr}" if texto_ocr else ""
        return (
            f"[IMAGEM:{imagem} | {sufixo} | timeout={timeout:g}s | conf={confianca:g} "
            f"| perfil:{perfil_busca} | area:{modo_area_busca}{janela}{ocr} | falha:{acao}{alvo}]"
        )

    if tipo_evento in TIPOS_CHROME:
        detalhes = metadados or {}
        alvo = detalhes.get("alvo") or {}
        acao = detalhes.get("acao") or {}
        metodo = str(alvo.get("metodo", "css") or "css")
        valor = str(alvo.get("valor", "") or "")
        valor_curto = Path(valor).name if metodo == "xpath" else valor
        acao_sucesso = str(detalhes.get("ao_sucesso", "continuar") or "continuar")
        bloco_sucesso = str(detalhes.get("bloco_sucesso", "") or "").strip()
        sufixo_sucesso = f" -> bloco {bloco_sucesso}" if acao_sucesso == "ir_para_bloco" and bloco_sucesso else ""
        return (
            f"[CHROME:{metodo}={valor_curto[:42]} | "
            f"acao:{str(acao.get('tipo', 'clicar') or 'clicar')} | "
            f"timeout={int(detalhes.get('timeout', 10000) or 10000)}ms | "
            f"sucesso:{acao_sucesso}{sufixo_sucesso}]"
        )

    if tipo_evento == "abrir_url":
        detalhes = metadados or {}
        url = str(detalhes.get("url", "") or "").strip()
        navegador = str(detalhes.get("navegador", "padrao") or "padrao")
        modo = str(detalhes.get("modo_abertura", "nova_janela") or "nova_janela")
        tela_cheia = bool(detalhes.get("tela_cheia", False))
        aguardar = float(detalhes.get("aguardar_segundos", 3) or 0)
        sufixo_tela = " | tela cheia" if tela_cheia else ""
        return f"[ABRIR URL:{url[:64]} | {navegador} | {modo}{sufixo_tela} | aguardar={aguardar:g}s]"

    if tipo_evento == "abrir_programa":
        detalhes = metadados or {}
        caminho = Path(str(detalhes.get("caminho", "") or "")).name or str(detalhes.get("caminho", "") or "").strip()
        parametros = str(detalhes.get("parametros", "") or "").strip()
        aguardar = float(detalhes.get("aguardar_segundos", 2) or 0)
        sufixo = f" | params:{parametros[:32]}" if parametros else ""
        return f"[ABRIR PROGRAMA:{caminho}{sufixo} | aguardar={aguardar:g}s]"

    if tipo_evento == "salvar_download":
        detalhes = metadados or {}
        extensao = str(detalhes.get("extensao", ".pdf") or ".pdf").strip() or ".pdf"
        destino = str(detalhes.get("pasta_destino", "") or "").strip() or "Destino manual"
        nome = str(detalhes.get("nome_arquivo", "") or "").strip() or "nome original"
        timeout = float(detalhes.get("timeout", 60) or 60)
        return f"[SALVAR DOWNLOAD:{extensao} -> {destino[:42]} | nome:{nome[:32]} | timeout={timeout:g}s]"

    if tipo_evento in TIPOS_VERIFICACAO:
        detalhes = metadados or {}
        ativo = Path(str(detalhes.get("imagem_ativo", ""))).name or "ativo"
        desativado = Path(str(detalhes.get("imagem_desativado", ""))).name or "desativado"
        return f"[ESTADO:{ativo}/{desativado}]"

    repeticoes = max(1, int((metadados or {}).get("repeticoes", 1) or 1))

    if tipo_evento in {"tecla", "atalho"} and not teclas:
        return "[TECLA]"

    if len(teclas) == 1:
        nome = teclas[0]
        exibicao = tecla_para_exibicao(nome)
        if nome == "enter":
            return f"[ENTER x{repeticoes}]\n" if repeticoes > 1 else "[ENTER]\n"
        if eh_texto_simples(nome):
            return f"{nome} x{repeticoes}" if repeticoes > 1 else nome
        return f"[{exibicao} x{repeticoes}]" if repeticoes > 1 else f"[{exibicao}]"

    partes = [tecla_para_exibicao(tecla) for tecla in teclas]
    atalho = f"[{'+'.join(partes)}]"
    return f"{atalho} x{repeticoes}" if repeticoes > 1 else atalho


def montar_comando_python(teclas, tipo_evento, texto_livre=None, metadados=None):
    if tipo_evento == "texto_livre":
        return f"digitar_texto({texto_livre!r})"

    if tipo_evento == "inicio_loop":
        return "# inicio_loop"

    if tipo_evento == "fim_loop":
        return "# fim_loop"

    if tipo_evento in TIPOS_SEQUENCIA_PLANILHA:
        detalhes = metadados or {}
        valores = detalhes.get("valores", [])
        acao_apos_item = detalhes.get("acao_apos_item", [])
        intervalo_entre_itens = detalhes.get(
            "intervalo_entre_itens",
            INTERVALO_ENTRE_ITENS_PLANILHA,
        )
        return (
            f"executar_sequencia({valores!r}, "
            f"acao_apos_item={acao_apos_item!r}, "
            f"intervalo_entre_itens={intervalo_entre_itens!r})"
        )

    if tipo_evento in TIPOS_ESCRITA_PLANILHA:
        detalhes = metadados or {}
        return (
            "escrever_planilha("
            f"aba={str(detalhes.get('aba', ''))!r}, "
            f"coluna={str(detalhes.get('coluna', ''))!r}, "
            f"valor={str(detalhes.get('valor', ''))!r}, "
            "linha='loop_atual')"
        )

    if tipo_evento in TIPOS_COMPARACAO_PLANILHA:
        detalhes = metadados or {}
        return (
            "comparar_planilha_com_ocr("
            f"coluna_esperada={str(detalhes.get('coluna_esperada', ''))!r}, "
            f"origem_ocr={str(detalhes.get('origem_tela', 'ultimo_ocr'))!r}, "
            f"status={str(detalhes.get('coluna_status', 'Status'))!r})"
        )

    if tipo_evento in TIPOS_MOUSE:
        detalhes = metadados or {}
        x = int(detalhes.get("x", 0))
        y = int(detalhes.get("y", 0))
        botao = str(detalhes.get("botao", "left"))
        cliques = max(1, int(detalhes.get("cliques", 1) or 1))
        return f'pyautogui.click(x={x}, y={y}, button="{botao}", clicks={cliques})'

    if tipo_evento in TIPOS_IMAGEM:
        detalhes = metadados or {}
        return (
            "clicar_por_imagem("
            f"imagem={str(detalhes.get('imagem', ''))!r}, "
            f"timeout={float(detalhes.get('timeout', 10) or 10)!r}, "
            f"confianca={float(detalhes.get('confianca', 0.8) or 0.8)!r}, "
            f"perfil_busca={str(detalhes.get('perfil_busca', 'balanced') or 'balanced')!r}, "
            f"modo_area_busca={normalizar_modo_area_busca_reconhecimento(detalhes.get('modo_area_busca', 'target_window'))!r}, "
            f"area_busca_x={int(detalhes.get('area_busca_x', 0) or 0)!r}, "
            f"area_busca_y={int(detalhes.get('area_busca_y', 0) or 0)!r}, "
            f"area_busca_largura={int(detalhes.get('area_busca_largura', 0) or 0)!r}, "
            f"area_busca_altura={int(detalhes.get('area_busca_altura', 0) or 0)!r}, "
            f"clicar={bool(detalhes.get('clicar', True))!r}, "
            f"clicks={max(0, int(detalhes.get('cliques', 1 if detalhes.get('clicar', True) else 0) or 0))!r}, "
            f"offset_x={int(detalhes.get('offset_x', 0) or 0)!r}, "
            f"offset_y={int(detalhes.get('offset_y', 0) or 0)!r}, "
            f"janela_alvo={str(detalhes.get('janela_alvo', ''))!r}, "
            f"texto_ocr_esperado={str(detalhes.get('texto_ocr_esperado', ''))!r}, "
            f"ao_nao_encontrar={str(detalhes.get('ao_nao_encontrar', 'parar'))!r}, "
            f"sequencia_nao_encontrar={str(detalhes.get('sequencia_nao_encontrar', ''))!r}, "
            f"bloco_nao_encontrar={detalhes.get('bloco_nao_encontrar', '')!r}, "
            f"quantidade_tentativas={int(detalhes.get('quantidade_tentativas', 1) or 1)!r}, "
            f"intervalo_tentativas={float(detalhes.get('intervalo_tentativas', 1) or 1)!r}"
            ")"
        )

    if tipo_evento in TIPOS_CHROME:
        detalhes = metadados or {}
        alvo = detalhes.get("alvo") or {}
        acao = detalhes.get("acao") or {}
        return (
            f"# chrome_elemento_html metodo={str(alvo.get('metodo', 'css'))!r} "
            f"valor={str(alvo.get('valor', ''))!r} acao={str(acao.get('tipo', 'clicar'))!r}"
        )

    if tipo_evento == "abrir_url":
        detalhes = metadados or {}
        return (
            "abrir_url("
            f"url={str(detalhes.get('url', ''))!r}, "
            f"navegador={str(detalhes.get('navegador', 'padrao'))!r}, "
            f"modo_abertura={str(detalhes.get('modo_abertura', 'nova_janela'))!r}, "
            f"aguardar_segundos={float(detalhes.get('aguardar_segundos', 3) or 0)!r}"
            ")"
        )

    if tipo_evento == "abrir_programa":
        detalhes = metadados or {}
        return (
            "abrir_programa("
            f"caminho={str(detalhes.get('caminho', ''))!r}, "
            f"parametros={str(detalhes.get('parametros', ''))!r}, "
            f"aguardar_segundos={float(detalhes.get('aguardar_segundos', 2) or 0)!r}"
            ")"
        )

    if tipo_evento == "salvar_download":
        detalhes = metadados or {}
        return (
            "salvar_download("
            f"pasta_origem={str(detalhes.get('pasta_origem', ''))!r}, "
            f"extensao={str(detalhes.get('extensao', '.pdf') or '.pdf')!r}, "
            f"pasta_destino={str(detalhes.get('pasta_destino', ''))!r}, "
            f"nome_arquivo={str(detalhes.get('nome_arquivo', ''))!r}, "
            f"timeout={float(detalhes.get('timeout', 60) or 60)!r}"
            ")"
        )

    if tipo_evento in TIPOS_VERIFICACAO:
        detalhes = metadados or {}
        return (
            "verificar_estado_botao("
            f"imagem_ativo={str(detalhes.get('imagem_ativo', ''))!r}, "
            f"imagem_desativado={str(detalhes.get('imagem_desativado', ''))!r}, "
            f"timeout={float(detalhes.get('timeout', 5) or 5)!r}, "
            f"confianca={float(detalhes.get('confianca', 0.85) or 0.85)!r}, "
            f"se_ativo={detalhes.get('se_ativo', {})!r}, "
            f"se_desativado={detalhes.get('se_desativado', {})!r}, "
            f"se_nao_encontrar={detalhes.get('se_nao_encontrar', {})!r}"
            ")"
        )

    repeticoes = max(1, int((metadados or {}).get("repeticoes", 1) or 1))
    return f"executar_teclas({teclas!r}, repeticoes={repeticoes})"


def montar_resumo(eventos):
    return "".join(evento.get("texto", "") for evento in eventos)


def montar_script_python(eventos):
    linhas = [
        "import pyautogui",
        "import time",
        "import ctypes",
        "",
        f"INTERVALO_TEXTO = {INTERVALO_TEXTO}",
        "",
        "",
        "def definir_texto_clipboard_windows(texto):",
        '    if not hasattr(ctypes, "windll"):',
        "        return False",
        "    user32 = ctypes.windll.user32",
        "    kernel32 = ctypes.windll.kernel32",
        "    kernel32.GlobalAlloc.restype = ctypes.c_void_p",
        "    kernel32.GlobalLock.restype = ctypes.c_void_p",
        "    user32.SetClipboardData.restype = ctypes.c_void_p",
        '    dados = (str(texto or "") + "\\0").encode("utf-16-le")',
        "    memoria = None",
        "    try:",
        "        aberto = False",
        "        for _ in range(6):",
        "            if user32.OpenClipboard(None):",
        "                aberto = True",
        "                break",
        "            time.sleep(0.01)",
        "        if not aberto:",
        "            return False",
        "        user32.EmptyClipboard()",
        "        memoria = kernel32.GlobalAlloc(0x0002, len(dados))",
        "        if not memoria:",
        "            return False",
        "        destino = kernel32.GlobalLock(memoria)",
        "        if not destino:",
        "            return False",
        "        ctypes.memmove(destino, dados, len(dados))",
        "        kernel32.GlobalUnlock(memoria)",
        "        if not user32.SetClipboardData(13, memoria):",
        "            return False",
        "        memoria = None",
        "        return True",
        "    except Exception:",
        "        return False",
        "    finally:",
        "        try:",
        "            user32.CloseClipboard()",
        "        except Exception:",
        "            pass",
        "        if memoria:",
        "            try:",
        "                kernel32.GlobalFree(memoria)",
        "            except Exception:",
        "                pass",
        "",
        "",
        "def digitar_texto(texto):",
        "    texto = str(texto)",
        "    if texto and definir_texto_clipboard_windows(texto):",
        '        pyautogui.hotkey("ctrl", "v")',
        "        time.sleep(0.008)",
        "        return",
        "    for caractere in texto:",
        '        if caractere == "\\r":',
        "            continue",
        '        if caractere == "\\n":',
        '            pyautogui.press("enter")',
        "            continue",
        '        if caractere == "\\t":',
        '            pyautogui.press("tab")',
        "            continue",
        "        pyautogui.write(caractere, interval=INTERVALO_TEXTO)",
        "",
        "",
        "def eh_texto_simples(nome):",
        '    return len(str(nome)) == 1 and str(nome) not in {"\\r", "\\n", "\\t"}',
        "",
        "",
        "def tecla_para_pyautogui(nome):",
        "    nome = str(nome)",
        "    if len(nome) == 1:",
        "        return nome",
        "    mapa = {",
        '        "ctrl": "ctrl",',
        '        "control": "ctrl",',
        '        "shift": "shift",',
        '        "alt": "alt",',
        '        "altgr": "altright",',
        '        "enter": "enter",',
        '        "tab": "tab",',
        '        "esc": "esc",',
        '        "escape": "esc",',
        '        "space": "space",',
        '        "backspace": "backspace",',
        '        "delete": "delete",',
        '        "del": "delete",',
        '        "insert": "insert",',
        '        "home": "home",',
        '        "end": "end",',
        '        "pgup": "pageup",',
        '        "pgdn": "pagedown",',
        '        "up": "up",',
        '        "down": "down",',
        '        "left": "left",',
        '        "right": "right",',
        "    }",
        "    return mapa.get(nome, nome)",
        "",
        "",
        "def executar_acao(acao):",
        "    if not acao:",
        "        return",
        "    if isinstance(acao, str):",
        '        acao = [parte.strip() for parte in acao.split("+") if parte.strip()]',
        "    if len(acao) == 1:",
        "        pyautogui.press(str(acao[0]))",
        "        return",
        "    pyautogui.hotkey(*[str(parte) for parte in acao])",
        "",
        "",
        "def executar_teclas(teclas, repeticoes=1):",
        "    teclas = [str(tecla) for tecla in (teclas or []) if str(tecla)]",
        "    if not teclas:",
        "        return",
        "    total = max(1, int(repeticoes))",
        "    teclas_pyautogui = [tecla_para_pyautogui(tecla) for tecla in teclas]",
        "    for _ in range(total):",
        "        if len(teclas_pyautogui) > 1:",
        "            pyautogui.hotkey(*teclas_pyautogui)",
        "            continue",
        "        tecla = teclas[0]",
        "        tecla_pyautogui = teclas_pyautogui[0]",
        "        if eh_texto_simples(tecla):",
        "            pyautogui.write(tecla)",
        "        else:",
        "            pyautogui.press(tecla_pyautogui)",
        "",
        "",
        "def executar_sequencia(valores, acao_apos_item=None, intervalo_entre_itens=0.5):",
        "    total = len(valores)",
        "    for indice, valor in enumerate(valores):",
        "        digitar_texto(str(valor))",
        "        if indice < total - 1 and acao_apos_item:",
        "            executar_acao(acao_apos_item)",
        "        if indice < total - 1 and intervalo_entre_itens > 0:",
        "            time.sleep(intervalo_entre_itens)",
        "",
        "",
        "def esperar_imagem(imagem, timeout=10, confianca=0.8):",
        "    confianca = float(confianca)",
        "    confiancas_busca = [confianca]",
        "    if confianca > 0.72:",
        "        confiancas_busca.append(round(max(0.55, confianca - 0.12), 2))",
        "    fim = time.perf_counter() + max(0, float(timeout))",
        "    while time.perf_counter() <= fim:",
        "        for confianca_atual in confiancas_busca:",
        "            for grayscale in (False, True):",
        "                local = pyautogui.locateOnScreen(imagem, confidence=confianca_atual, grayscale=grayscale)",
        "                if local:",
        "                    return local",
        "        time.sleep(0.25)",
        "    return None",
        "",
        "",
        "def clicar_por_imagem(",
        "    imagem,",
        "    timeout=10,",
        "    confianca=0.8,",
        "    clicar=True,",
        "    offset_x=0,",
        "    offset_y=0,",
        '    ao_nao_encontrar="parar",',
        '    sequencia_nao_encontrar="",',
        '    bloco_nao_encontrar="",',
        "    quantidade_tentativas=1,",
        "    intervalo_tentativas=1,",
        "):",
        "    tentativas = max(1, int(quantidade_tentativas))",
        "    for tentativa in range(tentativas):",
        "        local = esperar_imagem(imagem, timeout=timeout, confianca=confianca)",
        "        if local:",
        "            centro_x, centro_y = pyautogui.center(local)",
        "            destino_x = int(centro_x + int(offset_x))",
        "            destino_y = int(centro_y + int(offset_y))",
        "            if clicar:",
        "                pyautogui.click(destino_x, destino_y)",
        "            return True",
        '        if ao_nao_encontrar != "tentar_novamente" or tentativa >= tentativas - 1:',
        "            break",
        "        time.sleep(max(0, float(intervalo_tentativas)))",
        '    if ao_nao_encontrar == "parar":',
        '        raise RuntimeError(f"Imagem nao encontrada: {imagem}")',
        "    return False",
        "",
        "",
        "def verificar_estado_botao(",
        "    imagem_ativo,",
        "    imagem_desativado,",
        "    timeout=5,",
        "    confianca=0.85,",
        "    se_ativo=None,",
        "    se_desativado=None,",
        "    se_nao_encontrar=None,",
        "):",
        "    if esperar_imagem(imagem_ativo, timeout=timeout, confianca=confianca):",
        '        return se_ativo or {"acao": "continuar"}',
        "    if esperar_imagem(imagem_desativado, timeout=timeout, confianca=confianca):",
        '        return se_desativado or {"acao": "continuar"}',
        '    return se_nao_encontrar or {"acao": "parar"}',
        "",
        "",
        "def executar():",
    ]

    if not eventos:
        linhas.append("    pass")
    else:
        for evento in eventos:
            intervalo_antes = float(evento.get("intervalo_antes", 0) or 0)
            if intervalo_antes > 0:
                linhas.append(f"    time.sleep({intervalo_antes})")
            linhas.append(f"    {evento['comando_python']}")

    linhas.extend(
        [
            "",
            "",
            'if __name__ == "__main__":',
            "    executar()",
        ]
    )
    return "\n".join(linhas) + "\n"


def carregar_payload():
    if not ARQUIVO_SAIDA.exists():
        return criar_payload_vazio()

    payload = ler_json_arquivo(ARQUIVO_SAIDA)
    normalizar_sequencia_inicial_payload(payload)
    normalizar_configuracoes_sequencia(payload)
    return payload


def salvar_payload(payload):
    payload["criado_em"] = datetime.now().isoformat(timespec="seconds")
    payload["instrucoes"] = list(INSTRUCOES_PADRAO)
    normalizar_sequencia_inicial_payload(payload)
    normalizar_configuracoes_sequencia(payload)
    payload["resumo_digitado"] = montar_resumo(payload["eventos"])
    ARQUIVO_SAIDA.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    ARQUIVO_RESUMO.write_text(payload["resumo_digitado"], encoding="utf-8")
    ARQUIVO_SCRIPT.write_text(
        montar_script_python(payload["eventos"]),
        encoding="utf-8",
    )
    sincronizar_automacao_ativa(payload)


def normalizar_eventos(payload):
    payload.setdefault("configuracoes_sequencia", {})
    eventos = payload.get("eventos", [])
    for ordem, evento in enumerate(eventos, start=1):
        evento["ordem"] = ordem
        evento["intervalo_antes"] = round(float(evento.get("intervalo_antes", 0) or 0), 3)
        evento["sequencia"] = normalizar_nome_sequencia(evento.get("sequencia"))
        tipo = str(evento.get("tipo", "tecla") or "tecla").strip()
        teclas = evento.get("teclas", [])
        detalhes = evento.get("detalhes") or {}
        if tipo not in TIPOS_SUPORTADOS_RUNTIME:
            evento["detalhes"] = detalhes
            evento["texto"] = str(evento.get("texto") or f"[{tipo}]")
            evento["comando_python"] = str(evento.get("comando_python") or f"# bloco_nao_suportado: {tipo}")
            continue
        if tipo == "sequencia_planilha":
            sincronizar_evento_planilha(evento)
            continue
        elif tipo in TIPOS_IMAGEM:
            detalhes["imagem"] = normalizar_id_imagem(detalhes.get("imagem"))
            evento["detalhes"] = detalhes
        elif tipo in TIPOS_VERIFICACAO:
            detalhes["imagem_ativo"] = normalizar_id_imagem(detalhes.get("imagem_ativo"))
            detalhes["imagem_desativado"] = normalizar_id_imagem(detalhes.get("imagem_desativado"))
            evento["detalhes"] = detalhes
        elif tipo in TIPOS_CHROME:
            detalhes["fallback_imagem"] = normalizar_id_imagem(detalhes.get("fallback_imagem"))
            evento["detalhes"] = detalhes
        elif tipo in {"tecla", "atalho"}:
            detalhes["repeticoes"] = max(1, int(detalhes.get("repeticoes", 1) or 1))
            evento["detalhes"] = detalhes
        texto_livre = evento.get("texto") if tipo == "texto_livre" else None
        evento["texto"] = formatar_evento(teclas, tipo, texto_livre=texto_livre, metadados=detalhes)
        evento["comando_python"] = montar_comando_python(
            teclas,
            tipo,
            texto_livre=evento["texto"] if tipo == "texto_livre" else None,
            metadados=detalhes,
        )


def carregar_planilha_meta():
    workbooks = []
    for item in listar_arquivos_planilha():
        workbook_meta = dict(item)
        workbook_meta["sheets"] = []
        caminho = resolver_planilha(item["id"])
        if caminho.exists():
            try:
                workbook = load_workbook(caminho, read_only=True, data_only=True)
                try:
                    for sheet in workbook.worksheets:
                        headers = []
                        for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), []):
                            if value is None:
                                continue
                            header = str(value).strip()
                            if header:
                                headers.append(header)
                        workbook_meta["sheets"].append({"name": sheet.title, "columns": headers})
                finally:
                    workbook.close()
            except PermissionError:
                workbook_meta["erro"] = "Planilha em uso no Excel."
            except Exception as erro:
                workbook_meta["erro"] = f"Nao foi possivel ler esta planilha. Feche o Excel e salve novamente. Detalhe: {erro}"
        workbooks.append(workbook_meta)

    workbook_padrao = next((item for item in workbooks if item["id"] == ARQUIVO_PLANILHA.name), None)
    if workbook_padrao is None:
        workbook_padrao = next((item for item in workbooks if item.get("exists")), workbooks[0] if workbooks else None)
    return {
        "path": str(ARQUIVO_PLANILHA),
        "exists": bool(workbook_padrao["exists"]) if workbook_padrao else ARQUIVO_PLANILHA.exists(),
        "sheets": workbook_padrao.get("sheets", []) if workbook_padrao else [],
        "workbooks": workbooks,
        "default_workbook": workbook_padrao["id"] if workbook_padrao else ARQUIVO_PLANILHA.name,
    }


def carregar_planilha_preview(sheet_name=None, limit=200, workbook_file=None):
    caminho_planilha = resolver_planilha(workbook_file)
    workbook_id = normalizar_id_planilha(workbook_file)
    preview = {
        "path": str(caminho_planilha),
        "workbook_id": workbook_id,
        "workbook_name": Path(workbook_id).name,
        "exists": caminho_planilha.exists(),
        "sheet": sheet_name,
        "headers": [],
        "rows": [],
        "total_rows": 0,
        "shown_rows": 0,
    }
    if not caminho_planilha.exists():
        return preview

    try:
        workbook = load_workbook(caminho_planilha, read_only=True, data_only=True)
    except PermissionError:
        preview["erro"] = "Planilha em uso no Excel. Feche o arquivo e tente atualizar novamente."
        return preview
    except Exception as erro:
        preview["erro"] = f"Nao foi possivel ler a planilha. Feche o Excel, salve novamente e atualize. Detalhe: {erro}"
        return preview
    try:
        worksheets = list(workbook.worksheets)
        worksheet_names = [sheet.title for sheet in worksheets]
        if not worksheet_names:
            return preview

        selected_sheet = sheet_name if sheet_name in worksheet_names else worksheet_names[0]
        if selected_sheet not in worksheet_names:
            raise ValueError(f"Aba nao encontrada: {selected_sheet}")

        sheet = workbook[selected_sheet]
        header_row = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), []))
        headers = []
        for index, value in enumerate(header_row, start=1):
            texto = "" if value is None else str(value).strip()
            headers.append(texto or f"Coluna {index}")

        rows = []
        total_rows = 0
        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            values = []
            row_has_content = False
            for column_index in range(len(headers)):
                value = row[column_index] if column_index < len(row) else None
                texto = "" if value is None else str(value).strip()
                if texto:
                    row_has_content = True
                values.append(texto)

            if not row_has_content:
                continue

            total_rows += 1
            if len(rows) < limit:
                rows.append({"row_number": row_index, "values": values})

        preview.update(
            {
                "sheet": selected_sheet,
                "headers": headers,
                "rows": rows,
                "total_rows": total_rows,
                "shown_rows": len(rows),
            }
        )
        return preview
    finally:
        workbook.close()


def buscar_coluna(sheet, nome_coluna):
    for indice, value in enumerate(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), []), start=1):
        if value is None:
            continue
        if str(value).strip() == nome_coluna:
            return indice
    raise ValueError(f"Coluna nao encontrada: {nome_coluna}")


def valores_da_planilha(sheet_name, column_name, workbook_file=None):
    caminho_planilha = resolver_planilha(workbook_file)
    if not caminho_planilha.exists():
        raise ValueError(f"Arquivo de planilha nao encontrado: {caminho_planilha}")

    workbook = load_workbook(caminho_planilha, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Aba nao encontrada: {sheet_name}")
        sheet = workbook[sheet_name]
        column_index = buscar_coluna(sheet, column_name)
        valores = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            value = row[column_index - 1]
            if value is None:
                continue
            texto = str(value).strip()
            if not texto:
                continue
            valores.append(value)
        return valores
    finally:
        workbook.close()


def sincronizar_evento_planilha(evento):
    detalhes = evento.get("detalhes") or {}
    workbook_id = normalizar_id_planilha(detalhes.get("arquivo"))
    detalhes["arquivo"] = workbook_id
    detalhes["origem"] = detalhes.get("origem") or "coluna_inteira"
    detalhes["linhas_selecionadas"] = list(detalhes.get("linhas_selecionadas") or [])
    detalhes["acao_apos_item_texto"] = (
        str(detalhes.get("acao_apos_item_texto") or "nenhuma").strip().lower() or "nenhuma"
    )
    try:
        detalhes["valores"] = valores_da_planilha(
            detalhes.get("aba"),
            detalhes.get("coluna"),
            workbook_id,
        )
    except Exception:
        detalhes["valores"] = list(detalhes.get("valores") or [])
    evento["detalhes"] = detalhes
    evento["texto"] = formatar_evento(
        evento.get("teclas", []),
        evento.get("tipo", "sequencia_planilha"),
        metadados=detalhes,
    )
    evento["comando_python"] = montar_comando_python(
        evento.get("teclas", []),
        evento.get("tipo", "sequencia_planilha"),
        metadados=detalhes,
    )
    return evento


def normalizar_atraso_transicao(valor):
    return round(max(0.0, float(valor or 0)), 3)


def aplicar_fluxo_sucesso(
    detalhes=None,
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    base = dict(detalhes or {})
    base["ao_sucesso"] = str(ao_sucesso or "continuar")
    base["bloco_sucesso"] = int(bloco_sucesso) if str(bloco_sucesso or "").strip() else ""
    base["sequencia_sucesso"] = normalizar_nome_sequencia(sequencia_sucesso) if str(sequencia_sucesso or "").strip() else ""
    base["atraso_antes_sequencia_sucesso"] = normalizar_atraso_transicao(atraso_sequencia_sucesso)
    return base


def criar_bloco_texto(
    texto,
    delay,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
    editavel_execucao=False,
    rotulo_execucao="",
):
    detalhes = aplicar_fluxo_sucesso({}, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    detalhes["editavel_execucao"] = bool(editavel_execucao)
    detalhes["rotulo_execucao"] = str(rotulo_execucao or "").strip()
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "texto_livre",
        "teclas": ["texto_livre"],
        "texto": texto,
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["texto_livre"], "texto_livre", texto_livre=texto),
        "detalhes": detalhes,
    }


def parse_keys_string(keys_string):
    partes = [parte.strip().lower() for parte in (keys_string or "").replace(" ", "").split("+") if parte.strip()]
    if not partes:
        raise ValueError("Informe uma tecla ou atalho.")
    return partes


def parse_bool_value(value):
    if isinstance(value, bool):
        return value
    return str(value or "").strip().lower() in {"1", "true", "sim", "yes", "on"}


def criar_bloco_teclas(
    keys_string,
    delay,
    sequencia="principal",
    repeticoes=1,
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    teclas = parse_keys_string(keys_string) if str(keys_string or "").strip() else []
    tipo = "atalho" if len(teclas) > 1 else "tecla"
    detalhes = aplicar_fluxo_sucesso({
        "repeticoes": max(1, int(repeticoes or 1)),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": tipo,
        "teclas": teclas,
        "texto": formatar_evento(teclas, tipo, metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(teclas, tipo, metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_mouse(
    x,
    y,
    button,
    delay,
    clicks=1,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    detalhes = aplicar_fluxo_sucesso({
        "x": int(x),
        "y": int(y),
        "botao": str(button or "left").lower(),
        "cliques": max(1, int(clicks or 1)),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "mouse_click",
        "teclas": ["mouse"],
        "texto": formatar_evento(["mouse"], "mouse_click", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["mouse"], "mouse_click", metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_clicar_imagem(
    imagem,
    delay,
    timeout=10,
    confianca=0.8,
    perfil_busca="balanced",
    modo_area_busca="target_window",
    area_busca_x=0,
    area_busca_y=0,
    area_busca_largura=0,
    area_busca_altura=0,
    clicar=True,
    clicks=1,
    offset_x=0,
    offset_y=0,
    janela_alvo="",
    texto_ocr_esperado="",
    leitura_relativa_ativa=False,
    leitura_relativa_lado="livre",
    leitura_relativa_largura=140,
    leitura_relativa_altura=40,
    leitura_relativa_offset_x=0,
    leitura_relativa_offset_y=0,
    leitura_relativa_salvar_em="",
    ao_nao_encontrar="parar",
    sequencia_nao_encontrar="",
    atraso_sequencia_nao_encontrar=0,
    bloco_nao_encontrar="",
    quantidade_tentativas=1,
    intervalo_tentativas=1,
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
    sequencia="principal",
):
    detalhes = aplicar_fluxo_sucesso({
        "imagem": normalizar_id_imagem(imagem),
        "timeout": max(0.1, float(timeout)),
        "confianca": float(confianca),
        "perfil_busca": str(perfil_busca or "balanced").strip().lower() or "balanced",
        "modo_area_busca": normalizar_modo_area_busca_reconhecimento(modo_area_busca),
        "area_busca_x": int(area_busca_x or 0),
        "area_busca_y": int(area_busca_y or 0),
        "area_busca_largura": max(0, int(area_busca_largura or 0)),
        "area_busca_altura": max(0, int(area_busca_altura or 0)),
        "clicar": bool(clicar),
        "cliques": max(0, int(clicks if bool(clicar) else 0)),
        "offset_x": int(offset_x),
        "offset_y": int(offset_y),
        "janela_alvo": str(janela_alvo or "").strip(),
        "texto_ocr_esperado": str(texto_ocr_esperado or "").strip(),
        "leitura_relativa": {
            "ativo": parse_bool_value(leitura_relativa_ativa),
            "lado": str(leitura_relativa_lado or "livre").strip().lower() or "livre",
            "largura": max(16, int(leitura_relativa_largura or 140)),
            "altura": max(12, int(leitura_relativa_altura or 40)),
            "offset_x": int(leitura_relativa_offset_x or 0),
            "offset_y": int(leitura_relativa_offset_y or 0),
            "salvar_em": str(leitura_relativa_salvar_em or "").strip(),
        },
        "ao_nao_encontrar": str(ao_nao_encontrar or "parar"),
        "sequencia_nao_encontrar": str(sequencia_nao_encontrar or "").strip(),
        "atraso_antes_sequencia_nao_encontrar": normalizar_atraso_transicao(atraso_sequencia_nao_encontrar),
        "bloco_nao_encontrar": int(bloco_nao_encontrar) if str(bloco_nao_encontrar or "").strip() else "",
        "quantidade_tentativas": max(1, int(quantidade_tentativas or 1)),
        "intervalo_tentativas": max(0.0, float(intervalo_tentativas)),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "clicar_imagem",
        "teclas": ["imagem"],
        "texto": formatar_evento(["imagem"], "clicar_imagem", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["imagem"], "clicar_imagem", metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_verificar_estado_botao(
    nome,
    imagem_ativo,
    imagem_desativado,
    delay,
    timeout=5,
    confianca=0.85,
    se_ativo=None,
    se_desativado=None,
    se_nao_encontrar=None,
    sequencia="principal",
):
    detalhes = {
        "nome": str(nome or "Verificar estado do botao").strip() or "Verificar estado do botao",
        "imagem_ativo": normalizar_id_imagem(imagem_ativo),
        "imagem_desativado": normalizar_id_imagem(imagem_desativado),
        "timeout": max(0.1, float(timeout)),
        "confianca": float(confianca),
        "se_ativo": se_ativo or {"acao": "continuar"},
        "se_desativado": se_desativado or {"acao": "continuar"},
        "se_nao_encontrar": se_nao_encontrar or {"acao": "parar"},
    }
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "verificar_estado_botao",
        "teclas": ["estado_botao"],
        "texto": formatar_evento(["estado_botao"], "verificar_estado_botao", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["estado_botao"], "verificar_estado_botao", metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_chrome_elemento_html(
    nome,
    html_original,
    elemento,
    alvo,
    acao,
    timeout=10000,
    fallback_visual=False,
    fallback_imagem="",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
    ao_falhar="parar",
    bloco_falha="",
    sequencia_falha="",
    atraso_sequencia_falha=0,
    quantidade_tentativas=1,
    intervalo_tentativas=0.5,
    delay=0,
    sequencia="principal",
):
    html_original_limpo = str(html_original or "").strip()
    elemento_informado = elemento if isinstance(elemento, dict) else {}
    alvo_informado = alvo if isinstance(alvo, dict) else {}
    acao_informada = acao if isinstance(acao, dict) else {}

    if html_original_limpo:
        html_analisado = parse_html_elemento(html_original_limpo)
    else:
        # Bloco criado pelo painel rapido: o usuario captura/configura o elemento depois.
        html_analisado = {
            "html_original": "",
            "elemento": {
                "tag": str(elemento_informado.get("tag") or "elemento"),
                "type": str(elemento_informado.get("type") or ""),
                "id": str(elemento_informado.get("id") or ""),
                "name": str(elemento_informado.get("name") or ""),
                "placeholder": str(elemento_informado.get("placeholder") or ""),
                "aria_label": str(elemento_informado.get("aria_label") or ""),
                "role": str(elemento_informado.get("role") or ""),
                "text": str(elemento_informado.get("text") or ""),
                "classes": elemento_informado.get("classes") if isinstance(elemento_informado.get("classes"), list) else [],
                "attributes": elemento_informado.get("attributes") if isinstance(elemento_informado.get("attributes"), dict) else {},
            },
            "tipo_provavel": str(elemento_informado.get("tipo_provavel") or "element"),
            "seletores": [],
            "alvo_recomendado": {
                "metodo": str(alvo_informado.get("metodo") or "css"),
                "valor": str(alvo_informado.get("valor") or ""),
                "codigo_playwright": "",
                "confianca": "baixa",
                "seletor_css": str(alvo_informado.get("seletor_css") or ""),
            },
            "acao_sugerida": {
                "tipo": str(acao_informada.get("tipo") or "clicar"),
                "origem_valor": str(acao_informada.get("origem_valor") or "texto_fixo"),
                "valor": str(acao_informada.get("valor") or ""),
                "limpar_antes": bool(acao_informada.get("limpar_antes", True)),
                "pressionar_depois": str(acao_informada.get("pressionar_depois") or ""),
                "salvar_em": str(acao_informada.get("salvar_em") or ""),
            },
            "nome_sugerido": "Chrome/Web: capturar elemento",
            "confianca": "baixa",
        }

    elemento_base = dict(html_analisado.get("elemento") or {})
    elemento_base.update({chave: valor for chave, valor in elemento_informado.items() if valor not in (None, "")})

    chrome_sem_alvo = not html_original_limpo and not str(alvo_informado.get("valor") or "").strip()
    seletores = [] if chrome_sem_alvo else gerar_seletores(elemento_base)
    seletores_por_metodo = {str(item.get("metodo")): item for item in seletores}
    metodo_alvo = str(alvo_informado.get("metodo") or (html_analisado.get("alvo_recomendado") or {}).get("metodo") or "css")
    alvo_recomendado = dict(seletores_por_metodo.get(metodo_alvo) or html_analisado.get("alvo_recomendado") or {})
    alvo_final = {**alvo_recomendado, **alvo_informado}
    alvo_final["metodo"] = metodo_alvo
    if not str(alvo_final.get("valor", "")).strip():
        alvo_final["valor"] = str(alvo_recomendado.get("valor", "")).strip()
    if not alvo_final.get("confianca"):
        alvo_final["confianca"] = alvo_recomendado.get("confianca", "media")
    if not alvo_final.get("seletor_css"):
        alvo_final["seletor_css"] = alvo_recomendado.get("seletor_css", "")

    acao_sugerida = dict(html_analisado.get("acao_sugerida") or {})
    acao_final = {**acao_sugerida, **acao_informada}
    acao_final["tipo"] = str(acao_final.get("tipo") or acao_sugerida.get("tipo") or "clicar")
    acao_final["origem_valor"] = str(acao_final.get("origem_valor") or "texto_fixo")
    acao_final["valor"] = str(acao_final.get("valor", "") or "")
    acao_final["limpar_antes"] = bool(acao_final.get("limpar_antes", True))
    acao_final["pressionar_depois"] = str(acao_final.get("pressionar_depois", "") or "")
    acao_final["salvar_em"] = str(acao_final.get("salvar_em", "") or "")

    detalhes = {
        "nome": str(nome or html_analisado.get("nome_sugerido") or "Chrome: Elemento HTML").strip() or "Chrome: Elemento HTML",
        "html_original": str(html_original_limpo or html_analisado.get("html_original") or "").strip(),
        "elemento": elemento_base,
        "tipo_provavel": str(html_analisado.get("tipo_provavel") or classificar_tipo_elemento_html(elemento_base)),
        "seletores": seletores,
        "alvo": alvo_final,
        "acao": acao_final,
        "timeout": max(500, int(timeout or 10000)),
        "fallback_visual": bool(fallback_visual),
        "fallback_imagem": normalizar_id_imagem(fallback_imagem),
        "ao_sucesso": str(ao_sucesso or "continuar"),
        "bloco_sucesso": int(bloco_sucesso) if str(bloco_sucesso or "").strip() else "",
        "sequencia_sucesso": normalizar_nome_sequencia(sequencia_sucesso) if str(sequencia_sucesso or "").strip() else "",
        "atraso_antes_sequencia_sucesso": normalizar_atraso_transicao(atraso_sequencia_sucesso),
        "ao_falhar": str(ao_falhar or "parar"),
        "bloco_falha": int(bloco_falha) if str(bloco_falha or "").strip() else "",
        "sequencia_falha": normalizar_nome_sequencia(sequencia_falha) if str(sequencia_falha or "").strip() else "",
        "atraso_antes_sequencia_falha": normalizar_atraso_transicao(atraso_sequencia_falha),
        "quantidade_tentativas": max(1, int(quantidade_tentativas or 1)),
        "intervalo_tentativas": max(0.0, float(intervalo_tentativas or 0)),
    }
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "chrome_elemento_html",
        "teclas": ["chrome"],
        "texto": formatar_evento(["chrome"], "chrome_elemento_html", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["chrome"], "chrome_elemento_html", metadados=detalhes),
        "detalhes": detalhes,
    }


def normalizar_navegador_url(navegador):
    valor = str(navegador or "padrao").strip().lower()
    if valor in {"chrome", "edge"}:
        return valor
    return "padrao"


def normalizar_modo_abertura_url(modo):
    valor = str(modo or "nova_janela").strip().lower()
    if valor in {"nova_aba", "aba"}:
        return "nova_aba"
    return "nova_janela"


def criar_bloco_abrir_url(
    url,
    navegador="padrao",
    modo_abertura="nova_janela",
    tela_cheia=False,
    aguardar_segundos=3,
    delay=0,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    url_normalizada = str(url or "").strip()
    detalhes = aplicar_fluxo_sucesso({
        "url": url_normalizada,
        "navegador": normalizar_navegador_url(navegador),
        "modo_abertura": normalizar_modo_abertura_url(modo_abertura),
        "tela_cheia": bool(tela_cheia),
        "aguardar_segundos": max(0.0, float(aguardar_segundos or 0)),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "abrir_url",
        "teclas": ["abrir_url"],
        "texto": formatar_evento(["abrir_url"], "abrir_url", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["abrir_url"], "abrir_url", metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_abrir_programa(
    caminho,
    parametros="",
    aguardar_segundos=2,
    delay=0,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    caminho_normalizado = str(caminho or "").strip()
    detalhes = aplicar_fluxo_sucesso({
        "caminho": caminho_normalizado,
        "parametros": str(parametros or "").strip(),
        "aguardar_segundos": max(0.0, float(aguardar_segundos or 0)),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "abrir_programa",
        "teclas": ["abrir_programa"],
        "texto": formatar_evento(["abrir_programa"], "abrir_programa", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["abrir_programa"], "abrir_programa", metadados=detalhes),
        "detalhes": detalhes,
    }


def normalizar_extensao_download(extensao):
    valor = str(extensao or ".pdf").strip().lower()
    if not valor or valor == "*":
        return "*"
    return valor if valor.startswith(".") else f".{valor}"


def criar_bloco_salvar_download(
    pasta_origem="",
    extensao=".pdf",
    pasta_destino="",
    nome_arquivo="{empresa}_{competencia}_FGTS.pdf",
    timeout=60,
    delay=0,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    pasta_destino_normalizada = str(pasta_destino or "").strip()
    detalhes = aplicar_fluxo_sucesso({
        "pasta_origem": str(pasta_origem or "").strip(),
        "extensao": normalizar_extensao_download(extensao),
        "pasta_destino": pasta_destino_normalizada,
        "nome_arquivo": str(nome_arquivo or "").strip(),
        "timeout": max(1.0, float(timeout or 60)),
        "ultimo_arquivo_variavel": "ultimo_download",
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "salvar_download",
        "teclas": ["salvar_download"],
        "texto": formatar_evento(["salvar_download"], "salvar_download", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(["salvar_download"], "salvar_download", metadados=detalhes),
        "detalhes": detalhes,
    }


def criar_bloco_planilha(
    sheet,
    column,
    delay,
    after_action,
    interval_between,
    workbook_file=None,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    workbook_id = normalizar_id_planilha(workbook_file)
    valores = valores_da_planilha(sheet, column, workbook_id)
    acao_apos_item = normalizar_acao_apos_item(after_action)
    modo_loop = "somente_contexto" if str(after_action or "").strip().lower() in {"somente_contexto", "contexto", "so_contexto"} else "digitar_coluna"
    detalhes = aplicar_fluxo_sucesso({
        "arquivo": workbook_id,
        "aba": sheet,
        "coluna": column,
        "valores": valores,
        "origem": "coluna_inteira",
        "modo_loop": modo_loop,
        "linhas_selecionadas": [],
        "acao_apos_item": acao_apos_item,
        "acao_apos_item_texto": (after_action or "nenhuma").strip().lower() or "nenhuma",
        "intervalo_entre_itens": float(interval_between),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "sequencia_planilha",
        "teclas": [f"planilha:{column}"],
        "texto": formatar_evento([f"planilha:{column}"], "sequencia_planilha", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(
            [f"planilha:{column}"],
            "sequencia_planilha",
            metadados=detalhes,
        ),
        "detalhes": detalhes,
    }


def criar_bloco_escrever_planilha(
    sheet,
    column,
    value,
    delay,
    workbook_file=None,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    workbook_id = normalizar_id_planilha(workbook_file)
    detalhes = aplicar_fluxo_sucesso({
        "arquivo": workbook_id,
        "aba": str(sheet or "").strip(),
        "coluna": str(column or "").strip(),
        "valor": str(value or ""),
        "modo_linha": "loop_atual",
        "criar_coluna": True,
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "escrever_planilha",
        "teclas": [f"escrever_planilha:{column}"],
        "texto": formatar_evento([f"escrever_planilha:{column}"], "escrever_planilha", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(
            [f"escrever_planilha:{column}"],
            "escrever_planilha",
            metadados=detalhes,
        ),
        "detalhes": detalhes,
    }


def criar_bloco_comparar_planilha_tela(
    sheet,
    expected_column,
    delay,
    workbook_file=None,
    sequencia="principal",
    source="ultimo_ocr",
    comparison_mode="texto",
    status_column="",
    note_column="",
    ok_value="OK",
    fail_value="Divergente",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
    ao_divergir="continuar",
    bloco_divergir="",
    sequencia_divergir="",
    atraso_sequencia_divergir=0,
):
    workbook_id = normalizar_id_planilha(workbook_file)
    detalhes = aplicar_fluxo_sucesso({
        "arquivo": workbook_id,
        "aba": str(sheet or "").strip(),
        "coluna_esperada": str(expected_column or "").strip(),
        "origem_tela": normalizar_origem_ocr_comparacao(source),
        "modo_comparacao": str(comparison_mode or "texto").strip().lower() or "texto",
        "coluna_status": str(status_column or "").strip(),
        "coluna_observacao": str(note_column or "").strip(),
        "valor_ok": str(ok_value or "OK"),
        "valor_divergente": str(fail_value or "Divergente"),
        "ao_divergir": str(ao_divergir or "continuar").strip().lower() or "continuar",
        "bloco_divergir": bloco_divergir or "",
        "sequencia_divergir": normalizar_nome_sequencia(sequencia_divergir) if str(sequencia_divergir or "").strip() else "",
        "atraso_antes_sequencia_divergir": round(float(atraso_sequencia_divergir or 0), 3),
    }, ao_sucesso, bloco_sucesso, sequencia_sucesso, atraso_sequencia_sucesso)
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": "comparar_planilha_tela",
        "teclas": [f"comparar_planilha:{expected_column}"],
        "texto": formatar_evento([f"comparar_planilha:{expected_column}"], "comparar_planilha_tela", metadados=detalhes),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python(
            [f"comparar_planilha:{expected_column}"],
            "comparar_planilha_tela",
            metadados=detalhes,
        ),
        "detalhes": detalhes,
    }


def criar_bloco_marcador_loop(
    tipo,
    delay=0,
    sequencia="principal",
    ao_sucesso="continuar",
    bloco_sucesso="",
    sequencia_sucesso="",
    atraso_sequencia_sucesso=0,
):
    tipo_normalizado = "inicio_loop" if str(tipo or "").strip().lower() == "inicio_loop" else "fim_loop"
    detalhes = {}
    if tipo_normalizado == "fim_loop":
        detalhes = aplicar_fluxo_sucesso(
            {},
            ao_sucesso,
            bloco_sucesso,
            sequencia_sucesso,
            atraso_sequencia_sucesso,
        )
    return {
        "ordem": 0,
        "sequencia": normalizar_nome_sequencia(sequencia),
        "tipo": tipo_normalizado,
        "teclas": [tipo_normalizado],
        "texto": formatar_evento([tipo_normalizado], tipo_normalizado),
        "intervalo_antes": round(float(delay), 3),
        "comando_python": montar_comando_python([tipo_normalizado], tipo_normalizado),
        "detalhes": detalhes,
    }


def resumir_texto(texto, limite=84):
    texto = str(texto or "").replace("\n", " ").strip()
    if len(texto) <= limite:
        return texto or "-"
    return f"{texto[: limite - 3]}..."


def card_do_evento(evento):
    tipo = evento.get("tipo", "tecla")
    detalhes = evento.get("detalhes") or {}
    titulo = "Bloco"
    subtitulo = resumir_texto(evento.get("texto", ""))
    descricao = evento.get("comando_python", "")

    if tipo == "texto_livre":
        titulo = "Texto livre"
        subtitulo = resumir_texto(evento.get("texto", ""))
        descricao = f"{len(evento.get('texto', ''))} caracteres"
    elif tipo == "atalho":
        titulo = "Atalho"
        subtitulo = " + ".join(tecla_para_exibicao(tecla) for tecla in evento.get("teclas", [])) or "Configurar atalho"
        repeticoes = max(1, int(detalhes.get("repeticoes", 1) or 1))
        descricao = "Combinacao de teclas" if repeticoes == 1 else f"Repetir {repeticoes} vezes"
    elif tipo == "tecla":
        titulo = "Tecla"
        subtitulo = " + ".join(tecla_para_exibicao(tecla) for tecla in evento.get("teclas", [])) or "Configurar tecla"
        repeticoes = max(1, int(detalhes.get("repeticoes", 1) or 1))
        descricao = "Acao unica" if repeticoes == 1 else f"Repetir {repeticoes} vezes"
    elif tipo == "mouse_click":
        titulo = "Clique do mouse"
        cliques = max(1, int(detalhes.get("cliques", 1) or 1))
        tipo_clique = "Clique duplo" if cliques == 2 else "Clique simples"
        subtitulo = (
            f"{str(detalhes.get('botao', 'left')).upper()} em X {detalhes.get('x', 0)} | "
            f"Y {detalhes.get('y', 0)}"
        )
        descricao = f"{tipo_clique} absoluto na tela"
    elif tipo == "clicar_imagem":
        titulo = "Reconhecimento visual"
        subtitulo = Path(str(detalhes.get("imagem", ""))).name or "Referencia visual"
        descricao = (
            f"Timeout {float(detalhes.get('timeout', 10) or 10):g}s | "
            f"Confianca {float(detalhes.get('confianca', 0.8) or 0.8):g} | "
            f"Falha: {detalhes.get('ao_nao_encontrar', 'parar')}"
        )
    elif tipo == "chrome_elemento_html":
        titulo = detalhes.get("nome", "Chrome / Elemento HTML")
        elemento = detalhes.get("elemento") or {}
        alvo = detalhes.get("alvo") or {}
        acao = detalhes.get("acao") or {}
        subtitulo = (
            f"{str(elemento.get('tag', 'elemento')).lower()} | "
            f"{alvo.get('metodo', 'css')}: {str(alvo.get('valor', '') or '-')[:56]}"
        )
        descricao = (
            f"Acao: {acao.get('tipo', 'clicar')} | "
            f"Timeout {int(detalhes.get('timeout', 10000) or 10000)}ms"
        )
    elif tipo == "abrir_url":
        titulo = "Abrir URL"
        subtitulo = resumir_texto(detalhes.get("url", ""), 70)
        navegador = str(detalhes.get("navegador", "padrao") or "padrao").title()
        modo = "Nova aba" if detalhes.get("modo_abertura") == "nova_aba" else "Nova janela"
        descricao = f"{navegador} | {modo} | Aguarda {float(detalhes.get('aguardar_segundos', 3) or 0):g}s"
    elif tipo == "abrir_programa":
        titulo = "Abrir Programa"
        caminho = str(detalhes.get("caminho", "") or "").strip()
        subtitulo = Path(caminho).name or caminho or "Executavel"
        parametros = str(detalhes.get("parametros", "") or "").strip()
        descricao = (
            f"Parametros: {resumir_texto(parametros, 48)} | "
            f"Aguarda {float(detalhes.get('aguardar_segundos', 2) or 0):g}s"
        )
    elif tipo == "salvar_download":
        titulo = "Salvar download"
        destino = str(detalhes.get("pasta_destino", "") or "").strip()
        extensao = str(detalhes.get("extensao", ".pdf") or ".pdf")
        nome = str(detalhes.get("nome_arquivo", "") or "").strip() or "nome original"
        subtitulo = f"{extensao} -> {Path(destino).name or destino or 'Destino'}"
        descricao = f"Nome: {resumir_texto(nome, 44)} | Timeout {float(detalhes.get('timeout', 60) or 60):g}s"
    elif tipo == "verificar_estado_botao":
        titulo = detalhes.get("nome", "Verificar estado do botao")
        subtitulo = (
            f"Ativo: {Path(str(detalhes.get('imagem_ativo', ''))).name or '-'} | "
            f"Desativado: {Path(str(detalhes.get('imagem_desativado', ''))).name or '-'}"
        )
        descricao = (
            f"Ativo: {(detalhes.get('se_ativo') or {}).get('acao', 'continuar')} | "
            f"Desativado: {(detalhes.get('se_desativado') or {}).get('acao', 'continuar')} | "
            f"Nao encontrar: {(detalhes.get('se_nao_encontrar') or {}).get('acao', 'parar')}"
        )
    elif tipo == "sequencia_planilha":
        somente_contexto = detalhes.get("modo_loop") == "somente_contexto"
        titulo = "Loop de dados" if somente_contexto else "Loop da planilha"
        subtitulo = (
            f"{Path(detalhes.get('arquivo', ARQUIVO_PLANILHA.name)).name} • "
            f"{detalhes.get('aba', 'Sheet1')} -> {detalhes.get('coluna', 'Coluna')}"
        )
        descricao = f"{len(detalhes.get('valores', []))} linhas | apenas carrega a linha atual" if somente_contexto else (
            f"{len(detalhes.get('valores', []))} itens | apos item: "
            f"{detalhes.get('acao_apos_item_texto', 'nenhuma')}"
        )
    elif tipo == "escrever_planilha":
        titulo = "Escrever na planilha"
        subtitulo = (
            f"{Path(detalhes.get('arquivo', ARQUIVO_PLANILHA.name)).name} - "
            f"{detalhes.get('aba', 'Sheet1')} -> {detalhes.get('coluna', 'Coluna')}"
        )
        descricao = f"Linha atual do loop recebe: {detalhes.get('valor', '')}"
    elif tipo == "comparar_planilha_tela":
        titulo = "Conferir valor"
        subtitulo = (
            f"{Path(detalhes.get('arquivo', ARQUIVO_PLANILHA.name)).name} - "
            f"{detalhes.get('aba', 'Sheet1')} -> {detalhes.get('coluna_esperada', 'Coluna')}"
        )
        descricao = (
            f"Compara com OCR da tela | Resultado em "
            f"{detalhes.get('coluna_status', 'Status')}"
        )
    elif tipo == "inicio_loop":
        titulo = "Inicio do loop"
        subtitulo = "Daqui em diante o trecho pode repetir por linha da planilha."
        descricao = "Use junto com um bloco Fim do loop para limitar a repeticao."
    elif tipo == "fim_loop":
        titulo = "Fim do loop"
        subtitulo = "O trecho repetido termina imediatamente antes deste marcador."
        descricao = "Tudo depois daqui volta a executar apenas uma vez."
    elif tipo not in TIPOS_SUPORTADOS_RUNTIME:
        titulo = str(tipo or "Bloco").replace("_", " ").strip().title() or "Bloco"
        subtitulo = resumir_texto(evento.get("texto", "")) or "Bloco importado de uma versao mais recente do Studio."
        descricao = "Este bloco veio de uma versao mais nova e sera executado quando o runtime do Agent estiver atualizado."
    return {
        "order": evento.get("ordem", 0),
        "type": tipo,
        "title": titulo,
        "subtitle": subtitulo,
        "description": descricao,
        "sequence": evento.get("sequencia", "principal"),
        "delay": round(float(evento.get("intervalo_antes", 0) or 0), 3),
        "text": evento.get("texto", ""),
        "keys": evento.get("teclas", []),
        "command": evento.get("comando_python", ""),
        "details": detalhes,
    }


def payload_para_resposta(payload):
    normalizar_eventos(payload)
    normalizar_sequencia_inicial_payload(payload)
    configuracoes_sequencia = normalizar_configuracoes_sequencia(payload)
    eventos = payload.get("eventos", [])
    return {
        "created_at": payload.get("criado_em"),
        "initial_sequence": payload.get("sequencia_inicial", "principal"),
        "summary": {
            "total": len(eventos),
            "text_blocks": sum(1 for evento in eventos if evento.get("tipo") == "texto_livre"),
            "key_blocks": sum(1 for evento in eventos if evento.get("tipo") in {"tecla", "atalho"}),
            "mouse_blocks": sum(1 for evento in eventos if evento.get("tipo") == "mouse_click"),
            "image_blocks": sum(1 for evento in eventos if evento.get("tipo") == "clicar_imagem"),
            "chrome_blocks": sum(1 for evento in eventos if evento.get("tipo") == "chrome_elemento_html"),
            "branch_blocks": sum(1 for evento in eventos if evento.get("tipo") in TIPOS_VERIFICACAO),
            "spreadsheet_blocks": sum(1 for evento in eventos if evento.get("tipo") == "sequencia_planilha"),
            "unknown_blocks": sum(1 for evento in eventos if str(evento.get("tipo") or "").strip() not in TIPOS_SUPORTADOS_RUNTIME),
            "sequence_count": len({normalizar_nome_sequencia(evento.get("sequencia")) for evento in eventos}) or 1,
        },
        "events": [card_do_evento(evento) for evento in eventos],
        "sequences": sorted({normalizar_nome_sequencia(evento.get("sequencia")) for evento in eventos}) or ["principal"],
        "sequence_settings": configuracoes_sequencia,
        "workbook": carregar_planilha_meta(),
        "runtime_compatibility": gerar_resumo_compatibilidade_payload(payload),
    }


def carregar_e_normalizar():
    payload = carregar_payload()
    normalizar_eventos(payload)
    normalizar_sequencia_inicial_payload(payload)
    normalizar_configuracoes_sequencia(payload)
    return payload


def localizar_evento(payload, order):
    for indice, evento in enumerate(payload.get("eventos", [])):
        if int(evento.get("ordem", 0)) == int(order):
            return indice, evento
    raise ValueError(f"Bloco nao encontrado: {order}")


def mapa_ordens_eventos_existentes(eventos):
    mapa = {}
    for indice, evento in enumerate(eventos or [], start=1):
        try:
            ordem = int(evento.get("ordem", indice) or indice)
        except (TypeError, ValueError):
            ordem = indice
        mapa[id(evento)] = ordem
    return mapa


def remapear_referencias_de_blocos(payload, mapa_ordem_anterior):
    """Mantem atalhos de fluxo apontando para o mesmo bloco apos inserir/mover/remover."""
    eventos = payload.get("eventos", [])
    if not eventos or not mapa_ordem_anterior:
        return

    ordem_nova_por_antiga = {}
    for indice, evento in enumerate(eventos, start=1):
        ordem_antiga = mapa_ordem_anterior.get(id(evento))
        if ordem_antiga:
            ordem_nova_por_antiga[int(ordem_antiga)] = indice

    campos_destino = ("bloco_sucesso", "bloco_falha", "bloco_nao_encontrar")
    for evento in eventos:
        detalhes = evento.get("detalhes")
        if not isinstance(detalhes, dict):
            continue
        for campo in campos_destino:
            valor = detalhes.get(campo, "")
            if str(valor or "").strip() == "":
                continue
            try:
                ordem_destino = int(valor)
            except (TypeError, ValueError):
                continue
            nova_ordem = ordem_nova_por_antiga.get(ordem_destino)
            detalhes[campo] = int(nova_ordem) if nova_ordem else ""


def salvar_e_retornar(payload):
    normalizar_eventos(payload)
    normalizar_sequencia_inicial_payload(payload)
    salvar_payload(payload)
    return payload_para_resposta(payload)


def inserir_evento(payload, novo_evento, target_order, position):
    eventos = payload.get("eventos", [])
    mapa_ordem_anterior = mapa_ordens_eventos_existentes(eventos)
    indice, _ = localizar_evento(payload, target_order)
    indice_insercao = indice if position == "before" else indice + 1
    eventos.insert(indice_insercao, novo_evento)
    payload["eventos"] = eventos
    remapear_referencias_de_blocos(payload, mapa_ordem_anterior)
    return payload


def comando_listar(_args):
    return payload_para_resposta(carregar_e_normalizar())


def comando_listar_automacoes(_args):
    return listar_automacoes_resumidas()


def comando_criar_automacao(args):
    biblioteca = carregar_biblioteca()
    automacao_id = gerar_id_automacao()
    item = normalizar_item_biblioteca(
        {
            "id": automacao_id,
            "nome": args.nome,
            "descricao": args.descricao,
            "categoria": args.categoria,
            "criado_em": agora_iso(),
            "atualizado_em": agora_iso(),
        }
    )
    biblioteca["items"].append(item)
    salvar_biblioteca(biblioteca)
    salvar_payload_em_automacao(automacao_id, criar_payload_vazio())
    payload = carregar_automacao_na_raiz(automacao_id)
    resposta = payload_para_resposta(payload)
    resposta["library"] = listar_automacoes_resumidas()
    resposta["active_automation"] = item
    return resposta


def comando_abrir_automacao(args):
    biblioteca = carregar_biblioteca()
    _, item = localizar_item_biblioteca(biblioteca, args.id)
    payload = carregar_automacao_na_raiz(args.id)
    resposta = payload_para_resposta(payload)
    resposta["library"] = listar_automacoes_resumidas()
    resposta["active_automation"] = item
    return resposta


def comando_duplicar_automacao(args):
    biblioteca = carregar_biblioteca()
    _, original = localizar_item_biblioteca(biblioteca, args.id)
    novo_id = gerar_id_automacao()
    novo_item = normalizar_item_biblioteca(
        {
            **original,
            "id": novo_id,
            "nome": f"{original['nome']} - Copia",
            "criado_em": agora_iso(),
            "atualizado_em": agora_iso(),
            "vezes_executada": 0,
            "ultima_execucao": None,
        }
    )
    origem = diretorio_automacao(args.id)
    destino = diretorio_automacao(novo_id)
    if destino.exists():
        shutil.rmtree(destino, ignore_errors=True)
    if origem.exists():
        shutil.copytree(origem, destino, dirs_exist_ok=True)
    else:
        salvar_payload_em_automacao(novo_id, criar_payload_vazio())
    biblioteca["items"].append(novo_item)
    salvar_biblioteca(biblioteca)
    return {"library": listar_automacoes_resumidas(), "duplicated_id": novo_id}


def comando_renomear_automacao(args):
    biblioteca = carregar_biblioteca()
    indice, item = localizar_item_biblioteca(biblioteca, args.id)
    item["nome"] = str(args.nome or item["nome"]).strip() or item["nome"]
    if args.descricao is not None:
        item["descricao"] = str(args.descricao or "").strip()
    if args.categoria is not None:
        item["categoria"] = str(args.categoria or "Outros").strip() or "Outros"
    item["atualizado_em"] = agora_iso()
    biblioteca["items"][indice] = item
    salvar_biblioteca(biblioteca)
    return {"library": listar_automacoes_resumidas(), "renamed_id": args.id}


def comando_excluir_automacao(args):
    biblioteca = carregar_biblioteca()
    indice, item = localizar_item_biblioteca(biblioteca, args.id)
    pasta = diretorio_automacao(args.id)
    if pasta.exists():
        shutil.rmtree(pasta, ignore_errors=True)
    biblioteca["items"].pop(indice)
    salvar_biblioteca(biblioteca)
    if obter_automacao_ativa_id() == str(args.id):
        definir_automacao_ativa("")
    return {"library": listar_automacoes_resumidas(), "deleted_id": item["id"]}


def comando_marcar_execucao_automacao(args):
    biblioteca = carregar_biblioteca()
    indice, item = localizar_item_biblioteca(biblioteca, args.id)
    item["vezes_executada"] = max(0, int(item.get("vezes_executada") or 0)) + 1
    item["ultima_execucao"] = agora_iso()
    item["atualizado_em"] = agora_iso()
    biblioteca["items"][indice] = item
    salvar_biblioteca(biblioteca)
    return {"library": listar_automacoes_resumidas(), "executed_id": args.id}


def comando_exportar_automacao_zip(args):
    biblioteca = carregar_biblioteca()
    _, item = localizar_item_biblioteca(biblioteca, args.id)
    destino_zip = exportar_automacao_para_zip(args.id, item, args.destino)
    return {
        "zip_path": str(destino_zip),
        "automation": item,
        "exported_id": args.id,
    }


def exportar_runtime_para_zip(destino_zip=None):
    nome_base = "FlowDetect_Runtime_Update"
    if destino_zip:
        destino_zip = Path(destino_zip)
        if destino_zip.suffix.lower() != ".zip":
            destino_zip = destino_zip.with_suffix(".zip")
        destino_zip.parent.mkdir(parents=True, exist_ok=True)
    else:
        destino_zip = caminho_arquivo_unico(DIRETORIO_EXPORTACOES_AUTOMACAO, f"{nome_base}.zip")

    with tempfile.TemporaryDirectory(prefix="flowdetect_runtime_export_") as temp_dir:
        raiz_bundle = Path(temp_dir) / "bundle"
        raiz_bundle.mkdir(parents=True, exist_ok=True)
        runtime_exportado = exportar_runtime_para_pasta(raiz_bundle)
        (raiz_bundle / "COMO_ATUALIZAR_AGENT.txt").write_text(
            "No Agent, importe este ZIP pelo sistema de atualizacao para sincronizar os blocos e scripts mais recentes do Studio.",
            encoding="utf-8",
        )
        with zipfile.ZipFile(destino_zip, "w", compression=zipfile.ZIP_DEFLATED) as arquivo_zip:
            for caminho in sorted(raiz_bundle.rglob("*")):
                if caminho.is_file():
                    arquivo_zip.write(caminho, caminho.relative_to(raiz_bundle))

    return {
        "zip_path": str(destino_zip),
        "runtime": runtime_exportado["manifest"],
        "files": runtime_exportado["files"],
    }


def comando_exportar_runtime_zip(args):
    return exportar_runtime_para_zip(args.destino)


def comando_info_runtime(_args):
    return {
        "runtime": gerar_manifesto_runtime(),
    }


def comando_importar_runtime_zip(args):
    arquivo_zip = Path(args.arquivo)
    if not arquivo_zip.exists():
        raise ValueError(f"Arquivo ZIP nao encontrado: {arquivo_zip}")
    if arquivo_zip.suffix.lower() != ".zip":
        raise ValueError("Selecione um arquivo .zip valido para atualizar o runtime.")

    try:
        with tempfile.TemporaryDirectory(prefix="flowdetect_runtime_import_") as temp_dir:
            pasta_extraida = Path(temp_dir) / "conteudo"
            extrair_zip_com_seguranca(arquivo_zip, pasta_extraida)
            manifesto_runtime_path = localizar_manifesto_runtime_bundle(pasta_extraida)
            if not manifesto_runtime_path or not manifesto_runtime_path.exists():
                raise ValueError("Esse ZIP nao contem um pacote de atualizacao de runtime do FlowDetect.")
            manifesto_runtime = ler_json_arquivo(manifesto_runtime_path)
            if manifesto_runtime.get("bundle_type") not in {None, BUNDLE_RUNTIME_TIPO}:
                raise ValueError("Esse ZIP nao parece ser um pacote de runtime compativel do FlowDetect.")
            pasta_runtime_importado = localizar_pasta_runtime_importado(pasta_extraida, manifesto_runtime_path)
            atualizacao = aplicar_pacote_runtime_importado(
                pasta_runtime_importado,
                manifesto_runtime=manifesto_runtime,
            )
    except zipfile.BadZipFile as erro:
        raise ValueError(f"O arquivo selecionado nao e um ZIP valido: {erro}") from erro

    return {
        "source_zip": str(arquivo_zip),
        "runtime_update": atualizacao,
        "runtime": {
            "local": gerar_manifesto_runtime(),
            "bundle": manifesto_runtime,
        },
    }


def comando_importar_automacao_zip(args):
    arquivo_zip = Path(args.arquivo)
    if not arquivo_zip.exists():
        raise ValueError(f"Arquivo ZIP nao encontrado: {arquivo_zip}")
    if arquivo_zip.suffix.lower() != ".zip":
        raise ValueError("Selecione um arquivo .zip valido para importar a automacao.")

    try:
        with tempfile.TemporaryDirectory(prefix="flowdetect_import_") as temp_dir:
            pasta_extraida = Path(temp_dir) / "conteudo"
            extrair_zip_com_seguranca(arquivo_zip, pasta_extraida)

            manifesto_path = localizar_manifesto_bundle(pasta_extraida)
            manifesto = {}
            if manifesto_path and manifesto_path.exists():
                manifesto = ler_json_arquivo(manifesto_path)
                if manifesto.get("bundle_type") not in {None, BUNDLE_AUTOMACAO_TIPO}:
                    raise ValueError("Esse ZIP nao parece ser uma automacao compativel do FlowDetect.")

            manifesto_runtime_path = localizar_manifesto_runtime_bundle(pasta_extraida)
            manifesto_runtime = (
                ler_json_arquivo(manifesto_runtime_path)
                if manifesto_runtime_path and manifesto_runtime_path.exists()
                else {}
            )
            pasta_runtime_importado = localizar_pasta_runtime_importado(pasta_extraida, manifesto_runtime_path)

            pasta_automacao = localizar_pasta_automacao_importada(pasta_extraida, manifesto_path)
            if not pasta_automacao or not (pasta_automacao / ARQUIVO_SAIDA.name).exists():
                raise ValueError("Nao foi encontrado o arquivo principal da automacao dentro do ZIP.")

            payload = ler_json_arquivo(pasta_automacao / ARQUIVO_SAIDA.name)
            runtime_local_antes = gerar_manifesto_runtime()
            compatibilidade_antes = gerar_resumo_compatibilidade_payload(payload, runtime_local_antes)
            comparacao_runtime = comparar_runtime_manifestos(runtime_local_antes, manifesto_runtime or {})
            atualizacao_runtime = {
                "applied": False,
                "reason": "",
                "comparison": comparacao_runtime,
            }

            precisa_atualizar_runtime = bool(
                manifesto_runtime
                and pasta_runtime_importado
                and (
                    not comparacao_runtime.get("same_hash")
                    or compatibilidade_antes.get("unsupported_block_types")
                )
            )

            if precisa_atualizar_runtime:
                atualizacao_runtime = aplicar_pacote_runtime_importado(
                    pasta_runtime_importado,
                    manifesto_runtime=manifesto_runtime,
                )
                atualizacao_runtime["reason"] = (
                    "Runtime do Agent atualizado automaticamente para acompanhar os blocos mais recentes do Studio."
                )

            runtime_local_depois = gerar_manifesto_runtime()
            compatibilidade_depois = gerar_resumo_compatibilidade_payload(payload, runtime_local_depois)
            if compatibilidade_depois.get("unsupported_block_types"):
                raise ValueError(
                    "Esta automacao usa blocos que ainda nao sao suportados por este Agent: "
                    + ", ".join(compatibilidade_depois["unsupported_block_types"])
                    + ". Exporte novamente a automacao com o pacote de runtime do Studio ou atualize o Agent."
                )

            biblioteca = carregar_biblioteca()
            automacao_id = gerar_id_automacao()
            dados_item = manifesto.get("automation") or {}
            nome_importado = str(dados_item.get("nome") or arquivo_zip.stem or "Automacao importada").strip()
            item = normalizar_item_biblioteca(
                {
                    "id": automacao_id,
                    "nome": nome_importado or "Automacao importada",
                    "descricao": str(dados_item.get("descricao") or "Importada de arquivo ZIP.").strip(),
                    "categoria": str(dados_item.get("categoria") or "Importadas").strip() or "Importadas",
                    "criado_em": str(dados_item.get("criado_em") or payload.get("criado_em") or agora_iso()),
                    "atualizado_em": agora_iso(),
                    "sequencia_inicial": str(dados_item.get("sequencia_inicial") or "principal").strip() or "principal",
                    "vezes_executada": 0,
                    "ultima_execucao": None,
                }
            )

            destino = diretorio_automacao(automacao_id)
            if destino.exists():
                shutil.rmtree(destino, ignore_errors=True)
            shutil.copytree(pasta_automacao, destino, dirs_exist_ok=True)
            payload = escrever_payload_em_pasta(destino, payload)

            biblioteca["items"].append(item)
            salvar_biblioteca(biblioteca)

            payload_ativo = carregar_automacao_na_raiz(automacao_id)
    except zipfile.BadZipFile as erro:
        raise ValueError(f"O arquivo selecionado nao e um ZIP valido: {erro}") from erro

    resposta = payload_para_resposta(payload_ativo)
    resposta["library"] = listar_automacoes_resumidas()
    resposta["active_automation"] = item
    resposta["imported_bundle"] = {
        "source_zip": str(arquivo_zip),
        "imported_id": automacao_id,
        "manifest_found": bool(manifesto_path),
        "command_count": contar_blocos_payload(payload),
    }
    resposta["runtime_update"] = atualizacao_runtime
    resposta["runtime"] = {
        "local": runtime_local_depois,
        "bundle": manifesto_runtime,
        "compatibility": compatibilidade_depois,
        "manifest_found": bool(manifesto_runtime_path),
    }
    return resposta


def comando_meta(_args):
    return carregar_planilha_meta()


def comando_preview_planilha(args):
    return carregar_planilha_preview(args.sheet, args.limit, args.workbook)


def comando_adicionar_texto(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_texto(
            args.text,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            parse_bool_value(getattr(args, "editavel_execucao", False)),
            getattr(args, "rotulo_execucao", "") or "",
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_teclas(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(criar_bloco_teclas(args.keys, args.delay, args.sequencia, args.repeat, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso))
    return salvar_e_retornar(payload)


def comando_adicionar_mouse(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(criar_bloco_mouse(args.x, args.y, args.button, args.delay, args.clicks, args.sequencia, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso))
    return salvar_e_retornar(payload)


def comando_adicionar_imagem(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_clicar_imagem(
            args.imagem,
            args.delay,
            args.timeout,
            args.confianca,
            args.perfil_busca,
            args.modo_area_busca,
            args.area_busca_x,
            args.area_busca_y,
            args.area_busca_largura,
            args.area_busca_altura,
            args.clicar,
            args.clicks,
            args.offset_x,
            args.offset_y,
            args.janela_alvo,
            args.texto_ocr_esperado,
            args.leitura_relativa_ativa,
            args.leitura_relativa_lado,
            args.leitura_relativa_largura,
            args.leitura_relativa_altura,
            args.leitura_relativa_offset_x,
            args.leitura_relativa_offset_y,
            args.leitura_relativa_salvar_em,
            args.ao_nao_encontrar,
            args.sequencia_nao_encontrar,
            args.atraso_sequencia_nao_encontrar,
            args.bloco_nao_encontrar,
            args.quantidade_tentativas,
            args.intervalo_tentativas,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.sequencia,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_verificacao(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_verificar_estado_botao(
            args.nome,
            args.imagem_ativo,
            args.imagem_desativado,
            args.delay,
            args.timeout,
            args.confianca,
            {"acao": args.acao_ativo, "sequencia": args.sequencia_ativo},
            {"acao": args.acao_desativado, "sequencia": args.sequencia_desativado},
            {"acao": args.acao_nao_encontrar, "sequencia": args.sequencia_nao_encontrar},
            args.sequencia,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_chrome_html(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_chrome_elemento_html(
            args.nome,
            args.html_original,
            json.loads(args.elemento_json) if args.elemento_json else {},
            json.loads(args.alvo_json) if args.alvo_json else {},
            json.loads(args.acao_json) if args.acao_json else {},
            args.timeout,
            args.fallback_visual,
            args.fallback_imagem,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.ao_falhar,
            args.bloco_falha,
            args.sequencia_falha,
            args.atraso_sequencia_falha,
            args.quantidade_tentativas,
            args.intervalo_tentativas,
            args.delay,
            args.sequencia,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_abrir_url(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_abrir_url(
            args.url,
            args.navegador,
            args.modo_abertura,
            args.tela_cheia,
            args.aguardar_segundos,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_abrir_programa(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_abrir_programa(
            args.caminho,
            args.parametros,
            args.aguardar_segundos,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_salvar_download(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_salvar_download(
            args.pasta_origem,
            args.extensao,
            args.pasta_destino,
            args.nome_arquivo,
            args.timeout,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_planilha(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_planilha(
            args.sheet,
            args.column,
            args.delay,
            args.after_action,
            args.interval_between,
            args.workbook,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_escrita_planilha(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_escrever_planilha(
            args.sheet,
            args.column,
            args.value,
            args.delay,
            args.workbook,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_comparacao_planilha(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_comparar_planilha_tela(
            args.sheet,
            args.expected_column,
            args.delay,
            args.workbook,
            args.sequencia,
            args.source,
            args.comparison_mode,
            args.status_column,
            args.note_column,
            args.ok_value,
            args.fail_value,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.ao_divergir,
            args.bloco_divergir,
            args.sequencia_divergir,
            args.atraso_sequencia_divergir,
        )
    )
    return salvar_e_retornar(payload)


def comando_adicionar_inicio_loop(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(criar_bloco_marcador_loop("inicio_loop", args.delay, args.sequencia))
    return salvar_e_retornar(payload)


def comando_adicionar_fim_loop(args):
    payload = carregar_e_normalizar()
    payload["eventos"].append(
        criar_bloco_marcador_loop(
            "fim_loop",
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    )
    return salvar_e_retornar(payload)


def comando_inserir_texto(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_texto(
            args.text,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            parse_bool_value(getattr(args, "editavel_execucao", False)),
            getattr(args, "rotulo_execucao", "") or "",
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_teclas(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_teclas(args.keys, args.delay, args.sequencia, args.repeat, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_mouse(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_mouse(args.x, args.y, args.button, args.delay, args.clicks, args.sequencia, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_imagem(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_clicar_imagem(
            args.imagem,
            args.delay,
            args.timeout,
            args.confianca,
            args.perfil_busca,
            args.modo_area_busca,
            args.area_busca_x,
            args.area_busca_y,
            args.area_busca_largura,
            args.area_busca_altura,
            args.clicar,
            args.clicks,
            args.offset_x,
            args.offset_y,
            args.janela_alvo,
            args.texto_ocr_esperado,
            args.leitura_relativa_ativa,
            args.leitura_relativa_lado,
            args.leitura_relativa_largura,
            args.leitura_relativa_altura,
            args.leitura_relativa_offset_x,
            args.leitura_relativa_offset_y,
            args.leitura_relativa_salvar_em,
            args.ao_nao_encontrar,
            args.sequencia_nao_encontrar,
            args.atraso_sequencia_nao_encontrar,
            args.bloco_nao_encontrar,
            args.quantidade_tentativas,
            args.intervalo_tentativas,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.sequencia,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_verificacao(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_verificar_estado_botao(
            args.nome,
            args.imagem_ativo,
            args.imagem_desativado,
            args.delay,
            args.timeout,
            args.confianca,
            {"acao": args.acao_ativo, "sequencia": args.sequencia_ativo},
            {"acao": args.acao_desativado, "sequencia": args.sequencia_desativado},
            {"acao": args.acao_nao_encontrar, "sequencia": args.sequencia_nao_encontrar},
            args.sequencia,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_chrome_html(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_chrome_elemento_html(
            args.nome,
            args.html_original,
            json.loads(args.elemento_json) if args.elemento_json else {},
            json.loads(args.alvo_json) if args.alvo_json else {},
            json.loads(args.acao_json) if args.acao_json else {},
            args.timeout,
            args.fallback_visual,
            args.fallback_imagem,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.ao_falhar,
            args.bloco_falha,
            args.sequencia_falha,
            args.atraso_sequencia_falha,
            args.quantidade_tentativas,
            args.intervalo_tentativas,
            args.delay,
            args.sequencia,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_abrir_url(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_abrir_url(
            args.url,
            args.navegador,
            args.modo_abertura,
            args.tela_cheia,
            args.aguardar_segundos,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_abrir_programa(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_abrir_programa(
            args.caminho,
            args.parametros,
            args.aguardar_segundos,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_salvar_download(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_salvar_download(
            args.pasta_origem,
            args.extensao,
            args.pasta_destino,
            args.nome_arquivo,
            args.timeout,
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_planilha(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_planilha(
            args.sheet,
            args.column,
            args.delay,
            args.after_action,
            args.interval_between,
            args.workbook,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_escrita_planilha(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_escrever_planilha(
            args.sheet,
            args.column,
            args.value,
            args.delay,
            args.workbook,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_comparacao_planilha(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_comparar_planilha_tela(
            args.sheet,
            args.expected_column,
            args.delay,
            args.workbook,
            args.sequencia,
            args.source,
            args.comparison_mode,
            args.status_column,
            args.note_column,
            args.ok_value,
            args.fail_value,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
            args.ao_divergir,
            args.bloco_divergir,
            args.sequencia_divergir,
            args.atraso_sequencia_divergir,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_inicio_loop(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_marcador_loop("inicio_loop", args.delay, args.sequencia),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_inserir_fim_loop(args):
    payload = carregar_e_normalizar()
    inserir_evento(
        payload,
        criar_bloco_marcador_loop(
            "fim_loop",
            args.delay,
            args.sequencia,
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        ),
        args.target_order,
        args.position,
    )
    return salvar_e_retornar(payload)


def comando_atualizar_marcador_loop(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") not in {"inicio_loop", "fim_loop"}:
        raise ValueError("O bloco selecionado nao eh marcador de loop.")
    if args.delay is not None:
        evento["intervalo_antes"] = round(float(args.delay), 3)
    if args.sequencia:
        evento["sequencia"] = normalizar_nome_sequencia(args.sequencia)
    if evento.get("tipo") == "fim_loop":
        evento["detalhes"] = aplicar_fluxo_sucesso(
            evento.get("detalhes"),
            args.ao_sucesso,
            args.bloco_sucesso,
            args.sequencia_sucesso,
            args.atraso_sequencia_sucesso,
        )
    return salvar_e_retornar(payload)


def comando_atualizar_texto(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "texto_livre":
        raise ValueError("O bloco selecionado nao eh do tipo texto livre.")
    evento["texto"] = args.text
    detalhes = aplicar_fluxo_sucesso(evento.get("detalhes"), args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso)
    detalhes["editavel_execucao"] = parse_bool_value(getattr(args, "editavel_execucao", detalhes.get("editavel_execucao", False)))
    detalhes["rotulo_execucao"] = str(getattr(args, "rotulo_execucao", detalhes.get("rotulo_execucao", "")) or "").strip()
    evento["detalhes"] = detalhes
    if args.sequencia is not None:
        evento["sequencia"] = normalizar_nome_sequencia(args.sequencia)
    if args.delay is not None:
        evento["intervalo_antes"] = args.delay
    return salvar_e_retornar(payload)


def comando_atualizar_teclas(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") not in {"tecla", "atalho"}:
        raise ValueError("O bloco selecionado nao eh do tipo tecla/atalho.")
    teclas = parse_keys_string(args.keys)
    evento["tipo"] = "atalho" if len(teclas) > 1 else "tecla"
    evento["teclas"] = teclas
    detalhes = evento.get("detalhes") or {}
    detalhes["repeticoes"] = max(1, int(args.repeat or 1))
    detalhes = aplicar_fluxo_sucesso(detalhes, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso)
    evento["detalhes"] = detalhes
    if args.sequencia is not None:
        evento["sequencia"] = normalizar_nome_sequencia(args.sequencia)
    if args.delay is not None:
        evento["intervalo_antes"] = args.delay
    return salvar_e_retornar(payload)


def comando_atualizar_mouse(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "mouse_click":
        raise ValueError("O bloco selecionado nao eh do tipo mouse.")
    detalhes = evento.get("detalhes") or {}
    detalhes["x"] = int(args.x)
    detalhes["y"] = int(args.y)
    detalhes["botao"] = str(args.button).lower()
    detalhes["cliques"] = max(1, int(args.clicks or 1))
    detalhes = aplicar_fluxo_sucesso(detalhes, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso)
    evento["detalhes"] = detalhes
    if args.sequencia is not None:
        evento["sequencia"] = normalizar_nome_sequencia(args.sequencia)
    if args.delay is not None:
        evento["intervalo_antes"] = args.delay
    return salvar_e_retornar(payload)


def comando_atualizar_imagem(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "clicar_imagem":
        raise ValueError("O bloco selecionado nao eh do tipo clicar por imagem.")
    detalhes = evento.get("detalhes") or {}
    detalhes["imagem"] = normalizar_id_imagem(args.imagem)
    detalhes["timeout"] = max(0.1, float(args.timeout))
    detalhes["confianca"] = float(args.confianca)
    detalhes["perfil_busca"] = str(args.perfil_busca or "balanced").strip().lower() or "balanced"
    detalhes["modo_area_busca"] = normalizar_modo_area_busca_reconhecimento(args.modo_area_busca)
    detalhes["area_busca_x"] = int(args.area_busca_x or 0)
    detalhes["area_busca_y"] = int(args.area_busca_y or 0)
    detalhes["area_busca_largura"] = max(0, int(args.area_busca_largura or 0))
    detalhes["area_busca_altura"] = max(0, int(args.area_busca_altura or 0))
    detalhes["clicar"] = bool(args.clicar)
    detalhes["cliques"] = max(0, int(args.clicks if bool(args.clicar) else 0))
    detalhes["offset_x"] = int(args.offset_x)
    detalhes["offset_y"] = int(args.offset_y)
    detalhes["janela_alvo"] = str(args.janela_alvo or "").strip()
    detalhes["texto_ocr_esperado"] = str(args.texto_ocr_esperado or "").strip()
    detalhes["leitura_relativa"] = {
        "ativo": parse_bool_value(args.leitura_relativa_ativa),
        "lado": str(args.leitura_relativa_lado or "livre").strip().lower() or "livre",
        "largura": max(16, int(args.leitura_relativa_largura or 140)),
        "altura": max(12, int(args.leitura_relativa_altura or 40)),
        "offset_x": int(args.leitura_relativa_offset_x or 0),
        "offset_y": int(args.leitura_relativa_offset_y or 0),
        "salvar_em": str(args.leitura_relativa_salvar_em or "").strip(),
    }
    detalhes["ao_nao_encontrar"] = str(args.ao_nao_encontrar or "parar")
    detalhes["sequencia_nao_encontrar"] = str(args.sequencia_nao_encontrar or "").strip()
    detalhes["atraso_antes_sequencia_nao_encontrar"] = normalizar_atraso_transicao(args.atraso_sequencia_nao_encontrar)
    detalhes["bloco_nao_encontrar"] = int(args.bloco_nao_encontrar) if str(args.bloco_nao_encontrar or "").strip() else ""
    detalhes["quantidade_tentativas"] = max(1, int(args.quantidade_tentativas or 1))
    detalhes["intervalo_tentativas"] = max(0.0, float(args.intervalo_tentativas))
    detalhes = aplicar_fluxo_sucesso(detalhes, args.ao_sucesso, args.bloco_sucesso, args.sequencia_sucesso, args.atraso_sequencia_sucesso)
    evento["detalhes"] = detalhes
    if args.delay is not None:
        evento["intervalo_antes"] = args.delay
    if args.sequencia is not None:
        evento["sequencia"] = normalizar_nome_sequencia(args.sequencia)
    return salvar_e_retornar(payload)


def comando_atualizar_verificacao(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "verificar_estado_botao":
        raise ValueError("O bloco selecionado nao eh do tipo verificar estado do botao.")
    bloco = criar_bloco_verificar_estado_botao(
        args.nome,
        args.imagem_ativo,
        args.imagem_desativado,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.timeout,
        args.confianca,
        {"acao": args.acao_ativo, "sequencia": args.sequencia_ativo},
        {"acao": args.acao_desativado, "sequencia": args.sequencia_desativado},
        {"acao": args.acao_nao_encontrar, "sequencia": args.sequencia_nao_encontrar},
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_chrome_html(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "chrome_elemento_html":
        raise ValueError("O bloco selecionado nao eh do tipo Chrome / Elemento HTML.")
    bloco = criar_bloco_chrome_elemento_html(
        args.nome,
        args.html_original,
        json.loads(args.elemento_json) if args.elemento_json else (evento.get("detalhes") or {}).get("elemento", {}),
        json.loads(args.alvo_json) if args.alvo_json else (evento.get("detalhes") or {}).get("alvo", {}),
        json.loads(args.acao_json) if args.acao_json else (evento.get("detalhes") or {}).get("acao", {}),
        args.timeout,
        args.fallback_visual,
        args.fallback_imagem,
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
        args.ao_falhar,
        args.bloco_falha,
        args.sequencia_falha,
        args.atraso_sequencia_falha,
        args.quantidade_tentativas,
        args.intervalo_tentativas,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_abrir_url(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "abrir_url":
        raise ValueError("O bloco selecionado nao eh do tipo Abrir URL.")
    bloco = criar_bloco_abrir_url(
        args.url,
        args.navegador,
        args.modo_abertura,
        args.tela_cheia,
        args.aguardar_segundos,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_abrir_programa(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "abrir_programa":
        raise ValueError("O bloco selecionado nao eh do tipo Abrir Programa.")
    bloco = criar_bloco_abrir_programa(
        args.caminho,
        args.parametros,
        args.aguardar_segundos,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_salvar_download(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "salvar_download":
        raise ValueError("O bloco selecionado nao eh do tipo Salvar download.")
    detalhes_atuais = evento.get("detalhes") or {}
    bloco = criar_bloco_salvar_download(
        args.pasta_origem if args.pasta_origem is not None else detalhes_atuais.get("pasta_origem", ""),
        args.extensao if args.extensao is not None else detalhes_atuais.get("extensao", ".pdf"),
        args.pasta_destino if args.pasta_destino is not None else detalhes_atuais.get("pasta_destino", ""),
        args.nome_arquivo if args.nome_arquivo is not None else detalhes_atuais.get("nome_arquivo", ""),
        args.timeout if args.timeout is not None else detalhes_atuais.get("timeout", 60),
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_planilha(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "sequencia_planilha":
        raise ValueError("O bloco selecionado nao eh do tipo planilha.")
    bloco = criar_bloco_planilha(
        args.sheet,
        args.column,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.after_action,
        args.interval_between,
        args.workbook or (evento.get("detalhes") or {}).get("arquivo"),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_escrita_planilha(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "escrever_planilha":
        raise ValueError("O bloco selecionado nao eh do tipo escrever na planilha.")
    detalhes_atuais = evento.get("detalhes") or {}
    bloco = criar_bloco_escrever_planilha(
        args.sheet,
        args.column,
        args.value,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.workbook or detalhes_atuais.get("arquivo"),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_comparacao_planilha(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    if evento.get("tipo") != "comparar_planilha_tela":
        raise ValueError("O bloco selecionado nao eh do tipo conferir valor.")
    detalhes_atuais = evento.get("detalhes") or {}
    bloco = criar_bloco_comparar_planilha_tela(
        args.sheet,
        args.expected_column,
        args.delay if args.delay is not None else evento.get("intervalo_antes", 0),
        args.workbook or detalhes_atuais.get("arquivo"),
        args.sequencia if args.sequencia is not None else evento.get("sequencia", "principal"),
        args.source,
        args.comparison_mode,
        args.status_column,
        args.note_column,
        args.ok_value,
        args.fail_value,
        args.ao_sucesso,
        args.bloco_sucesso,
        args.sequencia_sucesso,
        args.atraso_sequencia_sucesso,
        args.ao_divergir,
        args.bloco_divergir,
        args.sequencia_divergir,
        args.atraso_sequencia_divergir,
    )
    evento.update(bloco)
    return salvar_e_retornar(payload)


def comando_atualizar_atraso(args):
    payload = carregar_e_normalizar()
    _, evento = localizar_evento(payload, args.order)
    evento["intervalo_antes"] = round(float(args.delay), 3)
    return salvar_e_retornar(payload)


def comando_atualizar_atraso_global(args):
    payload = carregar_e_normalizar()
    delay = round(float(args.delay), 3)
    for evento in payload.get("eventos", []):
        evento["intervalo_antes"] = delay
    return salvar_e_retornar(payload)


def comando_atualizar_configuracao_sequencia(args):
    payload = carregar_e_normalizar()
    sequencia = normalizar_nome_sequencia(args.sequencia)
    modo_execucao = normalizar_modo_execucao_sequencia(args.modo_execucao)
    sequencia_apos_loop = str(args.sequencia_apos_loop or "").strip()
    sequencia_apos_loop = normalizar_nome_sequencia(sequencia_apos_loop) if sequencia_apos_loop else ""
    if sequencia_apos_loop == sequencia:
        sequencia_apos_loop = ""

    configuracoes = normalizar_configuracoes_sequencia(payload)
    if modo_execucao == "normal":
        configuracoes.pop(sequencia, None)
    else:
        configuracoes[sequencia] = {
            "modo_execucao": modo_execucao,
            "sequencia_apos_loop": sequencia_apos_loop,
        }
    payload["configuracoes_sequencia"] = configuracoes
    normalizar_configuracoes_sequencia(payload)
    return salvar_e_retornar(payload)


def comando_atualizar_sequencia_inicial(args):
    payload = carregar_e_normalizar()
    payload["sequencia_inicial"] = normalizar_nome_sequencia(args.sequencia)
    return salvar_e_retornar(payload)


def comando_sincronizar_planilhas(_args):
    payload = carregar_e_normalizar()
    return salvar_e_retornar(payload)


def comando_mover(args):
    payload = carregar_e_normalizar()
    mapa_ordem_anterior = mapa_ordens_eventos_existentes(payload.get("eventos", []))
    indice, _ = localizar_evento(payload, args.order)
    novo_indice = indice - 1 if args.direction == "up" else indice + 1
    if novo_indice < 0 or novo_indice >= len(payload["eventos"]):
        return payload_para_resposta(payload)
    payload["eventos"][indice], payload["eventos"][novo_indice] = (
        payload["eventos"][novo_indice],
        payload["eventos"][indice],
    )
    remapear_referencias_de_blocos(payload, mapa_ordem_anterior)
    return salvar_e_retornar(payload)


def comando_reordenar(args):
    payload = carregar_e_normalizar()
    mapa_ordem_anterior = mapa_ordens_eventos_existentes(payload.get("eventos", []))
    indice_atual, evento = localizar_evento(payload, args.order)
    eventos = payload.get("eventos", [])
    evento_removido = eventos.pop(indice_atual)

    if not eventos:
        eventos.append(evento_removido)
        payload["eventos"] = eventos
        return salvar_e_retornar(payload)

    target_index = max(1, int(args.target_index))
    if indice_atual < target_index - 1:
        target_index -= 1
    target_index = min(target_index, len(eventos) + 1)
    eventos.insert(target_index - 1, evento_removido)
    payload["eventos"] = eventos
    remapear_referencias_de_blocos(payload, mapa_ordem_anterior)
    return salvar_e_retornar(payload)


def comando_remover(args):
    payload = carregar_e_normalizar()
    mapa_ordem_anterior = mapa_ordens_eventos_existentes(payload.get("eventos", []))
    indice, _ = localizar_evento(payload, args.order)
    payload["eventos"].pop(indice)
    remapear_referencias_de_blocos(payload, mapa_ordem_anterior)
    return salvar_e_retornar(payload)


def comando_limpar(_args):
    payload = carregar_e_normalizar()
    payload["eventos"] = []
    return salvar_e_retornar(payload)


def comando_analisar_html_elemento(args):
    return parse_html_elemento(args.html)


def build_parser():
    parser = argparse.ArgumentParser(description="Gerencia o roteiro visual da automacao.")
    subparsers = parser.add_subparsers(dest="command", required=True)

    def adicionar_argumentos_fluxo_sucesso(parser_destino):
        parser_destino.add_argument("--ao-sucesso", choices=["continuar", "ir_para_bloco", "ir_para_sequencia"], default="continuar")
        parser_destino.add_argument("--bloco-sucesso", default="")
        parser_destino.add_argument("--sequencia-sucesso", default="")
        parser_destino.add_argument("--atraso-sequencia-sucesso", type=float, default=0)

    def adicionar_argumentos_leitura_relativa(parser_destino):
        parser_destino.add_argument("--leitura-relativa-ativa", default="false")
        parser_destino.add_argument("--leitura-relativa-lado", choices=["direita", "esquerda", "abaixo", "acima", "centro", "livre"], default="livre")
        parser_destino.add_argument("--leitura-relativa-largura", type=int, default=140)
        parser_destino.add_argument("--leitura-relativa-altura", type=int, default=40)
        parser_destino.add_argument("--leitura-relativa-offset-x", type=int, default=0)
        parser_destino.add_argument("--leitura-relativa-offset-y", type=int, default=0)
        parser_destino.add_argument("--leitura-relativa-salvar-em", default="")

    def adicionar_argumentos_abrir_url(parser_destino):
        parser_destino.add_argument("--url", required=True)
        parser_destino.add_argument("--navegador", choices=["padrao", "chrome", "edge"], default="padrao")
        parser_destino.add_argument("--modo-abertura", choices=["nova_janela", "nova_aba"], default="nova_janela")
        parser_destino.add_argument("--tela-cheia", action="store_true")
        parser_destino.add_argument("--aguardar-segundos", type=float, default=3)

    def adicionar_argumentos_abrir_programa(parser_destino):
        parser_destino.add_argument("--caminho", required=True)
        parser_destino.add_argument("--parametros", default="")
        parser_destino.add_argument("--aguardar-segundos", type=float, default=2)

    def adicionar_argumentos_salvar_download(parser_destino):
        parser_destino.add_argument("--pasta-origem", default="")
        parser_destino.add_argument("--extensao", default=".pdf")
        parser_destino.add_argument("--pasta-destino", required=True)
        parser_destino.add_argument("--nome-arquivo", default="{empresa}_{competencia}_FGTS.pdf")
        parser_destino.add_argument("--timeout", type=float, default=60)

    def adicionar_argumentos_comparar_planilha(parser_destino):
        parser_destino.add_argument("--workbook")
        parser_destino.add_argument("--sheet", required=True)
        parser_destino.add_argument("--expected-column", required=True)
        parser_destino.add_argument("--source", default="ultimo_ocr")
        parser_destino.add_argument("--comparison-mode", choices=["texto", "numero", "data"], default="texto")
        parser_destino.add_argument("--status-column", default="")
        parser_destino.add_argument("--note-column", default="")
        parser_destino.add_argument("--ok-value", default="OK")
        parser_destino.add_argument("--fail-value", default="Divergente")
        parser_destino.add_argument("--ao-divergir", choices=["continuar", "parar", "ir_para_bloco", "ir_para_sequencia"], default="continuar")
        parser_destino.add_argument("--bloco-divergir", default="")
        parser_destino.add_argument("--sequencia-divergir", default="")
        parser_destino.add_argument("--atraso-sequencia-divergir", type=float, default=0)

    subparsers.add_parser("list")
    subparsers.add_parser("automation-library-list")
    subparsers.add_parser("workbook-meta")
    analyze_html = subparsers.add_parser("analyze-html-element")
    analyze_html.add_argument("--html", required=True)
    preview_sheet = subparsers.add_parser("workbook-preview")
    preview_sheet.add_argument("--workbook")
    preview_sheet.add_argument("--sheet")
    preview_sheet.add_argument("--limit", type=int, default=200)
    subparsers.add_parser("clear")

    create_automation = subparsers.add_parser("automation-create")
    create_automation.add_argument("--nome", required=True)
    create_automation.add_argument("--descricao", default="")
    create_automation.add_argument("--categoria", default="Outros")

    open_automation = subparsers.add_parser("automation-open")
    open_automation.add_argument("--id", required=True)

    duplicate_automation = subparsers.add_parser("automation-duplicate")
    duplicate_automation.add_argument("--id", required=True)

    rename_automation = subparsers.add_parser("automation-rename")
    rename_automation.add_argument("--id", required=True)
    rename_automation.add_argument("--nome", required=True)
    rename_automation.add_argument("--descricao")
    rename_automation.add_argument("--categoria")

    delete_automation = subparsers.add_parser("automation-delete")
    delete_automation.add_argument("--id", required=True)

    mark_run = subparsers.add_parser("automation-mark-run")
    mark_run.add_argument("--id", required=True)

    export_zip = subparsers.add_parser("automation-export-zip")
    export_zip.add_argument("--id", required=True)
    export_zip.add_argument("--destino")

    import_zip = subparsers.add_parser("automation-import-zip")
    import_zip.add_argument("--arquivo", required=True)

    runtime_info = subparsers.add_parser("runtime-info")

    export_runtime_zip = subparsers.add_parser("runtime-export-zip")
    export_runtime_zip.add_argument("--destino")

    import_runtime_zip = subparsers.add_parser("runtime-import-zip")
    import_runtime_zip.add_argument("--arquivo", required=True)

    add_text = subparsers.add_parser("add-text")
    add_text.add_argument("--text", required=True)
    add_text.add_argument("--delay", type=float, default=0)
    add_text.add_argument("--sequencia", default="principal")
    add_text.add_argument("--editavel-execucao", default="false")
    add_text.add_argument("--rotulo-execucao", default="")
    adicionar_argumentos_fluxo_sucesso(add_text)

    add_key = subparsers.add_parser("add-key")
    add_key.add_argument("--keys", required=True)
    add_key.add_argument("--repeat", type=int, default=1)
    add_key.add_argument("--delay", type=float, default=0)
    add_key.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_key)

    add_mouse = subparsers.add_parser("add-mouse")
    add_mouse.add_argument("--x", type=int, required=True)
    add_mouse.add_argument("--y", type=int, required=True)
    add_mouse.add_argument("--button", default="left")
    add_mouse.add_argument("--clicks", type=int, default=1)
    add_mouse.add_argument("--delay", type=float, default=0)
    add_mouse.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_mouse)

    add_image = subparsers.add_parser("add-image-click")
    add_image.add_argument("--imagem", required=True)
    add_image.add_argument("--timeout", type=float, default=10)
    add_image.add_argument("--confianca", type=float, default=0.8)
    add_image.add_argument("--perfil-busca", default="balanced")
    add_image.add_argument("--modo-area-busca", default="target_window")
    add_image.add_argument("--area-busca-x", type=int, default=0)
    add_image.add_argument("--area-busca-y", type=int, default=0)
    add_image.add_argument("--area-busca-largura", type=int, default=0)
    add_image.add_argument("--area-busca-altura", type=int, default=0)
    add_image.add_argument("--clicar", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=True)
    add_image.add_argument("--clicks", type=int, default=1)
    add_image.add_argument("--offset-x", type=int, default=0)
    add_image.add_argument("--offset-y", type=int, default=0)
    add_image.add_argument("--janela-alvo", default="")
    add_image.add_argument("--texto-ocr-esperado", default="")
    adicionar_argumentos_leitura_relativa(add_image)
    add_image.add_argument("--ao-nao-encontrar", choices=["parar", "continuar", "pular", "tentar_novamente", "ir_para_sequencia", "ir_para_bloco"], default="parar")
    add_image.add_argument("--sequencia-nao-encontrar", default="")
    add_image.add_argument("--atraso-sequencia-nao-encontrar", type=float, default=0)
    add_image.add_argument("--bloco-nao-encontrar", default="")
    add_image.add_argument("--quantidade-tentativas", type=int, default=1)
    add_image.add_argument("--intervalo-tentativas", type=float, default=1)
    add_image.add_argument("--delay", type=float, default=0)
    add_image.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_image)

    add_verify = subparsers.add_parser("add-verify-button-state")
    add_verify.add_argument("--nome", default="Verificar estado do botao")
    add_verify.add_argument("--imagem-ativo", required=True)
    add_verify.add_argument("--imagem-desativado", required=True)
    add_verify.add_argument("--timeout", type=float, default=5)
    add_verify.add_argument("--confianca", type=float, default=0.85)
    add_verify.add_argument("--acao-ativo", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="continuar")
    add_verify.add_argument("--sequencia-ativo", default="")
    add_verify.add_argument("--acao-desativado", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="ir_para_sequencia")
    add_verify.add_argument("--sequencia-desativado", default="")
    add_verify.add_argument("--acao-nao-encontrar", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="parar")
    add_verify.add_argument("--sequencia-nao-encontrar", default="")
    add_verify.add_argument("--delay", type=float, default=0)
    add_verify.add_argument("--sequencia", default="principal")

    add_chrome = subparsers.add_parser("add-chrome-html")
    add_chrome.add_argument("--nome", default="Chrome / Elemento HTML")
    add_chrome.add_argument("--html-original", required=True)
    add_chrome.add_argument("--elemento-json", default="{}")
    add_chrome.add_argument("--alvo-json", default="{}")
    add_chrome.add_argument("--acao-json", default="{}")
    add_chrome.add_argument("--timeout", type=int, default=10000)
    add_chrome.add_argument("--fallback-visual", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=False)
    add_chrome.add_argument("--fallback-imagem", default="")
    add_chrome.add_argument("--ao-sucesso", choices=["continuar", "ir_para_bloco", "ir_para_sequencia"], default="continuar")
    add_chrome.add_argument("--bloco-sucesso", default="")
    add_chrome.add_argument("--sequencia-sucesso", default="")
    add_chrome.add_argument("--atraso-sequencia-sucesso", type=float, default=0)
    add_chrome.add_argument("--ao-falhar", choices=["parar", "continuar", "tentar_novamente", "ir_para_bloco", "ir_para_sequencia"], default="parar")
    add_chrome.add_argument("--bloco-falha", default="")
    add_chrome.add_argument("--sequencia-falha", default="")
    add_chrome.add_argument("--atraso-sequencia-falha", type=float, default=0)
    add_chrome.add_argument("--quantidade-tentativas", type=int, default=1)
    add_chrome.add_argument("--intervalo-tentativas", type=float, default=0.5)
    add_chrome.add_argument("--delay", type=float, default=0)
    add_chrome.add_argument("--sequencia", default="principal")

    add_open_url = subparsers.add_parser("add-open-url")
    adicionar_argumentos_abrir_url(add_open_url)
    add_open_url.add_argument("--delay", type=float, default=0)
    add_open_url.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_open_url)

    add_open_program = subparsers.add_parser("add-open-program")
    adicionar_argumentos_abrir_programa(add_open_program)
    add_open_program.add_argument("--delay", type=float, default=0)
    add_open_program.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_open_program)

    add_save_download = subparsers.add_parser("add-save-download")
    adicionar_argumentos_salvar_download(add_save_download)
    add_save_download.add_argument("--delay", type=float, default=0)
    add_save_download.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_save_download)

    add_sheet = subparsers.add_parser("add-planilha")
    add_sheet.add_argument("--workbook")
    add_sheet.add_argument("--sheet", required=True)
    add_sheet.add_argument("--column", required=True)
    add_sheet.add_argument("--delay", type=float, default=0)
    add_sheet.add_argument("--after-action", default="nenhuma")
    add_sheet.add_argument("--interval-between", type=float, default=0.5)
    add_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_sheet)

    add_write_sheet = subparsers.add_parser("add-write-planilha")
    add_write_sheet.add_argument("--workbook")
    add_write_sheet.add_argument("--sheet", required=True)
    add_write_sheet.add_argument("--column", required=True)
    add_write_sheet.add_argument("--value", default="")
    add_write_sheet.add_argument("--delay", type=float, default=0)
    add_write_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_write_sheet)

    add_compare_sheet = subparsers.add_parser("add-compare-planilha")
    adicionar_argumentos_comparar_planilha(add_compare_sheet)
    add_compare_sheet.add_argument("--delay", type=float, default=0)
    add_compare_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_compare_sheet)

    add_loop_start = subparsers.add_parser("add-loop-start")
    add_loop_start.add_argument("--delay", type=float, default=0)
    add_loop_start.add_argument("--sequencia", default="principal")

    add_loop_end = subparsers.add_parser("add-loop-end")
    add_loop_end.add_argument("--delay", type=float, default=0)
    add_loop_end.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(add_loop_end)

    insert_text = subparsers.add_parser("insert-text")
    insert_text.add_argument("--target-order", type=int, required=True)
    insert_text.add_argument("--position", choices=["before", "after"], required=True)
    insert_text.add_argument("--text", required=True)
    insert_text.add_argument("--delay", type=float, default=0)
    insert_text.add_argument("--sequencia", default="principal")
    insert_text.add_argument("--editavel-execucao", default="false")
    insert_text.add_argument("--rotulo-execucao", default="")
    adicionar_argumentos_fluxo_sucesso(insert_text)

    insert_key = subparsers.add_parser("insert-key")
    insert_key.add_argument("--target-order", type=int, required=True)
    insert_key.add_argument("--position", choices=["before", "after"], required=True)
    insert_key.add_argument("--keys", required=True)
    insert_key.add_argument("--repeat", type=int, default=1)
    insert_key.add_argument("--delay", type=float, default=0)
    insert_key.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_key)

    insert_mouse = subparsers.add_parser("insert-mouse")
    insert_mouse.add_argument("--target-order", type=int, required=True)
    insert_mouse.add_argument("--position", choices=["before", "after"], required=True)
    insert_mouse.add_argument("--x", type=int, required=True)
    insert_mouse.add_argument("--y", type=int, required=True)
    insert_mouse.add_argument("--button", default="left")
    insert_mouse.add_argument("--clicks", type=int, default=1)
    insert_mouse.add_argument("--delay", type=float, default=0)
    insert_mouse.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_mouse)

    insert_image = subparsers.add_parser("insert-image-click")
    insert_image.add_argument("--target-order", type=int, required=True)
    insert_image.add_argument("--position", choices=["before", "after"], required=True)
    insert_image.add_argument("--imagem", required=True)
    insert_image.add_argument("--timeout", type=float, default=10)
    insert_image.add_argument("--confianca", type=float, default=0.8)
    insert_image.add_argument("--perfil-busca", default="balanced")
    insert_image.add_argument("--modo-area-busca", default="target_window")
    insert_image.add_argument("--area-busca-x", type=int, default=0)
    insert_image.add_argument("--area-busca-y", type=int, default=0)
    insert_image.add_argument("--area-busca-largura", type=int, default=0)
    insert_image.add_argument("--area-busca-altura", type=int, default=0)
    insert_image.add_argument("--clicar", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=True)
    insert_image.add_argument("--clicks", type=int, default=1)
    insert_image.add_argument("--offset-x", type=int, default=0)
    insert_image.add_argument("--offset-y", type=int, default=0)
    insert_image.add_argument("--janela-alvo", default="")
    insert_image.add_argument("--texto-ocr-esperado", default="")
    adicionar_argumentos_leitura_relativa(insert_image)
    insert_image.add_argument("--ao-nao-encontrar", choices=["parar", "continuar", "pular", "tentar_novamente", "ir_para_sequencia", "ir_para_bloco"], default="parar")
    insert_image.add_argument("--sequencia-nao-encontrar", default="")
    insert_image.add_argument("--atraso-sequencia-nao-encontrar", type=float, default=0)
    insert_image.add_argument("--bloco-nao-encontrar", default="")
    insert_image.add_argument("--quantidade-tentativas", type=int, default=1)
    insert_image.add_argument("--intervalo-tentativas", type=float, default=1)
    insert_image.add_argument("--delay", type=float, default=0)
    insert_image.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_image)

    insert_verify = subparsers.add_parser("insert-verify-button-state")
    insert_verify.add_argument("--target-order", type=int, required=True)
    insert_verify.add_argument("--position", choices=["before", "after"], required=True)
    insert_verify.add_argument("--nome", default="Verificar estado do botao")
    insert_verify.add_argument("--imagem-ativo", required=True)
    insert_verify.add_argument("--imagem-desativado", required=True)
    insert_verify.add_argument("--timeout", type=float, default=5)
    insert_verify.add_argument("--confianca", type=float, default=0.85)
    insert_verify.add_argument("--acao-ativo", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="continuar")
    insert_verify.add_argument("--sequencia-ativo", default="")
    insert_verify.add_argument("--acao-desativado", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="ir_para_sequencia")
    insert_verify.add_argument("--sequencia-desativado", default="")
    insert_verify.add_argument("--acao-nao-encontrar", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="parar")
    insert_verify.add_argument("--sequencia-nao-encontrar", default="")
    insert_verify.add_argument("--delay", type=float, default=0)
    insert_verify.add_argument("--sequencia", default="principal")

    insert_chrome = subparsers.add_parser("insert-chrome-html")
    insert_chrome.add_argument("--target-order", type=int, required=True)
    insert_chrome.add_argument("--position", choices=["before", "after"], required=True)
    insert_chrome.add_argument("--nome", default="Chrome / Elemento HTML")
    insert_chrome.add_argument("--html-original", required=True)
    insert_chrome.add_argument("--elemento-json", default="{}")
    insert_chrome.add_argument("--alvo-json", default="{}")
    insert_chrome.add_argument("--acao-json", default="{}")
    insert_chrome.add_argument("--timeout", type=int, default=10000)
    insert_chrome.add_argument("--fallback-visual", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=False)
    insert_chrome.add_argument("--fallback-imagem", default="")
    insert_chrome.add_argument("--ao-sucesso", choices=["continuar", "ir_para_bloco", "ir_para_sequencia"], default="continuar")
    insert_chrome.add_argument("--bloco-sucesso", default="")
    insert_chrome.add_argument("--sequencia-sucesso", default="")
    insert_chrome.add_argument("--atraso-sequencia-sucesso", type=float, default=0)
    insert_chrome.add_argument("--ao-falhar", choices=["parar", "continuar", "tentar_novamente", "ir_para_bloco", "ir_para_sequencia"], default="parar")
    insert_chrome.add_argument("--bloco-falha", default="")
    insert_chrome.add_argument("--sequencia-falha", default="")
    insert_chrome.add_argument("--atraso-sequencia-falha", type=float, default=0)
    insert_chrome.add_argument("--quantidade-tentativas", type=int, default=1)
    insert_chrome.add_argument("--intervalo-tentativas", type=float, default=0.5)
    insert_chrome.add_argument("--delay", type=float, default=0)
    insert_chrome.add_argument("--sequencia", default="principal")

    insert_open_url = subparsers.add_parser("insert-open-url")
    insert_open_url.add_argument("--target-order", type=int, required=True)
    insert_open_url.add_argument("--position", choices=["before", "after"], required=True)
    adicionar_argumentos_abrir_url(insert_open_url)
    insert_open_url.add_argument("--delay", type=float, default=0)
    insert_open_url.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_open_url)

    insert_open_program = subparsers.add_parser("insert-open-program")
    insert_open_program.add_argument("--target-order", type=int, required=True)
    insert_open_program.add_argument("--position", choices=["before", "after"], required=True)
    adicionar_argumentos_abrir_programa(insert_open_program)
    insert_open_program.add_argument("--delay", type=float, default=0)
    insert_open_program.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_open_program)

    insert_save_download = subparsers.add_parser("insert-save-download")
    insert_save_download.add_argument("--target-order", type=int, required=True)
    insert_save_download.add_argument("--position", choices=["before", "after"], required=True)
    adicionar_argumentos_salvar_download(insert_save_download)
    insert_save_download.add_argument("--delay", type=float, default=0)
    insert_save_download.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_save_download)

    insert_sheet = subparsers.add_parser("insert-planilha")
    insert_sheet.add_argument("--workbook")
    insert_sheet.add_argument("--target-order", type=int, required=True)
    insert_sheet.add_argument("--position", choices=["before", "after"], required=True)
    insert_sheet.add_argument("--sheet", required=True)
    insert_sheet.add_argument("--column", required=True)
    insert_sheet.add_argument("--delay", type=float, default=0)
    insert_sheet.add_argument("--after-action", default="nenhuma")
    insert_sheet.add_argument("--interval-between", type=float, default=0.5)
    insert_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_sheet)

    insert_write_sheet = subparsers.add_parser("insert-write-planilha")
    insert_write_sheet.add_argument("--workbook")
    insert_write_sheet.add_argument("--target-order", type=int, required=True)
    insert_write_sheet.add_argument("--position", choices=["before", "after"], required=True)
    insert_write_sheet.add_argument("--sheet", required=True)
    insert_write_sheet.add_argument("--column", required=True)
    insert_write_sheet.add_argument("--value", default="")
    insert_write_sheet.add_argument("--delay", type=float, default=0)
    insert_write_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_write_sheet)

    insert_compare_sheet = subparsers.add_parser("insert-compare-planilha")
    insert_compare_sheet.add_argument("--target-order", type=int, required=True)
    insert_compare_sheet.add_argument("--position", choices=["before", "after"], required=True)
    adicionar_argumentos_comparar_planilha(insert_compare_sheet)
    insert_compare_sheet.add_argument("--delay", type=float, default=0)
    insert_compare_sheet.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_compare_sheet)

    update_loop_marker = subparsers.add_parser("update-loop-marker")
    update_loop_marker.add_argument("--order", type=int, required=True)
    update_loop_marker.add_argument("--delay", type=float)
    update_loop_marker.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_loop_marker)

    insert_loop_start = subparsers.add_parser("insert-loop-start")
    insert_loop_start.add_argument("--target-order", type=int, required=True)
    insert_loop_start.add_argument("--position", choices=["before", "after"], required=True)
    insert_loop_start.add_argument("--delay", type=float, default=0)
    insert_loop_start.add_argument("--sequencia", default="principal")

    insert_loop_end = subparsers.add_parser("insert-loop-end")
    insert_loop_end.add_argument("--target-order", type=int, required=True)
    insert_loop_end.add_argument("--position", choices=["before", "after"], required=True)
    insert_loop_end.add_argument("--delay", type=float, default=0)
    insert_loop_end.add_argument("--sequencia", default="principal")
    adicionar_argumentos_fluxo_sucesso(insert_loop_end)

    update_text = subparsers.add_parser("update-text")
    update_text.add_argument("--order", type=int, required=True)
    update_text.add_argument("--text", required=True)
    update_text.add_argument("--delay", type=float)
    update_text.add_argument("--sequencia")
    update_text.add_argument("--editavel-execucao", default="false")
    update_text.add_argument("--rotulo-execucao", default="")
    adicionar_argumentos_fluxo_sucesso(update_text)

    update_key = subparsers.add_parser("update-key")
    update_key.add_argument("--order", type=int, required=True)
    update_key.add_argument("--keys", required=True)
    update_key.add_argument("--repeat", type=int, default=1)
    update_key.add_argument("--delay", type=float)
    update_key.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_key)

    update_mouse = subparsers.add_parser("update-mouse")
    update_mouse.add_argument("--order", type=int, required=True)
    update_mouse.add_argument("--x", type=int, required=True)
    update_mouse.add_argument("--y", type=int, required=True)
    update_mouse.add_argument("--button", default="left")
    update_mouse.add_argument("--clicks", type=int, default=1)
    update_mouse.add_argument("--delay", type=float)
    update_mouse.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_mouse)

    update_image = subparsers.add_parser("update-image-click")
    update_image.add_argument("--order", type=int, required=True)
    update_image.add_argument("--imagem", required=True)
    update_image.add_argument("--timeout", type=float, default=10)
    update_image.add_argument("--confianca", type=float, default=0.8)
    update_image.add_argument("--perfil-busca", default="balanced")
    update_image.add_argument("--modo-area-busca", default="target_window")
    update_image.add_argument("--area-busca-x", type=int, default=0)
    update_image.add_argument("--area-busca-y", type=int, default=0)
    update_image.add_argument("--area-busca-largura", type=int, default=0)
    update_image.add_argument("--area-busca-altura", type=int, default=0)
    update_image.add_argument("--clicar", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=True)
    update_image.add_argument("--clicks", type=int, default=1)
    update_image.add_argument("--offset-x", type=int, default=0)
    update_image.add_argument("--offset-y", type=int, default=0)
    update_image.add_argument("--janela-alvo", default="")
    update_image.add_argument("--texto-ocr-esperado", default="")
    adicionar_argumentos_leitura_relativa(update_image)
    update_image.add_argument("--ao-nao-encontrar", choices=["parar", "continuar", "pular", "tentar_novamente", "ir_para_sequencia", "ir_para_bloco"], default="parar")
    update_image.add_argument("--sequencia-nao-encontrar", default="")
    update_image.add_argument("--atraso-sequencia-nao-encontrar", type=float, default=0)
    update_image.add_argument("--bloco-nao-encontrar", default="")
    update_image.add_argument("--quantidade-tentativas", type=int, default=1)
    update_image.add_argument("--intervalo-tentativas", type=float, default=1)
    update_image.add_argument("--delay", type=float)
    update_image.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_image)

    update_verify = subparsers.add_parser("update-verify-button-state")
    update_verify.add_argument("--order", type=int, required=True)
    update_verify.add_argument("--nome", default="Verificar estado do botao")
    update_verify.add_argument("--imagem-ativo", required=True)
    update_verify.add_argument("--imagem-desativado", required=True)
    update_verify.add_argument("--timeout", type=float, default=5)
    update_verify.add_argument("--confianca", type=float, default=0.85)
    update_verify.add_argument("--acao-ativo", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="continuar")
    update_verify.add_argument("--sequencia-ativo", default="")
    update_verify.add_argument("--acao-desativado", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="ir_para_sequencia")
    update_verify.add_argument("--sequencia-desativado", default="")
    update_verify.add_argument("--acao-nao-encontrar", choices=["continuar", "parar", "pular_card", "ir_para_sequencia"], default="parar")
    update_verify.add_argument("--sequencia-nao-encontrar", default="")
    update_verify.add_argument("--delay", type=float)
    update_verify.add_argument("--sequencia")

    update_chrome = subparsers.add_parser("update-chrome-html")
    update_chrome.add_argument("--order", type=int, required=True)
    update_chrome.add_argument("--nome", default="Chrome / Elemento HTML")
    update_chrome.add_argument("--html-original", required=True)
    update_chrome.add_argument("--elemento-json", default="{}")
    update_chrome.add_argument("--alvo-json", default="{}")
    update_chrome.add_argument("--acao-json", default="{}")
    update_chrome.add_argument("--timeout", type=int, default=10000)
    update_chrome.add_argument("--fallback-visual", type=lambda value: str(value).lower() in {"1", "true", "sim", "yes"}, default=False)
    update_chrome.add_argument("--fallback-imagem", default="")
    update_chrome.add_argument("--ao-sucesso", choices=["continuar", "ir_para_bloco", "ir_para_sequencia"], default="continuar")
    update_chrome.add_argument("--bloco-sucesso", default="")
    update_chrome.add_argument("--sequencia-sucesso", default="")
    update_chrome.add_argument("--atraso-sequencia-sucesso", type=float, default=0)
    update_chrome.add_argument("--ao-falhar", choices=["parar", "continuar", "tentar_novamente", "ir_para_bloco", "ir_para_sequencia"], default="parar")
    update_chrome.add_argument("--bloco-falha", default="")
    update_chrome.add_argument("--sequencia-falha", default="")
    update_chrome.add_argument("--atraso-sequencia-falha", type=float, default=0)
    update_chrome.add_argument("--quantidade-tentativas", type=int, default=1)
    update_chrome.add_argument("--intervalo-tentativas", type=float, default=0.5)
    update_chrome.add_argument("--delay", type=float)
    update_chrome.add_argument("--sequencia")

    update_open_url = subparsers.add_parser("update-open-url")
    update_open_url.add_argument("--order", type=int, required=True)
    adicionar_argumentos_abrir_url(update_open_url)
    update_open_url.add_argument("--delay", type=float)
    update_open_url.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_open_url)

    update_open_program = subparsers.add_parser("update-open-program")
    update_open_program.add_argument("--order", type=int, required=True)
    adicionar_argumentos_abrir_programa(update_open_program)
    update_open_program.add_argument("--delay", type=float)
    update_open_program.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_open_program)

    update_save_download = subparsers.add_parser("update-save-download")
    update_save_download.add_argument("--order", type=int, required=True)
    adicionar_argumentos_salvar_download(update_save_download)
    update_save_download.add_argument("--delay", type=float)
    update_save_download.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_save_download)

    update_sheet = subparsers.add_parser("update-planilha")
    update_sheet.add_argument("--workbook")
    update_sheet.add_argument("--order", type=int, required=True)
    update_sheet.add_argument("--sheet", required=True)
    update_sheet.add_argument("--column", required=True)
    update_sheet.add_argument("--after-action", default="nenhuma")
    update_sheet.add_argument("--interval-between", type=float, default=0.5)
    update_sheet.add_argument("--delay", type=float)
    update_sheet.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_sheet)

    update_write_sheet = subparsers.add_parser("update-write-planilha")
    update_write_sheet.add_argument("--workbook")
    update_write_sheet.add_argument("--order", type=int, required=True)
    update_write_sheet.add_argument("--sheet", required=True)
    update_write_sheet.add_argument("--column", required=True)
    update_write_sheet.add_argument("--value", default="")
    update_write_sheet.add_argument("--delay", type=float)
    update_write_sheet.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_write_sheet)

    update_compare_sheet = subparsers.add_parser("update-compare-planilha")
    update_compare_sheet.add_argument("--order", type=int, required=True)
    adicionar_argumentos_comparar_planilha(update_compare_sheet)
    update_compare_sheet.add_argument("--delay", type=float)
    update_compare_sheet.add_argument("--sequencia")
    adicionar_argumentos_fluxo_sucesso(update_compare_sheet)

    update_delay = subparsers.add_parser("update-delay")
    update_delay.add_argument("--order", type=int, required=True)
    update_delay.add_argument("--delay", type=float, required=True)

    update_delay_all = subparsers.add_parser("update-delay-all")
    update_delay_all.add_argument("--delay", type=float, required=True)

    update_sequence_settings = subparsers.add_parser("update-sequence-settings")
    update_sequence_settings.add_argument("--sequencia", required=True)
    update_sequence_settings.add_argument("--modo-execucao", choices=sorted(MODOS_EXECUCAO_SEQUENCIA), default="normal")
    update_sequence_settings.add_argument("--sequencia-apos-loop", default="")

    update_initial_sequence = subparsers.add_parser("update-initial-sequence")
    update_initial_sequence.add_argument("--sequencia", required=True)

    subparsers.add_parser("sync-spreadsheets")

    move = subparsers.add_parser("move")
    move.add_argument("--order", type=int, required=True)
    move.add_argument("--direction", choices=["up", "down"], required=True)

    reorder = subparsers.add_parser("reorder")
    reorder.add_argument("--order", type=int, required=True)
    reorder.add_argument("--target-index", type=int, required=True)

    remove = subparsers.add_parser("remove")
    remove.add_argument("--order", type=int, required=True)

    return parser


COMANDOS = {
    "list": comando_listar,
    "automation-library-list": comando_listar_automacoes,
    "analyze-html-element": comando_analisar_html_elemento,
    "automation-create": comando_criar_automacao,
    "automation-open": comando_abrir_automacao,
    "automation-duplicate": comando_duplicar_automacao,
    "automation-rename": comando_renomear_automacao,
    "automation-delete": comando_excluir_automacao,
    "automation-mark-run": comando_marcar_execucao_automacao,
    "automation-export-zip": comando_exportar_automacao_zip,
    "automation-import-zip": comando_importar_automacao_zip,
    "runtime-info": comando_info_runtime,
    "runtime-export-zip": comando_exportar_runtime_zip,
    "runtime-import-zip": comando_importar_runtime_zip,
    "workbook-meta": comando_meta,
    "workbook-preview": comando_preview_planilha,
    "add-text": comando_adicionar_texto,
    "add-key": comando_adicionar_teclas,
    "add-mouse": comando_adicionar_mouse,
    "add-image-click": comando_adicionar_imagem,
    "add-verify-button-state": comando_adicionar_verificacao,
    "add-chrome-html": comando_adicionar_chrome_html,
    "add-open-url": comando_adicionar_abrir_url,
    "add-open-program": comando_adicionar_abrir_programa,
    "add-save-download": comando_adicionar_salvar_download,
    "add-planilha": comando_adicionar_planilha,
    "add-write-planilha": comando_adicionar_escrita_planilha,
    "add-compare-planilha": comando_adicionar_comparacao_planilha,
    "add-loop-start": comando_adicionar_inicio_loop,
    "add-loop-end": comando_adicionar_fim_loop,
    "insert-text": comando_inserir_texto,
    "insert-key": comando_inserir_teclas,
    "insert-mouse": comando_inserir_mouse,
    "insert-image-click": comando_inserir_imagem,
    "insert-verify-button-state": comando_inserir_verificacao,
    "insert-chrome-html": comando_inserir_chrome_html,
    "insert-open-url": comando_inserir_abrir_url,
    "insert-open-program": comando_inserir_abrir_programa,
    "insert-save-download": comando_inserir_salvar_download,
    "insert-planilha": comando_inserir_planilha,
    "insert-write-planilha": comando_inserir_escrita_planilha,
    "insert-compare-planilha": comando_inserir_comparacao_planilha,
    "insert-loop-start": comando_inserir_inicio_loop,
    "insert-loop-end": comando_inserir_fim_loop,
    "update-text": comando_atualizar_texto,
    "update-key": comando_atualizar_teclas,
    "update-mouse": comando_atualizar_mouse,
    "update-image-click": comando_atualizar_imagem,
    "update-verify-button-state": comando_atualizar_verificacao,
    "update-chrome-html": comando_atualizar_chrome_html,
    "update-open-url": comando_atualizar_abrir_url,
    "update-open-program": comando_atualizar_abrir_programa,
    "update-save-download": comando_atualizar_salvar_download,
    "update-planilha": comando_atualizar_planilha,
    "update-write-planilha": comando_atualizar_escrita_planilha,
    "update-compare-planilha": comando_atualizar_comparacao_planilha,
    "update-loop-marker": comando_atualizar_marcador_loop,
    "update-delay": comando_atualizar_atraso,
    "update-delay-all": comando_atualizar_atraso_global,
    "update-sequence-settings": comando_atualizar_configuracao_sequencia,
    "update-initial-sequence": comando_atualizar_sequencia_inicial,
    "sync-spreadsheets": comando_sincronizar_planilhas,
    "move": comando_mover,
    "reorder": comando_reordenar,
    "remove": comando_remover,
    "clear": comando_limpar,
}


def main():
    parser = build_parser()
    args = parser.parse_args()
    resposta = COMANDOS[args.command](args)
    payload = json.dumps(resposta, ensure_ascii=False)
    sys.stdout.buffer.write(payload.encode("utf-8"))


if __name__ == "__main__":
    main()
