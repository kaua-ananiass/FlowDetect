import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


ARQUIVO_MAPEAMENTO = Path("mapeamento_teclado.json")
ARQUIVO_PLANILHA = Path("db.xlsx")
PASTA_PLANILHAS = Path("planilhas")


def normalizar_id_planilha(arquivo=None):
    texto = str(arquivo or ARQUIVO_PLANILHA.name).strip().replace("\\", "/")
    return texto or ARQUIVO_PLANILHA.name


def resolver_planilha(arquivo=None):
    workbook_id = normalizar_id_planilha(arquivo)
    caminho = Path(workbook_id)
    if caminho.is_absolute():
        return caminho
    if caminho.exists():
        return caminho
    if "/" not in workbook_id and "\\" not in workbook_id and workbook_id != ARQUIVO_PLANILHA.name:
        alternativa = PASTA_PLANILHAS / workbook_id
        if alternativa.exists():
            return alternativa
    return caminho


def carregar_mapeamento(caminho):
    arquivo = Path(caminho)
    if not arquivo.exists():
        raise FileNotFoundError(f"Arquivo nao encontrado: {arquivo.resolve()}")
    return arquivo, json.loads(arquivo.read_text(encoding="utf-8"))


def identificar_coluna_empresa(colunas):
    candidatos_exatos = {"empresa", "razao social", "nome empresa", "cliente", "nome"}
    candidatos_parciais = ("empresa", "razao", "cliente")
    for coluna in colunas:
        nome = str(coluna).strip().lower()
        if nome in candidatos_exatos:
            return coluna
    for coluna in colunas:
        nome = str(coluna).strip().lower()
        if any(parte in nome for parte in candidatos_parciais):
            return coluna
    return None


def encontrar_evento_planilha(dados):
    for evento in dados.get("eventos", []):
        if evento.get("tipo") == "sequencia_planilha":
            return evento
    raise ValueError("Nenhuma sequencia de planilha foi encontrada no roteiro atual.")


def ler_cabecalhos(worksheet):
    return [
        str(valor).strip() if valor is not None else f"Coluna {indice + 1}"
        for indice, valor in enumerate(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), []))
    ]


def indice_coluna(colunas, nome_coluna):
    for indice, coluna in enumerate(colunas, start=1):
        if coluna == nome_coluna:
            return indice
    raise ValueError(f"Coluna nao encontrada na planilha: {nome_coluna}")


def listar_linhas_evento(arquivo_planilha, detalhes):
    workbook = load_workbook(arquivo_planilha, read_only=True, data_only=True)
    try:
        aba = detalhes.get("aba")
        coluna = detalhes.get("coluna")
        if not aba or not coluna:
            raise ValueError("Evento de planilha sem aba ou coluna configurada.")

        if aba not in workbook.sheetnames:
            raise ValueError(f"Aba nao encontrada na planilha: {aba}")

        worksheet = workbook[aba]
        colunas = ler_cabecalhos(worksheet)
        coluna_empresa = identificar_coluna_empresa(colunas)
        linhas_selecionadas = {int(linha) for linha in (detalhes.get("linhas_selecionadas") or [])}
        linhas = []

        for numero_linha, valores in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
            registro = {}
            for indice, nome_coluna in enumerate(colunas):
                registro[nome_coluna] = valores[indice] if indice < len(valores) else None

            valor = registro.get(coluna)
            if valor is None or str(valor).strip() == "":
                continue
            if linhas_selecionadas and numero_linha not in linhas_selecionadas:
                continue

            empresa = ""
            if coluna_empresa:
                valor_empresa = registro.get(coluna_empresa)
                if valor_empresa is not None and str(valor_empresa).strip():
                    empresa = str(valor_empresa).strip()

            linhas.append(
                {
                    "row_number": numero_linha,
                    "value": str(valor).strip(),
                    "company": empresa,
                }
            )

        return {
            "workbook_file": normalizar_id_planilha(detalhes.get("arquivo")),
            "sheet": aba,
            "column": coluna,
            "company_column": coluna_empresa or "",
            "origin": detalhes.get("origem", "coluna_inteira"),
            "rows": linhas,
        }
    finally:
        workbook.close()


def montar_texto_evento(coluna, quantidade, acao_texto):
    return f"[PLANILHA:{coluna} x{quantidade} | apos:{acao_texto}]"


def montar_comando_python(valores, acao_apos_item, intervalo_entre_itens):
    return (
        f"executar_sequencia({valores!r}, "
        f"acao_apos_item={acao_apos_item!r}, "
        f"intervalo_entre_itens={intervalo_entre_itens!r})"
    )


def montar_comando_texto(texto):
    return f"digitar_texto({texto!r})"


def sincronizar_eventos_planilha(arquivo_planilha, dados_mapeamento):
    workbook = load_workbook(arquivo_planilha, read_only=True, data_only=True)
    try:
        for evento in dados_mapeamento.get("eventos", []):
            if evento.get("tipo") != "sequencia_planilha":
                continue

            detalhes = evento.setdefault("detalhes", {})
            info = listar_linhas_evento(arquivo_planilha, detalhes)
            valores = []
            for linha in info["rows"]:
                texto = linha["value"]
                try:
                    valor_numero = int(texto)
                    valores.append(valor_numero)
                except ValueError:
                    try:
                        valor_decimal = float(texto.replace(",", "."))
                        valores.append(valor_decimal)
                    except ValueError:
                        valores.append(texto)

            detalhes["valores"] = valores
            coluna = detalhes.get("coluna", "coluna")
            acao_texto = detalhes.get("acao_apos_item_texto", "nenhuma")
            evento["texto"] = montar_texto_evento(coluna, len(valores), acao_texto)
            evento["comando_python"] = montar_comando_python(
                valores,
                detalhes.get("acao_apos_item", []),
                detalhes.get("intervalo_entre_itens", 0.5),
            )
    finally:
        workbook.close()


def ajustar_linhas_selecionadas_apos_remocao(dados_mapeamento, aba, linha_removida):
    for evento in dados_mapeamento.get("eventos", []):
        if evento.get("tipo") != "sequencia_planilha":
            continue
        detalhes = evento.get("detalhes", {})
        if detalhes.get("aba") != aba:
            continue
        linhas = []
        for linha in detalhes.get("linhas_selecionadas") or []:
            numero = int(linha)
            if numero == linha_removida:
                continue
            if numero > linha_removida:
                numero -= 1
            linhas.append(numero)
        detalhes["linhas_selecionadas"] = linhas


def salvar_mapeamento(arquivo_mapeamento, dados):
    arquivo_mapeamento.write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")


def listar_textos_livres(dados_mapeamento):
    itens = []
    for evento in dados_mapeamento.get("eventos", []):
        if evento.get("tipo") != "texto_livre":
            continue
        texto = str(evento.get("texto", ""))
        resumo = texto.replace("\r", " ").replace("\n", " / ").strip()
        itens.append(
            {
                "order": int(evento.get("ordem", 0)),
                "text": texto,
                "preview": resumo[:90],
            }
        )
    return {"items": itens}


def atualizar_texto_livre(dados_mapeamento, ordem, texto):
    for evento in dados_mapeamento.get("eventos", []):
        if evento.get("tipo") != "texto_livre":
            continue
        if int(evento.get("ordem", 0)) != int(ordem):
            continue
        evento["texto"] = texto
        evento["comando_python"] = montar_comando_texto(texto)
        return
    raise ValueError(f"Evento de texto livre nao encontrado para a ordem {ordem}.")


def adicionar_valor(arquivo_planilha, dados_mapeamento, valor, empresa=""):
    evento = encontrar_evento_planilha(dados_mapeamento)
    detalhes = evento.get("detalhes", {})
    aba = detalhes.get("aba")
    coluna = detalhes.get("coluna")

    workbook = load_workbook(arquivo_planilha)
    try:
        worksheet = workbook[aba]
        colunas = ler_cabecalhos(worksheet)
        indice_valor = indice_coluna(colunas, coluna)
        indice_empresa = None
        coluna_empresa = identificar_coluna_empresa(colunas)
        if coluna_empresa:
            indice_empresa = indice_coluna(colunas, coluna_empresa)

        nova_linha = worksheet.max_row + 1
        worksheet.cell(row=nova_linha, column=indice_valor, value=valor)
        if empresa and indice_empresa:
            worksheet.cell(row=nova_linha, column=indice_empresa, value=empresa)
        workbook.save(arquivo_planilha)
    finally:
        workbook.close()

    sincronizar_eventos_planilha(arquivo_planilha, dados_mapeamento)


def trocar_linhas_por_valor(worksheet, linha_a, linha_b):
    max_coluna = worksheet.max_column
    valores_a = [worksheet.cell(row=linha_a, column=coluna).value for coluna in range(1, max_coluna + 1)]
    valores_b = [worksheet.cell(row=linha_b, column=coluna).value for coluna in range(1, max_coluna + 1)]

    for coluna in range(1, max_coluna + 1):
        worksheet.cell(row=linha_a, column=coluna, value=valores_b[coluna - 1])
        worksheet.cell(row=linha_b, column=coluna, value=valores_a[coluna - 1])


def trocar_linhas_selecionadas(dados_mapeamento, aba, linha_a, linha_b):
    menor = min(int(linha_a), int(linha_b))
    maior = max(int(linha_a), int(linha_b))
    for evento in dados_mapeamento.get("eventos", []):
        if evento.get("tipo") != "sequencia_planilha":
            continue
        detalhes = evento.get("detalhes", {})
        if detalhes.get("aba") != aba:
            continue

        linhas_atuais = [int(linha) for linha in (detalhes.get("linhas_selecionadas") or [])]
        if not linhas_atuais:
            continue

        novas_linhas = []
        for numero in linhas_atuais:
            if numero == linha_a:
                novas_linhas.append(int(linha_b))
            elif numero == linha_b:
                novas_linhas.append(int(linha_a))
            else:
                novas_linhas.append(numero)
        detalhes["linhas_selecionadas"] = sorted(novas_linhas)


def remover_linha(arquivo_planilha, dados_mapeamento, row_number):
    evento = encontrar_evento_planilha(dados_mapeamento)
    detalhes = evento.get("detalhes", {})
    aba = detalhes.get("aba")

    workbook = load_workbook(arquivo_planilha)
    try:
        worksheet = workbook[aba]
        worksheet.delete_rows(int(row_number), 1)
        workbook.save(arquivo_planilha)
    finally:
        workbook.close()

    ajustar_linhas_selecionadas_apos_remocao(dados_mapeamento, aba, int(row_number))
    sincronizar_eventos_planilha(arquivo_planilha, dados_mapeamento)


def mover_linha(arquivo_planilha, dados_mapeamento, row_number, direcao):
    evento = encontrar_evento_planilha(dados_mapeamento)
    detalhes = evento.get("detalhes", {})
    aba = detalhes.get("aba")
    info = listar_linhas_evento(arquivo_planilha, detalhes)
    linhas = [item["row_number"] for item in info["rows"]]

    if int(row_number) not in linhas:
        raise ValueError("A linha informada nao faz parte da sequencia atual do loop.")

    indice_atual = linhas.index(int(row_number))
    deslocamento = -1 if direcao == "subir" else 1
    indice_destino = indice_atual + deslocamento
    if indice_destino < 0 or indice_destino >= len(linhas):
        raise ValueError("Nao e possivel mover este item nessa direcao.")

    linha_destino = linhas[indice_destino]

    workbook = load_workbook(arquivo_planilha)
    try:
        worksheet = workbook[aba]
        trocar_linhas_por_valor(worksheet, int(row_number), int(linha_destino))
        workbook.save(arquivo_planilha)
    finally:
        workbook.close()

    trocar_linhas_selecionadas(dados_mapeamento, aba, int(row_number), int(linha_destino))
    sincronizar_eventos_planilha(arquivo_planilha, dados_mapeamento)


def criar_parser():
    parser = argparse.ArgumentParser(description="Gerencia os dados usados pela automacao.")
    parser.add_argument("acao", choices=("listar", "adicionar", "remover", "mover", "listar_textos", "atualizar_texto"))
    parser.add_argument("--arquivo", default=str(ARQUIVO_MAPEAMENTO))
    parser.add_argument("--planilha", default=str(ARQUIVO_PLANILHA))
    parser.add_argument("--valor", default="")
    parser.add_argument("--empresa", default="")
    parser.add_argument("--linha", type=int, default=0)
    parser.add_argument("--direcao", choices=("subir", "descer"), default="subir")
    parser.add_argument("--ordem", type=int, default=0)
    parser.add_argument("--texto", default="")
    return parser


def main():
    args = criar_parser().parse_args()
    arquivo_mapeamento, dados_mapeamento = carregar_mapeamento(args.arquivo)
    arquivo_planilha = resolver_planilha(args.planilha)
    if not arquivo_planilha.exists():
        raise FileNotFoundError(f"Arquivo nao encontrado: {arquivo_planilha.resolve()}")

    evento_planilha = None
    arquivo_evento = None
    if args.acao in {"listar", "adicionar", "remover", "mover"}:
        evento_planilha = encontrar_evento_planilha(dados_mapeamento)
        arquivo_evento = resolver_planilha((evento_planilha.get("detalhes") or {}).get("arquivo") or args.planilha)

    if args.acao == "listar":
        print(json.dumps(listar_linhas_evento(arquivo_evento, evento_planilha.get("detalhes", {})), ensure_ascii=False))
        return 0

    if args.acao == "listar_textos":
        print(json.dumps(listar_textos_livres(dados_mapeamento), ensure_ascii=False))
        return 0

    if args.acao == "adicionar":
        if not str(args.valor).strip():
            raise ValueError("Informe um valor para adicionar ao loop.")
        adicionar_valor(arquivo_evento, dados_mapeamento, args.valor, args.empresa)
        salvar_mapeamento(arquivo_mapeamento, dados_mapeamento)
        evento = encontrar_evento_planilha(dados_mapeamento)
        arquivo_evento = resolver_planilha((evento.get("detalhes") or {}).get("arquivo") or args.planilha)
        print(json.dumps(listar_linhas_evento(arquivo_evento, evento.get("detalhes", {})), ensure_ascii=False))
        return 0

    if args.acao == "atualizar_texto":
        if args.ordem <= 0:
            raise ValueError("Informe a ordem do texto que deve ser alterado.")
        atualizar_texto_livre(dados_mapeamento, args.ordem, args.texto)
        salvar_mapeamento(arquivo_mapeamento, dados_mapeamento)
        print(json.dumps(listar_textos_livres(dados_mapeamento), ensure_ascii=False))
        return 0

    if args.linha <= 0:
        raise ValueError("Informe a linha da planilha que deve ser alterada.")

    if args.acao == "remover":
        remover_linha(arquivo_evento, dados_mapeamento, args.linha)
    else:
        mover_linha(arquivo_evento, dados_mapeamento, args.linha, args.direcao)

    salvar_mapeamento(arquivo_mapeamento, dados_mapeamento)
    evento = encontrar_evento_planilha(dados_mapeamento)
    arquivo_evento = resolver_planilha((evento.get("detalhes") or {}).get("arquivo") or args.planilha)
    print(json.dumps(listar_linhas_evento(arquivo_evento, evento.get("detalhes", {})), ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
