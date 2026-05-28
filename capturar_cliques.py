import ctypes
import json
import queue
import shutil
import threading
import time
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk

import pyautogui
from openpyxl import load_workbook
from pynput import keyboard, mouse as pynput_mouse


ARQUIVO_SAIDA = Path("mapeamento_teclado.json")
ARQUIVO_RESUMO = Path("resumo_teclado.txt")
ARQUIVO_SCRIPT = Path("comandos_teclado.py")
ARQUIVO_PLANILHA = Path("db.xlsx")
ARQUIVO_EXECUTOR_SEPARADO = Path("executar_automacao.py")
ARQUIVO_BAT_EXECUTOR = Path("executar_automacao.bat")
ARQUIVO_INTERFACE_EXECUTOR = Path("interface_automacao.py")
ARQUIVO_DADOS_PLANILHA = Path("gerenciar_dados_planilha.py")
PASTA_BUILD = Path("build_automacao")
PASTA_AUTOMACAO_EXPORTADA = Path("Automacao_Iss")
ARQUIVO_INICIAR_EXPORTACAO = "iniciar.bat"
INTERVALO_TEXTO = 0.01
PAUSA_ANTES_DE_DIGITAR = 0.3
INTERVALO_ENTRE_ITENS_PLANILHA = 0.5
TIPOS_SEQUENCIA_PLANILHA = {"sequencia_planilha"}
TIPOS_MOUSE = {"mouse_click"}
TECLAS_MODIFICADORAS = {
    "shift",
    "shift_r",
    "ctrl",
    "ctrl_r",
    "alt",
    "alt_r",
    "cmd",
    "cmd_r",
}
MAPA_NORMALIZACAO = {
    "shift_l": "shift",
    "shift_r": "shift_r",
    "ctrl_l": "ctrl",
    "ctrl_r": "ctrl_r",
    "alt_l": "alt",
    "alt_r": "alt_r",
    "cmd_l": "cmd",
    "cmd_r": "cmd_r",
}
MAPA_EXIBICAO = {
    "space": "SPACE",
    "enter": "ENTER",
    "tab": "TAB",
    "backspace": "BACKSPACE",
    "esc": "ESC",
    "shift": "SHIFT",
    "shift_r": "SHIFT_R",
    "ctrl": "CTRL",
    "ctrl_r": "CTRL_R",
    "alt": "ALT",
    "alt_r": "ALT_R",
    "cmd": "WIN",
    "cmd_r": "WIN_R",
    "up": "UP",
    "down": "DOWN",
    "left": "LEFT",
    "right": "RIGHT",
    "delete": "DELETE",
    "insert": "INSERT",
    "home": "HOME",
    "end": "END",
    "page_up": "PAGE_UP",
    "page_down": "PAGE_DOWN",
    "f8": "F8",
    "f9": "F9",
}
MAPA_PYAUTOGUI = {
    "space": "space",
    "enter": "enter",
    "tab": "tab",
    "backspace": "backspace",
    "esc": "esc",
    "shift": "shift",
    "shift_r": "shift",
    "ctrl": "ctrl",
    "ctrl_r": "ctrl",
    "alt": "alt",
    "alt_r": "alt",
    "cmd": "win",
    "cmd_r": "win",
    "up": "up",
    "down": "down",
    "left": "left",
    "right": "right",
    "delete": "delete",
    "insert": "insert",
    "home": "home",
    "end": "end",
    "page_up": "pageup",
    "page_down": "pagedown",
    "f8": "f8",
    "f9": "f9",
}


class GravadorTeclado:
    def __init__(self, arquivo_saida, arquivo_planilha=ARQUIVO_PLANILHA):
        self.arquivo_saida = Path(arquivo_saida)
        self.arquivo_resumo = ARQUIVO_RESUMO
        self.arquivo_script = ARQUIVO_SCRIPT
        self.arquivo_planilha = Path(arquivo_planilha)
        self.lock = threading.Lock()
        self.eventos = []
        self.teclas_pressionadas = set()
        self.teclas_principais_por_vk = {}
        self.modificadores_ativos = []
        self.modificadores_da_sequencia = []
        self.houve_tecla_principal = False
        self.captura_pausada = False
        self.encerrando = False
        self.fila_acoes = queue.Queue()
        self.listener = None
        self.mouse_listener = None
        self.root = None
        self.notebook = None
        self.janela_texto = None
        self.campo_texto = None
        self.janela_anterior = None
        self.digitando = threading.Lock()
        self.user32 = ctypes.windll.user32
        self.ultimo_evento_em = None
        self.instantes_eventos = []
        self.planilha_cache = {}
        self.planilha_ativa = None
        self.planilha_colunas = []
        self.planilha_linhas = []
        self.executando_roteiro = False
        self.cancelar_execucao = threading.Event()
        self.evento_selecionado_ordem = None
        self.evento_execucao_atual_ordem = None
        self.pausa_manual = False
        self.capturando_mouse = False

    def normalizar_nome_tecla(self, key):
        if hasattr(key, "char") and key.char is not None:
            if ord(key.char) < 32 and hasattr(key, "vk") and 32 <= key.vk <= 126:
                return chr(key.vk).lower()
            return key.char

        if hasattr(key, "name"):
            return MAPA_NORMALIZACAO.get(key.name, key.name)

        texto = str(key).replace("Key.", "")
        return MAPA_NORMALIZACAO.get(texto, texto)

    def tecla_para_exibicao(self, nome):
        if len(nome) == 1:
            if nome == " ":
                return "SPACE"
            return nome.upper()

        return MAPA_EXIBICAO.get(nome, nome.upper())

    def tecla_para_pyautogui(self, nome):
        if len(nome) == 1:
            return nome

        return MAPA_PYAUTOGUI.get(nome, nome)

    def eh_modificadora(self, nome):
        return nome in TECLAS_MODIFICADORAS

    def eh_texto_simples(self, nome):
        return len(nome) == 1 and nome not in {"\r", "\n", "\t"}

    def normalizar_tecla_principal(self, key, nome):
        if self.modificadores_ativos and hasattr(key, "vk") and 32 <= key.vk <= 126:
            return chr(key.vk).lower()

        return nome

    def formatar_evento(self, teclas, tipo_evento, texto_livre=None, metadados=None):
        if tipo_evento == "texto_livre":
            return texto_livre or ""

        if tipo_evento in TIPOS_SEQUENCIA_PLANILHA:
            coluna = (metadados or {}).get("coluna", "Coluna")
            quantidade = len((metadados or {}).get("valores", []))
            acao = (metadados or {}).get("acao_apos_item_texto") or "nenhuma"
            return f"[PLANILHA:{coluna} x{quantidade} | apos:{acao}]"

        if tipo_evento in TIPOS_MOUSE:
            detalhes = metadados or {}
            x = detalhes.get("x", 0)
            y = detalhes.get("y", 0)
            botao = str(detalhes.get("botao", "left")).upper()
            return f"[MOUSE:{botao} x={x} y={y}]"

        if len(teclas) == 1:
            nome = teclas[0]
            exibicao = self.tecla_para_exibicao(nome)
            if nome == "enter":
                return "[ENTER]\n"
            if self.eh_texto_simples(nome):
                return nome
            return f"[{exibicao}]"

        partes = [self.tecla_para_exibicao(tecla) for tecla in teclas]
        return f"[{'+'.join(partes)}]"

    def montar_resumo(self):
        return "".join(evento["texto"] for evento in self.eventos)

    def montar_script_python(self):
        linhas = [
            "import pyautogui",
            "import time",
            "",
            f"INTERVALO_TEXTO = {INTERVALO_TEXTO}",
            "",
            "",
            "def digitar_texto(texto):",
            "    for caractere in texto:",
            "        if caractere == \"\\r\":",
            "            continue",
            "        if caractere == \"\\n\":",
            "            pyautogui.press(\"enter\")",
            "            continue",
            "        if caractere == \"\\t\":",
            "            pyautogui.press(\"tab\")",
            "            continue",
            "        pyautogui.write(caractere, interval=INTERVALO_TEXTO)",
            "",
            "",
            "def executar_acao(acao):",
            "    if not acao:",
            "        return",
            "    if isinstance(acao, str):",
            "        acao = [parte.strip() for parte in acao.split(\"+\") if parte.strip()]",
            "    if len(acao) == 1:",
            "        pyautogui.press(str(acao[0]))",
            "        return",
            "    pyautogui.hotkey(*[str(parte) for parte in acao])",
            "",
            "",
            "def executar_sequencia(valores, acao_apos_item=None, intervalo_entre_itens=0.5):",
            "    total = len(valores)",
            "    for indice, valor in enumerate(valores):",
            "        digitar_texto(str(valor))",
            "        if indice < total - 1 and acao_apos_item:",
            "            executar_acao(acao_apos_item)",
            "        if indice < total - 1 and intervalo_entre_itens > 0:",
            "            time.sleep(intervalo_entre_itens)",
            "",
            "",
            "def executar():",
        ]

        if not self.eventos:
            linhas.append("    pass")
        else:
            for evento in self.eventos:
                intervalo_antes = evento.get("intervalo_antes", 0)
                if intervalo_antes > 0:
                    linhas.append(f"    time.sleep({intervalo_antes})")
                linhas.append(f"    {evento['comando_python']}")

        linhas.extend(
            [
                "",
                "",
                'if __name__ == "__main__":',
                "    executar()",
            ]
        )
        return "\n".join(linhas) + "\n"

    def salvar_arquivo(self):
        dados = {
            "criado_em": datetime.now().isoformat(timespec="seconds"),
            "instrucoes": [
                "Teclas simples e atalhos sao gravados individualmente.",
                "F8 abre uma caixa de texto grande e salva o bloco como um unico comando.",
                "O botao Capturar mouse salva a coordenada exata do proximo clique.",
                "Cada evento tambem salva o intervalo em segundos desde a acao anterior.",
                "A aba Planilha pode adicionar uma sequencia inteira ou so linhas selecionadas ao roteiro.",
                "Backspace remove a ultima entrada salva e F9 ou ESC encerram a captura.",
            ],
            "resumo_digitado": self.montar_resumo(),
            "eventos": self.eventos,
        }
        self.arquivo_saida.write_text(
            json.dumps(dados, indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        self.arquivo_resumo.write_text(
            self.montar_resumo(),
            encoding="utf-8",
        )
        self.arquivo_script.write_text(
            self.montar_script_python(),
            encoding="utf-8",
        )

    def montar_comando_python(self, teclas, tipo_evento, texto_livre=None, metadados=None):
        if tipo_evento == "texto_livre":
            return f"digitar_texto({texto_livre!r})"

        if tipo_evento in TIPOS_SEQUENCIA_PLANILHA:
            detalhes = metadados or {}
            valores = detalhes.get("valores", [])
            acao_apos_item = detalhes.get("acao_apos_item", [])
            intervalo_entre_itens = detalhes.get(
                "intervalo_entre_itens",
                INTERVALO_ENTRE_ITENS_PLANILHA,
            )
            return (
                f"executar_sequencia({valores!r}, "
                f"acao_apos_item={acao_apos_item!r}, "
                f"intervalo_entre_itens={intervalo_entre_itens!r})"
            )

        if tipo_evento in TIPOS_MOUSE:
            detalhes = metadados or {}
            x = int(detalhes.get("x", 0))
            y = int(detalhes.get("y", 0))
            botao = str(detalhes.get("botao", "left"))
            return f'pyautogui.click(x={x}, y={y}, button="{botao}")'

        teclas_pyautogui = [self.tecla_para_pyautogui(tecla) for tecla in teclas]

        if len(teclas_pyautogui) > 1:
            argumentos = ", ".join(f'"{tecla}"' for tecla in teclas_pyautogui)
            return f"pyautogui.hotkey({argumentos})"

        tecla = teclas[0]
        tecla_pyautogui = teclas_pyautogui[0]

        if self.eh_texto_simples(tecla):
            return f"pyautogui.write({tecla!r})"

        return f'pyautogui.press("{tecla_pyautogui}")'

    def registrar_evento(self, teclas, tipo_evento, texto_livre=None, metadados=None):
        texto = self.formatar_evento(teclas, tipo_evento, texto_livre, metadados)
        comando_python = self.montar_comando_python(
            teclas,
            tipo_evento,
            texto_livre,
            metadados,
        )
        agora = time.perf_counter()
        intervalo_antes = 0 if self.ultimo_evento_em is None else round(agora - self.ultimo_evento_em, 3)

        with self.lock:
            ordem = len(self.eventos) + 1
            item = {
                "ordem": ordem,
                "tipo": tipo_evento,
                "teclas": teclas,
                "texto": texto,
                "intervalo_antes": intervalo_antes,
                "comando_python": comando_python,
            }
            if metadados:
                item["detalhes"] = metadados
            self.eventos.append(item)
            self.instantes_eventos.append(agora)
            self.ultimo_evento_em = agora
            self.salvar_arquivo()

        self.enfileirar_ui(
            "evento_registrado",
            {
                "mensagem": f"[{ordem}] {comando_python} | intervalo: {intervalo_antes}s",
                "tag": "ok",
            },
        )

    def remover_ultima_entrada(self):
        with self.lock:
            if not self.eventos:
                self.enfileirar_ui(
                    "log",
                    {"mensagem": "Nenhuma entrada para remover.", "tag": "alerta"},
                )
                return

            removida = self.eventos.pop()
            if self.instantes_eventos:
                self.instantes_eventos.pop()
            self.ultimo_evento_em = self.instantes_eventos[-1] if self.instantes_eventos else None
            self.salvar_arquivo()

        self.enfileirar_ui(
            "evento_removido",
            {
                "mensagem": f"Entrada deletada: {removida['comando_python']}",
                "tag": "removido",
            },
        )

    def mover_evento_selecionado(self, direcao):
        if self.evento_selecionado_ordem is None:
            self.adicionar_log("Selecione um item do roteiro para mover.", "alerta")
            return

        with self.lock:
            indice = next(
                (i for i, e in enumerate(self.eventos) if e.get("ordem") == self.evento_selecionado_ordem),
                None,
            )
            if indice is None:
                return

            novo_indice = indice + direcao
            if novo_indice < 0 or novo_indice >= len(self.eventos):
                return

            self.eventos[indice], self.eventos[novo_indice] = self.eventos[novo_indice], self.eventos[indice]

            for nova_ordem, evento in enumerate(self.eventos, start=1):
                evento["ordem"] = nova_ordem

            self.evento_selecionado_ordem = self.eventos[novo_indice]["ordem"]
            self.salvar_arquivo()

        nova_ordem = self.evento_selecionado_ordem
        self.enfileirar_ui("evento_registrado", {
            "mensagem": f"Evento movido para posicao {nova_ordem}.",
            "tag": "info",
        })
        self.enfileirar_ui("selecionar_evento", {"ordem": nova_ordem})

    def remover_evento_selecionado(self):
        if self.evento_selecionado_ordem is None:
            self.adicionar_log("Selecione um item do roteiro para apagar.", "alerta")
            return

        with self.lock:
            indice = next(
                (i for i, e in enumerate(self.eventos) if e.get("ordem") == self.evento_selecionado_ordem),
                None,
            )
            if indice is None:
                self.adicionar_log("Evento selecionado nao encontrado.", "alerta")
                return

            removido = self.eventos.pop(indice)
            if indice < len(self.instantes_eventos):
                self.instantes_eventos.pop(indice)
            self.ultimo_evento_em = self.instantes_eventos[-1] if self.instantes_eventos else None

            for nova_ordem, evento in enumerate(self.eventos, start=1):
                evento["ordem"] = nova_ordem

            self.salvar_arquivo()

        self.evento_selecionado_ordem = None
        self.enfileirar_ui(
            "evento_removido",
            {
                "mensagem": f"Evento apagado: {removido['comando_python']}",
                "tag": "removido",
            },
        )

    def adicionar_modificadora(self, nome):
        if nome not in self.modificadores_ativos:
            self.modificadores_ativos.append(nome)

        if nome not in self.modificadores_da_sequencia:
            self.modificadores_da_sequencia.append(nome)

    def remover_modificadora(self, nome):
        if nome in self.modificadores_ativos:
            self.modificadores_ativos.remove(nome)

    def resetar_estado_captura(self):
        self.teclas_pressionadas.clear()
        self.teclas_principais_por_vk.clear()
        self.modificadores_ativos = []
        self.modificadores_da_sequencia = []
        self.houve_tecla_principal = False

    def finalizar_sequencia(self):
        if not self.houve_tecla_principal and self.modificadores_da_sequencia:
            self.registrar_evento(list(self.modificadores_da_sequencia), "atalho")

        self.modificadores_ativos = []
        self.modificadores_da_sequencia = []
        self.houve_tecla_principal = False

    def capturar_janela_ativa(self):
        return self.user32.GetForegroundWindow()

    def restaurar_janela_ativa(self, hwnd):
        if not hwnd:
            return

        self.user32.ShowWindow(hwnd, 5)
        self.user32.SetForegroundWindow(hwnd)

    def janela_do_app_esta_ativa(self):
        if self.root is None:
            return False

        try:
            hwnd_atual = self.capturar_janela_ativa()
            janelas = {self.root.winfo_id()}
            if self.janela_texto is not None and self.janela_texto.winfo_exists():
                janelas.add(self.janela_texto.winfo_id())
            return hwnd_atual in janelas
        except tk.TclError:
            return False

    def solicitar_abertura_texto(self):
        self.fila_acoes.put(("abrir_texto", None))

    def solicitar_encerramento(self):
        self.fila_acoes.put(("encerrar", None))

    def enfileirar_ui(self, acao, payload=None):
        if self.root is not None:
            self.fila_acoes.put((acao, payload))

    def processar_fila(self):
        while True:
            try:
                acao, payload = self.fila_acoes.get_nowait()
            except queue.Empty:
                break

            if acao == "abrir_texto":
                self.abrir_janela_texto()
            elif acao == "encerrar":
                self.encerrar()
            elif acao == "evento_registrado":
                self.atualizar_interface_eventos()
                self.adicionar_log(payload["mensagem"], payload.get("tag", "ok"))
            elif acao == "evento_removido":
                self.atualizar_interface_eventos()
                self.adicionar_log(payload["mensagem"], payload.get("tag", "removido"))
            elif acao == "log":
                self.adicionar_log(payload["mensagem"], payload.get("tag", "info"))
            elif acao == "selecionar_evento":
                iid = str(payload["ordem"])
                if self.lista_eventos.exists(iid):
                    self.lista_eventos.selection_set(iid)
                    self.lista_eventos.see(iid)
            elif acao == "status_execucao":
                self.atualizar_status_execucao(payload["mensagem"])
            elif acao == "detalhe_execucao":
                self.atualizar_detalhe_execucao(
                    payload["mensagem"],
                    payload.get("ordem"),
                )
            elif acao == "restaurar_janela":
                try:
                    self.root.deiconify()
                    self.root.lift()
                except tk.TclError:
                    pass
            elif acao == "mouse_capturado":
                self.concluir_captura_mouse(payload)
            elif acao == "mouse_erro":
                self.finalizar_captura_mouse_com_erro(payload["mensagem"])

        if not self.encerrando and self.root is not None:
            self.root.after(100, self.processar_fila)

    def carregar_eventos_existentes(self):
        if not self.arquivo_saida.exists():
            return

        try:
            dados = json.loads(self.arquivo_saida.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return

        eventos = dados.get("eventos")
        if not isinstance(eventos, list):
            return

        self.eventos = []
        for ordem, evento in enumerate(eventos, start=1):
            tipo = evento.get("tipo", "tecla")
            item = {
                "ordem": ordem,
                "tipo": tipo,
                "teclas": evento.get("teclas", []),
                "texto": evento.get("texto", ""),
                "intervalo_antes": evento.get("intervalo_antes", 0),
                "comando_python": evento.get("comando_python", "pass"),
            }
            if "detalhes" in evento and isinstance(evento["detalhes"], dict):
                item["detalhes"] = evento["detalhes"]
            self.eventos.append(item)

    def resumir_para_lista(self, texto, limite=60):
        texto = str(texto).replace("\n", "\\n")
        if len(texto) <= limite:
            return texto
        return f"{texto[: limite - 3]}..."

    def normalizar_acao_apos_item(self, acao_texto):
        texto = (acao_texto or "").strip().lower()
        if not texto or texto == "nenhuma":
            return []
        return [parte.strip() for parte in texto.replace(" ", "").split("+") if parte.strip()]

    def montar_launcher_exportacao(self):
        return (
            "@echo off\n"
            "setlocal\n"
            "cd /d \"%~dp0\"\n\n"
            "if not exist \"Automacao.exe\" (\n"
            "    echo Arquivo Automacao.exe nao encontrado nesta pasta.\n"
            "    pause\n"
            "    endlocal\n"
            "    exit /b 1\n"
            ")\n\n"
            "start \"\" \"%~dp0Automacao.exe\"\n"
            "endlocal\n"
        )

    def montar_arquivo_guia_build(self):
        return (
            "EXECUTOR SEPARADO DA AUTOMACAO\n"
            "==============================\n\n"
            "Arquivos desta pasta:\n"
            "- executar_automacao.bat: abre a interface da automacao.\n"
            "- Automacao.exe: painel visual em Tauri com HTML/CSS, quando disponivel.\n"
            "- interface_automacao.py: fallback visual em Python com progresso e empresa atual.\n"
            "- executar_automacao.py: executor separado do roteiro gravado.\n"
            "- gerenciar_dados_planilha.py: utilitario para listar e editar os itens do loop na planilha.\n"
            "- mapeamento_teclado.json: roteiro completo salvo pela interface.\n"
            "- db.xlsx: planilha usada para ler os dados mais recentes.\n\n"
            "Como usar:\n"
            "1. Edite a planilha db.xlsx se precisar atualizar os dados.\n"
            "2. Clique duas vezes em executar_automacao.bat.\n"
            "3. Escolha entre execucao normal ou em loop e ajuste o atraso inicial.\n"
            "4. Acompanhe a barra de porcentagem e a empresa atual durante a execucao.\n"
            "5. Clique em Iniciar ou Parar quando precisar.\n"
        )

    def montar_arquivo_guia_exportacao(self):
        return (
            "AUTOMACAO EXPORTADA SEPARADA DO MAIN\n"
            "====================================\n\n"
            "Esta pasta contem apenas os arquivos necessarios para rodar a automacao gerada.\n"
            "O programa principal do projeto continua separado no main.py.\n\n"
            "Arquivos principais desta pasta:\n"
            "- iniciar.bat: abre a interface principal da automacao.\n"
            "- Automacao.exe: interface visual da automacao.\n"
            "- executar_automacao.py: executor do roteiro.\n"
            "- gerenciar_dados_planilha.py: edita dados do loop e textos livres.\n"
            "- mapeamento_teclado.json: roteiro atual salvo.\n"
            "- db.xlsx: base de dados usada na execucao.\n\n"
            "Como usar:\n"
            "1. Clique duas vezes em iniciar.bat.\n"
            "2. Ajuste o atraso inicial, se precisar.\n"
            "3. Use Dados do loop para adicionar, remover ou reordenar itens.\n"
            "4. Use Textos livres para alterar os textos digitados.\n\n"
            "Observacao:\n"
            "- O Python precisa estar instalado no Windows para a parte de execucao funcionar.\n"
        )

    def exportar_arquivos_automacao(self, pasta_destino, arquivos_para_copiar, executaveis_tauri):
        pasta_destino.mkdir(exist_ok=True)

        copiados = []
        faltando = []
        for arquivo in arquivos_para_copiar:
            if arquivo.exists():
                shutil.copy2(arquivo, pasta_destino / arquivo.name)
                copiados.append(arquivo.name)
            else:
                faltando.append(arquivo.name)

        for executavel in executaveis_tauri:
            if executavel.exists():
                destino = pasta_destino / "Automacao.exe"
                shutil.copy2(executavel, destino)
                copiados.append(destino.name)
                break

        return copiados, faltando

    def atualizar_cache_planilha_silenciosamente(self):
        if not self.arquivo_planilha.exists():
            return

        try:
            self.carregar_planilha()
        except Exception:
            pass

    def gerar_build_automacao(self):
        self.atualizar_cache_planilha_silenciosamente()
        self.salvar_arquivo()
        pasta_build = PASTA_BUILD
        pasta_exportacao = PASTA_AUTOMACAO_EXPORTADA

        arquivos_para_copiar = [
            self.arquivo_saida,
            self.arquivo_resumo,
            self.arquivo_script,
            self.arquivo_planilha,
            ARQUIVO_EXECUTOR_SEPARADO,
            ARQUIVO_BAT_EXECUTOR,
            ARQUIVO_INTERFACE_EXECUTOR,
            ARQUIVO_DADOS_PLANILHA,
            Path("requirements.txt"),
        ]
        executaveis_tauri = [
            Path("tauri_ui/src-tauri/target/release/Automacao.exe"),
            Path("tauri_ui/src-tauri/target/release/automacao_tauri_ui.exe"),
        ]

        copiados_build, faltando_build = self.exportar_arquivos_automacao(
            pasta_build,
            arquivos_para_copiar,
            executaveis_tauri,
        )
        copiados_exportacao, faltando_exportacao = self.exportar_arquivos_automacao(
            pasta_exportacao,
            arquivos_para_copiar,
            executaveis_tauri,
        )

        (pasta_build / "COMO_EXECUTAR_BUILD.txt").write_text(
            self.montar_arquivo_guia_build(),
            encoding="utf-8",
        )
        (pasta_exportacao / ARQUIVO_INICIAR_EXPORTACAO).write_text(
            self.montar_launcher_exportacao(),
            encoding="utf-8",
        )
        (pasta_exportacao / "COMO_EXECUTAR.txt").write_text(
            self.montar_arquivo_guia_exportacao(),
            encoding="utf-8",
        )

        mensagem = (
            "Automacao exportada separada do main.py.\n\n"
            f"Pasta tecnica de build:\n{pasta_build.resolve()}\n\n"
            f"Pasta final separada da automacao:\n{pasta_exportacao.resolve()}\n\n"
            f"Arquivos copiados para a exportacao: {', '.join(copiados_exportacao) if copiados_exportacao else 'nenhum'}"
        )
        faltando_total = sorted(set(faltando_build + faltando_exportacao))
        if faltando_total:
            mensagem += f"\n\nArquivos nao encontrados: {', '.join(faltando_total)}"

        self.adicionar_log(
            f"Automacao separada gerada em {pasta_exportacao.resolve()}.",
            "ok",
        )
        messagebox.showinfo("Build da automacao", mensagem)

    def adicionar_log(self, mensagem, tag="info"):
        if self.log_texto is None:
            return

        horario = datetime.now().strftime("%H:%M:%S")
        self.log_texto.configure(state="normal")
        self.log_texto.insert("end", f"[{horario}] {mensagem}\n", tag)
        self.log_texto.see("end")
        self.log_texto.configure(state="disabled")

    def atualizar_status_execucao(self, mensagem):
        if self.status_execucao_var is not None:
            self.status_execucao_var.set(mensagem)

    def atualizar_detalhe_execucao(self, mensagem, ordem=None):
        if self.execucao_detalhe_var is not None:
            self.execucao_detalhe_var.set(mensagem)
        self.destacar_evento_execucao(ordem)

    def destacar_evento_execucao(self, ordem):
        if self.lista_eventos is None:
            self.evento_execucao_atual_ordem = ordem
            return

        if self.evento_execucao_atual_ordem is not None:
            iid_anterior = str(self.evento_execucao_atual_ordem)
            if self.lista_eventos.exists(iid_anterior):
                self.lista_eventos.item(iid_anterior, tags=())

        self.evento_execucao_atual_ordem = ordem
        if ordem is None:
            return

        iid = str(ordem)
        if self.lista_eventos.exists(iid):
            self.lista_eventos.item(iid, tags=("executando",))
            self.lista_eventos.selection_set(iid)
            self.lista_eventos.see(iid)

    def descrever_evento_em_execucao(
        self,
        evento,
        indice_item=None,
        total_itens=None,
        valor_atual=None,
        iteracao_atual=None,
        total_iteracoes=None,
        aguardando_intervalo=False,
    ):
        ordem = evento.get("ordem", "?")
        tipo = evento.get("tipo", "tecla")
        detalhes = evento.get("detalhes", {})
        partes = []

        if iteracao_atual is not None and total_iteracoes is not None:
            partes.append(f"Loop {iteracao_atual}/{total_iteracoes}")

        partes.append(f"Evento #{ordem}")

        if tipo == "texto_livre":
            conteudo = self.resumir_para_lista(evento.get("texto", ""), 95)
            detalhe = f"Texto livre: {conteudo}"
        elif tipo in TIPOS_SEQUENCIA_PLANILHA:
            coluna = detalhes.get("coluna", "Coluna")
            total_valores = total_itens if total_itens is not None else len(detalhes.get("valores", []))
            if valor_atual is not None and indice_item is not None:
                valor = self.resumir_para_lista(valor_atual, 80)
                detalhe = f"Planilha {coluna} item {indice_item + 1}/{total_valores}: {valor}"
            else:
                detalhe = f"Planilha {coluna} com {total_valores} itens"
        elif tipo in TIPOS_MOUSE:
            x = int(detalhes.get("x", 0))
            y = int(detalhes.get("y", 0))
            botao = str(detalhes.get("botao", "left")).upper()
            detalhe = f"Mouse {botao} em X={x} Y={y}"
        else:
            conteudo = evento.get("texto") or self.formatar_evento(evento.get("teclas", []), tipo)
            detalhe = self.resumir_para_lista(conteudo, 95)

        if aguardando_intervalo:
            detalhe += f" | aguardando {evento.get('intervalo_antes', 0)}s"

        partes.append(detalhe)
        return " | ".join(partes)

    def sinalizar_evento_em_execucao(
        self,
        evento,
        indice_item=None,
        total_itens=None,
        valor_atual=None,
        iteracao_atual=None,
        total_iteracoes=None,
        aguardando_intervalo=False,
        registrar_log=False,
    ):
        descricao = self.descrever_evento_em_execucao(
            evento,
            indice_item=indice_item,
            total_itens=total_itens,
            valor_atual=valor_atual,
            iteracao_atual=iteracao_atual,
            total_iteracoes=total_iteracoes,
            aguardando_intervalo=aguardando_intervalo,
        )
        status = "Aguardando proximo passo..." if aguardando_intervalo else "Executando passo do roteiro..."
        self.enfileirar_ui("status_execucao", {"mensagem": status})
        self.enfileirar_ui(
            "detalhe_execucao",
            {"mensagem": f"Passo atual: {descricao}", "ordem": evento.get("ordem")},
        )
        if registrar_log:
            self.enfileirar_ui("log", {"mensagem": f"Executando {descricao}", "tag": "info"})

    def atualizar_interface_eventos(self):
        self.lista_eventos.delete(*self.lista_eventos.get_children())
        for evento in self.eventos:
            detalhes = evento.get("detalhes", {})
            conteudo = evento["texto"]
            if evento["tipo"] in TIPOS_SEQUENCIA_PLANILHA:
                conteudo = self.formatar_evento(
                    evento.get("teclas", []),
                    evento["tipo"],
                    metadados=detalhes,
                )

            self.lista_eventos.insert(
                "",
                "end",
                iid=str(evento["ordem"]),
                values=(
                    evento["ordem"],
                    evento["tipo"],
                    self.resumir_para_lista(conteudo, 42),
                    evento.get("intervalo_antes", 0),
                    self.resumir_para_lista(evento["comando_python"], 58),
                ),
            )
            if evento["ordem"] == self.evento_execucao_atual_ordem:
                self.lista_eventos.item(str(evento["ordem"]), tags=("executando",))

        resumo = self.montar_resumo()
        self.resumo_texto.configure(state="normal")
        self.resumo_texto.delete("1.0", "end")
        self.resumo_texto.insert("1.0", resumo if resumo else "Nenhum evento gravado ainda.")
        self.resumo_texto.configure(state="disabled")

        total = len(self.eventos)
        total_texto_grande = sum(1 for evento in self.eventos if evento["tipo"] == "texto_livre")
        total_planilha = sum(1 for evento in self.eventos if evento["tipo"] in TIPOS_SEQUENCIA_PLANILHA)
        total_mouse = sum(1 for evento in self.eventos if evento["tipo"] in TIPOS_MOUSE)
        sufixo_status = " | captura pausada" if self.pausa_manual else " | captura ativa"
        self.status_var.set(
            f"{total} eventos no roteiro | {total_texto_grande} blocos de texto | {total_planilha} sequencias de planilha | {total_mouse} cliques do mouse{sufixo_status}"
        )
        self.atualizar_editor_tempo()

    def obter_evento_por_ordem(self, ordem):
        for evento in self.eventos:
            if evento.get("ordem") == ordem:
                return evento
        return None

    def ao_selecionar_evento(self, event=None):
        selecao = self.lista_eventos.selection()
        if not selecao:
            self.evento_selecionado_ordem = None
            self.atualizar_editor_tempo()
            return

        try:
            self.evento_selecionado_ordem = int(selecao[0])
        except ValueError:
            self.evento_selecionado_ordem = None
        self.atualizar_editor_tempo()

    def atualizar_editor_tempo(self):
        evento = self.obter_evento_por_ordem(self.evento_selecionado_ordem) if self.evento_selecionado_ordem else None
        if evento is None:
            self.evento_tempo_info_var.set("Selecione um item do roteiro para editar os tempos.")
            self.intervalo_evento_var.set(0.0)
            self.intervalo_planilha_evento_var.set(0.5)
            self.entrada_intervalo_planilha_evento.configure(state="disabled")
            return

        self.intervalo_evento_var.set(float(evento.get("intervalo_antes", 0)))
        if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA:
            detalhes = evento.get("detalhes", {})
            self.intervalo_planilha_evento_var.set(float(detalhes.get("intervalo_entre_itens", INTERVALO_ENTRE_ITENS_PLANILHA)))
            self.entrada_intervalo_planilha_evento.configure(state="normal")
            self.evento_tempo_info_var.set(
                f"Evento #{evento['ordem']} selecionado. Ajuste o tempo antes dele e o tempo interno da sequencia."
            )
        else:
            self.intervalo_planilha_evento_var.set(0.5)
            self.entrada_intervalo_planilha_evento.configure(state="disabled")
            self.evento_tempo_info_var.set(
                f"Evento #{evento['ordem']} selecionado. Ajuste o tempo antes deste comando."
            )

    def aplicar_tempo_evento_selecionado(self):
        evento = self.obter_evento_por_ordem(self.evento_selecionado_ordem) if self.evento_selecionado_ordem else None
        if evento is None:
            messagebox.showwarning("Tempo", "Selecione um item do roteiro antes de alterar o tempo.")
            return

        try:
            novo_intervalo = round(float(self.intervalo_evento_var.get()), 3)
        except (tk.TclError, ValueError):
            messagebox.showwarning("Tempo", "Informe um intervalo valido para o evento.")
            return

        if novo_intervalo < 0:
            messagebox.showwarning("Tempo", "O intervalo antes do evento nao pode ser negativo.")
            return

        evento["intervalo_antes"] = novo_intervalo

        if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA:
            try:
                novo_intervalo_interno = round(float(self.intervalo_planilha_evento_var.get()), 3)
            except (tk.TclError, ValueError):
                messagebox.showwarning("Tempo", "Informe um intervalo interno valido para a sequencia da planilha.")
                return

            if novo_intervalo_interno < 0:
                messagebox.showwarning("Tempo", "O intervalo interno da sequencia nao pode ser negativo.")
                return

            detalhes = evento.setdefault("detalhes", {})
            detalhes["intervalo_entre_itens"] = novo_intervalo_interno
            evento["comando_python"] = self.montar_comando_python(
                evento.get("teclas", []),
                evento.get("tipo", "tecla"),
                texto_livre=evento.get("texto", ""),
                metadados=detalhes,
            )

        self.salvar_arquivo()
        self.atualizar_interface_eventos()
        self.adicionar_log(
            f"Tempo do evento #{evento['ordem']} atualizado para {novo_intervalo}s.",
            "ok",
        )

    def configurar_estilo(self):
        self.root.title("Captura de Atalhos e Automacao")
        self.root.geometry("1380x820")
        self.root.minsize(1120, 720)
        self.root.configure(bg="#EEF2F7")

        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.configure("TFrame", background="#EEF2F7")
        style.configure("Card.TFrame", background="#FFFFFF")
        style.configure("TNotebook", background="#EEF2F7", borderwidth=0)
        style.configure("TNotebook.Tab", padding=(18, 10), font=("Segoe UI Semibold", 10))
        style.configure("Title.TLabel", background="#EEF2F7", foreground="#16324F", font=("Segoe UI Semibold", 17))
        style.configure("Sub.TLabel", background="#EEF2F7", foreground="#55657A", font=("Segoe UI", 10))
        style.configure("Section.TLabelframe", background="#FFFFFF")
        style.configure("Section.TLabelframe.Label", background="#FFFFFF", foreground="#16324F", font=("Segoe UI Semibold", 10))
        style.configure("Primary.TButton", font=("Segoe UI Semibold", 10))
        style.configure("Treeview", font=("Consolas", 10), rowheight=28)
        style.configure("Treeview.Heading", font=("Segoe UI Semibold", 10))

    def construir_interface(self):
        self.status_var = tk.StringVar(value="Carregando interface...")
        self.status_execucao_var = tk.StringVar(value="Automacao pronta para executar.")
        self.execucao_detalhe_var = tk.StringVar(value="Passo atual: nenhum evento em execucao ainda.")
        self.evento_tempo_info_var = tk.StringVar(value="Selecione um item do roteiro para editar os tempos.")
        self.botao_pausa_var = tk.StringVar(value="Pausar captura")
        self.mouse_posicao_var = tk.StringVar(value="Nenhum clique do mouse capturado ainda.")
        self.planilha_info_var = tk.StringVar(value="Planilha ainda nao carregada.")
        self.coluna_info_var = tk.StringVar(value="Selecione uma coluna para ver os valores.")
        self.planilha_aba_var = tk.StringVar()
        self.coluna_planilha_var = tk.StringVar()
        self.acao_apos_item_var = tk.StringVar(value="enter")
        self.planilha_intervalo_var = tk.DoubleVar(value=INTERVALO_ENTRE_ITENS_PLANILHA)
        self.atraso_execucao_var = tk.DoubleVar(value=3.0)
        self.intervalo_evento_var = tk.DoubleVar(value=0.0)
        self.intervalo_planilha_evento_var = tk.DoubleVar(value=INTERVALO_ENTRE_ITENS_PLANILHA)

        self.configurar_estilo()

        container = ttk.Frame(self.root, padding=18)
        container.pack(fill="both", expand=True)

        topo = ttk.Frame(container)
        topo.pack(fill="x", pady=(0, 14))

        ttk.Label(topo, text="Automacao com Captura de Teclas e Mouse", style="Title.TLabel").pack(anchor="w")
        ttk.Label(
            topo,
            text="A captura roda em segundo plano, mostra o historico visualmente, captura cliques do mouse e conversa com a planilha db.xlsx.",
            style="Sub.TLabel",
        ).pack(anchor="w", pady=(4, 0))

        acoes = ttk.Frame(container)
        acoes.pack(fill="x", pady=(0, 12))

        ttk.Button(acoes, text="Executar roteiro", command=self.executar_roteiro_pelo_botao, style="Primary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Executar em loop", command=self.executar_roteiro_em_loop_pelo_botao, style="Primary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Parar execucao", command=self.parar_execucao_roteiro).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, textvariable=self.botao_pausa_var, command=self.alternar_pausa_captura).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Texto grande (F8)", command=self.abrir_texto_pelo_botao, style="Primary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="↑ Mover acima", command=lambda: self.mover_evento_selecionado(-1)).pack(side="left", padx=(0, 4))
        ttk.Button(acoes, text="↓ Mover abaixo", command=lambda: self.mover_evento_selecionado(1)).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Apagar selecionado", command=self.remover_evento_selecionado).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Apagar ultima entrada", command=self.remover_ultima_entrada).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Mostrar informacoes da planilha", command=self.recarregar_planilha).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Gerar build", command=self.gerar_build_automacao, style="Primary.TButton").pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Salvar arquivos", command=self.salvar_e_refrescar).pack(side="left", padx=(0, 8))
        ttk.Button(acoes, text="Fechar programa", command=self.encerrar).pack(side="right")

        status_card = ttk.Frame(container, style="Card.TFrame", padding=12)
        status_card.pack(fill="x", pady=(0, 12))
        ttk.Label(status_card, textvariable=self.status_var, background="#FFFFFF", foreground="#17324D", font=("Segoe UI Semibold", 10)).pack(anchor="w")
        linha_execucao = ttk.Frame(status_card, style="Card.TFrame")
        linha_execucao.pack(fill="x", pady=(8, 0))
        ttk.Label(linha_execucao, textvariable=self.status_execucao_var, background="#FFFFFF", foreground="#5B677A", font=("Segoe UI", 9)).pack(side="left")
        ttk.Label(linha_execucao, text="Atraso antes de executar:", background="#FFFFFF", foreground="#5B677A", font=("Segoe UI", 9)).pack(side="left", padx=(16, 6))
        ttk.Spinbox(linha_execucao, from_=0, to=15, increment=0.5, textvariable=self.atraso_execucao_var, width=6).pack(side="left")
        ttk.Label(
            status_card,
            textvariable=self.execucao_detalhe_var,
            background="#FFFFFF",
            foreground="#35506B",
            font=("Consolas", 9),
            wraplength=980,
            justify="left",
        ).pack(anchor="w", pady=(6, 0))
        ttk.Label(
            status_card,
            text=f"Arquivos: {self.arquivo_saida.resolve()} | {self.arquivo_script.resolve()}",
            background="#FFFFFF",
            foreground="#5B677A",
            font=("Segoe UI", 9),
        ).pack(anchor="w", pady=(4, 0))

        self.notebook = ttk.Notebook(container)
        self.notebook.pack(fill="both", expand=True)

        self.aba_captura = ttk.Frame(self.notebook, padding=12)
        self.aba_planilha = ttk.Frame(self.notebook, padding=12)
        self.notebook.add(self.aba_captura, text="Captura")
        self.notebook.add(self.aba_planilha, text="Planilha")

        self.construir_aba_captura()
        self.construir_aba_planilha()

        self.root.protocol("WM_DELETE_WINDOW", self.encerrar)

    def construir_aba_captura(self):
        painel = ttk.PanedWindow(self.aba_captura, orient="horizontal")
        painel.pack(fill="both", expand=True)

        esquerda = ttk.Frame(painel)
        direita = ttk.Frame(painel)
        painel.add(esquerda, weight=3)
        painel.add(direita, weight=2)

        quadro_eventos = ttk.LabelFrame(esquerda, text="Roteiro gravado", style="Section.TLabelframe", padding=10)
        quadro_eventos.pack(fill="both", expand=True)

        colunas = ("ordem", "tipo", "conteudo", "intervalo", "comando")
        self.lista_eventos = ttk.Treeview(quadro_eventos, columns=colunas, show="headings")
        self.lista_eventos.heading("ordem", text="#")
        self.lista_eventos.heading("tipo", text="Tipo")
        self.lista_eventos.heading("conteudo", text="Conteudo")
        self.lista_eventos.heading("intervalo", text="Intervalo(s)")
        self.lista_eventos.heading("comando", text="Comando")
        self.lista_eventos.column("ordem", width=55, anchor="center")
        self.lista_eventos.column("tipo", width=140, anchor="w")
        self.lista_eventos.column("conteudo", width=300, anchor="w")
        self.lista_eventos.column("intervalo", width=95, anchor="center")
        self.lista_eventos.column("comando", width=410, anchor="w")

        scroll_y = ttk.Scrollbar(quadro_eventos, orient="vertical", command=self.lista_eventos.yview)
        scroll_x = ttk.Scrollbar(quadro_eventos, orient="horizontal", command=self.lista_eventos.xview)
        self.lista_eventos.configure(yscrollcommand=scroll_y.set, xscrollcommand=scroll_x.set)
        self.lista_eventos.tag_configure("executando", background="#DBFCE7", foreground="#123728")
        self.lista_eventos.bind("<<TreeviewSelect>>", self.ao_selecionar_evento)
        self.lista_eventos.bind("<Delete>", lambda _: self.remover_evento_selecionado())
        self.lista_eventos.bind("<Alt-Up>", lambda _: self.mover_evento_selecionado(-1))
        self.lista_eventos.bind("<Alt-Down>", lambda _: self.mover_evento_selecionado(1))
        self.lista_eventos.pack(fill="both", expand=True, side="left")
        scroll_y.pack(fill="y", side="right")
        scroll_x.pack(fill="x", side="bottom")

        editor_tempo = ttk.LabelFrame(esquerda, text="Editar tempos do roteiro", style="Section.TLabelframe", padding=10)
        editor_tempo.pack(fill="x", pady=(10, 0))

        ttk.Label(editor_tempo, textvariable=self.evento_tempo_info_var, style="Sub.TLabel").pack(anchor="w")

        linha_tempo = ttk.Frame(editor_tempo)
        linha_tempo.pack(fill="x", pady=(10, 0))
        ttk.Label(linha_tempo, text="Tempo antes do evento:").pack(side="left")
        ttk.Spinbox(
            linha_tempo,
            from_=0,
            to=300,
            increment=0.1,
            textvariable=self.intervalo_evento_var,
            width=8,
        ).pack(side="left", padx=(8, 6))
        ttk.Label(linha_tempo, text="segundos").pack(side="left")

        linha_tempo_planilha = ttk.Frame(editor_tempo)
        linha_tempo_planilha.pack(fill="x", pady=(8, 0))
        ttk.Label(linha_tempo_planilha, text="Tempo entre itens da sequencia:").pack(side="left")
        self.entrada_intervalo_planilha_evento = ttk.Spinbox(
            linha_tempo_planilha,
            from_=0,
            to=300,
            increment=0.1,
            textvariable=self.intervalo_planilha_evento_var,
            width=8,
        )
        self.entrada_intervalo_planilha_evento.pack(side="left", padx=(8, 6))
        ttk.Label(linha_tempo_planilha, text="segundos").pack(side="left")

        ttk.Button(
            editor_tempo,
            text="Salvar tempo do item selecionado",
            command=self.aplicar_tempo_evento_selecionado,
            style="Primary.TButton",
        ).pack(anchor="e", pady=(10, 0))

        quadro_mouse = ttk.LabelFrame(direita, text="Captura de mouse", style="Section.TLabelframe", padding=10)
        quadro_mouse.pack(fill="x", pady=(0, 10))

        ttk.Label(
            quadro_mouse,
            text="Clique no botao abaixo e depois clique no ponto exato da tela que deve ser repetido no roteiro.",
            style="Sub.TLabel",
            wraplength=360,
            justify="left",
        ).pack(anchor="w")
        ttk.Label(
            quadro_mouse,
            textvariable=self.mouse_posicao_var,
            background="#FFFFFF",
            foreground="#17324D",
            font=("Consolas", 10),
        ).pack(anchor="w", pady=(8, 10))
        ttk.Button(
            quadro_mouse,
            text="Capturar mouse",
            command=self.capturar_mouse_pelo_botao,
            style="Primary.TButton",
        ).pack(anchor="w")

        quadro_resumo = ttk.LabelFrame(direita, text="Resumo digitado", style="Section.TLabelframe", padding=10)
        quadro_resumo.pack(fill="both", expand=True, pady=(0, 10))

        self.resumo_texto = tk.Text(
            quadro_resumo,
            wrap="word",
            height=10,
            font=("Consolas", 10),
            bg="#FCFCFD",
            fg="#1D2939",
            relief="flat",
        )
        self.resumo_texto.pack(fill="both", expand=True)
        self.resumo_texto.configure(state="disabled")

        quadro_log = ttk.LabelFrame(direita, text="Atividade e exclusoes", style="Section.TLabelframe", padding=10)
        quadro_log.pack(fill="both", expand=True)

        self.log_texto = tk.Text(
            quadro_log,
            wrap="word",
            font=("Consolas", 10),
            bg="#0F172A",
            fg="#E2E8F0",
            relief="flat",
        )
        self.log_texto.pack(fill="both", expand=True)
        self.log_texto.tag_configure("info", foreground="#CBD5E1")
        self.log_texto.tag_configure("ok", foreground="#93C5FD")
        self.log_texto.tag_configure("removido", foreground="#FCA5A5")
        self.log_texto.tag_configure("alerta", foreground="#FDE68A")
        self.log_texto.configure(state="disabled")

    def construir_aba_planilha(self):
        topo = ttk.Frame(self.aba_planilha)
        topo.pack(fill="x", pady=(0, 10))

        ttk.Label(topo, text=f"Arquivo da planilha: {self.arquivo_planilha.resolve()}", style="Sub.TLabel").pack(anchor="w")
        ttk.Label(topo, textvariable=self.planilha_info_var, style="Sub.TLabel").pack(anchor="w", pady=(4, 0))

        filtros = ttk.Frame(self.aba_planilha)
        filtros.pack(fill="x", pady=(0, 10))

        ttk.Label(filtros, text="Aba da planilha:").pack(side="left")
        self.combo_abas = ttk.Combobox(filtros, textvariable=self.planilha_aba_var, state="readonly", width=24)
        self.combo_abas.pack(side="left", padx=(8, 14))
        self.combo_abas.bind("<<ComboboxSelected>>", self.alterar_aba_planilha)

        ttk.Label(filtros, text="Coluna:").pack(side="left")
        self.combo_colunas = ttk.Combobox(filtros, textvariable=self.coluna_planilha_var, state="readonly", width=28)
        self.combo_colunas.pack(side="left", padx=(8, 8))
        self.combo_colunas.bind("<<ComboboxSelected>>", self.mostrar_valores_coluna_evento)

        ttk.Button(filtros, text="Mostrar coluna", command=self.mostrar_valores_coluna).pack(side="left", padx=(0, 8))
        ttk.Button(filtros, text="Adicionar coluna toda", command=self.adicionar_coluna_ao_roteiro).pack(side="left", padx=(0, 8))
        ttk.Button(filtros, text="Adicionar linhas selecionadas", command=self.adicionar_linhas_selecionadas_ao_roteiro).pack(side="left")

        painel = ttk.PanedWindow(self.aba_planilha, orient="horizontal")
        painel.pack(fill="both", expand=True)

        esquerda = ttk.Frame(painel)
        direita = ttk.Frame(painel)
        painel.add(esquerda, weight=3)
        painel.add(direita, weight=2)

        quadro_linhas = ttk.LabelFrame(esquerda, text="Informacoes disponiveis na planilha", style="Section.TLabelframe", padding=10)
        quadro_linhas.pack(fill="both", expand=True)

        self.tabela_planilha = ttk.Treeview(quadro_linhas, show="headings")
        scroll_planilha_y = ttk.Scrollbar(quadro_linhas, orient="vertical", command=self.tabela_planilha.yview)
        scroll_planilha_x = ttk.Scrollbar(quadro_linhas, orient="horizontal", command=self.tabela_planilha.xview)
        self.tabela_planilha.configure(yscrollcommand=scroll_planilha_y.set, xscrollcommand=scroll_planilha_x.set)
        self.tabela_planilha.pack(fill="both", expand=True, side="left")
        scroll_planilha_y.pack(fill="y", side="right")
        scroll_planilha_x.pack(fill="x", side="bottom")

        quadro_coluna = ttk.LabelFrame(direita, text="Percorrer linhas em sequencia", style="Section.TLabelframe", padding=10)
        quadro_coluna.pack(fill="both", expand=True)

        ttk.Label(quadro_coluna, textvariable=self.coluna_info_var, style="Sub.TLabel").pack(anchor="w")

        opcoes = ttk.Frame(quadro_coluna)
        opcoes.pack(fill="x", pady=(10, 10))

        linha_acao = ttk.Frame(opcoes)
        linha_acao.pack(fill="x")
        ttk.Label(linha_acao, text="Tecla ou atalho apos cada item:").pack(side="left")
        self.combo_acao_apos_item = ttk.Combobox(
            linha_acao,
            textvariable=self.acao_apos_item_var,
            width=18,
            values=("enter", "tab", "down", "up", "ctrl+f", "nenhuma"),
        )
        self.combo_acao_apos_item.pack(side="left", padx=(8, 0))

        linha_intervalo = ttk.Frame(opcoes)
        linha_intervalo.pack(fill="x", pady=(10, 0))
        ttk.Label(linha_intervalo, text="Intervalo entre itens da sequencia:").pack(side="left")
        self.entrada_intervalo = ttk.Spinbox(
            linha_intervalo,
            from_=0,
            to=60,
            increment=0.1,
            textvariable=self.planilha_intervalo_var,
            width=8,
        )
        self.entrada_intervalo.pack(side="left", padx=(8, 0))
        ttk.Label(linha_intervalo, text="segundos").pack(side="left", padx=(6, 0))

        ttk.Label(
            quadro_coluna,
            text="Voce pode adicionar a coluna toda ou selecionar linhas especificas com Ctrl/Shift.",
            style="Sub.TLabel",
        ).pack(anchor="w", pady=(0, 10))

        botoes_selecao = ttk.Frame(quadro_coluna)
        botoes_selecao.pack(fill="x", pady=(0, 8))
        ttk.Button(botoes_selecao, text="Selecionar tudo", command=self.selecionar_todas_as_linhas_coluna).pack(side="left", padx=(0, 8))
        ttk.Button(botoes_selecao, text="Limpar selecao", command=self.limpar_selecao_linhas_coluna).pack(side="left")

        self.lista_coluna = ttk.Treeview(
            quadro_coluna,
            columns=("linha", "valor"),
            show="headings",
            height=18,
            selectmode="extended",
        )
        self.lista_coluna.heading("linha", text="Linha")
        self.lista_coluna.heading("valor", text="Valor da coluna")
        self.lista_coluna.column("linha", width=70, anchor="center")
        self.lista_coluna.column("valor", width=260, anchor="w")
        scroll_coluna = ttk.Scrollbar(quadro_coluna, orient="vertical", command=self.lista_coluna.yview)
        self.lista_coluna.configure(yscrollcommand=scroll_coluna.set)
        self.lista_coluna.pack(fill="both", expand=True, side="left")
        scroll_coluna.pack(fill="y", side="right")

    def salvar_e_refrescar(self):
        self.salvar_arquivo()
        self.atualizar_interface_eventos()
        self.adicionar_log("Arquivos atualizados manualmente.", "info")

    def alternar_pausa_captura(self):
        self.pausa_manual = not self.pausa_manual
        if self.pausa_manual:
            self.botao_pausa_var.set("Retomar captura")
            self.adicionar_log("Captura pausada manualmente. Nada sera gravado ate retomar.", "alerta")
        else:
            self.botao_pausa_var.set("Pausar captura")
            self.adicionar_log("Captura retomada.", "ok")
        self.atualizar_interface_eventos()

    def abrir_texto_pelo_botao(self):
        if self.captura_pausada:
            return
        self.captura_pausada = True
        self.resetar_estado_captura()
        self.abrir_janela_texto()

    def capturar_mouse_pelo_botao(self):
        if self.capturando_mouse:
            self.adicionar_log("Ja existe uma captura de mouse aguardando clique.", "alerta")
            return

        if self.executando_roteiro:
            self.adicionar_log("Aguarde a automacao terminar antes de capturar um clique do mouse.", "alerta")
            return

        self.capturando_mouse = True
        self.captura_pausada = True
        self.resetar_estado_captura()
        self.mouse_posicao_var.set("Aguardando clique na tela...")
        self.adicionar_log(
            "Captura de mouse iniciada. Clique no ponto exato da tela para salvar esse clique no roteiro.",
            "info",
        )

        try:
            self.root.iconify()
        except tk.TclError:
            pass

        threading.Thread(target=self.capturar_posicao_mouse, daemon=True).start()

    def capturar_posicao_mouse(self):
        time.sleep(0.2)

        try:
            def ao_clicar(x, y, button, pressed):
                if not pressed:
                    return True

                nome_botao = getattr(button, "name", str(button).replace("Button.", ""))
                self.enfileirar_ui(
                    "mouse_capturado",
                    {
                        "x": int(x),
                        "y": int(y),
                        "botao": nome_botao,
                    },
                )
                return False

            self.mouse_listener = pynput_mouse.Listener(on_click=ao_clicar)
            self.mouse_listener.start()
            self.mouse_listener.join()
        except Exception as erro:
            self.enfileirar_ui("mouse_erro", {"mensagem": f"Falha ao capturar o mouse: {erro}"})
        finally:
            self.mouse_listener = None

    def concluir_captura_mouse(self, payload):
        self.capturando_mouse = False
        self.captura_pausada = False

        x = int(payload["x"])
        y = int(payload["y"])
        botao = str(payload.get("botao", "left"))
        self.mouse_posicao_var.set(f"Ultimo clique salvo: {botao.upper()} em X={x} | Y={y}")
        self.registrar_evento(["mouse"], "mouse_click", metadados={"x": x, "y": y, "botao": botao})
        self.adicionar_log(
            f"Clique do mouse salvo no roteiro: botao {botao.upper()} em X={x}, Y={y}.",
            "ok",
        )

        try:
            self.root.deiconify()
            self.root.lift()
        except tk.TclError:
            pass

    def finalizar_captura_mouse_com_erro(self, mensagem):
        self.capturando_mouse = False
        self.captura_pausada = False
        self.mouse_posicao_var.set("Nenhum clique do mouse capturado ainda.")
        self.adicionar_log(mensagem, "removido")

        try:
            self.root.deiconify()
            self.root.lift()
        except tk.TclError:
            pass

    def executar_roteiro_pelo_botao(self):
        if self.executando_roteiro:
            self.adicionar_log("Ja existe uma automacao em andamento.", "alerta")
            return

        if not self.eventos:
            messagebox.showwarning("Automacao", "Nenhum evento foi gravado ainda.")
            return

        try:
            atraso = max(0, round(float(self.atraso_execucao_var.get()), 2))
        except (tk.TclError, ValueError):
            messagebox.showwarning("Automacao", "Informe um atraso valido antes de executar.")
            return

        self.atualizar_cache_planilha_silenciosamente()
        self.executando_roteiro = True
        self.captura_pausada = True
        self.cancelar_execucao.clear()
        self.atualizar_status_execucao("Preparando automacao...")
        self.adicionar_log(
            f"Automacao iniciada. Voce tem {atraso}s para focar a janela onde o roteiro sera executado.",
            "ok",
        )

        try:
            self.root.iconify()
        except tk.TclError:
            pass

        threading.Thread(
            target=self.executar_roteiro,
            args=(atraso,),
            daemon=True,
        ).start()

    def parar_execucao_roteiro(self):
        if not self.executando_roteiro:
            self.adicionar_log("Nenhuma automacao esta em andamento.", "alerta")
            return

        self.cancelar_execucao.set()
        self.adicionar_log("Pedido de parada da automacao enviado.", "alerta")
        self.atualizar_status_execucao("Parando automacao...")

    def esperar_com_cancelamento(self, segundos):
        fim = time.perf_counter() + max(0, segundos)
        while time.perf_counter() < fim:
            if self.cancelar_execucao.is_set() or self.encerrando:
                return False
            time.sleep(0.05)
        return True

    def executar_acao_automacao(self, acao):
        if not acao:
            return

        if isinstance(acao, str):
            partes = self.normalizar_acao_apos_item(acao)
        else:
            partes = [str(parte).strip() for parte in acao if str(parte).strip()]

        if not partes:
            return

        if len(partes) == 1:
            pyautogui.press(partes[0])
            return

        pyautogui.hotkey(*partes)

    def executar_evento_roteiro(self, evento, iteracao_atual=None, total_iteracoes=None):
        tipo = evento.get("tipo")

        if tipo == "texto_livre":
            self.escrever_texto(evento.get("texto", ""))
            return True

        if tipo in TIPOS_SEQUENCIA_PLANILHA:
            detalhes = evento.get("detalhes", {})
            valores = detalhes.get("valores", [])
            acao_apos_item = detalhes.get("acao_apos_item")
            if acao_apos_item is None and detalhes.get("pressionar_enter"):
                acao_apos_item = ["enter"]
            intervalo_entre_itens = detalhes.get(
                "intervalo_entre_itens",
                INTERVALO_ENTRE_ITENS_PLANILHA,
            )

            total = len(valores)
            for indice, valor in enumerate(valores):
                self.sinalizar_evento_em_execucao(
                    evento,
                    indice_item=indice,
                    total_itens=total,
                    valor_atual=str(valor),
                    iteracao_atual=iteracao_atual,
                    total_iteracoes=total_iteracoes,
                    registrar_log=True,
                )
                if self.cancelar_execucao.is_set() or self.encerrando:
                    return False
                self.escrever_texto(str(valor))
                if indice < total - 1 and acao_apos_item:
                    self.executar_acao_automacao(acao_apos_item)
                if indice < total - 1 and intervalo_entre_itens > 0:
                    if not self.esperar_com_cancelamento(intervalo_entre_itens):
                        return False
            return True

        if tipo in TIPOS_MOUSE:
            detalhes = evento.get("detalhes", {})
            x = int(detalhes.get("x", 0))
            y = int(detalhes.get("y", 0))
            botao = str(detalhes.get("botao", "left"))
            pyautogui.click(x=x, y=y, button=botao)
            return True

        teclas = evento.get("teclas", [])
        if not teclas:
            return True

        if len(teclas) > 1:
            pyautogui.hotkey(*[self.tecla_para_pyautogui(tecla) for tecla in teclas])
            return True

        tecla = teclas[0]
        if self.eh_texto_simples(tecla):
            pyautogui.write(tecla, interval=INTERVALO_TEXTO)
            return True

        pyautogui.press(self.tecla_para_pyautogui(tecla))
        return True

    def finalizar_execucao_roteiro(self, mensagem, tag="info"):
        self.executando_roteiro = False
        self.captura_pausada = False
        self.cancelar_execucao.clear()
        self.enfileirar_ui("log", {"mensagem": mensagem, "tag": tag})
        self.enfileirar_ui("status_execucao", {"mensagem": mensagem})
        self.enfileirar_ui("restaurar_janela", None)

    def executar_roteiro(self, atraso):
        try:
            if atraso > 0 and not self.esperar_com_cancelamento(atraso):
                self.finalizar_execucao_roteiro("Automacao cancelada antes de iniciar.", "alerta")
                return

            self.enfileirar_ui("status_execucao", {"mensagem": "Executando roteiro..."})
            self.enfileirar_ui(
                "log",
                {"mensagem": "Roteiro em execucao na janela selecionada.", "tag": "ok"},
            )

            for evento in self.eventos:
                if self.cancelar_execucao.is_set() or self.encerrando:
                    self.finalizar_execucao_roteiro("Automacao interrompida pelo usuario.", "alerta")
                    return

                intervalo_antes = evento.get("intervalo_antes", 0)
                if intervalo_antes > 0:
                    self.sinalizar_evento_em_execucao(evento, aguardando_intervalo=True)
                if intervalo_antes > 0 and not self.esperar_com_cancelamento(intervalo_antes):
                    self.finalizar_execucao_roteiro("Automacao interrompida durante a espera.", "alerta")
                    return

                if evento.get("tipo") not in TIPOS_SEQUENCIA_PLANILHA:
                    self.sinalizar_evento_em_execucao(evento, registrar_log=True)

                resultado = self.executar_evento_roteiro(evento)
                if resultado is False:
                    self.finalizar_execucao_roteiro("Automacao interrompida pelo usuario.", "alerta")
                    return

            self.finalizar_execucao_roteiro("Automacao concluida com sucesso.", "ok")
        except Exception as erro:
            self.finalizar_execucao_roteiro(f"Falha na automacao: {erro}", "removido")

    def executar_roteiro_em_loop_pelo_botao(self):
        if self.executando_roteiro:
            self.adicionar_log("Ja existe uma automacao em andamento.", "alerta")
            return

        if not self.eventos:
            messagebox.showwarning("Automacao", "Nenhum evento foi gravado ainda.")
            return

        evento_planilha = next(
            (e for e in self.eventos if e.get("tipo") in TIPOS_SEQUENCIA_PLANILHA),
            None,
        )
        if evento_planilha is None:
            messagebox.showwarning(
                "Loop",
                "Nenhuma sequencia de planilha encontrada no roteiro.\n"
                "Adicione uma coluna da planilha ao roteiro antes de usar o loop.",
            )
            return

        total_iteracoes = len(evento_planilha.get("detalhes", {}).get("valores", []))
        if total_iteracoes == 0:
            messagebox.showwarning("Loop", "A sequencia de planilha no roteiro nao possui valores.")
            return

        try:
            atraso = max(0, round(float(self.atraso_execucao_var.get()), 2))
        except (tk.TclError, ValueError):
            messagebox.showwarning("Automacao", "Informe um atraso valido antes de executar.")
            return

        self.atualizar_cache_planilha_silenciosamente()
        self.executando_roteiro = True
        self.captura_pausada = True
        self.cancelar_execucao.clear()
        self.atualizar_status_execucao(f"Preparando loop: {total_iteracoes} iteracoes...")
        self.adicionar_log(
            f"Loop iniciado com {total_iteracoes} iteracoes. Voce tem {atraso}s para focar a janela.",
            "ok",
        )

        try:
            self.root.iconify()
        except tk.TclError:
            pass

        threading.Thread(
            target=self.executar_roteiro_em_loop,
            args=(atraso, total_iteracoes),
            daemon=True,
        ).start()

    def executar_roteiro_em_loop(self, atraso, total_iteracoes):
        try:
            if atraso > 0 and not self.esperar_com_cancelamento(atraso):
                self.finalizar_execucao_roteiro("Loop cancelado antes de iniciar.", "alerta")
                return

            for indice in range(total_iteracoes):
                if self.cancelar_execucao.is_set() or self.encerrando:
                    self.finalizar_execucao_roteiro("Loop interrompido pelo usuario.", "alerta")
                    return

                self.enfileirar_ui(
                    "status_execucao",
                    {"mensagem": f"Loop: iteracao {indice + 1} de {total_iteracoes}..."},
                )
                self.enfileirar_ui(
                    "log",
                    {"mensagem": f"--- Iteracao {indice + 1} de {total_iteracoes} ---", "tag": "ok"},
                )

                for evento in self.eventos:
                    if self.cancelar_execucao.is_set() or self.encerrando:
                        self.finalizar_execucao_roteiro("Loop interrompido pelo usuario.", "alerta")
                        return

                    intervalo_antes = evento.get("intervalo_antes", 0)
                    if intervalo_antes > 0:
                        self.sinalizar_evento_em_execucao(
                            evento,
                            iteracao_atual=indice + 1,
                            total_iteracoes=total_iteracoes,
                            aguardando_intervalo=True,
                        )
                    if intervalo_antes > 0 and not self.esperar_com_cancelamento(intervalo_antes):
                        self.finalizar_execucao_roteiro("Loop interrompido durante a espera.", "alerta")
                        return

                    if evento.get("tipo") in TIPOS_SEQUENCIA_PLANILHA:
                        detalhes = evento.get("detalhes", {})
                        valores = detalhes.get("valores", [])
                        if indice < len(valores):
                            acao_apos_item = detalhes.get("acao_apos_item")
                            if acao_apos_item is None and detalhes.get("pressionar_enter"):
                                acao_apos_item = ["enter"]
                            self.sinalizar_evento_em_execucao(
                                evento,
                                indice_item=indice,
                                total_itens=len(valores),
                                valor_atual=str(valores[indice]),
                                iteracao_atual=indice + 1,
                                total_iteracoes=total_iteracoes,
                                registrar_log=True,
                            )
                            self.escrever_texto(str(valores[indice]))
                            if acao_apos_item:
                                self.executar_acao_automacao(acao_apos_item)
                        continue

                    self.sinalizar_evento_em_execucao(
                        evento,
                        iteracao_atual=indice + 1,
                        total_iteracoes=total_iteracoes,
                        registrar_log=True,
                    )
                    resultado = self.executar_evento_roteiro(
                        evento,
                        iteracao_atual=indice + 1,
                        total_iteracoes=total_iteracoes,
                    )
                    if resultado is False:
                        self.finalizar_execucao_roteiro("Loop interrompido pelo usuario.", "alerta")
                        return

            self.finalizar_execucao_roteiro(
                f"Loop concluido com sucesso: {total_iteracoes} iteracoes realizadas.", "ok"
            )
        except Exception as erro:
            self.finalizar_execucao_roteiro(f"Falha no loop: {erro}", "removido")

    def abrir_janela_texto(self):
        if self.digitando.locked():
            self.adicionar_log("Aguarde a digitacao atual terminar.", "alerta")
            self.captura_pausada = False
            return

        if self.janela_texto is not None and self.janela_texto.winfo_exists():
            self.janela_texto.lift()
            self.janela_texto.focus_force()
            self.campo_texto.focus_set()
            return

        self.janela_anterior = self.capturar_janela_ativa()
        self.janela_texto = tk.Toplevel(self.root)
        self.janela_texto.title("Texto grande")
        self.janela_texto.geometry("660x420")
        self.janela_texto.minsize(520, 320)
        self.janela_texto.configure(bg="#F8FAFC")
        self.janela_texto.attributes("-topmost", True)
        self.janela_texto.transient(self.root)
        self.janela_texto.protocol("WM_DELETE_WINDOW", self.cancelar_texto)

        titulo = tk.Label(
            self.janela_texto,
            text="Digite o texto grande que deve ser salvo e depois reproduzido",
            anchor="w",
            padx=14,
            pady=14,
            bg="#F8FAFC",
            fg="#17324D",
            font=("Segoe UI Semibold", 12),
        )
        titulo.pack(fill="x")

        self.campo_texto = tk.Text(
            self.janela_texto,
            wrap="word",
            font=("Consolas", 11),
            undo=True,
            bg="#FFFFFF",
            fg="#111827",
            relief="solid",
            borderwidth=1,
        )
        self.campo_texto.pack(fill="both", expand=True, padx=14, pady=(0, 14))

        rodape = tk.Frame(self.janela_texto, bg="#F8FAFC", padx=14, pady=14)
        rodape.pack(fill="x")

        dica = tk.Label(
            rodape,
            text="Ctrl+Enter confirma | Esc cancela",
            anchor="w",
            bg="#F8FAFC",
            fg="#55657A",
            font=("Segoe UI", 10),
        )
        dica.pack(side="left")

        tk.Button(
            rodape,
            text="Cancelar",
            command=self.cancelar_texto,
            bg="#E5E7EB",
            fg="#111827",
            relief="flat",
            padx=12,
            pady=8,
        ).pack(side="right")
        tk.Button(
            rodape,
            text="Confirmar e digitar",
            command=self.confirmar_texto,
            bg="#1D4ED8",
            fg="#FFFFFF",
            relief="flat",
            padx=12,
            pady=8,
        ).pack(side="right", padx=(0, 8))

        self.janela_texto.bind("<Control-Return>", self.confirmar_texto)
        self.janela_texto.bind("<Escape>", self.cancelar_texto)
        self.janela_texto.after(50, self.campo_texto.focus_set)

    def fechar_janela_texto(self):
        if self.janela_texto is not None and self.janela_texto.winfo_exists():
            self.janela_texto.destroy()

        self.janela_texto = None
        self.campo_texto = None

    def cancelar_texto(self, event=None):
        self.fechar_janela_texto()
        self.captura_pausada = False
        self.adicionar_log("Entrada de texto grande cancelada.", "alerta")
        return "break"

    def confirmar_texto(self, event=None):
        if self.janela_texto is None or not self.janela_texto.winfo_exists():
            return "break"

        texto = self.campo_texto.get("1.0", "end-1c")
        if not texto.strip():
            self.adicionar_log("Nenhum texto informado para o bloco grande.", "alerta")
            self.campo_texto.focus_set()
            return "break"

        self.registrar_evento(["texto_livre"], "texto_livre", texto_livre=texto)
        self.fechar_janela_texto()
        threading.Thread(
            target=self.digitar_texto_livre,
            args=(texto,),
            daemon=True,
        ).start()
        return "break"

    def digitar_texto_livre(self, texto):
        if not self.digitando.acquire(blocking=False):
            self.enfileirar_ui("log", {"mensagem": "Ja existe uma digitacao em andamento.", "tag": "alerta"})
            self.captura_pausada = False
            return

        try:
            self.restaurar_janela_ativa(self.janela_anterior)
            time.sleep(PAUSA_ANTES_DE_DIGITAR)
            self.escrever_texto(texto)
            self.enfileirar_ui("log", {"mensagem": "Texto grande digitado com sucesso.", "tag": "ok"})
        finally:
            self.captura_pausada = False
            self.digitando.release()

    def escrever_texto(self, texto):
        for caractere in texto:
            if caractere == "\r":
                continue
            if caractere == "\n":
                pyautogui.press("enter")
                continue
            if caractere == "\t":
                pyautogui.press("tab")
                continue
            pyautogui.write(caractere, interval=INTERVALO_TEXTO)

    def carregar_planilha(self):
        if not self.arquivo_planilha.exists():
            self.planilha_cache = {}
            self.planilha_ativa = None
            self.planilha_colunas = []
            self.planilha_linhas = []
            return False, f"Arquivo nao encontrado: {self.arquivo_planilha.resolve()}"

        try:
            workbook = load_workbook(self.arquivo_planilha, read_only=True, data_only=True)
        except Exception as erro:
            return False, f"Falha ao abrir a planilha: {erro}"

        cache = {}
        for nome_aba in workbook.sheetnames:
            worksheet = workbook[nome_aba]
            linhas = list(worksheet.iter_rows(values_only=True))
            if not linhas:
                cache[nome_aba] = {"colunas": [], "linhas": []}
                continue

            cabecalhos = [str(valor).strip() if valor is not None else f"Coluna {indice + 1}" for indice, valor in enumerate(linhas[0])]
            registros = []
            for numero_linha, linha in enumerate(linhas[1:], start=2):
                registro = {"__linha__": numero_linha}
                for indice, coluna in enumerate(cabecalhos):
                    registro[coluna] = linha[indice] if indice < len(linha) else None
                registros.append(registro)
            cache[nome_aba] = {"colunas": cabecalhos, "linhas": registros}

        workbook.close()
        self.planilha_cache = cache
        self.planilha_ativa = workbook.sheetnames[0] if workbook.sheetnames else None
        if self.planilha_ativa:
            self.planilha_colunas = cache[self.planilha_ativa]["colunas"]
            self.planilha_linhas = cache[self.planilha_ativa]["linhas"]
        else:
            self.planilha_colunas = []
            self.planilha_linhas = []
        return True, f"{len(self.planilha_linhas)} linhas encontradas em {self.planilha_ativa or 'nenhuma aba'}."

    def recarregar_planilha(self):
        sucesso, mensagem = self.carregar_planilha()
        if sucesso:
            self.preencher_planilha_ui()
            self.adicionar_log(f"Planilha carregada: {mensagem}", "info")
        else:
            self.planilha_info_var.set(mensagem)
            self.adicionar_log(mensagem, "removido")
            messagebox.showerror("Planilha", mensagem)

    def preencher_planilha_ui(self):
        abas = list(self.planilha_cache.keys())
        self.combo_abas["values"] = abas

        if not abas:
            self.planilha_info_var.set("Nenhuma aba disponivel na planilha.")
            self.combo_colunas["values"] = []
            self.atualizar_tabela_planilha()
            self.atualizar_lista_coluna([])
            return

        if self.planilha_ativa not in abas:
            self.planilha_ativa = abas[0]

        self.planilha_aba_var.set(self.planilha_ativa)
        dados_aba = self.planilha_cache[self.planilha_ativa]
        self.planilha_colunas = dados_aba["colunas"]
        self.planilha_linhas = dados_aba["linhas"]
        self.combo_colunas["values"] = self.planilha_colunas

        if self.coluna_planilha_var.get() not in self.planilha_colunas:
            coluna_padrao = "Id Alterdata" if "Id Alterdata" in self.planilha_colunas else (self.planilha_colunas[0] if self.planilha_colunas else "")
            self.coluna_planilha_var.set(coluna_padrao)

        self.planilha_info_var.set(
            f"Aba atual: {self.planilha_ativa} | {len(self.planilha_linhas)} linhas | {len(self.planilha_colunas)} colunas"
        )
        self.atualizar_tabela_planilha()
        self.mostrar_valores_coluna()

    def alterar_aba_planilha(self, event=None):
        aba = self.planilha_aba_var.get()
        if aba not in self.planilha_cache:
            return
        self.planilha_ativa = aba
        self.preencher_planilha_ui()

    def atualizar_tabela_planilha(self):
        colunas = ["__linha__"] + self.planilha_colunas
        self.tabela_planilha.delete(*self.tabela_planilha.get_children())
        self.tabela_planilha["columns"] = colunas

        for coluna in colunas:
            titulo = "Linha" if coluna == "__linha__" else coluna
            largura = 75 if coluna == "__linha__" else 180
            self.tabela_planilha.heading(coluna, text=titulo)
            self.tabela_planilha.column(coluna, width=largura, anchor="w")

        for linha in self.planilha_linhas:
            valores = [linha.get("__linha__", "")]
            valores.extend("" if linha.get(coluna) is None else linha.get(coluna) for coluna in self.planilha_colunas)
            self.tabela_planilha.insert("", "end", values=valores)

    def obter_valores_coluna(self, coluna):
        valores = []
        for linha in self.planilha_linhas:
            valor = linha.get(coluna)
            if valor is None or str(valor).strip() == "":
                continue
            valores.append((linha.get("__linha__", ""), valor))
        return valores

    def atualizar_lista_coluna(self, valores):
        self.lista_coluna.delete(*self.lista_coluna.get_children())
        for linha, valor in valores:
            self.lista_coluna.insert("", "end", values=(linha, valor))

    def selecionar_todas_as_linhas_coluna(self):
        itens = self.lista_coluna.get_children()
        if itens:
            self.lista_coluna.selection_set(itens)

    def limpar_selecao_linhas_coluna(self):
        self.lista_coluna.selection_remove(self.lista_coluna.selection())

    def obter_configuracao_sequencia_planilha(self, coluna, valores, origem, linhas_selecionadas=None):
        try:
            intervalo_entre_itens = round(float(self.planilha_intervalo_var.get()), 3)
        except (tk.TclError, ValueError):
            messagebox.showwarning("Planilha", "Informe um intervalo valido para a sequencia.")
            return None

        if intervalo_entre_itens < 0:
            messagebox.showwarning("Planilha", "O intervalo entre itens nao pode ser negativo.")
            return None

        acao_texto = self.acao_apos_item_var.get().strip()
        acao_apos_item = self.normalizar_acao_apos_item(acao_texto)

        return {
            "arquivo": str(self.arquivo_planilha),
            "aba": self.planilha_ativa,
            "coluna": coluna,
            "valores": valores,
            "origem": origem,
            "linhas_selecionadas": linhas_selecionadas or [],
            "acao_apos_item": acao_apos_item,
            "acao_apos_item_texto": acao_texto or "nenhuma",
            "intervalo_entre_itens": intervalo_entre_itens,
        }

    def mostrar_valores_coluna_evento(self, event=None):
        self.mostrar_valores_coluna()

    def mostrar_valores_coluna(self):
        coluna = self.coluna_planilha_var.get()
        if not coluna:
            self.coluna_info_var.set("Selecione uma coluna para ver os valores.")
            self.atualizar_lista_coluna([])
            return

        valores = self.obter_valores_coluna(coluna)
        self.atualizar_lista_coluna(valores)
        self.coluna_info_var.set(
            f"Coluna {coluna} com {len(valores)} valores prontos para percorrer em sequencia."
        )

    def adicionar_coluna_ao_roteiro(self):
        coluna = self.coluna_planilha_var.get()
        if not coluna:
            messagebox.showwarning("Planilha", "Selecione uma coluna antes de adicionar ao roteiro.")
            return

        itens_coluna = self.obter_valores_coluna(coluna)
        valores = [valor for _, valor in itens_coluna]
        if not valores:
            messagebox.showwarning("Planilha", f"A coluna {coluna} nao possui valores para adicionar.")
            return

        metadados = self.obter_configuracao_sequencia_planilha(
            coluna,
            valores,
            origem="coluna_inteira",
        )
        if metadados is None:
            return

        self.registrar_evento([f"planilha:{coluna}"], "sequencia_planilha", metadados=metadados)
        self.adicionar_log(
            f"Coluna {coluna} adicionada ao roteiro com {len(valores)} itens em sequencia.",
            "ok",
        )

    def adicionar_linhas_selecionadas_ao_roteiro(self):
        coluna = self.coluna_planilha_var.get()
        if not coluna:
            messagebox.showwarning("Planilha", "Selecione uma coluna antes de montar a sequencia.")
            return

        selecao = self.lista_coluna.selection()
        if not selecao:
            messagebox.showwarning(
                "Planilha",
                "Selecione uma ou mais linhas na lista da coluna para adicionar separadamente.",
            )
            return

        valores = []
        linhas = []
        for item_id in selecao:
            linha, valor = self.lista_coluna.item(item_id, "values")
            linhas.append(int(linha))
            valores.append(valor)

        metadados = self.obter_configuracao_sequencia_planilha(
            coluna,
            valores,
            origem="linhas_selecionadas",
            linhas_selecionadas=linhas,
        )
        if metadados is None:
            return

        self.registrar_evento([f"planilha:{coluna}:selecionadas"], "sequencia_planilha", metadados=metadados)
        self.adicionar_log(
            f"{len(valores)} linhas selecionadas da coluna {coluna} foram adicionadas ao roteiro.",
            "ok",
        )

    def on_press(self, key):
        if self.pausa_manual or self.captura_pausada or self.janela_do_app_esta_ativa():
            return True

        if key == keyboard.Key.f8 and not self.teclas_pressionadas:
            self.captura_pausada = True
            self.resetar_estado_captura()
            self.solicitar_abertura_texto()
            return True

        if key in {keyboard.Key.f9, keyboard.Key.esc} and not self.teclas_pressionadas:
            self.enfileirar_ui("log", {"mensagem": "Encerrando captura.", "tag": "info"})
            self.solicitar_encerramento()
            return False

        if key == keyboard.Key.backspace and not self.teclas_pressionadas:
            self.remover_ultima_entrada()
            return True

        nome = self.normalizar_nome_tecla(key)
        if nome in self.teclas_pressionadas:
            return True

        if self.eh_modificadora(nome):
            self.teclas_pressionadas.add(nome)
            self.adicionar_modificadora(nome)
            return True

        nome = self.normalizar_tecla_principal(key, nome)
        self.teclas_pressionadas.add(nome)
        if hasattr(key, "vk"):
            self.teclas_principais_por_vk[key.vk] = nome
        teclas_evento = [*self.modificadores_ativos, nome]
        tipo_evento = "atalho" if len(teclas_evento) > 1 else "tecla"
        self.registrar_evento(teclas_evento, tipo_evento)
        self.houve_tecla_principal = True
        return True

    def on_release(self, key):
        if self.pausa_manual or self.captura_pausada or self.janela_do_app_esta_ativa():
            return True

        if key in {keyboard.Key.f8, keyboard.Key.f9, keyboard.Key.esc}:
            return True

        nome = self.normalizar_nome_tecla(key)
        if hasattr(key, "vk") and key.vk in self.teclas_principais_por_vk:
            nome = self.teclas_principais_por_vk.pop(key.vk)
        self.teclas_pressionadas.discard(nome)

        if self.eh_modificadora(nome):
            self.remover_modificadora(nome)

        if self.teclas_pressionadas:
            return True

        self.finalizar_sequencia()
        return True

    def encerrar(self):
        if self.encerrando:
            return

        self.encerrando = True
        self.captura_pausada = True
        self.cancelar_execucao.set()
        self.fechar_janela_texto()

        if self.listener is not None:
            self.listener.stop()
        if self.mouse_listener is not None:
            self.mouse_listener.stop()

        if self.root is not None:
            self.root.quit()

    def iniciar(self):
        self.carregar_eventos_existentes()

        self.root = tk.Tk()
        self.construir_interface()
        self.atualizar_interface_eventos()

        sucesso, mensagem = self.carregar_planilha()
        if sucesso:
            self.preencher_planilha_ui()
            self.adicionar_log(f"Planilha carregada: {mensagem}", "info")
        else:
            self.planilha_info_var.set(mensagem)
            self.adicionar_log(mensagem, "removido")

        self.adicionar_log("Captura iniciada. Teclas, atalhos e cliques do mouse podem entrar no roteiro.", "info")
        self.adicionar_log("F8 abre o texto grande. O botao Capturar mouse salva o proximo clique. Backspace remove a ultima entrada. F9 fecha o programa.", "info")

        self.listener = keyboard.Listener(
            on_press=self.on_press,
            on_release=self.on_release,
        )
        self.listener.start()
        self.root.after(100, self.processar_fila)
        self.root.mainloop()
        self.listener.join()


def executar_captura_teclado():
    gravador = GravadorTeclado(ARQUIVO_SAIDA)
    gravador.iniciar()


def iniciar_atalhos_texto():
    executar_captura_teclado()


def main():
    executar_captura_teclado()


if __name__ == "__main__":
    main()
